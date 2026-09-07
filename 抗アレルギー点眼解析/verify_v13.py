# -*- coding: utf-8 -*-
"""v13 の成果物を検証する。

v13 は数値の再計算を伴わないため、gini_cv_bounds.csv 等に対する数値の不変条件は
verify_v12.py に委ね、ここでは
  1. 原稿から消えているべき記述・入っているべき記述
  2. 原稿の数値が Table 2 / Suppl Table S8 / Table 3 と整合すること
  3. 図表ファイルのキャプションに古い数値が残っていないこと
  4. まとめファイルが v13 の名前で揃っていること
を確認する。

実行: python verify_v13.py
"""
import io
import os
import re
import sys

import pandas as pd
from docx import Document

sys.stdout.reconfigure(encoding="utf-8")
BASE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, BASE)
import docx_finalize as F                                # noqa: E402

NEW = os.path.join(BASE, "05_論文成果物", "公費含めない_new")
V13 = os.path.join(NEW, "投稿改定ver", "投稿用v13")
CLEAN = os.path.join(V13, "NDB抗アレルギー解析_v13_clean.docx")
TRACKED = os.path.join(V13, "NDB抗アレルギー解析_v13.docx")
MD = os.path.join(V13, "論文下書き_総量2014_シェア2022_区間解析_投稿改定ver13.md")

ok, ng = [], []


def check(cond, label, detail=""):
    (ok if cond else ng).append(label + (" — " + detail if detail else ""))


def main():
    # ---------------------------------------------------------- 原稿
    a = F.audit(CLEAN)
    check(a["w:ins"] == 0 and a["w:del"] == 0, "clean版に変更履歴なし", str(a))
    check(a["commentReference"] == 0 and not a["comment_parts"], "clean版にコメントなし", str(a))
    check(F.audit(TRACKED)["w:ins"] > 0, "履歴つき版には変更履歴がある")

    doc = Document(CLEAN)
    paras = [p.text for p in doc.paragraphs]
    txt = "\n".join(paras)

    must_go = [
        "同方向である14,15",                       # D 引用番号
        "0.39〜0.43",                              # L 男／女比
        "主要3成分で85%前後",                      # K
        "後発医薬品比率は全期間を図示",            # E
        "最大36ポイント",                          # O
        "上限の過小評価にはつながらない",          # C
        "2023年度は上限未到達",                    # H
        "後発医薬品が2021年度以前非掲載",          # I
        "15歳以上では全",                          # F
        "15歳以上で一貫して",                      # F
        "PMDA薬価基準収載品目表",                  # J
        "セル秘匿の影響は2022年度以降",            # K
        "セル秘匿は2022年度以降の全国集計の結論",  # K
        "収録品目数の多い年度で大きく",            # G
        "処方機会の多さ19",                        # M
        "処方閾値の低下20",                        # M
        "分布しない地域であった4,5,14",            # N（Results 末尾の孤立文）
        "該当はクロモグリク酸ナトリウムのユニットドーズ製剤のみ））",  # A
    ]
    for s in must_go:
        check(s not in txt, f"削除済み: {s}")

    must_have = [
        "ケトチフェンのユニットドーズ製剤も含まれ",
        "補完的秘匿の対象となる構造が存在しない",
        "総計列に補完的秘匿が及ばないという公表仕様の読み",
        "過大評価にはならないが、過小評価となりうる",
        "同方向である15,16",
        "Figure 6Aに2015年度以降を",
        "15〜94歳では全年齢階級で",
        "15〜94歳では全階級で",
        "15〜94歳で一貫して女性優位",
        "処方量の小さい成分が掲載された年度で大きく",
        "適用された閾値を特定できず",
        "2015〜2021年度は一部の品目が",
        "前節に示した突合",
        "総計秘匿による最大0.111%",
        "主要3成分が83〜85%",
        "0.39〜0.42",
        "病態の変化19",
        "有効性が示されているエピナスチン20",
        "関東・東海に位置し、最小の沖縄県",
        "最大35.7ポイント",
        "品目非掲載に関しては点識別",           # J: Methods 第一 に1回は残る
    ]
    for s in must_have:
        check(s in txt, f"追加済み: {s}")

    # 重複文の除去（J）：「品目非掲載に関しては点識別」は1回だけ
    check(txt.count("品目非掲載に関しては点識別") == 1,
          "「品目非掲載に関しては点識別」が1回のみ", str(txt.count("品目非掲載に関しては点識別")))

    # 引用番号の初出順（14 が 15,16 より先に現れる）
    def first(pat):
        m = re.search(pat, txt)
        return m.start() if m else -1
    i14, i15 = first(r"4,5,14"), first(r"15,16")
    check(0 <= i14 < i15, "文献14の初出が15,16より前", f"{i14} < {i15}")

    # ---------------------------------------------------------- 数値の整合
    t2 = pd.read_csv(os.path.join(V13, "Table2.csv"))
    r = t2[t2["Age Group"].isin(["50-54", "55-59"])]
    lo = round(r.MF_per100k_ratio_lower.min(), 2)
    hi = round(r.MF_per100k_ratio_upper.max(), 2)
    check(f"{lo:.2f}〜{hi:.2f}" == "0.39〜0.42", "50〜59歳の男／女比の最低水準", f"{lo:.2f}〜{hi:.2f}")
    r95 = t2[t2["Age Group"].isin(["95-99", "100+"])]
    check((r95.MF_per100k_ratio_upper > 1).all() and (r95.MF_per100k_ratio_lower < 1).all(),
          "95歳以上は男／女比の区間が1をまたぐ（「15〜94歳」限定の根拠）")

    s8 = pd.read_excel(os.path.join(V13, "SupplTableS8.xlsx"), sheet_name="Fig4B_top3")
    w14 = s8[s8["Fiscal year"] == 2014].Share_width_pp.max()
    w21 = s8[s8["Fiscal year"] == 2021].Share_width_pp.max()
    check(f"{w14:.1f}" == "35.7" and f"{w21:.1f}" == "19.0",
          "Fig 4B の最大区間幅（2014: 35.7、2021: 19.0）", f"{w14:.2f}, {w21:.2f}")
    s1c = pd.read_excel(os.path.join(V13, "SupplTableS8.xlsx"), sheet_name="Fig1C_age")
    w04 = s1c[s1c["Age group"] == "0-4"].Share_width_pp
    check(f"{w04.min():.1f}" == "5.0" and f"{w04.max():.1f}" == "8.2",
          "Fig 1C 0〜4歳の区間幅（5.0〜8.2 pp）", f"{w04.min():.2f}〜{w04.max():.2f}")

    t3 = pd.read_csv(os.path.join(V13, "Table3.csv"))
    c = [c for c in t3.columns if c.startswith("Top 3 total")]
    if c:
        lo3, hi3 = t3[c[0]].min(), t3[[x for x in c][-1]].max()
        check(lo3 >= 83.0 and hi3 <= 85.6, "主要3成分シェア 2022〜2024 は 83〜85% の範囲",
              f"{lo3:.1f}〜{hi3:.1f}")

    s2 = pd.read_csv(os.path.join(V13, "SupplTableS2.csv"), comment="#")
    lev = s2[(s2.code == "LEVOCASTINE") & (s2.year <= 2021)]
    all14 = lev[lev.year == 2014].brand_unit__listed_generic.iloc[0] == 0
    part = (lev[lev.year >= 2015].brand_unit__listed_generic > 0).all() and \
           (lev[lev.year >= 2015].brand_unit__unlisted_generic > 0).all()
    check(all14 and part, "レボカバスチン後発: 2014年度は全品目、2015〜2021年度は一部が非掲載")

    # ---------------------------------------------------------- md
    check(os.path.exists(MD), "v13 の md がある")
    if os.path.exists(MD):
        md = io.open(MD, encoding="utf-8").read()
        for s in ["付記15", "15〜94歳", "0.39〜0.42", "1.8倍", "5.0〜8.2ポイント",
                  "最大35.7ポイント", "83〜85%", "少なくとも2.6倍"]:
            check(s in md, f"md に反映済み: {s}")
        body = md.split("## 付記1：")[0]
        for s in ["0.39〜0.43", "1.9倍", "4.9〜8.2ポイント", "85%前後", "3.1倍の開きがある"]:
            check(s not in body, f"md本文に旧表現なし: {s}")

    # ---------------------------------------------------------- 図表ファイル
    from pptx import Presentation
    from openpyxl import load_workbook
    bad = []
    BAD = ("1.9x", "37.7 pp", "4.9-8.2")
    for f in os.listdir(V13):
        p = os.path.join(V13, f)
        if f.endswith(".pptx"):
            for si, s in enumerate(Presentation(p).slides):
                for sh in s.shapes:
                    if sh.has_text_frame and any(b in sh.text_frame.text for b in BAD):
                        bad.append(f"{f}#{si + 1}")
        elif f.endswith(".xlsx"):
            for ws in load_workbook(p, read_only=True).worksheets:
                for row in ws.iter_rows(values_only=True):
                    if any(isinstance(c, str) and any(b in c for b in BAD) for c in row):
                        bad.append(f"{f}/{ws.title}")
    check(not bad, "図表ファイルに古いキャプションなし", ", ".join(bad))

    # まとめファイルと原稿の突合（付記15-3）
    s1 = pd.read_csv(os.path.join(V13, "SupplTableS1.csv"), comment="#")
    check(f"全{s1.product_name.nunique()}品目" in txt and f"計{len(s1)}行" in txt,
          "S1 凡例の品目数・行数がファイルと一致",
          f"{s1.product_name.nunique()}品目 / {len(s1)}行")
    s10 = pd.read_csv(os.path.join(V13, "SupplTableS10.csv"), comment="#").iloc[:44]
    import math
    gap = pd.to_numeric(s10.gini_max_gap, errors="coerce").max()
    ncert = int((s10.gini_max_certified.astype(str) == "True").sum())
    want = f"{math.ceil(gap*1e4)/1e4:.4f}"
    check(f"{len(s10)}セル中{ncert}セル" in txt and f"両者の差は{want}以下" in txt,
          "ジニ最大の証明ギャップの記述がS10と一致（上界として真）",
          f"certified={ncert}/{len(s10)} gap={gap:.6f} → {want}")
    s8t1 = pd.read_excel(os.path.join(V13, "SupplTableS8.xlsx"), sheet_name="Table1_9agents")
    tbl1 = {r.cells[0].text.strip(): [c.text.strip() for c in r.cells] for r in doc.tables[0].rows}
    JA = {"エピナスチン": "Epinastine", "オロパタジン": "Olopatadine", "レボカバスチン": "Levocabastine",
          "ケトチフェン": "Ketotifen", "クロモグリク酸ナトリウム": "Sodium cromoglicate",
          "トラニラスト": "Tranilast", "ペミロラスト": "Pemirolast", "イブジラスト": "Ibudilast",
          "アシタザノラスト": "Acitazanolast"}
    bad_share = []
    for ja, en in JA.items():
        r = s8t1[s8t1.Drug == en].iloc[0]
        lo, hi = f"{r.Share_of_9_agents_pct_min:.1f}", f"{r.Share_of_9_agents_pct_max:.1f}"
        want_s = lo if lo == hi else f"{lo}–{hi}"
        if tbl1[ja][3] != want_s:
            bad_share.append(f"{ja}: docx={tbl1[ja][3]} S8={want_s}")
    check(not bad_share, "Table 1 の9成分中シェアが Suppl Table S8 と同じ丸めで一致",
          "; ".join(bad_share))
    fb = os.path.join(V13, "SupplFigureSまとめ_v13_提出.xlsx")
    notes = "\n".join(str(c) for row in load_workbook(fb, read_only=True)["Notes"]
                      .iter_rows(values_only=True) for c in row if c)
    check("manuscript v13" in notes and "1st-11th releases" in notes
          and "build_supplfig_en_v11.py" not in notes,
          "SupplFigureSまとめ Notes の版数・公表回が最新",
          notes.split("\n")[0][:60])
    for f in ["figureまとめ_v13_提出.pptx", "figureまとめ_v13_提出.xlsx",
              "SupplTableSまとめ_v13_提出.xlsx", "Tableまとめ_v13_提出.xlsx",
              "SupplFigureSまとめ_v13_提出.xlsx", "SupplFigureSまとめ_v13_提出.pptx"]:
        check(os.path.exists(os.path.join(V13, f)), f"まとめファイルがある: {f}")
    stale = [f for f in os.listdir(V13) if "_v12" in f]
    check(not stale, "v12 名のファイルが残っていない", ", ".join(stale))
    # ネイティブグラフが残っているか（zip 内に chart パートがあるか）
    import zipfile
    with zipfile.ZipFile(os.path.join(V13, "figureまとめ_v13_提出.xlsx")) as z:
        n_chart = sum(1 for nme in z.namelist() if nme.startswith("xl/charts/chart"))
    check(n_chart >= 13, "figureまとめ xlsx にネイティブグラフが残っている", f"{n_chart} charts")

    print("\n=== OK (%d) ===" % len(ok))
    for s in ok:
        print("  [OK]", s)
    print("\n=== NG (%d) ===" % len(ng))
    for s in ng:
        print("  [NG]", s)
    sys.exit(1 if ng else 0)


if __name__ == "__main__":
    main()
