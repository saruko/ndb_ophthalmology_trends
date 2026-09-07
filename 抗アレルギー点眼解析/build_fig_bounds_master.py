# -*- coding: utf-8 -*-
"""マスター `figureまとめnew.pptx` の Fig 1A・1B・2・3A・4A に識別区間（上限）を描き足す。

方針（原稿 §2.8 改訂版）: 棒・点＝下限（公表値）、ひげ＝識別区間の上限。
Excel ネイティブのエラーバー（カスタム値、上側のみ）として付けるので、投稿用 pptx でも
編集できる。区間が極端に広い 100歳以上（Fig 1A・2）は値軸を固定して図外に逃がし、
図中に注記を置く。再実行しても二重には入らない。

入力: 論文図表/csv/Fig1_age_sex_profile_bounds__Fig1A_TOP3.csv、…__Fig1B_{drug}.csv、
      …__Fig1B_TOP3.csv、Fig3_prefecture_map_bounds__Fig3_2024_per100k.csv、
      Fig4_trends_shares_bounds__Fig4A_per100k.csv
出力: figureまとめnew.pptx（上書き。差し替え前の状態は figureまとめnew_backup_20260818.pptx）
"""
import csv
import os

import numpy as np
from pptx import Presentation
from pptx.util import Emu, Pt

from pptx_bounds import add_error_bars, set_value_axis_max

BASE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(BASE, "05_論文成果物", "公費含めない_new")
CSV = os.path.join(OUT, "論文図表", "csv")
MASTER = os.path.join(OUT, "figureまとめnew.pptx")
NOTE_PREFIX = "Whiskers"          # 注記テキストボックスの識別子（再実行時に消す）
AXIS_MAX = {"1A": 350000, "2": 400000}   # 100歳以上の上限を図外に逃がすための固定値


def rows(name):
    return list(csv.DictReader(open(os.path.join(CSV, name), encoding="utf-8-sig")))


def chart_by_title(slide, key):
    for sh in slide.shapes:
        if getattr(sh, "has_chart", False) and key in sh.chart.chart_title.text_frame.text:
            return sh
    raise KeyError(key)


def drop_notes(slide):
    for sh in list(slide.shapes):
        if sh.has_text_frame and sh.text_frame.text.startswith(NOTE_PREFIX):
            sh._element.getparent().remove(sh._element)


def add_note(slide, gf, text, gap=40000):
    """グラフ枠 gf の直下（プロットと重ならない空き領域）に注記を置く。"""
    tb = slide.shapes.add_textbox(Emu(int(gf.left + gf.width * 0.06)),
                                  Emu(int(gf.top + gf.height + gap)),
                                  Emu(int(gf.width * 0.92)), Emu(500000))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.text = text
    for para in tf.paragraphs:
        for run in para.runs:
            run.font.size = Pt(12)
            run.font.name = "Arial"
    return tb


def plus_from(cats, lower_by_cat, upper_by_cat):
    out = []
    for c in cats:
        lo, hi = float(lower_by_cat[c]), float(upper_by_cat[c])
        out.append(max(hi - lo, 0.0))
    return out


def fig1a(prs):
    s = prs.slides[0]
    gf = chart_by_title(s, "Fig 1A")
    r = rows("Fig1_age_sex_profile_bounds__Fig1A_TOP3.csv")
    lo = {x["Age Group"]: x["per100k_lower"] for x in r}
    hi = {x["Age Group"]: x["per100k_upper"] for x in r}
    cats = [str(c) for c in gf.chart.plots[0].categories]
    ser = gf.chart.plots[0].series[0]
    add_error_bars(ser, plus_from(cats, lo, hi))
    set_value_axis_max(gf.chart, AXIS_MAX["1A"], 0)
    off = [c for c in cats if float(hi[c]) > AXIS_MAX["1A"]]
    txt = ("Whiskers: upper bound of the identification interval "
           "(bars = lower bound = published value).")
    if off:
        txt += " Off scale: " + "; ".join(
            "%s upper %s (+%.0f%%)" % (c, format(float(hi[c]), ",.0f"),
                                        (float(hi[c]) / float(lo[c]) - 1) * 100) for c in off)
    drop_notes(s)
    add_note(s, gf, txt)
    return "Fig 1A: whiskers on 21 bars, axis max %d, off-scale %s" % (AXIS_MAX["1A"], off)


def fig1b(prs):
    s = prs.slides[0]
    gf = chart_by_title(s, "Fig 1B")
    cats = [str(c) for c in gf.chart.plots[0].categories]
    tables = {}
    for drug in ("Epinastine", "Olopatadine", "Levocabastine"):
        r = rows("Fig1_age_sex_profile_bounds__Fig1B_%s.csv" % drug)
        tables[drug] = ({x["Age Group"]: x["Total_lower"] for x in r},
                        {x["Age Group"]: x["Total_upper"] for x in r})
    done = []
    for ser in gf.chart.plots[0].series:
        vals = np.array([v if v is not None else np.nan for v in ser.values], float)
        # 系列名が空なので値で成分を同定する
        match = None
        for drug, (lo, hi) in tables.items():
            ref = np.array([float(lo[c]) for c in cats])
            if np.nanmax(np.abs(vals - ref) / np.maximum(ref, 1e-9)) < 1e-4:
                match = drug
                break
        if match is None:
            raise RuntimeError("Fig 1B: 系列を成分に対応づけられません")
        lo, hi = tables[match]
        add_error_bars(ser, plus_from(cats, lo, hi))
        done.append(match)
    add_note(s, gf, "Whiskers: upper bound of the identification interval "
                    "(points = lower bound = published value).")
    return "Fig 1B: whiskers on %s" % done


def fig2(prs):
    s = prs.slides[1]
    gf = chart_by_title(s, "Sex-specific")
    r = rows("Fig1_age_sex_profile_bounds__Fig1B_TOP3.csv")
    cats = [str(c) for c in gf.chart.plots[0].categories]
    off = []
    for ser in gf.chart.plots[0].series:
        sx = ser.name            # Male / Female
        lo = {x["Age Group"]: x["%s_lower" % sx] for x in r}
        hi = {x["Age Group"]: x["%s_upper" % sx] for x in r}
        add_error_bars(ser, plus_from(cats, lo, hi))
        off += ["%s %s upper %s (+%.0f%%)" % (c, sx, format(float(hi[c]), ",.0f"),
                                              (float(hi[c]) / float(lo[c]) - 1) * 100)
                for c in cats if float(hi[c]) > AXIS_MAX["2"]]
    set_value_axis_max(gf.chart, AXIS_MAX["2"], 0)
    txt = ("Whiskers: upper bound of the identification interval "
           "(bars = lower bound = published value).")
    if off:
        txt += " Off scale: " + "; ".join(off)
    drop_notes(s)
    add_note(s, gf, txt)
    return "Fig 2: whiskers on Male/Female, axis max %d, off-scale %d" % (AXIS_MAX["2"], len(off))


def fig3a(prs):
    s = prs.slides[2]
    gf = chart_by_title(s, "Fig 3A")
    r = rows("Fig3_prefecture_map_bounds__Fig3_2024_per100k.csv")
    lows = np.array([float(x["TOP3_lower"]) for x in r])
    ups = np.array([float(x["TOP3_upper"]) for x in r])
    ser = gf.chart.plots[0].series[0]
    plus = []
    for v in ser.values:          # 県名が英語なので値で対応づける（値は全県異なる）
        i = int(np.argmin(np.abs(lows - float(v))))
        if abs(lows[i] - float(v)) / lows[i] > 1e-4:
            raise RuntimeError("Fig 3A: 値 %r に対応する県が見つかりません" % v)
        plus.append(ups[i] - lows[i])
    add_error_bars(ser, plus)
    drop_notes(s)
    add_note(s, gf, "Whiskers: upper bound of the identification interval "
                    "(bars = lower bound = published value). Prefectures are ordered by "
                    "the lower bound; ranks of prefectures with overlapping intervals "
                    "are not identified.")
    return "Fig 3A: whiskers on 47 bars"


def fig4a(prs):
    s = prs.slides[3]
    gf = chart_by_title(s, "Fig 4A")
    r = rows("Fig4_trends_shares_bounds__Fig4A_per100k.csv")
    cats = [str(c) for c in gf.chart.plots[0].categories]
    cats = [c[:-2] if c.endswith(".0") else c for c in cats]
    for ser in gf.chart.plots[0].series:
        lo = {x["Year"]: x[ser.name + "_lower"] for x in r}
        hi = {x["Year"]: x[ser.name + "_upper"] for x in r}
        add_error_bars(ser, plus_from(cats, lo, hi))
    drop_notes(s)
    add_note(s, gf, "Whiskers: upper bound of the identification interval "
                    "(points = lower bound = published value; includes products not "
                    "listed in NDB before FY2022).")
    return "Fig 4A: whiskers on 4 series"


def main():
    prs = Presentation(MASTER)
    for f in (fig1a, fig1b, fig2, fig3a, fig4a):
        print(f(prs))
    prs.save(MASTER)
    print("-> %s" % MASTER)


if __name__ == "__main__":
    main()
