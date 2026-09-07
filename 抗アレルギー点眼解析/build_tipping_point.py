# -*- coding: utf-8 -*-
"""転覆点解析（tipping point / quantitative bias analysis）。

査読指摘①「2014〜2021年度は集計されていない薬剤があるのに、3成分の処方シェアが
多いと本当に言えるのか」への定量的回答。

## 考え方

3成分シェア s = T / (T + O)（T=3成分の量、O=他6成分の量）。
公表されているTは確定値に近い（エピナスチン・オロパタジンは2020年度まで
未収載品目ゼロ＝公表値が真値）。一方Oは未収載品目のぶん過小評価されている。

そこで「**結論が覆るには、隠れた他成分がどれだけ存在しなければならないか**」を逆算する。

    s < 閾値 となる条件:  O > T × (1 − 閾値) / 閾値

    閾値50%   → O > T
    閾値2/3   → O > T / 2
    閾値80%   → O > T / 4

この必要量を、足切りが撤廃されて全品目が見えた2022年度の他成分実測値と比べる。
必要量が実測の数倍に達するなら、「3成分が主要でなかった」というシナリオは
**未収載の旧世代薬が2022年時点の数倍存在し、足切り撤廃と同時に消滅した**
ことを意味し、観測された市場動向（他6成分は2022→2024年度も減少）と矛盾する。

これは疫学の quantitative bias analysis と同じ論法であり、
「結論を覆すのに必要なバイアスの大きさを示し、それが非現実的であることを示す」。

出力: 05_論文成果物/公費含めない_new/論文図表/
  csv/tipping_point_analysis.csv        本体（先頭にREADMEコメント行）
  SupplTable_tipping_point.xlsx         READMEシート＋データシート
  Fig6_tipping_point.png                必要量と実測・最悪ケースの対比（対数軸）
また 査読回答_根拠数値.csv に転覆点の行を追記する（再実行しても重複しない）。
"""
import os
import sys

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np               # noqa: E402
import pandas as pd              # noqa: E402

BASE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(BASE, "src"))

from paths import BASE_DIR  # noqa: E402

NEW_DIR = os.path.join(BASE_DIR, "05_論文成果物", "公費含めない_new")
OUT_DIR = os.path.join(NEW_DIR, "論文図表")
CSV_DIR = os.path.join(OUT_DIR, "csv")
SHARE_CSV = os.path.join(
    CSV_DIR, "Fig4_trends_shares_bounds__Fig4C_top3_share_bounds.csv")
EVIDENCE_CSV = os.path.join(OUT_DIR, "査読回答_根拠数値.csv")

FULL_COVERAGE_FROM = 2022     # 足切り撤廃＝全品目が公表される最初の年度
THRESHOLDS = [(0.50, "50"), [2 / 3, "67"], (0.80, "80")]


def build(share):
    ref = share[share.Year == FULL_COVERAGE_FROM].iloc[0]
    others_ref = ref.others_published_mL      # 2022年度の他6成分実測（mL）

    rows = []
    for _, r in share.iterrows():
        T = r.top3_published_mL
        rec = {
            "Year": int(r.Year),
            "top3_published_mL": T,
            "others_published_mL": r.others_published_mL,
            "others_rank_upper_mL": r.others_upper_mL,
            "top3_share_published_pct": r.top3_share_published_pct,
            "top3_share_lower_pct": r.top3_share_lower_pct,
        }
        for thr, tag in THRESHOLDS:
            need = T * (1 - thr) / thr
            rec[f"required_others_below{tag}pct_mL"] = need
            rec[f"ratio_to_2022actual_below{tag}pct"] = need / others_ref
        # 最悪ケース（順位ベース上限）が転覆点に届くか
        rec["worstcase_crosses_50pct"] = bool(
            r.others_upper_mL > T)
        rec["ratio_worstcase_to_2022actual"] = r.others_upper_mL / others_ref
        rows.append(rec)
    d = pd.DataFrame(rows)

    # 他6成分の市場トレンド（フルカバレッジ期間）— 「消滅シナリオ」の反証材料
    full = share[share.Year >= FULL_COVERAGE_FROM]
    n = len(full) - 1
    cagr = ((full.others_published_mL.iloc[-1] / full.others_published_mL.iloc[0])
            ** (1 / n) - 1) * 100 if n > 0 else np.nan
    return d, others_ref, cagr


README = [
    "# ========== README（この行はデータではない。読込時は comment='#' を指定） ==========",
    "# 転覆点解析（tipping point / quantitative bias analysis）。数量はすべてmL換算。",
    "# 目的: 「2014〜2021年度は未収載の薬剤があるのに3成分のシェアが多いと言えるのか」に対し、",
    "#       結論が覆るために必要な『隠れた他成分の量』を逆算し、その非現実性を示す。",
    "# ",
    "# Year                          : 年度",
    "# top3_published_mL             : 3成分（エピナスチン・オロパタジン・レボカバスチン）の公表総量",
    "# others_published_mL           : 他6成分の公表総量（0はその年度に全成分がNDB非掲載）",
    "# others_rank_upper_mL          : 他6成分の順位ベース上限（未収載品目を最大限見積もった最悪ケース）",
    "# top3_share_published_pct      : 公表値どうしで計算した3成分シェア",
    "# top3_share_lower_pct          : 識別区間の下限（他6成分を上限に置いた最悪ケースのシェア）",
    "# required_others_below50pct_mL : 3成分シェアが50%を割るために必要な他6成分の真の量",
    "#                                 （s < 閾値 ⇔ O > T×(1−閾値)/閾値。50%ならO>T）",
    "# required_others_below67pct_mL : 同、閾値2/3の場合（O > T/2）",
    "# required_others_below80pct_mL : 同、閾値80%の場合（O > T/4）",
    "# ratio_to_2022actual_below**pct: 上記必要量が、全品目公開後の2022年度実測の何倍か",
    "#                                 → この倍率だけ隠れた薬剤が存在し、足切り撤廃と同時に",
    "#                                    消滅した、というシナリオが必要になる",
    "# worstcase_crosses_50pct       : 順位ベース最悪ケースでも50%を割るか（論理的に否定できるか）",
    "# ratio_worstcase_to_2022actual : 最悪ケースの他6成分量が2022年度実測の何倍か",
    "# ",
    "# 解釈: 倍率が大きいほど『3成分が主要でなかった』シナリオは非現実的。",
    "#       他6成分は1990年代以前発売の縮小市場であり、2022年度以降も減少が続いている。",
    "# ================================================================================",
]


def write_csv(d):
    os.makedirs(CSV_DIR, exist_ok=True)
    p = os.path.join(CSV_DIR, "tipping_point_analysis.csv")
    with open(p, "w", encoding="utf-8-sig", newline="") as f:
        f.write("\n".join(README) + "\n")
        d.to_csv(f, index=False)
    print(f"-> {os.path.relpath(p, BASE)} ({len(d)} rows)")
    return p


def write_xlsx(d, others_ref, cagr):
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    notes = [
        "転覆点解析（tipping point / quantitative bias analysis）",
        "査読指摘①「未収載の薬剤があるのに3成分のシェアが多いと言えるのか」への定量的回答",
        "",
        "条件: 3成分シェア s = T/(T+O) が閾値を割るには O > T×(1−閾値)/閾値 が必要",
        "  閾値50% → O > T ／ 閾値2/3 → O > T/2 ／ 閾値80% → O > T/4",
        "",
        f"比較基準: 全品目が公開された{FULL_COVERAGE_FROM}年度の他6成分実測 = {others_ref:,.0f} mL",
        f"他6成分の{FULL_COVERAGE_FROM}〜2024年度の年平均変化率 = {cagr:+.1f}%/年（縮小市場）",
        "",
        "結論: 2014〜2021年度に3成分が過半数でなかったと言うためには、未収載の旧世代薬が"
        f"{FULL_COVERAGE_FROM}年度実測の数倍存在し、足切り撤廃と同時に消滅した必要がある。"
        "これは観測された市場動向と矛盾するため、3成分の代表性はplausibleと判断できる。",
        "",
        "ただし無仮定で厳密にシェアを主張できるのは2022年度以降（識別区間が点に縮む）である。",
        "2014〜2021年度は本解析を感度分析として併記し、シェアの断定は避けること。",
    ]
    p = os.path.join(OUT_DIR, "SupplTable_tipping_point.xlsx")
    with pd.ExcelWriter(p, engine="openpyxl") as xw:
        pd.DataFrame({"Notes": notes}).to_excel(xw, sheet_name="README", index=False)
        d.to_excel(xw, sheet_name="TippingPoint", index=False)
        for ws in xw.book.worksheets:
            for cell in ws[1]:
                cell.font = Font(name="Arial", bold=True)
            for i, col in enumerate(ws.iter_cols(min_row=1, max_row=1), 1):
                ws.column_dimensions[get_column_letter(i)].width = max(
                    12, min(34, len(str(col[0].value or "")) + 3))
            ws.freeze_panes = "A2"
    print(f"-> {os.path.relpath(p, BASE)}")


def figure(d, others_ref):
    plt.rcParams["font.family"] = "MS Gothic"
    plt.rcParams["axes.unicode_minus"] = False
    fig, ax = plt.subplots(figsize=(11, 6))
    y = d.Year

    ax.plot(y, d.required_others_below50pct_mL / 1e6, "o-", color="tab:red", lw=2.2,
            label="3成分シェアが50%を割るのに必要な他6成分の量")
    ax.plot(y, d.required_others_below80pct_mL / 1e6, "^--", color="tab:orange",
            lw=1.6, label="同 80%を割るのに必要な量")
    ax.plot(y, d.others_rank_upper_mL / 1e6, "s--", color="tab:purple", lw=1.6,
            label="他6成分の順位ベース上限（最悪ケース）")
    ax.plot(y, d.others_published_mL.replace(0, np.nan) / 1e6, "D-",
            color="tab:blue", lw=2.2, label="他6成分の公表値（実測）")
    ax.axhline(others_ref / 1e6, color="gray", ls=":", lw=1.5,
               label=f"{FULL_COVERAGE_FROM}年度の他6成分実測 = {others_ref/1e6:.1f}百万mL")
    ax.axvspan(FULL_COVERAGE_FROM - .5, y.max() + .4, color="tab:green", alpha=.07)

    ax.set_yscale("log")
    ax.set_xlabel("年度")
    ax.set_ylabel("処方数量（百万mL、対数軸）")
    ax.set_title("転覆点解析: 「3成分が主要でない」と言うために必要な他6成分の量\n"
                 "（必要量は実測を大きく上回り、そのシナリオは非現実的）",
                 fontsize=12, fontweight="bold")
    ax.set_xticks(list(y))
    ax.grid(alpha=.3, which="both")
    ax.text(FULL_COVERAGE_FROM + .5, ax.get_ylim()[1] * .75, "全品目公開\n（足切り撤廃）",
            ha="center", va="top", fontsize=9, color="tab:green")
    # 凡例は系列と重ならないよう軸の下に置く
    ax.legend(fontsize=9, loc="upper center", bbox_to_anchor=(.5, -.12), ncol=2,
              frameon=False)
    fig.tight_layout()
    p = os.path.join(OUT_DIR, "Fig6_tipping_point.png")
    fig.savefig(p, dpi=150)
    plt.close(fig)
    print(f"-> {os.path.relpath(p, BASE)}")


def update_evidence(d, others_ref, cagr):
    """査読回答_根拠数値.csv に転覆点の行を追記する（再実行で重複しない）。"""
    if not os.path.exists(EVIDENCE_CSV):
        print("   （査読回答_根拠数値.csv が無いため追記をスキップ）")
        return
    ev = pd.read_csv(EVIDENCE_CSV)
    ev = ev[~ev.項目.astype(str).str.startswith("転覆点")]
    src = "論文図表/csv/tipping_point_analysis.csv"
    rows = []
    for _, r in d.iterrows():
        rows += [
            {"指摘": "指摘①", "項目": "転覆点_50%割れに必要な他6成分量",
             "年度": r.Year, "値": round(r.required_others_below50pct_mL),
             "単位": "mL", "出典ファイル": src,
             "備考": "3成分シェアが50%を割るには他6成分がこの量必要"},
            {"指摘": "指摘①", "項目": "転覆点_50%割れ必要量の2022年実測比",
             "年度": r.Year, "値": round(r.ratio_to_2022actual_below50pct, 2),
             "単位": "倍", "出典ファイル": src,
             "備考": f"全品目公開後の2022年度実測{others_ref:,.0f}mLの何倍必要か"},
            {"指摘": "指摘①", "項目": "転覆点_80%割れ必要量の2022年実測比",
             "年度": r.Year, "値": round(r.ratio_to_2022actual_below80pct, 2),
             "単位": "倍", "出典ファイル": src, "備考": ""},
        ]
    rows.append({"指摘": "指摘①", "項目": "転覆点_他6成分の年平均変化率(2022-2024)",
                 "年度": "2022-2024", "値": round(cagr, 1), "単位": "%/年",
                 "出典ファイル": src,
                 "備考": "縮小市場であり『消滅シナリオ』の反証材料"})
    out = pd.concat([ev, pd.DataFrame(rows)], ignore_index=True)
    out.to_csv(EVIDENCE_CSV, index=False, encoding="utf-8-sig")
    print(f"-> {os.path.relpath(EVIDENCE_CSV, BASE)} "
          f"({len(ev)} + {len(rows)} = {len(out)} rows)")


def main():
    share = pd.read_csv(SHARE_CSV)
    d, others_ref, cagr = build(share)
    write_csv(d)
    write_xlsx(d, others_ref, cagr)
    figure(d, others_ref)
    update_evidence(d, others_ref, cagr)

    print()
    print(f"他6成分 {FULL_COVERAGE_FROM}年度実測: {others_ref:,.0f} mL "
          f"／ {FULL_COVERAGE_FROM}〜2024年度 CAGR {cagr:+.1f}%/年")
    pre = d[d.Year < FULL_COVERAGE_FROM]
    print(f"2014〜2021年度に50%割れへ必要な倍率: "
          f"{pre.ratio_to_2022actual_below50pct.min():.1f}〜"
          f"{pre.ratio_to_2022actual_below50pct.max():.1f}倍")


if __name__ == "__main__":
    main()
