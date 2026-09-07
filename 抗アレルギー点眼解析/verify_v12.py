# -*- coding: utf-8 -*-
"""v12 の成果物を検証する。数値の不変条件と、原稿から消えているべき記述を確認する。

実行: C:\\Users\\goodt\\anaconda3\\python.exe verify_v12.py
"""
import os
import re
import sys

import pandas as pd
from docx import Document

sys.stdout.reconfigure(encoding="utf-8")

BASE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, BASE)
import docx_finalize as F                                # noqa: E402

NEW = os.path.join(BASE, "05_論文成果物", "公費含めない_new")
V12 = os.path.join(NEW, "投稿改定ver", "投稿用v12")
GINI = os.path.join(NEW, "ジニ係数あり", "gini_cv_bounds.csv")
CLEAN = os.path.join(V12, "NDB抗アレルギー解析_v12_clean.docx")
TRACKED = os.path.join(V12, "NDB抗アレルギー解析_v12.docx")

ok, ng = [], []


def check(cond, label, detail=""):
    (ok if cond else ng).append(label + (" — " + detail if detail else ""))


def main():
    # ---------------------------------------------------------- 数値
    d = pd.read_csv(GINI)
    check(len(d) == 44, "gini_cv_bounds.csv は44行", f"{len(d)}行")

    for m in ("gini", "cv"):
        bad = d[d[f"{m}_min"] > d[f"{m}_max"] + 1e-12]
        check(bad.empty, f"{m}: min ≦ max", f"{len(bad)}件違反")
        # 公表値（全県が下限）は実行可能なので区間に入るはず
        bad = d[(d[f"{m}_published"] < d[f"{m}_min"] - 1e-6)
                | (d[f"{m}_published"] > d[f"{m}_max"] + 1e-6)]
        check(bad.empty, f"{m}: 公表値が区間内",
              ", ".join(f"{r.series}{r.year}" for r in bad.itertuples()))
        # 制約つき区間は箱制約の区間に含まれる
        bad = d[(d[f"{m}_min"] < d[f"{m}_min_boxonly"] - 1e-6)
                | (d[f"{m}_max"] > d[f"{m}_max_boxonly"] + 1e-6)]
        check(bad.empty, f"{m}: 制約つき区間 ⊆ 箱制約区間",
              ", ".join(f"{r.series}{r.year}" for r in bad.itertuples()))
        # 証明付き上界は達成値以上
        bad = d[d[f"{m}_max_upper_certificate"] < d[f"{m}_max"] - 1e-9]
        check(bad.empty, f"{m}: 上界 ≧ 達成値",
              ", ".join(f"{r.series}{r.year}" for r in bad.itertuples()))

    check(d["gini_max_gap"].max() < 1e-3,
          "ジニ最大：証明付き上界と達成値の差が全セルで0.001未満",
          f"厳密一致 {int(d['gini_max_certified'].sum())}/{len(d)}、"
          f"最大ギャップ {d['gini_max_gap'].max():.2e}")

    chk = pd.read_csv(os.path.join(NEW, "ジニ係数あり", "gini_cv_consistency.csv"))
    exact = chk[chk.n_censored_prefectures == 0]
    check(len(exact) > 0 and (exact.ratio_lower.sub(1).abs() < 1e-6).all(),
          "秘匿セル0の成分・年度で Σ県別 = 全国総計",
          ", ".join(f"{r.series}{r.year}:{r.ratio_lower:.8f}"
                    for r in exact.itertuples()))
    check((chk.ratio_lower <= 1 + 1e-9).all(),
          "Σ県別下限 ≦ 全国上限（制約が実行可能）")

    comp = d[((d.series == "TOP3") & (d.year >= 2022))
             | ((d.series == "Epinastine") & (d.year >= 2014))
             | ((d.series == "Olopatadine") & (d.year >= 2014))
             | ((d.series == "Levocabastine") & (d.year >= 2022))]
    red = comp["gini_width_reduction_pct"].dropna()
    print(f"\n比較可能年度のジニ区間幅の縮小: {red.min():.1f}〜{red.max():.1f}% "
          f"(中央値 {red.median():.1f}%)")

    def row(s, y):
        return d[(d.series == s) & (d.year == y)].iloc[0]

    o14, o24 = row("Olopatadine", 2014), row("Olopatadine", 2024)
    print(f"オロパタジン ジニ 2014 [{o14.gini_min:.4f},{o14.gini_max:.4f}] "
          f"vs 2024 [{o24.gini_min:.4f},{o24.gini_max:.4f}] "
          f"→ {'分離（縮小が確定）' if o14.gini_min > o24.gini_max else '重複'}")
    print(f"  箱制約のみ: 2014 [{o14.gini_min_boxonly:.4f},{o14.gini_max_boxonly:.4f}] "
          f"vs 2024 [{o24.gini_min_boxonly:.4f},{o24.gini_max_boxonly:.4f}] "
          f"→ {'分離' if o14.gini_min_boxonly > o24.gini_max_boxonly else '重複'}")

    # ---------------------------------------------------------- 原稿
    if not os.path.exists(CLEAN):
        ng.append("clean docx が未生成")
        report()
        return

    a = F.audit(CLEAN)
    check(a["w:ins"] == 0 and a["w:del"] == 0, "clean版に変更履歴なし", str(a))
    check(a["commentReference"] == 0 and not a["comment_parts"],
          "clean版にコメントなし", str(a))
    at = F.audit(TRACKED)
    check(at["w:ins"] > 0, "履歴つき版には変更履歴がある", str(at))

    doc = Document(CLEAN)
    txt = "\n".join(p.text for p in doc.paragraphs)

    must_go = [
        "主要な抗ヒスタミン点眼薬の全国処方動向",
        "Major Antihistamine Eye Drops",
        "ジニ係数0.1前後という水準は",
        "少数の外れ値の影響を受けにくい",
        "処方患者数ベースの増加は、処方数量が示す以上に大きい",
        "当該3年度は上限と下限が一致し、点識別となる",
        "n(d,y) × Σ_s m(y,s) × c",
        "Published online 20",
    ]
    for s in must_go:
        check(s not in txt, f"削除済み: {s[:34]}")

    must_have = [
        "抗アレルギー点眼薬の全国処方動向 2014–2024年度",
        "National Prescription Trends for Anti-allergic Eye Drops in Japan",
        "集計整合性",
        "Σ_{i=1}^{n(d,y)} c_i × Σ_s m(y,s)",
        "品目非掲載に関しては点識別",
        "長期収載品の選定療養",
        "中学校卒業（15歳到達年度末）",
        "残薬の発生",
        "valid ではあるが sharp ではない",
        "2018;73(12):2395-2397",
        "2017;172(4):224-235",
        "2011;118(12):2361-2367",
        "2014;113(4):476-481",
        "2020;16(1):5",
        "https://www.mhlw.go.jp/stf/newpage_39830.html",
    ]
    for s in must_have:
        check(s in txt, f"追加済み: {s[:34]}")

    # 本文の指標値が CSV と一致するか
    for s, key in (("TOP3", "主要3成分合計"), ("Epinastine", "エピナスチン"),
                   ("Olopatadine", "オロパタジン"),
                   ("Levocabastine", "レボカバスチン")):
        r = row(s, 2024)
        import math
        want = (f"{math.floor(r.gini_min * 1000) / 1000:.3f}〜"
                f"{math.ceil(r.gini_max * 1000) / 1000:.3f}")
        check(want in txt, f"本文のジニ2024（{key}）が CSV と一致（外向き丸め）", want)
        # 表示区間は計算区間を包含していなければならない
        lo_s, hi_s = want.split("〜")
        check(float(lo_s) <= r.gini_min + 1e-12 and float(hi_s) >= r.gini_max - 1e-12,
              f"表示区間が計算区間を包含（{key}）", want)

    # 模擬査読で報告された「誤植」が実在しないことの再確認
    phantom = ["男／女比都道府県", "後発医薬品後発医薬品", "開きがある倍の開き",
               "0.0842024年度"]
    for s in phantom:
        check(s not in txt, f"幻の誤植は存在しない: {s}")

    verify_md(d, row)
    verify_bundles(d)
    report()


def verify_md(d, row):
    """作業原稿（.md）が docx と同じ内容になっているか。"""
    import io
    import math

    md_path = os.path.join(V12, "論文下書き_総量2014_シェア2022_区間解析_投稿改定ver12.md")
    check(os.path.exists(md_path), "v12 の md がある")
    check(not os.path.exists(
        os.path.join(V12, "論文下書き_総量2014_シェア2022_区間解析_投稿改定ver11.md")),
        "v11 の md は v12 フォルダに残っていない")
    if not os.path.exists(md_path):
        return
    md = io.open(md_path, encoding="utf-8").read()
    lines = md.split("\n")
    end = [i for i, l in enumerate(lines) if l.startswith("## 付記1：")]
    body = "\n".join(lines[:end[0]]) if end else md

    for s in ["0.084", "0.083〜0.130", "0.154〜0.230", "0.111〜0.139",
              "ジニ係数0.1前後", "数量が示す以上に大きい", "極端な偏在では"]:
        check(s not in body, f"md本文に旧v11表現なし: {s}")
    for s in ["部分識別解析", "集計整合性", "付記14", "品目非掲載に関しては",
              "中学校卒業", "選定療養", "gini_cv_consistency.csv"]:
        check(s in md, f"md に追加済み: {s}")
    for s, ja in (("TOP3", "主要3成分合計"), ("Epinastine", "エピナスチン"),
                  ("Olopatadine", "オロパタジン"),
                  ("Levocabastine", "レボカバスチン")):
        r = row(s, 2024)
        want = (f"{math.floor(r.gini_min * 1000) / 1000:.3f}〜"
                f"{math.ceil(r.gini_max * 1000) / 1000:.3f}")
        check(want in body, f"md本文のジニ2024（{ja}）が CSV と一致", want)


def verify_bundles(d):
    """まとめファイルが v12 の名前で、Fig 7 と S10 の値が生成元と一致するか。"""
    from openpyxl import load_workbook

    names = ["figureまとめ_v12_提出.pptx", "figureまとめ_v12_提出.xlsx",
             "SupplTableSまとめ_v12_提出.xlsx", "Tableまとめ_v12_提出.xlsx",
             "SupplFigureSまとめ_v12_提出.xlsx", "SupplFigureSまとめ_v12_提出.pptx"]
    for n in names:
        check(os.path.exists(os.path.join(V12, n)), f"まとめファイルがある: {n}")
    stale = [f for f in os.listdir(V12) if "_v11" in f]
    check(not stale, "v11 名のファイルが残っていない", ", ".join(stale))

    p = os.path.join(V12, "figureまとめ_v12_提出.xlsx")
    if os.path.exists(p):
        for sheet, met in (("Fig7A", "gini"), ("Fig7B", "cv")):
            ws = load_workbook(p)[sheet]
            rows = list(ws.iter_rows(values_only=True))
            hi = [i for i, r in enumerate(rows) if r[0] == "category"]
            if not hi:
                check(False, f"{sheet} の見出し行が見つかる")
                continue
            hdr = list(rows[hi[0]])
            bad = 0
            for r in rows[hi[0] + 1:]:
                if r[0] is None:
                    continue
                y = int(r[0])
                for s in ("TOP3", "Epinastine", "Olopatadine", "Levocabastine"):
                    src = d[(d.series == s) & (d.year == y)].iloc[0]
                    mid = (src[met + "_min"] + src[met + "_max"]) / 2
                    for col, exp in ((hdr.index(s), mid),
                                     (hdr.index(s + " lower"), src[met + "_min"]),
                                     (hdr.index(s + " upper"), src[met + "_max"])):
                        if r[col] is None or abs(float(r[col]) - exp) > 1e-9:
                            bad += 1
            check(bad == 0, f"まとめ {sheet} が生成元と一致", f"{bad}件不一致")
            note = rows[1][0] or ""
            # 7A は "aggregation consistency"、7B は "aggregation-consistency"
            check(re.search(r"aggregation[ -]consisten", note) is not None,
                  f"まとめ {sheet} の脚注が v12 の文言")

    p = os.path.join(V12, "SupplTableSまとめ_v12_提出.xlsx")
    if os.path.exists(p):
        ws = load_workbook(p)["S10"]
        rows = list(ws.iter_rows(values_only=True))
        hi = [i for i, r in enumerate(rows) if r[0] == "series"]
        if hi:
            hdr = list(rows[hi[0]])
            check("gini_min_boxonly" in hdr, "まとめ S10 に箱制約のみの列がある")
            gi, yi = hdr.index("gini_min"), hdr.index("year")
            n = 0
            for r in rows[hi[0] + 1:hi[0] + 45]:
                src = d[(d.series == r[0]) & (d.year == r[yi])]
                if len(src) == 1 and abs(float(r[gi]) - src.iloc[0].gini_min) < 1e-9:
                    n += 1
            check(n == 44, "まとめ S10 の44行が生成元と一致", f"{n}/44")
            heads = [r for r in rows[:hi[0]] if r[0] and str(r[0]).startswith("#")]
            check(all(r[1] is None for r in heads),
                  "まとめ S10 の説明行が A 列に収まっている")
        else:
            check(False, "まとめ S10 の見出し行が見つかる")


def report():
    print("\n=== OK (%d) ===" % len(ok))
    for s in ok:
        print("  [OK]", s)
    print("\n=== NG (%d) ===" % len(ng))
    for s in ng:
        print("  [NG]", s)
    sys.exit(1 if ng else 0)


if __name__ == "__main__":
    main()
