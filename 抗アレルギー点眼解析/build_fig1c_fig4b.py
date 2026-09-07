# -*- coding: utf-8 -*-
"""マスター `figureまとめnew.pptx` に構成比の識別区間を描き込む。

- **Fig 4B**: 100%積み上げ棒 → 折れ線（マーカー）に差し替え、上下のひげ（識別区間）を付ける。
  積み上げ棒ではひげが累積位置に付いてしまい、成分ごとのシェアの区間を表せないため。
- **Fig 1C**: 新規パネル。2024年度の年齢階級別・主要3成分の構成比（%）を上下ひげつきで描く
  （§3.2 の 71.4%・24.6% などの数値そのもの）。Fig 1B（量）の下に置く。

いずれも Fig 1B の折れ線グラフを雛形として複製するので、系列色・フォントは既存図と揃う。
シェアの区間は build_share_bounds.bounds（Table 1・Table 3・Suppl Table S8 と同一の式）。

再実行しても二重には入らない。
出力: figureまとめnew.pptx（上書き。差し替え前は figureまとめnew_backup_20260819.pptx）
"""
import copy
import csv
import os
import shutil

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from build_share_bounds import bounds
from pptx_bounds import add_error_bars

BASE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(BASE, "05_論文成果物", "公費含めない_new")
CSV = os.path.join(OUT, "論文図表", "csv")
MASTER = os.path.join(OUT, "figureまとめnew.pptx")
BACKUP = os.path.join(OUT, "figureまとめnew_backup_20260819.pptx")

DRUGS = ("Epinastine", "Olopatadine", "Levocabastine")
NOTE = ("Whiskers: identification interval of the share "
        "(points = published value, which lies between the two bounds).")


def chart_shape(slide, key):
    for sh in slide.shapes:
        if getattr(sh, "has_chart", False) and sh.chart.has_title \
                and key in sh.chart.chart_title.text_frame.text:
            return sh
    raise KeyError(key)


def set_text(el, text):
    """a:t を1つだけ残して差し替える（フォント設定を保ったままテキストだけ変える）。"""
    ts = list(el.iter(qn("a:t")))
    ts[0].text = text
    for t in ts[1:]:
        r = t.getparent()
        r.getparent().remove(r)


def set_axis_title(chart, tag, text):
    for ax in chart._chartSpace.iter(qn(tag)):
        ttl = ax.find(qn("c:title"))
        if ttl is not None:
            set_text(ttl, text)


def set_valax_numfmt(chart, code):
    for ax in chart._chartSpace.iter(qn("c:valAx")):
        nf = ax.find(qn("c:numFmt"))
        if nf is not None:
            nf.set("formatCode", code)
            nf.set("sourceLinked", "0")


FILLS = ("a:noFill", "a:solidFill", "a:gradFill", "a:blipFill", "a:pattFill", "a:grpFill")


def set_transparent(chart):
    """グラフ領域・プロット領域の塗りを透明にする（背面の縦破線を隠さないため。Fig 4A と同じ）。"""
    cs = chart._chartSpace
    targets = [cs.find(qn("c:spPr"))]
    pa = cs.find(qn("c:chart")).find(qn("c:plotArea"))
    targets.append(pa.find(qn("c:spPr")))
    for sp in targets:
        if sp is None:
            continue
        for tag in FILLS:
            for el in sp.findall(qn(tag)):
                sp.remove(el)
        sp.insert(0, etree.Element(qn("a:noFill")))


def share_rows(values):
    """values: {category: {drug: (lower, upper)}} -> (plotted, minus, plus) の dict。"""
    out = {}
    for cat, d in values.items():
        lo = {k: v[0] for k, v in d.items()}
        hi = {k: v[1] for k, v in d.items()}
        b = bounds(lo, hi)
        tot = sum(lo.values())
        out[cat] = {k: (lo[k] / tot * 100, b[k][0], b[k][1]) for k in lo}
    return out


def clone_chart(slide, src_shape, left, top, width, height, cd):
    """Fig 1B のグラフを雛形にして、同じ書式の折れ線グラフを新しく作る。

    グラフの XML だけを複製する（グラフ部品自体は add_chart で新規に作る）ので、
    埋め込みブックは複製先が独自に持つ。
    """
    gf = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, left, top, width, height, cd)
    new = gf.chart._chartSpace
    ext = new.find(qn("c:externalData"))
    src = copy.deepcopy(src_shape.chart._chartSpace)
    old_ext = src.find(qn("c:externalData"))
    if old_ext is not None:
        src.remove(old_ext)
    for child in list(new):
        new.remove(child)
    for child in list(src):
        new.append(child)
    if ext is not None:
        new.append(ext)
    gf.chart.replace_data(cd)
    return gf


def add_note(slide, gf, text):
    tb = slide.shapes.add_textbox(gf.left + int(gf.width * 0.06), gf.top + gf.height + 40000,
                                  int(gf.width * 0.92), 500000)
    tb.text_frame.word_wrap = True
    tb.text_frame.text = text
    for para in tb.text_frame.paragraphs:
        for run in para.runs:
            run.font.size = Pt(12)
            run.font.name = "Arial"
    return tb


def drop_note(slide, gf):
    for sh in list(slide.shapes):
        if sh.has_text_frame and sh.text_frame.text.startswith("Whiskers") \
                and abs(sh.top - (gf.top + gf.height)) < gf.height * 0.2:
            sh._element.getparent().remove(sh._element)


def drop_shape(sh):
    sh._element.getparent().remove(sh._element)


def build(chart_slide, src_shape, cats, data, left, top, width, height, title,
          cat_title, val_title):
    """data: {category: {drug: (plotted, min, max)}}"""
    cd = CategoryChartData()
    cd.categories = cats
    for d in DRUGS:
        cd.add_series(d, [round(data[c][d][0], 2) for c in cats])
    gf = clone_chart(chart_slide, src_shape, left, top, width, height, cd)
    ch = gf.chart
    set_text(ch.chart_title._element, title)
    set_axis_title(ch, "c:catAx", cat_title)
    set_axis_title(ch, "c:valAx", val_title)
    set_valax_numfmt(ch, "0")
    for ser, d in zip(ch.plots[0].series, DRUGS):
        add_error_bars(ser,
                       plus=[max(data[c][d][2] - data[c][d][0], 0.0) for c in cats],
                       minus=[max(data[c][d][0] - data[c][d][1], 0.0) for c in cats])
    return gf


def fig1c(prs, src_shape):
    """2024年度 年齢階級別の主要3成分構成比（100歳以上は除く。Fig 1A・1B と同じ扱い）。"""
    vals = {}
    for drug in DRUGS:
        for r in csv.DictReader(open(os.path.join(
                CSV, "Fig1_age_sex_profile_bounds__Fig1B_%s.csv" % drug), encoding="utf-8-sig")):
            if r["Age Group"] == "100+":
                continue
            vals.setdefault(r["Age Group"], {})[drug] = (float(r["Total_lower"]),
                                                         float(r["Total_upper"]))
    cats = list(vals)
    data = share_rows(vals)
    s = prs.slides[0]
    for sh in list(s.shapes):                      # 再実行時は古い 1C を消す
        if getattr(sh, "has_chart", False) and sh.chart.has_title \
                and "Fig 1C" in sh.chart.chart_title.text_frame.text:
            drop_note(s, sh)
            drop_shape(sh)
    b = chart_shape(s, "Fig 1B")
    gf = build(s, src_shape, cats, data, b.left, b.top + b.height + Emu(1100000),
               b.width, b.height,
               "Fig 1C. Age-specific composition of the three major agents (%, FY2024)",
               "Age group (years)", "Share among the three major agents (%)")
    add_note(s, gf, NOTE + " Age 100+ is omitted (see Fig 1A/1B).")


def fig4b(prs, src_shape):
    """主要3成分内シェアの推移。積み上げ棒 → 折れ線＋上下ひげ。"""
    vals = {}
    for r in csv.DictReader(open(os.path.join(
            CSV, "Fig4_trends_shares_bounds__Fig4A_per100k.csv"), encoding="utf-8-sig")):
        vals[r["Year"]] = {d: (float(r[d + "_lower"]), float(r[d + "_upper"])) for d in DRUGS}
    cats = list(vals)
    data = share_rows(vals)
    s = prs.slides[3]
    old = chart_shape(s, "Fig 4B")
    left, top, width, height = old.left, old.top, old.width, old.height
    title = old.chart.chart_title.text_frame.text
    drop_note(s, old)
    tree = old._element.getparent()
    z = list(tree).index(old._element)      # 縦破線より下に置く（前面に出すと線が隠れる）
    drop_shape(old)
    gf = build(s, src_shape, cats, data, left, top, width, height,
               title, "Fiscal year", "Prescription volume share (%)")
    set_transparent(gf.chart)
    tree.remove(gf._element)
    tree.insert(z, gf._element)
    add_note(s, gf, NOTE + " The intervals for FY2014 and FY2021 are wide because of "
                           "item-level non-listing; see Suppl Table S8.")


def main():
    if not os.path.exists(BACKUP):
        shutil.copyfile(MASTER, BACKUP)
        print("-> backup %s" % BACKUP)
    prs = Presentation(MASTER)
    src = chart_shape(prs.slides[0], "Fig 1B")     # 雛形（色・フォント）
    fig4b(prs, src)
    fig1c(prs, src)
    prs.save(MASTER)
    print("-> %s" % MASTER)


if __name__ == "__main__":
    main()
