# -*- coding: utf-8 -*-
"""現行の _v11.docx に、ジニ係数・CV の再計算値を変更履歴つきで反映する。

注意
----
`apply_docx_v11_gini.py` が作った `_v11rev*.docx` は 2026-09-04 15:xx 時点の
`_v11.docx` を元にしている。その後 `_v11.docx` は著者が編集しており
（「後発品」→「後発医薬品」ほか、Methods にジニ係数の最適化手法の記述を追加）、
rev ファイルを上書きすると著者の編集が失われる。本スクリプトは
**現行の _v11.docx を入力**として数値だけを差し替える。

反映する内容
------------
1. Results（都道府県）・Results（経年）・Discussion の 2024年度のジニ係数／CV を
   再計算値に更新（最小側が下がったため）。
2. 丸めの訂正。以下は v11 以前からの誤りで、今回の再計算とは独立:
     TOP3   CV上限   0.23049 → 0.231 と書かれていた → 0.230
     Epinastine ジニ上限 0.13946 → 0.140 と書かれていた → 0.139
     Epinastine CV公表値 0.23348 → 0.234 と書かれていた → 0.233
3. Figure 7 legend に、識別できていない年度を破線＋白抜きで描き分けた旨を追記。
4. Methods の手法記述を、実装に合わせて具体化（探索の初期値数、ジニ最大の
   証明付き上界が全44セルで達成値と一致したこと）。

出力:
  投稿用v11/NDB抗アレルギー解析_verジニ+改定_v11_num.docx        （変更履歴つき）
  投稿用v11/NDB抗アレルギー解析_verジニ+改定_v11_num_マークなし.docx
  （_v11.docx 自体は書き換えない。バックアップは _v11_backup_YYYYMMDD.docx）
"""
import os
import shutil
import sys

from docx import Document

BASE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, BASE)
from docx_track import TrackedEditor                      # noqa: E402

V11 = os.path.join(BASE, "05_論文成果物", "公費含めない_new", "投稿改定ver",
                   "投稿用v11")
SRC = os.path.join(V11, "NDB抗アレルギー解析_verジニ+改定_v11.docx")
BACKUP = os.path.join(V11, "NDB抗アレルギー解析_verジニ+改定_v11_backup_20260904.docx")
DEST = os.path.join(V11, "NDB抗アレルギー解析_verジニ+改定_v11_num.docx")
DEST_CLEAN = os.path.join(V11,
                          "NDB抗アレルギー解析_verジニ+改定_v11_num_マークなし.docx")
DATE = "2026-09-04T19:00:00Z"

REPLACEMENTS = [
    # --- Results（都道府県別の2024年度） ---
    ("2024年度のジニ係数は、主要3成分合計0.084〜0.130（公表値のみで計算すると0.117）、"
     "エピナスチン0.111〜0.140（同0.133）、オロパタジン0.079〜0.139（同0.118）、"
     "レボカバスチン0.107〜0.231（同0.180）、CVはそれぞれ0.155〜0.231（同0.209）、"
     "0.198〜0.245（同0.234）、0.152〜0.248（同0.218）、0.190〜0.408（同0.322）であった。",
     "2024年度のジニ係数は、主要3成分合計0.083〜0.130（公表値のみで計算すると0.117）、"
     "エピナスチン0.111〜0.139（同0.133）、オロパタジン0.078〜0.139（同0.118）、"
     "レボカバスチン0.083〜0.231（同0.180）、CVはそれぞれ0.154〜0.230（同0.209）、"
     "0.198〜0.245（同0.233）、0.152〜0.248（同0.218）、0.161〜0.408（同0.322）であった。"),

    # --- Results（経年） ---
    ("2024年度（ジニ係数0.111〜0.140、CV 0.198〜0.245）への縮小が識別区間を含めて確定した。"
     "オロパタジンは2014年度（0.129〜0.131）と2024年度（0.079〜0.139）の区間が重なるため",
     "2024年度（ジニ係数0.111〜0.139、CV 0.198〜0.245）への縮小が識別区間を含めて確定した。"
     "オロパタジンは2014年度（0.129〜0.131）と2024年度（0.078〜0.139）の区間が重なるため"),

    # --- Discussion ---
    ("2024年度の主要3成分合計のジニ係数は0.084〜0.130、CVは0.155〜0.231であった（Figure 7）。",
     "2024年度の主要3成分合計のジニ係数は0.083〜0.130、CVは0.154〜0.230であった（Figure 7）。"),
    ("エピナスチンの県間格差は2014年度（ジニ係数0.235）から2024年度（0.111〜0.140）への縮小が確定した。",
     "エピナスチンの県間格差は2014年度（ジニ係数0.235）から2024年度（0.111〜0.139）への縮小が確定した。"),

    # --- Methods（実装に合わせて具体化） ---
    ("そこで最大値については、各県を下限または上限に置く組み合わせを座標降下（複数初期値）で"
     "探索した達成値を真の最大値の下界とし、分子を上限・分母を下限で評価した緩和上界を併記した。",
     "そこで最大値については、各県を下限または上限に置く組み合わせを座標上昇（5通りの並べ方"
     "それぞれについて分割点を全通り試す構造化初期値と乱数初期値の計368通り）で探索した達成値を"
     "真の最大値の下界とした。さらにジニ係数については、ジニ係数がΣ_r w_r·x_(r)/(n·Σx)"
     "（w_r＝2r−n−1、rは昇順の順位）と書けることを用い、県と順位の対応を自由にしてよいと緩めた"
     "割当問題とDinkelbach型の二分探索により valid な上界を求めたところ、"
     "全44セル（4系列×11年度）で達成値と一致し、最大値が大域最適であることを確認した。"),

    # --- Figure 7 legend ---
    ("レボカバスチン（後発医薬品が2021年度以前非掲載）とそれを含む主要3成分合計は2022年度以降に"
     "限定する（それ以前は公表品目のみに基づく参考値）。レボカバスチンは処方量が小さく"
     "秘匿セルの比率が高いため区間が最も広い（2015–2017年度はCVの上限が軸の範囲外に達する）。",
     "レボカバスチン（後発医薬品が2021年度以前非掲載）とそれを含む主要3成分合計は2022年度以降に"
     "限定する。識別できていない年度（レボカバスチンと主要3成分合計の2014–2021年度）は破線と"
     "白抜きマーカーで描き分けた。当該年度は区間がほぼ0から最大値まで開くため中点は水準の情報を"
     "持たず、2015–2017年度に見えるピークは格差の上昇ではなく欠測の産物である"
     "（レボカバスチン2015年度のジニ係数の区間は0.000–0.733、CVは0.000–3.171）。"
     "実線と塗りつぶしマーカーが比較可能な年度を示す。"),
]


def find(paras, needle):
    for p in paras:
        if needle in p.text:
            return p
    raise KeyError(needle[:60])


def main():
    if not os.path.exists(BACKUP):
        shutil.copy2(SRC, BACKUP)
        print("backup -> %s" % BACKUP)
    doc = Document(SRC)
    ed = TrackedEditor(DATE)
    P = doc.paragraphs
    for old, new in REPLACEMENTS:
        ed.replace(find(P, old), old, new)
    doc.save(DEST)
    print("-> %s （変更履歴つき、%d 箇所）" % (DEST, len(REPLACEMENTS)))
    return DEST


def make_clean(path):
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    doc = Document(path)
    body = doc.element.body
    for para in list(body.iter("{%s}p" % W)):
        pPr = para.find("{%s}pPr" % W)
        if pPr is None:
            continue
        rPr = pPr.find("{%s}rPr" % W)
        if rPr is None or rPr.find("{%s}del" % W) is None:
            continue
        parent = para.getparent()
        if parent is not None:
            parent.remove(para)
    for el in list(body.iter("{%s}del" % W)):
        parent = el.getparent()
        if parent is not None:
            parent.remove(el)
    for el in list(body.iter("{%s}ins" % W)):
        parent = el.getparent()
        if parent is None:
            continue
        idx = list(parent).index(el)
        for k, ch in enumerate(list(el)):
            el.remove(ch)
            parent.insert(idx + k, ch)
        parent.remove(el)
    doc.save(DEST_CLEAN)
    print("-> %s （履歴確定版）" % DEST_CLEAN)


if __name__ == "__main__":
    make_clean(main())
