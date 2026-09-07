# -*- coding: utf-8 -*-
"""投稿用v13 フォルダを v12 の図表・表・補遺から組み立てる。

v13 は数値の再計算を伴わないため、図表・表・補遺ファイルは v12 のものをそのまま
複製し、まとめファイルの名前だけ v13 に改める。原稿（docx・md）は
apply_docx_v13.py / apply_md_v13.py が生成する。

PPT の英文キャプション（Fig 1C・4A・4B）は v12 のファイル上で訂正済み
（付記15-1）。本スクリプトは複製後にその訂正が入っていることを検証する。

実行: python build_v13_folder.py
"""
import os
import shutil
import sys

BASE = os.path.dirname(os.path.abspath(__file__))
NEW = os.path.join(BASE, "05_論文成果物", "公費含めない_new")
V12 = os.path.join(NEW, "投稿改定ver", "投稿用v12")
V13 = os.path.join(NEW, "投稿改定ver", "投稿用v13")

SKIP_PREFIX = ("NDB抗アレルギー解析_v12", "論文下書き_", "~$", "_v12_base")
BAD_CAPTIONS = ("1.9x", "37.7 pp", "4.9-8.2")


def main():
    os.makedirs(V13, exist_ok=True)
    n = 0
    for f in sorted(os.listdir(V12)):
        if f.startswith(SKIP_PREFIX):
            continue
        src = os.path.join(V12, f)
        if not os.path.isfile(src):
            continue
        dst = os.path.join(V13, f.replace("_v12_", "_v13_"))
        shutil.copy2(src, dst)
        n += 1
    print(f"{n} files copied -> {V13}")

    # PPT キャプションの検証
    from pptx import Presentation
    bad = []
    for f in os.listdir(V13):
        if not f.endswith(".pptx"):
            continue
        prs = Presentation(os.path.join(V13, f))
        for si, s in enumerate(prs.slides):
            for sh in s.shapes:
                if sh.has_text_frame and any(b in sh.text_frame.text for b in BAD_CAPTIONS):
                    bad.append((f, si + 1))
    from openpyxl import load_workbook
    for f in os.listdir(V13):
        if not f.endswith(".xlsx"):
            continue
        wb = load_workbook(os.path.join(V13, f), read_only=True)
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                if any(isinstance(c, str) and any(b in c for b in BAD_CAPTIONS) for c in row):
                    bad.append((f, ws.title))
    if bad:
        raise SystemExit("古いキャプションが残っている: %s" % bad)
    print("PPT/XLSX captions OK")

    # README
    with open(os.path.join(V12, "README.md"), encoding="utf-8") as fh:
        readme = fh.read()
    readme = readme.replace("投稿用ファイル 対応表（原稿 v12・投稿改定ver）",
                            "投稿用ファイル 対応表（原稿 v13・投稿改定ver）")
    readme = readme.replace("_v12_提出", "_v13_提出")
    readme += """

---

## v13 での変更（v12 からの差分）

数値の再計算はない。2件の外部レビュー（図表ファイル間の突合、原稿内部の論理・記述の
突合）で指摘された矛盾を解消した。内容は原稿 md の付記15 を参照。

| ファイル | 生成 | v12 からの変更 |
|---|---|---|
| `NDB抗アレルギー解析_v13.docx` / `_v13_clean.docx` | `apply_docx_v13.py` | 本文・脚注・補遺凡例の16項目（付記15-2） |
| `論文下書き_…_投稿改定ver13.md` | `apply_md_v13.py` | 同上のうち作業原稿に対応箇所がある分＋付記15 |
| `Figure4.pptx`・`figureまとめ_v13_提出.pptx` | v12 の複製 | Fig 1C・4A・4B の英文キャプションの数値（付記15-1）。生成元 `build_submission_v11.py`・`build_figure_bundle.py` も訂正済み |
| 上記以外の図表・表・補遺 | v12 の複製（`build_v13_folder.py`） | 変更なし |

**投稿には `NDB抗アレルギー解析_v13_clean.docx` を使うこと。**

再生成: `python apply_docx_v13.py` → `python apply_md_v13.py` →
`python build_v13_folder.py` → `python verify_v13.py`
"""
    with open(os.path.join(V13, "README.md"), "w", encoding="utf-8", newline="\n") as fh:
        fh.write(readme)
    print("README written")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
