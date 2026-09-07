# -*- coding: utf-8 -*-
"""Figure 3B（都道府県別 日本地図）を主要3成分・現行データで描き、マスター pptx に差し込む。

背景: 旧 Figure 3B は旧版凍結フォルダの 9成分合計（Fig3_prefecture_map_k.xlsx Fig3_data）で
描かれており、Figure 3A（主要3成分・現行）と集計単位・データ版の両方が食い違っていた
（原稿 付記2）。本スクリプトは
  05_論文成果物/公費含めない_new/論文図表/Fig3_prefecture_map_bounds.xlsx
  の Fig3_2024_per100k シート TOP3_lower（＝Figure 3A と同じ値）
から地図本体とカラーバーの PNG を作り、`figureまとめnew.pptx` スライド3 の
グループ内の 2 枚の画像（地図・カラーバー）を同じ寸法で差し替える。

出力:
  論文図表/Fig3B_map_top3.png        地図本体（3300×4458 px, 背景透明）
  論文図表/Fig3B_colorbar_top3.png   カラーバー（505×1123 px, 背景透明）
  figureまとめnew.pptx               スライド3 の画像を差し替え（実行前に手動でバックアップ済み）

地図データ: data/japan_prefectures.geojson（dataofjapan/land, パブリックドメイン）
"""
import json
import os

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib.collections import PathCollection
from matplotlib.path import Path as MplPath
from pptx import Presentation

BASE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(BASE, "05_論文成果物", "公費含めない_new")
FIG = os.path.join(OUT, "論文図表")
GEOJSON = os.path.join(BASE, "data", "japan_prefectures.geojson")
XLSX = os.path.join(FIG, "Fig3_prefecture_map_bounds.xlsx")
MASTER = os.path.join(OUT, "figureまとめnew.pptx")
MAP_PNG = os.path.join(FIG, "Fig3B_map_top3.png")
CB_PNG = os.path.join(FIG, "Fig3B_colorbar_top3.png")

# 差し替え先の画像と同じピクセル寸法（pptx 内の枠のアスペクト比を保つ）
MAP_PX = (3300, 4458)
CB_PX = (505, 1123)
COLUMN = "TOP3_lower"   # Figure 3A と同じ値（識別区間の下限＝公表値）
UPPER = "TOP3_upper"
HATCH_WIDTH_PCT = 10.0  # 区間幅 (上限−下限)/下限 がこれを超える県にハッチングを重ねる
HATCH_NOTE = ("Hatched: identification interval width >%d%% of the lower bound "
              "(colour = lower bound = published value)" % HATCH_WIDTH_PCT)


def load_geometry():
    gj = json.load(open(GEOJSON, encoding="utf-8"))
    geoms = {}
    for feat in gj["features"]:
        name = feat["properties"]["nam_ja"]
        geom = feat["geometry"]
        polys = [geom["coordinates"]] if geom["type"] == "Polygon" else geom["coordinates"]
        paths = []
        for poly in polys:
            verts, codes = [], []
            for ring in poly:
                verts.extend(ring)
                codes.extend([MplPath.MOVETO] + [MplPath.LINETO] * (len(ring) - 1))
            paths.append(MplPath(np.asarray(verts), codes))
        geoms[name] = paths
    return geoms


def load_values():
    df = pd.read_excel(XLSX, sheet_name="Fig3_2024_per100k")
    names = df["Prefecture"].astype(str).str.strip()
    width = (df[UPPER] / df[COLUMN] - 1) * 100
    return (dict(zip(names, df[COLUMN].astype(float))),
            set(names[width > HATCH_WIDTH_PCT]))


def draw_map(geoms, values, hatched, cmap, norm):
    dpi = 300
    fig = plt.figure(figsize=(MAP_PX[0] / dpi, MAP_PX[1] / dpi), dpi=dpi)
    ax = fig.add_axes([0, 0, 1, 1])
    paths, colors = [], []
    for name, plist in geoms.items():
        c = cmap(norm(values[name]))
        for p in plist:
            paths.append(p)
            colors.append(c)
    ax.add_collection(PathCollection(paths, facecolors=colors, edgecolors="white",
                                     linewidths=0.3))
    # 区間幅が広い県にハッチング（色は下限のまま。順位・水準が確定しないことを示す）
    hp = [p for name in hatched for p in geoms[name]]
    plt.rcParams["hatch.linewidth"] = 0.4
    ax.add_collection(PathCollection(hp, facecolors="none", edgecolors="#404040",
                                     linewidths=0.0, hatch="////"))
    ax.set_xlim(127.5, 146.5)
    ax.set_ylim(25.5, 46.0)
    ax.set_aspect(1.0 / np.cos(np.radians(37)))
    ax.axis("off")
    fig.savefig(MAP_PNG, dpi=dpi, transparent=True)
    plt.close(fig)


def draw_colorbar(cmap, norm, vmin, vmax):
    dpi = 300
    fig = plt.figure(figsize=(CB_PX[0] / dpi, CB_PX[1] / dpi), dpi=dpi)
    # 旧画像と同じ配置: 左に細長いバー、上に見出し、右に最大・最小値
    cax = fig.add_axes([0.02, 0.05, 0.29, 0.77])
    cb = plt.colorbar(plt.cm.ScalarMappable(norm=norm, cmap=cmap), cax=cax)
    cb.outline.set_visible(False)
    cb.set_ticks([vmin, vmax])
    cb.ax.set_yticklabels([f"{vmin:,.0f}", f"{vmax:,.0f}"], fontsize=15)
    cb.ax.tick_params(length=0, pad=2)
    fig.text(0.02, 0.99, "Prescriptions\nper 100,000", fontsize=15, va="top", ha="left",
             linespacing=1.1)
    fig.savefig(CB_PNG, dpi=dpi, transparent=True)
    plt.close(fig)


def replace_pictures():
    """スライド3 のグループ内の画像 2 枚（地図・カラーバー）を寸法で見分けて差し替える。"""
    prs = Presentation(MASTER)
    slide = prs.slides[2]
    done = []
    for sh in slide.shapes:
        if sh.shape_type != 6:      # GROUP
            continue
        for pic in sh.shapes:
            if pic.shape_type != 13:    # PICTURE
                continue
            size = pic.image.size
            if size == MAP_PX:
                src = MAP_PNG
            elif size == CB_PX:
                src = CB_PNG
            else:
                continue
            part = slide.part.related_part(pic._element.blip_rId)
            part._blob = open(src, "rb").read()
            done.append((pic.shape_id, os.path.basename(src)))
        # 同じグループ内の「Fig 3B.」テキストにハッチングの説明を1段落足す（再実行で二重にしない）
        for tb in sh.shapes:
            if tb.has_text_frame and tb.text_frame.text.startswith("Fig 3B"):
                tf = tb.text_frame
                if not any(p.text.startswith("Hatched") for p in tf.paragraphs):
                    ref = tf.paragraphs[-1]
                    para = tf.add_paragraph()
                    para.text = HATCH_NOTE
                    if ref.runs:
                        para.runs[0].font.size = ref.runs[0].font.size
                        para.runs[0].font.bold = False
                        para.runs[0].font.name = ref.runs[0].font.name
    if len(done) != 2:
        raise RuntimeError("差し替え対象の画像が2枚見つかりません: %r" % done)
    prs.save(MASTER)
    return done


def main():
    geoms = load_geometry()
    values, hatched = load_values()
    missing = set(geoms) - set(values)
    if missing:
        raise KeyError("都道府県名が一致しません -> %s" % sorted(missing))
    vmin, vmax = min(values.values()), max(values.values())
    cmap = plt.get_cmap("Blues")
    norm = plt.Normalize(vmin, vmax)
    draw_map(geoms, values, hatched, cmap, norm)
    print("hatched (%d): %s" % (len(hatched), " ".join(sorted(hatched))))
    draw_colorbar(cmap, norm, vmin, vmax)
    print("map  : %s" % MAP_PNG)
    print("cbar : %s  (%s = %s .. %s)" % (CB_PNG, COLUMN, f"{vmin:,.1f}", f"{vmax:,.1f}"))
    for sid, name in replace_pictures():
        print("pptx : slide 3 shape %d <- %s" % (sid, name))


if __name__ == "__main__":
    main()
