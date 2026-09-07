# -*- coding: utf-8 -*-
"""v14 の成果物を検証する。

v14 は数値の再計算を伴わない。数値そのものの検証は verify_v13.py（および
そこから委ねている verify_v12.py）に委ね、ここでは

  1. v13 から引き継いだ図表・表・補遺がバイト単位で同一であること
     （＝ v13 の数値検証がそのまま v14 にも当てはまること）
  2. Suppl Fig S1 の削除と S2〜S5 → S1〜S4 の繰り上げが、原稿・図表・
     まとめファイル・ファイル名のすべてで整合していること
  3. 補遺が本文からもれなく引用されていること
  4. キャプションを直した本文図の PNG が pptx より新しいこと
  5. 作業原稿（md）の数値が docx と揃っていること

を確認する。

実行: python verify_v14.py
"""
import hashlib
import io
import os
import re
import sys
import zipfile

from docx import Document

sys.stdout.reconfigure(encoding="utf-8")
BASE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, BASE)
import docx_finalize as F                                # noqa: E402

NEW = os.path.join(BASE, "05_論文成果物", "公費含めない_new")
V13 = os.path.join(NEW, "投稿改定ver", "投稿用v13")
V14 = os.path.join(NEW, "投稿改定ver", "投稿用v14")
CLEAN = os.path.join(V14, "NDB抗アレルギー解析_v14_clean.docx")
TRACKED = os.path.join(V14, "NDB抗アレルギー解析_v14.docx")
MD = os.path.join(V14, "論文下書き_総量2014_シェア2022_区間解析_投稿改定ver14.md")

FIGMAP = {"2": "1", "3": "2", "4": "3", "5": "4"}
# v14 で中身を変えたファイル（1. のバイト同一チェックから除く）
CHANGED = {"Figure2.png", "Figure2.pptx", "Figure4.png", "Figure6.png",
           "Figure6.pptx", "SupplFigureSまとめ_v14_提出.xlsx",
           "SupplFigureSまとめ_v14_提出.pptx", "README.md",
           "SupplFigure_en_README.md"}

ok, ng = [], []


def check(cond, label, detail=""):
    (ok if cond else ng).append(label + (" — " + detail if detail else ""))


def sha(path):
    return hashlib.sha256(io.open(path, "rb").read()).hexdigest()


def oldname(f):
    """v14 のファイル名 → v13 のファイル名。"""
    m = re.match(r"SupplFigureS([1-4])(.*)$", f)
    if m:
        back = {v: k for k, v in FIGMAP.items()}
        return "SupplFigureS%s%s" % (back[m.group(1)], m.group(2))
    return f.replace("_v14_", "_v13_")


def main():
    # ------------------------------------------------ 1. v13 からの引き継ぎ
    same = diff = 0
    for f in sorted(os.listdir(V14)):
        p = os.path.join(V14, f)
        if not os.path.isfile(p) or f in CHANGED:
            continue
        if f.startswith(("NDB抗アレルギー解析_v14", "論文下書き_")):
            continue
        if f.startswith("SupplFigureS") and f.endswith(".csv"):
            continue                       # 説明行の図番号だけ書き換えている
        src = os.path.join(V13, oldname(f))
        if not os.path.exists(src):
            ng.append("v13 に対応ファイルがない: %s" % f)
            continue
        if sha(p) == sha(src):
            same += 1
        else:
            diff += 1
            ng.append("v13 と中身が違う: %s" % f)
    check(diff == 0, "図表・表・補遺は v13 とバイト単位で同一",
          "%d ファイル" % same)

    # 消えているべきファイル・あるべきファイル
    for f in ("SupplFigureS5.png", "SupplFigureS5_en.png",
              "SupplFigureS5_data_en.csv"):
        check(not os.path.exists(os.path.join(V14, f)), "S5 のファイルがない: %s" % f)
    for n in (1, 2, 3, 4):
        check(os.path.exists(os.path.join(V14, "SupplFigureS%d_en.png" % n)),
              "SupplFigureS%d_en.png がある" % n)
    check(not os.path.exists(os.path.join(V14, "SupplFigureS1.pptx")),
          "模式図の pptx を複製していない")
    for f in ("figureまとめ_v14_提出.pptx", "figureまとめ_v14_提出.xlsx",
              "SupplTableSまとめ_v14_提出.xlsx", "Tableまとめ_v14_提出.xlsx",
              "SupplFigureSまとめ_v14_提出.xlsx", "SupplFigureSまとめ_v14_提出.pptx"):
        check(os.path.exists(os.path.join(V14, f)), "まとめファイルがある: %s" % f)
    check(not [f for f in os.listdir(V14) if "_v13_" in f],
          "v13 名のファイルが残っていない")

    # ------------------------------------------------------------ 2. 原稿
    a = F.audit(CLEAN)
    check(a["w:ins"] == 0 and a["w:del"] == 0, "clean版に変更履歴なし", str(a))
    check(a["commentReference"] == 0 and not a["comment_parts"],
          "clean版にコメントなし", str(a))
    check(F.audit(TRACKED)["w:ins"] > 0, "履歴つき版には変更履歴がある")

    doc = Document(CLEAN)
    paras = [p.text for p in doc.paragraphs]
    txt = "\n".join(paras)
    check("Supplementary Figure S5" not in txt, "docx に Suppl Fig S5 が残っていない")
    check("NDBオープンデータの欠測の3層構造" not in txt,
          "docx から模式図の legend が消えている")
    heads = [p.split(".")[0] for p in paras if p.startswith("Supplementary Figure S")]
    check(heads == ["Supplementary Figure S1", "Supplementary Figure S2",
                    "Supplementary Figure S3", "Supplementary Figure S4"],
          "補遺図の legend が S1〜S4 の順で並ぶ", str(heads))
    check(len(doc.tables) == 3, "本文表が3つ", str(len(doc.tables)))

    # 抄録の圧縮（英訳で約300語に収まる長さ）
    ab = [p for p in paras if p.startswith(("Background:", "Methods: NDB",
                                            "Results: 2024", "Conclusions: エピ"))]
    n_ab = sum(len(p.rstrip()) for p in ab)
    check(len(ab) == 4, "抄録が4段落", str(len(ab)))
    check(850 <= n_ab <= 950, "抄録が圧縮されている（旧1,293字）", "%d字" % n_ab)
    check("人口10万対189,671" not in txt.split("INTRODUCTION")[0],
          "抄録から人口10万対の値を落としてある")

    # ------------------------------------------- 3. 補遺の引用（本文のみ）
    body = "\n".join(paras[:paras.index("Supplementary figure and table legends")])
    for i in range(1, 12):
        check(re.search(r"Supplementary Table S%d(?![0-9])" % i, body) is not None,
              "本文が Supplementary Table S%d を引用" % i)
    order = []
    for m in re.finditer(r"Supplementary Figure S([1-4])", body):
        if m.group(1) not in order:
            order.append(m.group(1))
    check(order == ["1", "2", "3", "4"], "補遺図の初出順が昇順", "".join(order))

    # --------------------------------------- 4. まとめファイルと本文図 pptx
    z = zipfile.ZipFile(os.path.join(V14, "SupplFigureSまとめ_v14_提出.xlsx"))
    wbx = z.read("xl/workbook.xml").decode("utf-8")
    names = re.findall(r'<sheet[^>]*\bname="([^"]+)"', wbx)
    check(names == ["Notes", "S1A", "S1B", "S2", "S3A", "S3B", "S4"],
          "まとめ xlsx のシート名が繰り上がっている", str(names))
    charts = [n for n in z.namelist() if "/charts/chart" in n and n.endswith(".xml")]
    check(len(charts) == 14, "まとめ xlsx にネイティブグラフが残っている",
          "%d charts" % len(charts))
    refs = set()
    for n in charts:
        x = z.read(n).decode("utf-8")
        refs |= {m.split("!")[0] for m in re.findall(r"<f>([^<]*)</f>", x)}
    check(refs == {"'S1A'", "'S1B'", "'S2'", "'S3A'", "'S3B'", "'S4'"},
          "グラフの参照式が新しいシート名を指す", str(sorted(refs)))

    z = zipfile.ZipFile(os.path.join(V14, "SupplFigureSまとめ_v14_提出.pptx"))
    titles = []
    for n in sorted(x for x in z.namelist()
                    if re.match(r"ppt/slides/slide[0-9]+\.xml$", x)):
        x = z.read(n).decode("utf-8")
        titles += [t[:26] for t in re.findall(r"<a:t>([^<]*)</a:t>", x)
                   if t.startswith("Supplementary Figure")]
    check(titles == ["Supplementary Figure S1. G", "Supplementary Figure S2. M",
                     "Supplementary Figure S3. N", "Supplementary Figure S4. S"],
          "まとめ pptx のスライドタイトルが繰り上がっている", str(titles))

    for f, want, bad in (("Figure2.pptx", "Supplementary Figure S2",
                          "Supplementary Figure S3"),
                         ("Figure6.pptx", "Supplementary Figure S1",
                          "Supplementary Figure S2")):
        x = zipfile.ZipFile(os.path.join(V14, f)).read(
            "ppt/slides/slide1.xml").decode("utf-8")
        check(want in x and bad not in x, "%s の note が繰り上がっている" % f)

    # 補遺図 CSV の説明行
    for n in (1, 2, 3, 4):
        p = os.path.join(V14, "SupplFigureS%d_data_en.csv" % n)
        head = io.open(p, encoding="utf-8-sig").readline()
        check("Supplementary Figure S%d." % n in head,
              "SupplFigureS%d_data_en.csv の説明行が一致" % n, head.strip()[:60])

    # ----------------------------------------- 5. 書き出し直した PNG の鮮度
    for n in (2, 4, 6):
        png = os.path.join(V14, "Figure%d.png" % n)
        ppt = os.path.join(V14, "Figure%d.pptx" % n)
        check(os.path.getmtime(png) >= os.path.getmtime(ppt),
              "Figure%d.png が pptx より新しい" % n)
    x = zipfile.ZipFile(os.path.join(V14, "Figure4.pptx")).read(
        "ppt/slides/slide1.xml").decode("utf-8")
    check("1.8x" in x and "37.7 pp" not in x,
          "Figure4.pptx のキャプションが訂正済み（PNG はこれを書き出したもの）")

    # ------------------------------------------------------------ 6. 作業原稿
    md = io.open(MD, encoding="utf-8").read()
    mdbody = md[md.index("## 構造化抄録"):md.index("## 付記1：編集方針")]
    for s in ("15〜94歳", "+30〜+983%", "+17,550%超"):
        check(s in mdbody, "md に反映済み: %s" % s)
    for s in ("15歳以上の全階級で女性優位", "+50〜+430%", "+16,000%超"):
        check(s not in mdbody, "md本文に旧表現なし: %s" % s)
    # 抄録は docx と同じ本文に差し替え（見出しも Purpose → Background）
    mdab = md[md.index("## 構造化抄録"):md.index("**Keywords**:")]
    check("**Background**" in mdab and "**Purpose（目的）**" not in mdab,
          "md の抄録が docx と同じ構成に差し替わっている")
    check("※" in mdab, "md の抄録に250語版の削除候補（※）がある")
    for s in ("Suppl Fig S5", "Supplementary Figure S5",
              "Supplementary Figure S1.** NDB"):
        check(s not in mdbody, "md本文に旧番号なし: %s" % s)
    check("付記16" in md, "md に付記16 がある")

    print("=== OK (%d) ===" % len(ok))
    for s in ok:
        print("  [OK] %s" % s)
    print("\n=== NG (%d) ===" % len(ng))
    for s in ng:
        print("  [NG] %s" % s)
    return 1 if ng else 0


if __name__ == "__main__":
    sys.exit(main())
