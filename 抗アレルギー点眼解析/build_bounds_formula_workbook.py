# -*- coding: utf-8 -*-
"""識別区間（value_lower / value_upper）の計算式を Excel の数式そのままで示す監査用ブック。

投稿用の各 `_data.csv` は数字の羅列で、下限・上限がどう出たのかが読めない。本ブックは
生データ（公表総計・Σ開示セル・欠落量）から下限・上限までを **Excel の数式** で組み立て、
最後に投稿用ファイルの値と突き合わせる。数式セルを追えば導出が全部たどれる。

計算式の根拠は原稿§2.8 と `bounded_outputs_report.md`（アルゴリズム本体は
`build_bounded_outputs.py`）。本スクリプトは既存の出力を再計算せず、同じ入力から
数式を組み立てて一致することだけを確認する。

出力: 05_論文成果物/公費含めない_new/論文図表/識別区間の計算式.xlsx
      （`build_submission_files.py` が投稿用/ へコピーする）
"""
import os
import sys

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(BASE, "src"))

from ndb_reader import CENSOR_THRESHOLD  # noqa: E402
from paths import COVARIATE_PATH, INPUT_DIR  # noqa: E402

OUT = os.path.join(BASE, "05_論文成果物", "公費含めない_new")
FIG = os.path.join(OUT, "論文図表")
CSV = os.path.join(FIG, "csv")
DEST = os.path.join(FIG, "識別区間の計算式.xlsx")
CAP = CENSOR_THRESHOLD - 1              # 999
TOP3 = ["EPINASTINE", "OLOPATADINE", "LEVOCASTINE"]
YEAR = 2024
EXAMPLE_PREFS = ["沖縄県", "静岡県"]     # 下限が最小の県と、区間が最も広い県

S01 = "01_品目行の上限"
HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")
CALC_FILL = PatternFill("solid", fgColor="FFF2CC")   # 数式セル
CHK_FILL = PatternFill("solid", fgColor="E2EFDA")    # 検算セル
THIN = Side(style="thin", color="AAAAAA")


def note(ws, row, text, bold=False):
    ws.cell(row=row, column=1, value=text).font = Font(bold=bold, size=11)
    return row + 1


def header(ws, row, cols, widths):
    for i, name in enumerate(cols, start=1):
        c = ws.cell(row=row, column=i, value=name)
        c.font = Font(bold=True)
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = Border(bottom=THIN)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 32
    return row + 1


def put(ws, row, values, numfmt=None, fills=None):
    for i, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=i, value=v)
        if numfmt and numfmt.get(i):
            c.number_format = numfmt[i]
        if fills and fills.get(i):
            c.fill = fills[i]
    return row + 1


def read_csv_utf8(path):
    return pd.read_csv(path, encoding="utf-8-sig")


def load_inputs():
    lv = read_csv_utf8(os.path.join(INPUT_DIR, "product_level_censoring.csv"))
    lg = read_csv_utf8(os.path.join(INPUT_DIR, "product_pref_long.csv"))
    lv = lv[(lv.year == YEAR) & (lv.code.isin(TOP3))].reset_index(drop=True)
    lg = lg[(lg.year == YEAR) & (lg.code.isin(TOP3))]
    lv["per_cell_cap"] = np.where(lv.total_censored,
                                  (CAP - lv.sum_disclosed).clip(lower=0),
                                  lv.missing.clip(lower=0))
    return lv, lg


def pref_population(pref):
    cov = read_csv_utf8(COVARIATE_PATH)
    return float(cov[(cov.year == YEAR) & (cov.prefecture == pref)].population_total.iloc[0])


# ----------------------------------------------------------------- 00
README_LINES = [
    ("識別区間（value_lower / value_upper）の計算式", True),
    ("", False),
    ("NDBオープンデータは、処方数量が1,000未満のセルを「-」で伏せて公表している（セル秘匿）。"
     "伏せられた値そのものは不明だが、公表されている「総計」列から範囲を絞り込める。"
     "図表の下限・上限は統計的な信頼区間ではなく、この「真値が必ず入る範囲」（識別区間）である。", False),
    ("", False),
    ("【記号】", True),
    ("　total　　　　　… 公表されている品目行の総計", False),
    ("　sum_disclosed … 同じ行で開示されているセルの合計", False),
    ("　missing　　　 … total − sum_disclosed（伏せられたセルの合計。行内で確定している）", False),
    ("　総計秘匿　　　… 総計そのものが「-」の行。missing は不明だが total < 1,000 は確定している", False),
    ("", False),
    ("【基本式】", True),
    ("　各セルの下限 ＝ 秘匿なら 0、開示なら公表値", False),
    ("　各セルの上限 ＝ 秘匿なら per_cell_cap、開示なら公表値", False),
    ("　per_cell_cap ＝ IF(総計秘匿, MAX(0, 999 − sum_disclosed), MAX(0, missing))", False),
    ("", False),
    ("　上限に一律 999 を使わない理由: NDB は補完的秘匿を行っており、1,000 以上のセルでも"
     "「単独で伏せると総計から逆算できてしまう」場合に巻き添えで伏せられている。実データでも "
     "missing が 秘匿セル数×999 を超える行が都道府県軸に8行・年齢性別軸に41行ある。"
     "そのため「秘匿セルは必ず1,000未満」は前提にできず、「各セルは非負」「行内の秘匿セルの合計は "
     "missing にちょうど一致する」の2つの事実だけから各セルを [0, missing] と押さえる。", False),
    ("", False),
    ("【集計】独立な行の下限どうし・上限どうしを単純合算してよい（各行の真値が区間内のどこにあっても、"
     "合計の真値は Σ下限〜Σ上限に必ず入る）。成分をまたぐ合算は必ず mL 換算後に行う（07シート）。", False),
    ("", False),
    ("【割合】分子と分母がどちらも区間を持つため、割合の区間は次のとおり（比は分子で増え分母で減る）。", False),
    ("　割合の下限 ＝ 自分の下限 ÷（自分の下限 ＋ 相手の上限）", False),
    ("　割合の上限 ＝ 自分の上限 ÷（自分の上限 ＋ 相手の下限）", False),
    ("　このため割合の panel（Fig 5A・5B・6A・6B）だけは、下限が図に描いた公表値を下回りうる。", False),
    ("", False),
    ("【シート一覧】", True),
    ("　01_品目行の上限　　… 秘匿セル1個あたりの上限 per_cell_cap を品目行ごとに算出（2024年度・主要3成分）", False),
    ("　02_県別集計_沖縄県 … 上の per_cell_cap から1県の下限・上限・人口10万対を出し、Figure 3A と照合", False),
    ("　02_県別集計_静岡県 … 同上（区間が最も広い県）", False),
    ("　04_人口10万対_Fig1A … 数量の区間を人口で割って人口10万対にする（Figure 1A と照合）", False),
    ("　05_全国総量の上限　… 全国総量の上限＝公表総量＋総計秘匿の加算＋未収載品目の加算（Suppl Table S6 と照合）", False),
    ("　06_割合の区間_Fig5A … 割合の区間の式（Figure 5A と照合）", False),
    ("　07_単位換算　　　　… 成分をまたぐ合算に使う mL 換算係数", False),
    ("", False),
    ("黄色のセルが数式、緑色のセルが投稿用ファイルとの差（すべて 0 になる）。", False),
    ("本ブックは build_bounds_formula_workbook.py が自動生成する（直接編集しないこと）。", False),
]


def sheet_readme(wb):
    ws = wb.create_sheet("00_読み方")
    ws.column_dimensions["A"].width = 118
    r = 1
    for text, bold in README_LINES:
        r = note(ws, r, text, bold)


# ----------------------------------------------------------------- 01
def sheet_level(wb, lv):
    ws = wb.create_sheet(S01)
    r = note(ws, 1, "秘匿セル1個あたりの上限 per_cell_cap（2024年度・主要3成分・都道府県軸）", True)
    r = note(ws, r, "K列・M列・O列が数式。M列が以降のシートの入力になる。"
                    "年齢性別軸も式は同じだが missing の値が違うので、軸ごとに別の表を使うこと。")
    r += 1
    top = header(ws, r,
                 ["キー", "年度", "処方区分", "成分", "品目名", "単位", "mL係数",
                  "公表総計 total", "Σ開示セル sum_disclosed", "秘匿県数",
                  "欠落量 missing（式）", "総計秘匿", "セル上限 per_cell_cap（式）",
                  "参照値（pipeline）", "差（式）"],
                 [26, 6, 12, 14, 34, 6, 7, 15, 17, 8, 16, 9, 20, 16, 10])
    num = {8: "#,##0.00", 9: "#,##0.00", 11: "#,##0.00", 13: "#,##0.00",
           14: "#,##0.00", 15: "0.000000"}
    fills = {11: CALC_FILL, 13: CALC_FILL, 15: CHK_FILL}
    row = top
    for _, x in lv.iterrows():
        put(ws, row, [
            "%s|%s" % (x.sheet, x.product_name), int(x.year), x.sheet, x.code,
            x.product_name, x.unit, float(x.ml_factor),
            None if pd.isna(x.total) else float(x.total), float(x.sum_disclosed),
            int(x.n_pref_censored),
            '=IF(L{0},"—",H{0}-I{0})'.format(row), bool(x.total_censored),
            "=IF(L{0},MAX(0,{1}-I{0}),MAX(0,K{0}))".format(row, CAP),
            float(x.per_cell_cap), "=M{0}-N{0}".format(row),
        ], num, fills)
        row += 1
    ws.freeze_panes = "F%d" % top
    return top, row - 1


# ----------------------------------------------------------------- 02
def sheet_pref(wb, lv, lg, pref, lv_rows, fig3):
    ws = wb.create_sheet("02_県別集計_%s" % pref)
    m = (lg[lg.prefecture == pref]
         .merge(lv[["sheet", "code", "product_name", "per_cell_cap"]],
                on=["sheet", "code", "product_name"], how="left")
         .reset_index(drop=True))
    pub = fig3[fig3.Prefecture == pref].iloc[0]
    pop = pref_population(pref)

    r = note(ws, 1, "%s の主要3成分合計（2024年度）— Figure 3A の棒1本ぶん" % pref, True)
    r = note(ws, r, "G列は 01_品目行の上限 のM列を INDEX/MATCH で引いている。"
                    "H・I列がセルごとの下限・上限、下の合計行が棒（下限）とひげ（上限）の値。")
    r += 1
    top = header(ws, r,
                 ["キー", "処方区分", "成分", "品目名", "公表値（この県のセル）", "秘匿",
                  "セル上限（01シート参照・式）", "下限（式）", "上限（式）"],
                 [26, 12, 14, 34, 18, 7, 24, 16, 16])
    num = {5: "#,##0.00", 7: "#,##0.00", 8: "#,##0.00", 9: "#,##0.00"}
    fills = {7: CALC_FILL, 8: CALC_FILL, 9: CALC_FILL}
    row = top
    for _, x in m.iterrows():
        put(ws, row, [
            "%s|%s" % (x.sheet, x.product_name), x.sheet, x.code, x.product_name,
            None if pd.isna(x.value) else float(x.value), bool(x.censored),
            "=INDEX('{0}'!$M${1}:$M${2},MATCH(A{3},'{0}'!$A${1}:$A${2},0))".format(
                S01, lv_rows[0], lv_rows[1], row),
            "=IF(F{0},0,E{0})".format(row),
            "=IF(F{0},G{0},E{0})".format(row),
        ], num, fills)
        row += 1
    last = row - 1

    r_sum, r_pop, r_p100 = last + 2, last + 3, last + 4
    put(ws, r_sum, ["合計（mL）", None, None, None, None, None, None,
                    "=SUM(H{0}:H{1})".format(top, last),
                    "=SUM(I{0}:I{1})".format(top, last)],
        {8: "#,##0.00", 9: "#,##0.00"}, {8: CALC_FILL, 9: CALC_FILL})
    put(ws, r_pop, ["人口（2024年度）", None, None, None, None, None, None, pop],
        {8: "#,##0"})
    put(ws, r_p100, ["人口10万対 ＝ 合計 ÷ 人口 × 100,000", None, None, None, None, None, None,
                     "=H{0}/$H${1}*100000".format(r_sum, r_pop),
                     "=I{0}/$H${1}*100000".format(r_sum, r_pop)],
        {8: "#,##0.000000", 9: "#,##0.000000"}, {8: CALC_FILL, 9: CALC_FILL})
    put(ws, r_p100 + 1, ["投稿用ファイルの値（Figure3_data.csv／Suppl Table S7）",
                         None, None, None, None, None, None,
                         float(pub.TOP3_lower), float(pub.TOP3_upper)],
        {8: "#,##0.000000", 9: "#,##0.000000"})
    put(ws, r_p100 + 2, ["差（式・0になる）", None, None, None, None, None, None,
                         "=H{0}-H{1}".format(r_p100, r_p100 + 1),
                         "=I{0}-I{1}".format(r_p100, r_p100 + 1)],
        {8: "0.000000", 9: "0.000000"}, {8: CHK_FILL, 9: CHK_FILL})
    for rr in (r_sum, r_p100):
        ws.cell(row=rr, column=1).font = Font(bold=True)
    ws.freeze_panes = "E%d" % top


# ----------------------------------------------------------------- 04
def sheet_per100k(wb):
    d = read_csv_utf8(os.path.join(CSV, "Fig1_age_sex_profile_bounds__Fig1A_TOP3.csv"))
    ws = wb.create_sheet("04_人口10万対_Fig1A")
    r = note(ws, 1, "数量の区間 → 人口10万対（Figure 1A・2024年度・主要3成分合計）", True)
    r = note(ws, r, "分母の人口は確定値なので、下限・上限をそれぞれ同じ人口で割るだけ。"
                    "E・F列が数式、H・I列が Figure1_data.csv との差。")
    r += 1
    top = header(ws, r,
                 ["年齢階級", "数量 下限（mL）", "数量 上限（mL）", "人口",
                  "人口10万対 下限（式）", "人口10万対 上限（式）",
                  "Figure1_data.csv の下限", "差 下限（式）", "差 上限（式）"],
                 [12, 18, 18, 14, 20, 20, 24, 14, 14])
    row = top
    for _, x in d.iterrows():
        put(ws, row, [
            x["Age Group"], float(x.count_lower), float(x.count_upper), float(x.population),
            "=B{0}/$D{0}*100000".format(row), "=C{0}/$D{0}*100000".format(row),
            float(x.per100k_lower), "=E{0}-G{0}".format(row),
            "=F{0}-{1}".format(row, float(x.per100k_upper)),
        ], {2: "#,##0.00", 3: "#,##0.00", 4: "#,##0", 5: "#,##0.000000",
            6: "#,##0.000000", 7: "#,##0.000000", 8: "0.000000", 9: "0.000000"},
            {5: CALC_FILL, 6: CALC_FILL, 8: CHK_FILL, 9: CHK_FILL})
        row += 1
    ws.freeze_panes = "B%d" % top


# ----------------------------------------------------------------- 05
def sheet_national(wb):
    d = pd.read_csv(os.path.join(OUT, "top3_full_range_2014_2024.csv"), encoding="utf-8-sig", comment="#")
    ws = wb.create_sheet("05_全国総量の上限")
    r = note(ws, 1, "全国総量の上限 ＝ 公表総量 ＋ 総計秘匿の加算 ＋ 未収載品目の加算", True)
    r = note(ws, r, "全国総量では都道府県セルの秘匿は効かない（総計列が公表されているため）。"
                    "上限が動くのは (1) 総計そのものが秘匿の行（999 − Σ開示セル を足す）と "
                    "(2) NDB に品目自体が載っていない未収載品目（順位ベースの上限を足す）だけ。")
    r += 1
    top = header(ws, r,
                 ["成分", "年度", "公表総量（mL）", "総計秘匿の加算（mL）",
                  "未収載品目の加算（mL）", "下限（式）", "上限（式）",
                  "Suppl Table S6 の上限", "差（式）"],
                 [22, 7, 20, 20, 20, 20, 20, 20, 14])
    row = top
    for _, x in d.iterrows():
        put(ws, row, [
            x.drug, int(x.year), float(x.published_total_mL),
            float(x.censored_total_cap_add_mL), float(x.unlisted_cap_add_rank_mL),
            "=C{0}".format(row), "=C{0}+D{0}+E{0}".format(row),
            float(x.max_total_mL_rank), "=G{0}-H{0}".format(row),
        ], {3: "#,##0.00", 4: "#,##0.00", 5: "#,##0.00", 6: "#,##0.00",
            7: "#,##0.00", 8: "#,##0.00", 9: "0.000000"},
            {6: CALC_FILL, 7: CALC_FILL, 9: CHK_FILL})
        row += 1
    ws.freeze_panes = "C%d" % top


# ----------------------------------------------------------------- 06
def sheet_share(wb):
    d = read_csv_utf8(os.path.join(
        CSV, "Fig5_ge_formulation_bounds__Fig5A_epinastine_LX.csv")).fillna(0.0)
    f5 = read_csv_utf8(os.path.join(OUT, "投稿用", "Figure5_data.csv"))
    f5 = f5[(f5.panel == "5A") & (f5.series == "0.1% share (%)")]
    ref = {str(int(float(x.category))): (float(x.value_lower), float(x.value_upper))
           for _, x in f5.iterrows()}

    ws = wb.create_sheet("06_割合の区間_Fig5A")
    r = note(ws, 1, "割合の区間（Figure 5A: エピナスチン0.1%製剤の割合）", True)
    r = note(ws, r, "分子・分母がどちらも区間を持つので、割合の下限は「自分が最小・相手が最大」、"
                    "上限は「自分が最大・相手が最小」で取る。F・G列が数式。")
    r += 1
    top = header(ws, r,
                 ["年度", "0.1% 下限（mL）", "0.1% 上限（mL）",
                  "0.05%+GE 下限（mL）", "0.05%+GE 上限（mL）",
                  "0.1%割合 下限%（式）", "0.1%割合 上限%（式）",
                  "Figure5_data.csv の下限", "差 下限（式）", "差 上限（式）"],
                 [7, 18, 18, 18, 18, 20, 20, 22, 13, 13])
    row = top
    for _, x in d.iterrows():
        lo, hi = ref.get(str(int(x.Year)), (0.0, 0.0))
        put(ws, row, [
            int(x.Year), float(x.LX_lower), float(x.LX_upper),
            float(x["Standard+GE_lower"]), float(x["Standard+GE_upper"]),
            "=IF(B{0}+E{0}=0,0,B{0}/(B{0}+E{0})*100)".format(row),
            "=IF(C{0}+D{0}=0,0,C{0}/(C{0}+D{0})*100)".format(row),
            lo, "=ROUND(F{0},4)-H{0}".format(row),
            "=ROUND(G{0},4)-{1}".format(row, hi),
        ], {2: "#,##0.00", 3: "#,##0.00", 4: "#,##0.00", 5: "#,##0.00",
            6: "0.0000", 7: "0.0000", 8: "0.0000", 9: "0.0000", 10: "0.0000"},
            {6: CALC_FILL, 7: CALC_FILL, 9: CHK_FILL, 10: CHK_FILL})
        row += 1
    ws.freeze_panes = "B%d" % top


# ----------------------------------------------------------------- 07
def sheet_units(wb):
    # 先頭にREADMEコメント行があるため comment='#' で読む（unit_conversion_audit.csv）
    d = pd.read_csv(os.path.join(FIG, "unit_conversion_audit.csv"),
                    encoding="utf-8-sig", comment="#")
    d = d.drop_duplicates(subset=["code", "unit", "ml_factor"])
    ws = wb.create_sheet("07_単位換算")
    r = note(ws, 1, "mL 換算係数（成分をまたぐ合算の前に必ず掛ける）", True)
    r = note(ws, r, "9成分の処方数量の単位は ｍＬ／瓶／個 が混在する。換算せずに足すと"
                    "「mL＋瓶」という無意味な量になる。品目別の全一覧は Suppl Table S1。")
    r += 1
    top = header(ws, r, ["成分コード", "成分", "単位", "mL係数", "代表品目", "根拠"],
                 [18, 22, 8, 10, 40, 60])
    row = top
    for _, x in d.iterrows():
        put(ws, row, [x.code, x.drug, x.unit, float(x.ml_factor),
                      x.product_name, x.ml_factor_basis], {4: "0.00"})
        row += 1
    ws.freeze_panes = "C%d" % top


def verify(lv, lg, fig3):
    """openpyxl は数式を計算しないので、同じ式を Python で再現して一致を確認する。"""
    for pref in EXAMPLE_PREFS:
        m = lg[lg.prefecture == pref].merge(
            lv[["sheet", "code", "product_name", "per_cell_cap"]],
            on=["sheet", "code", "product_name"], how="left")
        pop = pref_population(pref)
        lo = m.value.fillna(0).sum() / pop * 1e5
        hi = np.where(m.censored, m.per_cell_cap, m.value.fillna(0)).sum() / pop * 1e5
        pub = fig3[fig3.Prefecture == pref].iloc[0]
        assert abs(lo - float(pub.TOP3_lower)) < 1e-6, (pref, lo, pub.TOP3_lower)
        assert abs(hi - float(pub.TOP3_upper)) < 1e-6, (pref, hi, pub.TOP3_upper)
        print("   検算 %s: 下限 %.4f / 上限 %.4f（投稿用ファイルと一致）" % (pref, lo, hi))

    d = pd.read_csv(os.path.join(OUT, "top3_full_range_2014_2024.csv"), encoding="utf-8-sig", comment="#")
    w = (d.published_total_mL + d.censored_total_cap_add_mL
         + d.unlisted_cap_add_rank_mL - d.max_total_mL_rank).abs().max()
    assert w < 1e-6, w
    print("   検算 全国総量: 公表+秘匿加算+未収載加算 と Suppl Table S6 の最大差 %.1e" % w)


def main():
    lv, lg = load_inputs()
    fig3 = read_csv_utf8(os.path.join(
        CSV, "Fig3_prefecture_map_bounds__Fig3_2024_per100k.csv"))

    wb = Workbook()
    wb.remove(wb.active)
    sheet_readme(wb)
    lv_rows = sheet_level(wb, lv)
    for pref in EXAMPLE_PREFS:
        sheet_pref(wb, lv, lg, pref, lv_rows, fig3)
    sheet_per100k(wb)
    sheet_national(wb)
    sheet_share(wb)
    sheet_units(wb)
    wb.save(DEST)
    print("-> %s（%d シート）" % (DEST, len(wb.sheetnames)))
    verify(lv, lg, fig3)


if __name__ == "__main__":
    main()
