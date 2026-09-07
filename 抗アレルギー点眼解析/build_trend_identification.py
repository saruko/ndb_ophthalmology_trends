# -*- coding: utf-8 -*-
"""総量トレンドの「増加が識別できるか」判定表と Fig 1（総量推移・識別区間つき）を生成する。

書き分け（教室決定 2026-08-14）:
  - エピナスチン単独      : 2014年度から一貫して増加（2014〜2020年度は区間なしで確定）
  - 主要3成分合計         : 2015年度起点の増加は上限を含めても確定。2014年度起点は不定
  - 9成分合計             : 増加は主張しない（2014年度は3成分しか掲載なし）。2022年度以降の水準のみ
  - 3成分シェア           : 2022年度以降のみ比較（本スクリプトの対象外）

上限 = 公表値 + 総計秘匿の上限加算（national_trends_bounds） + 未収載品目の順位ベース上限加算
       （Fig4D_totals_full_period / top3_full_range_2014_2024）
  ※未収載品目の加算は足切りが働いていた 2021年度以前のみ（原稿 §2.5 の定義）。
    2022年度以降は総計秘匿の加算のみとし、Table 1・Suppl Table S1 と同じ上限に揃える。

出力（05_論文成果物/公費含めない_new/）:
  trend_identification_table.csv           成分別・起点別の増加判定（Suppl Table S4）
  論文図表/Fig1_trend_bounds_editable.xlsx   Fig 1 データ＋Excelネイティブ折れ線（エラーバー=上限）
  論文図表/Fig1_trend_bounds.png             参照用PNG（網掛け=識別区間、縦破線=2021→2022公開仕様変更）
"""
import os

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference, Series
from openpyxl.chart.data_source import NumDataSource, NumRef
from openpyxl.chart.error_bar import ErrorBars

plt.rcParams["font.family"] = ["Meiryo", "Yu Gothic", "MS Gothic"]
plt.rcParams["axes.unicode_minus"] = False

BASE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(BASE, "05_論文成果物", "公費含めない_new")
FIG = os.path.join(OUT, "論文図表")
CSV = os.path.join(FIG, "csv")

LABEL = {"EPINASTINE": "エピナスチン", "OLOPATADINE": "オロパタジン",
         "LEVOCASTINE": "レボカバスチン", "TOP3_TOTAL": "主要3成分合計",
         "ALLERGY_EYE_TOTAL": "9成分合計"}
LABEL_EN = {"EPINASTINE": "Epinastine", "OLOPATADINE": "Olopatadine",
            "LEVOCASTINE": "Levocabastine", "TOP3_TOTAL": "Top 3 total",
            "ALLERGY_EYE_TOTAL": "All 9 agents"}
ORDER = list(LABEL)


def load():
    t3 = pd.read_csv(os.path.join(OUT, "top3_full_range_2014_2024.csv"), comment="#")
    nat = pd.read_csv(os.path.join(OUT, "national_trends_bounds.csv"))
    d = pd.read_csv(os.path.join(CSV, "Fig4_trends_shares_bounds__Fig4D_totals_full_period.csv"))

    # 2022年度以降は足切りなし → 未収載加算を外し、総計秘匿の加算のみ
    t3 = t3.copy()
    post = t3.year >= 2022
    t3.loc[post, "max_total_mL_rank"] = t3.loc[post, "min_total_mL"] + t3.loc[post, "censored_total_cap_add_mL"]

    rows = []
    for _, r in t3.iterrows():
        rows.append({"code": r.code, "year": r.year,
                     "published_mL": r.published_total_mL,
                     "lower_mL": r.min_total_mL, "upper_mL": r.max_total_mL_rank})
    # 主要3成分合計 = 3成分の下限・上限の和
    g = t3.groupby("year").agg(published_mL=("published_total_mL", "sum"),
                               lower_mL=("min_total_mL", "sum"),
                               upper_mL=("max_total_mL_rank", "sum")).reset_index()
    for _, r in g.iterrows():
        rows.append({"code": "TOP3_TOTAL", **r.to_dict()})
    # 9成分合計: 未収載加算は Fig4D（2021年度以前のみ）、総計秘匿加算は national_trends_bounds から
    n9 = nat[nat.code == "ALLERGY_EYE_TOTAL"].set_index("year")
    for _, r in d.iterrows():
        y = int(r.Year)
        tc_add = float(n9.loc[y, "count_upper"] - n9.loc[y, "count_lower"]) if y in n9.index else 0.0
        unl_add = (r.all9_upper_mL - r.all9_published_mL) if y <= 2021 else 0.0
        rows.append({"code": "ALLERGY_EYE_TOTAL", "year": y,
                     "published_mL": r.all9_published_mL,
                     "lower_mL": r.all9_published_mL,
                     "upper_mL": r.all9_published_mL + unl_add + tc_add})
    df = pd.DataFrame(rows)
    df["label"] = df.code.map(LABEL)
    df["label_en"] = df.code.map(LABEL_EN)
    return df.sort_values(["code", "year"], key=lambda s: s.map(
        {c: i for i, c in enumerate(ORDER)}) if s.name == "code" else s)


def identification_table(df):
    """起点年 y0 → 2024 の増加が識別できるか: upper(y0) < lower(2024)."""
    out = []
    for code in ORDER:
        s = df[df.code == code].set_index("year")
        lo24, hi24 = s.loc[2024, "lower_mL"], s.loc[2024, "upper_mL"]
        for y0 in (2014, 2015, 2022):
            if y0 not in s.index:
                continue
            lo0, hi0 = s.loc[y0, "lower_mL"], s.loc[y0, "upper_mL"]
            if hi0 < lo24:
                verdict, direction = "確定", "増加"
            elif lo0 > hi24:
                verdict, direction = "確定", "減少"
            else:
                verdict, direction = "不定（区間が重なる）", "—"
            out.append({
                "code": code, "label": LABEL[code], "start_year": y0,
                "start_lower_mL": lo0, "start_upper_mL": hi0,
                "start_width_pct": (hi0 - lo0) / lo0 * 100 if lo0 else np.nan,
                "fy2024_lower_mL": lo24, "fy2024_upper_mL": hi24,
                "change_identified": verdict, "direction": direction,
                "published_change_pct": (s.loc[2024, "published_mL"] / s.loc[y0, "published_mL"] - 1) * 100
                if s.loc[y0, "published_mL"] else np.nan,
            })
    t = pd.DataFrame(out)
    p = os.path.join(OUT, "trend_identification_table.csv")
    t.to_csv(p, index=False, encoding="utf-8-sig")
    print(f"-> {p}")
    return t


def excel_fig1(df):
    wb = Workbook()
    ws = wb.active
    ws.title = "Fig1_data"
    years = sorted(df.year.unique())
    hdr = ["Year"]
    for c in ORDER:
        hdr += [f"{LABEL_EN[c]}_lower", f"{LABEL_EN[c]}_add(upper-lower)",
                f"{LABEL_EN[c]}_upper"]
    hdr += ["zero"]
    ws.append(hdr)
    for y in years:
        row = [y]
        for c in ORDER:
            s = df[(df.code == c) & (df.year == y)]
            if s.empty:
                row += [None, None, None]
            else:
                lo, hi = float(s.lower_mL.iloc[0]), float(s.upper_mL.iloc[0])
                row += [lo / 1e6, (hi - lo) / 1e6, hi / 1e6]
        row += [0.0]
        ws.append(row)
    n = len(years)
    zero_col = len(hdr)

    def num(col):
        L = ws.cell(row=1, column=col).column_letter
        return NumDataSource(NumRef(f=f"'{ws.title}'!${L}$2:${L}${n + 1}"))

    def chart(title, codes, anchor, colors):
        ch = LineChart()
        ch.title = title
        ch.y_axis.title = "処方量（百万mL）"
        ch.x_axis.title = "年度"
        ch.x_axis.delete = False
        ch.y_axis.delete = False
        ch.height, ch.width = 12, 26
        for c, col in zip(codes, colors):
            i = ORDER.index(c)
            lo_col = 2 + i * 3
            add_col = lo_col + 1
            s = Series(Reference(ws, min_col=lo_col, min_row=2, max_row=n + 1),
                       title=f"{LABEL_EN[c]}（下限=公表値）")
            s.graphicalProperties.line.solidFill = col
            s.graphicalProperties.line.width = 22000
            s.marker.symbol = "circle"
            s.errBars = ErrorBars(errBarType="both", errValType="cust",
                                  plus=num(add_col), minus=num(zero_col),
                                  noEndCap=False)
            ch.series.append(s)
        ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=n + 1))
        ws.add_chart(ch, anchor)

    chart("Fig 1A 主要3成分の全国処方量 2014〜2024（エラーバー=秘匿・未収載を含む上限）",
          ["EPINASTINE", "OLOPATADINE", "LEVOCASTINE"], "R2",
          ["1F77B4", "FF7F0E", "2CA02C"])
    chart("Fig 1B 主要3成分合計・9成分合計（同上）",
          ["TOP3_TOTAL", "ALLERGY_EYE_TOTAL"], "R28",
          ["7F7F7F", "9467BD"])
    ws["R54"] = "※2021→2022年度の間に公開仕様の変更（上位品目の足切り緩和・瓶単位6成分の一斉掲載）あり。図中に縦破線＋注記を手動で追加すること。"
    ws["R55"] = "※9成分合計の2014〜2021年度の上限は極端に大きい（未収載成分の順位ベース上限）。軸を切るか対数軸にすること。"
    ws["R56"] = "※エラーバー: プラス側=上限−下限（add列）、マイナス側=0（zero列）。"
    p = os.path.join(FIG, "Fig1_trend_bounds_editable.xlsx")
    wb.save(p)
    print(f"-> {p}")


def png_fig1(df):
    from matplotlib.patches import Patch
    fig, axes = plt.subplots(1, 2, figsize=(13, 5.8))
    cols = {"EPINASTINE": "#1f77b4", "OLOPATADINE": "#ff7f0e",
            "LEVOCASTINE": "#2ca02c", "TOP3_TOTAL": "#555555",
            "ALLERGY_EYE_TOTAL": "#9467bd"}
    for ax, codes, ttl in (
        (axes[0], ["EPINASTINE", "OLOPATADINE", "LEVOCASTINE"], "A. 主要3成分（成分別）"),
        (axes[1], ["TOP3_TOTAL", "ALLERGY_EYE_TOTAL"], "B. 主要3成分合計・9成分合計（対数軸）"),
    ):
        for c in codes:
            s = df[df.code == c].sort_values("year")
            ax.plot(s.year, s.lower_mL / 1e6, "o-", color=cols[c], lw=1.8, ms=4,
                    label=f"{LABEL[c]}（下限=公表値）")
            ax.fill_between(s.year, s.lower_mL / 1e6, s.upper_mL / 1e6,
                            color=cols[c], alpha=0.18, lw=0)
        ax.axvline(2021.5, color="black", ls="--", lw=1)
        ax.set_xlabel("年度")
        ax.set_ylabel("処方量（百万mL）")
        ax.set_title(ttl, fontsize=11, loc="left")
        ax.set_xticks(range(2014, 2025))
        ax.tick_params(axis="x", labelsize=8)
        ax.grid(alpha=0.3)
        # 凡例は軸の下に置き、2014年度の区間の頂点を隠さない。網掛けも凡例に登録する
        h, l = ax.get_legend_handles_labels()
        h.append(Patch(facecolor="#888888", alpha=0.3, lw=0))
        l.append("網掛け＝識別区間（下限〜上限）")
        ax.legend(h, l, fontsize=8, loc="upper center", bbox_to_anchor=(0.5, -0.13),
                  ncol=2, frameon=False)
    axes[0].set_ylim(bottom=0)
    axes[1].set_yscale("log")
    axes[1].set_ylabel("処方量（百万mL、対数軸）")
    axes[1].set_ylim(top=3000)
    for ax in axes:
        ax.annotate("公開仕様の変更\n（足切り緩和・6成分一斉掲載）",
                    xy=(2021.5, ax.get_ylim()[0]), xycoords="data",
                    xytext=(0.50, 0.30) if ax is axes[0] else (0.62, 0.06),
                    textcoords="axes fraction",
                    fontsize=8, ha="left",
                    arrowprops=dict(arrowstyle="-", color="black", lw=0.6))
    axes[1].set_ylim(bottom=80)
    axes[1].annotate("2014年度は3成分のみ掲載\n（両系列の下限が一致）",
                     xy=(2014, 109.06), xytext=(2015.2, 88), fontsize=7,
                     va="center", color="#444444",
                     arrowprops=dict(arrowstyle="-", color="#444444", lw=0.6))
    fig.suptitle("全国処方量の推移 2014〜2024年度（網掛け＝セル秘匿・総計秘匿・未収載品目を含む識別区間）",
                 fontsize=11)
    fig.tight_layout()
    p = os.path.join(FIG, "Fig1_trend_bounds.png")
    fig.savefig(p, dpi=300)
    plt.close(fig)
    print(f"-> {p}")


if __name__ == "__main__":
    df = load()
    df.to_csv(os.path.join(CSV, "Fig1_trend_bounds_data.csv"),
              index=False, encoding="utf-8-sig")
    t = identification_table(df)
    print(t[["label", "start_year", "start_upper_mL", "fy2024_lower_mL",
             "change_identified", "direction"]].to_string(index=False))
    excel_fig1(df)
    png_fig1(df)
