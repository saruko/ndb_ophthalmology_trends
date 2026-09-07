# -*- coding: utf-8 -*-
"""Word 文書に「変更履歴つき」で置換・追記するための補助。

原稿 docx は査読者の変更履歴（w:ins / w:del）が入ったまま回っているため、
本ツールの編集も履歴として残す（既存の履歴を確定・破棄しない）。
既に author="Claude" の履歴が入っているので、同じ著者名を使う。
"""
import copy

from docx.oxml.ns import qn
from lxml import etree

AUTHOR = "Claude"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


class TrackedEditor:
    def __init__(self, date, author=AUTHOR):
        self.date = date
        self.author = author
        self.n = 0

    def _attrs(self, el):
        self.n += 1
        el.set(qn("w:id"), str(900000 + self.n))
        el.set(qn("w:author"), self.author)
        el.set(qn("w:date"), self.date)
        return el

    def _run(self, template, text, deleted=False):
        r = etree.SubElement(etree.Element("dummy"), qn("w:r"))
        rPr = template.find(qn("w:rPr")) if template is not None else None
        if rPr is not None:
            r.append(copy.deepcopy(rPr))
        t = etree.SubElement(r, qn("w:delText") if deleted else qn("w:t"))
        t.text = text
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        return r

    def _del(self, run):
        d = self._attrs(etree.Element(qn("w:del")))
        d.append(run)
        return d

    def _ins(self, run):
        i = self._attrs(etree.Element(qn("w:ins")))
        i.append(run)
        return i

    def _splice(self, r, start, end, ins_text):
        """run r の [start:end) を削除履歴にし、必要なら ins_text を挿入履歴として続ける。"""
        t = r.find(qn("w:t"))
        text = t.text or ""
        pre, mid, post = text[:start], text[start:end], text[end:]
        nodes = []
        if pre:
            nodes.append(self._run(r, pre))
        if mid:
            nodes.append(self._del(self._run(r, mid, deleted=True)))
        if ins_text:
            nodes.append(self._ins(self._run(r, ins_text)))
        if post:
            nodes.append(self._run(r, post))
        parent = r.getparent()
        if parent.tag == qn("w:ins"):
            # 既存の挿入履歴の中の文字を消す → その著者の w:ins の中に自分の w:del を入れ子にする
            holder = copy.deepcopy(parent)
            for ch in list(holder):
                holder.remove(ch)
            out = []
            for nd in nodes:
                if nd.tag == qn("w:ins"):
                    out.append(nd)                 # 自分の挿入は独立した w:ins
                else:
                    h = copy.deepcopy(holder)
                    h.append(nd)
                    out.append(h)
            gp = parent.getparent()
            idx = list(gp).index(parent)
            sib = list(parent)
            pos = sib.index(r)
            head, tail = sib[:pos], sib[pos + 1:]
            for ch in list(parent):
                parent.remove(ch)
            seq = []
            if head:
                for ch in head:
                    parent.append(ch)
                seq.append(parent)
            else:
                gp.remove(parent)
            seq += out
            if tail:
                h = copy.deepcopy(holder)
                for ch in tail:
                    h.append(ch)
                seq.append(h)
            for k, nd in enumerate(seq):
                if nd is parent:
                    continue
                gp.insert(idx + k, nd)
        else:
            idx = list(parent).index(r)
            parent.remove(r)
            for k, nd in enumerate(nodes):
                parent.insert(idx + k, nd)

    def replace(self, par, old, new):
        """段落 par の中の old を new に置き換える（old は削除履歴、new は挿入履歴）。

        old が複数の run にまたがっていてもよい。挿入は最初の run の位置に置く。
        """
        runs, spans, acc = [], [], 0
        for r in par._p.iter(qn("w:r")):
            t = r.find(qn("w:t"))
            if t is None or t.text is None:
                continue
            runs.append(r)
            spans.append((acc, acc + len(t.text)))
            acc += len(t.text)
        text = "".join((r.find(qn("w:t")).text or "") for r in runs)
        at = text.find(old)
        if at < 0:
            raise KeyError(old[:40])
        end = at + len(old)
        hits = [(r, s, e) for r, (s, e) in zip(runs, spans) if s < end and e > at]
        first = True
        for r, s, e in hits:                       # 後ろから触ると位置がずれないが、
            lo, hi = max(at, s) - s, min(end, e) - s   # run ごとに独立なので順序は不問
            self._splice(r, lo, hi, new if first else None)
            first = False
        return True

    def add_paragraph_after(self, par, text):
        """par の直後に、挿入履歴として段落を1つ追加する（書式は par を踏襲）。"""
        p = copy.deepcopy(par._p)
        for ch in list(p):
            if ch.tag not in (qn("w:pPr"),):
                p.remove(ch)
        pPr = p.find(qn("w:pPr"))
        if pPr is None:
            pPr = etree.SubElement(p, qn("w:pPr"))
            p.insert(0, pPr)
        rPr = pPr.find(qn("w:rPr"))
        if rPr is None:
            rPr = etree.SubElement(pPr, qn("w:rPr"))
        rPr.insert(0, self._attrs(etree.Element(qn("w:ins"))))   # 段落記号も挿入扱い
        template = None
        for r in par._p.iter(qn("w:r")):
            template = r
            break
        p.append(self._ins(self._run(template, text)))
        par._p.addnext(p)
        return p

    def delete_paragraph(self, par):
        """段落を丸ごと削除履歴にする（本文の run と段落記号の両方）。"""
        p = par._p
        pPr = p.find(qn("w:pPr"))
        if pPr is None:
            pPr = etree.Element(qn("w:pPr"))
            p.insert(0, pPr)
        rPr = pPr.find(qn("w:rPr"))
        if rPr is None:
            rPr = etree.SubElement(pPr, qn("w:rPr"))
        rPr.insert(0, self._attrs(etree.Element(qn("w:del"))))   # 段落記号
        for r in list(p.iter(qn("w:r"))):
            parent = r.getparent()
            if parent is None or parent.tag == qn("w:del"):
                continue
            t = r.find(qn("w:t"))
            if t is None:
                continue
            dt = etree.Element(qn("w:delText"))
            dt.text = t.text or ""
            dt.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            r.replace(t, dt)
            idx = list(parent).index(r)
            parent.remove(r)
            parent.insert(idx, self._del(r))
        return p

    def insert_paragraphs_after(self, par, texts, template=None):
        """par の直後に複数段落を挿入履歴として追加する。書式は template（既定は par）。"""
        anchor = par
        out = []
        for txt in texts:
            src = template if template is not None else anchor
            newp = self.add_paragraph_after_styled(anchor, txt, src)
            out.append(newp)
            anchor = _ParaLike(newp)
        return out

    def add_paragraph_after_styled(self, par, text, style_src):
        """par の直後に、style_src の書式を持つ段落を挿入履歴として追加する。"""
        p = copy.deepcopy(style_src._p if hasattr(style_src, "_p") else style_src)
        for ch in list(p):
            if ch.tag != qn("w:pPr"):
                p.remove(ch)
        pPr = p.find(qn("w:pPr"))
        if pPr is None:
            pPr = etree.Element(qn("w:pPr"))
            p.insert(0, pPr)
        rPr = pPr.find(qn("w:rPr"))
        if rPr is None:
            rPr = etree.SubElement(pPr, qn("w:rPr"))
        # 書式見本が「削除済み段落」だった場合、その w:del を引き継がない
        for stale in rPr.findall(qn("w:del")) + rPr.findall(qn("w:ins")):
            rPr.remove(stale)
        rPr.insert(0, self._attrs(etree.Element(qn("w:ins"))))
        template = None
        for r in (style_src._p if hasattr(style_src, "_p") else style_src).iter(qn("w:r")):
            template = r
            break
        p.append(self._ins(self._run(template, text)))
        (par._p if hasattr(par, "_p") else par).addnext(p)
        return p


class _ParaLike:
    """add_paragraph_after 系に渡すための最小ラッパ（_p だけ持つ）。"""

    def __init__(self, p):
        self._p = p
