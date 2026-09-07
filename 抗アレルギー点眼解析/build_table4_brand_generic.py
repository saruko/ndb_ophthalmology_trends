# -*- coding: utf-8 -*-
"""Table4: 先発品/後発品を分け、後発品を「NDB掲載」「上市済み未掲載」に区別した表。

Tables_123_bounds.xlsx と同じ様式（READMEシート＋データシート＋csv/への同内容CSV）。
後発品の上市メーカー情報は drug_master.py の MARKETED_GENERIC_MAKERS
（NDB公開構造と秘匿実態_9成分_統合版.md ③章のファクトチェック済み情報）に基づく。

シート:
  Table4A_summary   成分×年度: 先発（掲載/上市品目数・数量区間）、
                    後発（掲載/上市/未掲載メーカー数・掲載分数量区間・
                    未掲載分の順位ベース上限）、GE比率区間とA〜D分類
  Table4B_makers    成分×メーカー×年度: ○=NDB掲載 / ×=上市済み未掲載 / —=未上市

出力: 05_論文成果物/公費含めない_new/論文図表/Table4_brand_generic_listing_bounds.xlsx
"""
import os
import re
import sys

import numpy as np
import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(BASE, "src"))

from drug_master import (DRUG_NAME_JA, DRUG_ORDER,            # noqa: E402
                         MARKETED_GENERIC_MAKERS)
from paths import BASE_DIR, INPUT_DIR                          # noqa: E402

NEW_DIR = os.path.join(BASE_DIR, "05_論文成果物", "公費含めない_new")
OUT_DIR = os.path.join(NEW_DIR, "論文図表")
CSV_DIR = os.path.join(OUT_DIR, "csv")

YEARS = list(range(2014, 2025))

_ZEN = str.maketrans(
    "ＡＢＣＤＥＦＧＨＩＪＫＬＭＮＯＰＱＲＳＴＵＶＷＸＹＺ０１２３４５６７８９",
    "ABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789")


def norm(s: str) -> str:
    return str(s).translate(_ZEN).strip()


def listed_maker_of(product_name: str):
    """品目名の「」内をメーカー名として取り出す（銘柄別収載の後発品）。"""
    m = re.search(r"「(.+?)」", str(product_name))
    return norm(m.group(1)) if m else None


def marketed_makers(code, year):
    """上市メーカー名の集合と（名簿が無い期間の）注記を返す。"""
    for (y0, y1, lo, hi, note) in [(r[0], r[1], r[2], r[3], r[4])
                                   for r in MARKETED_GENERIC_MAKERS[code]]:
        if y0 <= year <= y1:
            if lo == 0:
                return set(), "未発売", lo
            m = re.search(r"[:：]\s*(.+)$", note)
            names = {norm(x) for x in m.group(1).split(",")} if m else set()
            return names, note, lo
    return set(), "", 0


def main():
    os.makedirs(CSV_DIR, exist_ok=True)
    level = pd.read_csv(os.path.join(INPUT_DIR, "product_level_censoring.csv"))
    inv = pd.read_csv(os.path.join(NEW_DIR, "unlisted_products_inventory.csv"))
    ge = pd.read_csv(os.path.join(NEW_DIR, "brand_generic_share_bounds.csv"))
    geann = pd.read_csv(os.path.join(BASE_DIR, "03_解析結果", "後発品_剤形",
                                     "ge_share_annotated.csv"))
    uns = pd.read_csv(os.path.join(NEW_DIR, "unlisted_sensitivity.csv"),
                      comment="#")

    # ------------------------------------------------------------------
    # Table4A: 成分×年度サマリ
    # ------------------------------------------------------------------
    ge_i = ge.set_index(["code", "year"])
    ann_i = geann.set_index(["code", "year"])
    uns_i = uns.set_index(["code", "year"])
    rows = []
    for code in DRUG_ORDER:
        for year in YEARS:
            iv = inv[(inv.code == code) & (inv.year == year)]
            if iv.empty:
                continue
            iv = iv.iloc[0]
            rec = {
                "Year": year, "Drug": DRUG_NAME_JA[code],
                # ---- 先発品 ----
                "先発_NDB掲載品目数": int(iv.listed_brand_products),
                "先発_上市品目数": int(iv.marketed_brand),
                "先発_未掲載品目数": int(iv.unlisted_brand),
            }
            if (code, year) in ge_i.index:
                g = ge_i.loc[(code, year)]
                rec |= {
                    "先発_数量下限_mL": g.quantity_brand_lower,
                    "先発_数量上限_mL": g.quantity_brand_upper,
                    "後発_掲載分数量下限_mL": g.quantity_generic_lower,
                    "後発_掲載分数量上限_mL": g.quantity_generic_upper,
                    "GE比率下限_%": g.share_pct_generic_lower,
                    "GE比率上限_%": g.share_pct_generic_upper,
                }
            # ---- 後発品（メーカー数）----
            hi = iv.marketed_generic_hi
            rec |= {
                "後発_NDB掲載メーカー数": int(iv.listed_generic_makers),
                "後発_上市メーカー数": (f"{int(iv.marketed_generic_lo)}以上"
                                if pd.isna(hi) else
                                f"{int(iv.marketed_generic_lo)}"
                                if iv.marketed_generic_lo == hi else
                                f"{int(iv.marketed_generic_lo)}〜{int(hi)}"),
                "後発_上市済み未掲載メーカー数(下限)": int(iv.unlisted_generic_lo),
            }
            # ---- 未掲載品目分の順位ベース上限 ----
            if (code, year) in uns_i.index:
                u = uns_i.loc[(code, year)]
                if u.listed_in_ndb and pd.notna(u.upper1):
                    rec["未掲載品目分の上限_mL"] = u.upper1 - u.published_total
            if (code, year) in ann_i.index:
                a = ann_i.loc[(code, year)]
                rec["GE比率の解釈区分"] = f"{a.ge_class}: {a.ge_class_label}"
            rows.append(rec)
    t4a = pd.DataFrame(rows)

    # ------------------------------------------------------------------
    # Table4B: メーカー別の掲載状況マトリクス
    # ------------------------------------------------------------------
    lst = level[level.product_type == "generic_maker"].copy()
    lst["maker"] = lst.product_name.map(listed_maker_of)
    listed_by = lst.groupby(["code", "year"])["maker"].agg(
        lambda s: {x for x in s if x})

    rows_b, unmatched = [], set()
    for code in DRUG_ORDER:
        all_names = set()
        for year in YEARS:
            names, _, _ = marketed_makers(code, year)
            all_names |= names
        # 名簿が無い期間（レボカバスチン2014-2021の概数表記）に備え、
        # NDBに実際に掲載されたメーカーも候補に加える
        for year in YEARS:
            all_names |= listed_by.get((code, year), set())
        if not all_names:
            continue
        for name in sorted(all_names):
            rec = {"Drug": DRUG_NAME_JA[code], "Maker": name}
            for year in YEARS:
                marketed, note, lo = marketed_makers(code, year)
                listed = listed_by.get((code, year), set())
                if name in listed:
                    rec[str(year)] = "○"        # NDB掲載
                elif name in marketed:
                    rec[str(year)] = "×"        # 上市済み・未掲載
                elif lo > 0 and not marketed:
                    rec[str(year)] = "×?"       # 上市数のみ判明（名簿なし）
                else:
                    rec[str(year)] = "—"        # 未上市
                if name in listed and marketed and name not in marketed:
                    unmatched.add((code, name, year))
            rows_b.append(rec)
    t4b = pd.DataFrame(rows_b)

    # ------------------------------------------------------------------
    # 出力（Tables_123_bounds.xlsx と同じ様式）
    # ------------------------------------------------------------------
    notes = [
        "抗アレルギー点眼薬9成分（NDBオープンデータ第1〜11回、公費レセプトを含まない集計）",
        "数量はすべてmL換算。下限/上限は秘匿識別区間（bounded_outputs_report.md 参照）",
        "後発品の上市メーカー名簿は NDB公開構造と秘匿実態_9成分_統合版.md ③章"
        "（src/drug_master.py MARKETED_GENERIC_MAKERS）に基づく",
        "Table4A: 後発_掲載分数量は NDB掲載品目のみの合算（未掲載分を含まない下限）。"
        "未掲載品目分の上限_mL は順位ベース（未掲載品目数×最小公表総計、mL換算）",
        "Table4A: 統一名収載品はメーカー数に数えない（1銘柄として数量には含む）",
        "Table4B: ○=NDB掲載 / ×=上市済み未掲載 / ×?=上市済み未掲載と推定"
        "（レボカバスチン2014〜2021年度は出典が「約10〜11社」の概数でメーカー名簿なし）/ —=未上市",
        "GE比率の解釈区分: A=真値0%（後発未発売）, B=算出不能（足切りで非掲載）, "
        "C=比較不能（収載初年度・部分年）, D=信頼可（2022年度以降）",
    ]
    path = os.path.join(OUT_DIR, "Table4_brand_generic_listing_bounds.xlsx")
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        pd.DataFrame({"Notes": notes}).to_excel(xw, sheet_name="README", index=False)
        for name, df in [("Table4A_summary", t4a), ("Table4B_makers", t4b)]:
            df.to_excel(xw, sheet_name=name, index=False)
            df.to_csv(os.path.join(
                CSV_DIR, f"Table4_brand_generic_listing_bounds__{name}.csv"),
                index=False, encoding="utf-8-sig")
        for ws in xw.book.worksheets:
            for cell in ws[1]:
                cell.font = Font(name="Arial", bold=True)
            for i, col in enumerate(ws.iter_cols(min_row=1, max_row=1), 1):
                ws.column_dimensions[get_column_letter(i)].width = max(
                    8, min(30, len(str(col[0].value or "")) + 4))
            ws.freeze_panes = "A2"
    print(f"-> {os.path.relpath(path, BASE)} "
          f"(Table4A {len(t4a)} rows / Table4B {len(t4b)} rows)")
    if unmatched:
        print("★名簿と照合できなかった掲載メーカー:", sorted(unmatched))


if __name__ == "__main__":
    main()
