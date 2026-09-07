# -*- coding: utf-8 -*-
"""原稿 docx（NDB抗アレルギー解析_ver公費なしnew.docx）を、md 原稿の改訂に合わせて更新する。

内容（いずれも 2026-08-19 の改訂。詳細は md の付記6・6-2・6-3）:
  1. シェアの識別区間の式を sharp なものに統一（Table 1・Table 3 の legend と数値）
  2. Fig 4B に上下エラーバー、Fig 1C（年齢階級別構成比）を新設したことの反映
  3. Supplementary Table S8（構成比の識別区間）の追加

編集はすべて変更履歴（author="Claude"）として入れる。既存の履歴は触らない。
"""
import os
import shutil

import docx
from docx.oxml.ns import qn

from docx_track import TrackedEditor

BASE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(BASE, "05_論文成果物", "公費含めない_new")
SRC = os.path.join(OUT, "NDB抗アレルギー解析_ver公費なしnew.docx")
BACKUP = os.path.join(OUT, "NDB抗アレルギー解析_ver公費なしnew_backup_20260819.docx")
DATE = "2026-08-19T00:00:00Z"

SHARP = ("［成分の下限／（成分の下限＋他成分の上限の和），"
         "成分の上限／（成分の上限＋他成分の下限の和）］")
NAIVE = "［成分の下限／9成分合計の上限，成分の上限／9成分合計の下限］"


def text_of(el):
    """変更履歴（w:ins）の中の文字も拾う（python-docx の .text は拾わない）。"""
    return "".join(n.text or "" for n in el.iter(qn("w:t")))


def para_by(doc, needle, start=0):
    for i, p in enumerate(doc.paragraphs):
        if i < start:
            continue
        if needle in text_of(p._p):
            return p
    raise KeyError(needle[:40])


def main():
    if not os.path.exists(BACKUP):
        shutil.copyfile(SRC, BACKUP)
        print("-> backup %s" % os.path.basename(BACKUP))
    d = docx.Document(SRC)
    ed = TrackedEditor(DATE)

    # --- 1. Table 3 を sharp な式で作り直したことによる数値の変更 -----------------
    ed.replace(para_by(d, "Results: 2024年度の9成分合計"), "2024年度85.7〜85.8%", "2024年度85.8%")
    ed.replace(para_by(d, "これらの主要3成分の合計は9成分全体の"), "85.7〜85.8%を占めた", "85.8%を占めた")
    ed.replace(para_by(d, "9成分合計に占めるシェアは2022〜2024年度で"), "→85.7〜85.8%と漸増", "→85.8%と漸増")

    t3 = [t for t in d.tables if text_of(t.rows[0].cells[0]._tc).strip() == "年度"][0]
    cells = {(text_of(r.cells[0]._tc).strip(), i): r.cells[i]
             for r in t3.rows[1:] for i in range(7)}
    for (year, col), old, new in ((("2024", 1), "85.7–85.8", "85.8"),
                                  (("2022", 5), "87.3–87.4", "87.3"),
                                  (("2022", 6), "12.6–12.7", "12.7"),
                                  (("2023", 5), "87.9–88.1", "88.0")):
        ed.replace(cells[(year, col)].paragraphs[0], old, new)

    for key in ("Table 1. 2024年度の成分別処方量", "Table 3. 9成分合計に占める成分別シェア"):
        ed.replace(para_by(d, key), NAIVE, SHARP)

    # --- 2. Fig 1C の新設・Fig 4B のエラーバー ----------------------------------
    p62 = para_by(d, "区間の提示：本研究は記述研究であり")
    ed.replace(p62, "（Figure 1A・1B・2・3A・4A・5B）。",
               "（Figure 1A・1B・2・3A・4A）。割合を描いた図（Figure 1C・4B・5B）では、"
               "描いた値が公表値どうしの比であり識別区間の中間に位置するため、エラーバーを"
               "上下両側に描いた。割合・シェアの識別区間は、分子が分母の一部であること"
               "（分子と分母が同時に最悪値を取れないこと）を用いて、本研究を通じて"
               + SHARP + "として算出し（Table 1・Table 3、Figure 1C・4B・5A・5B・6A・6B で共通）、"
               "構成比の数値は Supplementary Table S8 に示した。")
    ed.replace(p62, "年齢階級別の図（Figure 1A・1B・2・5B）", "年齢階級別の図（Figure 1A・1B・1C・2・5B）")
    ed.replace(p62,
               "全国総計に基づく割合（Figure 4B・5A・6A・6B）は総計が公表値であるため"
               "区間幅が0.01ポイント未満であり、エラーバーは描かず legend に幅を記した。",
               "全国総計に基づく割合のうち Figure 5A・6A・6B は総計が公表値であるため区間幅が"
               "全年度で0.01ポイント未満であり、エラーバーは描かず legend に幅を記した。"
               "Figure 4B も2022年度以降は0.01ポイント未満であるが、2021年度以前は品目非掲載の"
               "影響で最大38ポイントに達するため、上下のエラーバーを描いた。")

    p74 = para_by(d, "主要3成分の内訳も年齢とともに系統的に変化した")
    ed.replace(p74, "（Figure 1B）", "（Figure 1C）")
    ed.replace(p74, "エピナスチンの占める割合は0〜34歳で約70%", "エピナスチンの占める割合は5〜34歳で約70%")
    ed.replace(p74, "35歳以降は加齢とともに単調に低下し", "35歳以降は加齢とともに概ね単調に低下し")
    ed.replace(p74, "これらは下限（公表値）に基づく構成比であり、95歳以上は除いて読む。",
               "点は公表値どうしの比、エラーバーは識別区間である（数値は Supplementary Table S8）。"
               "区間を考慮すると、0〜4歳のエピナスチンは64.1〜72.3%と幅が広く「約70%」とは"
               "断定できない。また70〜74歳・75〜79歳・80〜84歳（それぞれ47.9〜48.8%、"
               "47.7〜48.4%、47.0〜47.9%）は区間が重なるため、この3階級間の順序は確定しない。"
               "全体としての低下、およびオロパタジン・レボカバスチンの上昇は、区間を考慮しても"
               "確定する。95歳以上は除いて読む。")

    ed.replace(para_by(d, "同図は視覚的な連続性のため2014年度から描いているが"),
               "2021年度以前は上位品目の足切りの影響を受けるため参考値であり、数値の解釈は行わない。",
               "2021年度以前は上位品目の足切りの影響を受けるため参考値であり、数値の解釈は行わない。"
               "実際、構成比の識別区間は2014年度で最大37.7ポイント、2021年度で最大19.1ポイントに"
               "達する一方、2022年度以降は0.01ポイント未満である（Supplementary Table S8）。")

    ed.replace(para_by(d, "(A) 年齢階級別の人口10万対の主要3成分合計処方量"),
               "(B) 同じ集計を成分別に分けたもの（エピナスチン、オロパタジン、レボカバスチン）。"
               "棒・点は識別区間の下限（公表値）、エラーバーは上限を示す。",
               "(B) 同じ集計を成分別に分けたもの（エピナスチン、オロパタジン、レボカバスチン）。"
               "(C) 主要3成分合計に占める成分別の構成比（%）。(A)(B) の棒・点は識別区間の下限"
               "（公表値）でエラーバーは上限、(C) の点は公表値どうしの比でエラーバーは識別区間の"
               "下限・上限（上下両側）を示す。")
    ed.replace(para_by(d, "(A) 年齢階級別の人口10万対の主要3成分合計処方量"),
               "95〜99歳は区間が広いため、水準および成分別構成の比較には用いない。",
               "95〜99歳は区間が広いため、水準および成分別構成の比較には用いない。"
               "(C) の区間幅は0〜4歳で4.9〜8.2ポイント、5〜89歳で0.6〜5.0ポイント、"
               "90〜94歳で4.9〜5.9ポイント、95〜99歳で23.0〜26.9ポイントであり、"
               "数値は Supplementary Table S8 に示した。")

    ed.replace(para_by(d, "(B) 主要3成分内の成分別構成比（%）"),
               "(B) 主要3成分内の成分別構成比（%）。全国総計は公表値であるためセル秘匿による"
               "区間幅は0.01ポイント未満であり、エラーバーは描いていない。",
               "(B) 主要3成分内の成分別構成比（%）。点は公表値どうしの比、エラーバーは識別区間の"
               "下限・上限（上下両側）を示す。区間幅は2022年度以降は0.01ポイント未満であるが、"
               "品目非掲載の影響を受ける2021年度以前は広い（2014年度で最大37.7ポイント、"
               "2021年度で最大19.1ポイント）。数値は Supplementary Table S8 に示した。")

    # --- 3. Supplementary Table S8 の legend を追加 ------------------------------
    ed.add_paragraph_after(
        para_by(d, "Supplementary Table S7. 2024年度の都道府県別"),
        "Supplementary Table S8. 構成比（シェア）の識別区間。分子と分母が同じ秘匿セルを"
        "共有するため、シェアにも理論上の最小・最大が存在する。各系列について"
        + SHARP + "として算出した（本研究のシェアはすべてこの式による）。"
        "シートは Table1_9agents（2024年度の9成分構成比。Table 1 の該当列。先発／後発シェアの"
        "区間を併記）、Fig4B_top3（2014〜2024年度の主要3成分内シェア。Figure 4B の数値版）、"
        "Fig1C_age（2024年度の年齢階級別の主要3成分構成比。Figure 1C および本文の数値版）である。")

    d.save(SRC)
    print("-> %s（変更履歴 %d 件）" % (os.path.basename(SRC), ed.n))


if __name__ == "__main__":
    main()
