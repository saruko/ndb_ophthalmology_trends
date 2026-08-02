"""NDBオープンデータ「注射」から薬効分類131（眼科用剤）を抽出してCSV化する。

出力（いずれもロング形式）:
  ophthalmic_injection_prefecture.csv : 都道府県別
  ophthalmic_injection_agesex.csv     : 年齢階級・性別

--nokouhi を付けると2024年度のみ「公費レセプトを含まない」版を読み、
「公費含まない」サブフォルダに出力する。第1回〜第10回（2014〜2023年度）は
もともと公費レセプトを含まない集計のみが公表されているため、差し替えは2024年度だけでよい。
"""

import argparse
import csv
import re
from pathlib import Path

import openpyxl

BASE = Path(__file__).resolve().parents[1]
PREF_DIR = BASE / "data" / "raw"
AGESEX_DIR = BASE / "data" / "raw" / "ndb_age_sex"
NOKOUHI_DIR = BASE / "data" / "raw" / "ndb_2024_nokouhi"
OUT_DIR = BASE / "抗VEFG薬"

YEARS = range(2014, 2025)
TARGET_CLASS = "131"  # 眼科用剤

# メタ列の見出し → 出力列名（2014・2015年度は「単位」列が存在しない等、年度で構成が異なる）
META_MAP = {
    "薬効分類": "薬効分類",
    "薬効分類名称": "薬効分類名称",
    "医薬品コード": "医薬品コード",
    "医薬品名": "医薬品名",
    "単位": "単位",
    "薬価基準収載医薬品コード": "薬価基準収載医薬品コード",
    "薬価": "薬価",
    "後発品区分": "後発品区分",
}


def clean(v):
    if v is None:
        return ""
    return re.sub(r"\s+", "", str(v))


def num(v):
    """秘匿('-')は空文字、それ以外は数値文字列を返す。"""
    s = clean(v)
    if s in ("", "-", "－"):
        return ""
    return s


def load_sheet(path, sheet_name):
    """ヘッダ2段を読み、131のデータ行を (meta, values) で返す。"""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet_name]
    rows = ws.iter_rows(values_only=True)

    # 見出し行（1列目に「薬効分類」）を探す。次の行が下段見出し。
    for row in rows:
        if clean(row[0]).startswith("薬効分類"):
            top = [clean(c) for c in row]
            break
    else:
        raise ValueError(f"header not found: {path.name} / {sheet_name}")
    sub = [clean(c) for c in next(rows)]

    meta_idx = {META_MAP[c]: i for i, c in enumerate(top) if c in META_MAP}
    total_idx = next(i for i, c in enumerate(top) if c.startswith("総計"))
    first_val = total_idx + 1

    # 上段（男/女。都道府県版では空）を前方補完
    filled_top, last = [], ""
    for c in top:
        if c:
            last = c
        filled_top.append(last)
    header = list(zip(filled_top, sub))

    out = []
    cls = cls_name = ""
    for row in rows:
        if clean(row[meta_idx["薬効分類"]]):
            cls = clean(row[meta_idx["薬効分類"]])
            cls_name = clean(row[meta_idx["薬効分類名称"]])
        if cls != TARGET_CLASS:
            continue
        meta = {out_col: clean(row[i]) for out_col, i in meta_idx.items()}
        meta["薬効分類"], meta["薬効分類名称"] = cls, cls_name
        meta.setdefault("単位", "")
        meta["総計_処方数量"] = num(row[total_idx])
        vals = [(header[i], row[i]) for i in range(first_val, len(row))]
        out.append((meta, vals))
    wb.close()
    return out


def extract(src_dir, prefix, kind, nokouhi=False):
    records = []
    for year in YEARS:
        if nokouhi and year == 2024:
            # 2024年度のみ「公費レセプトを含まない」版に差し替える
            path = NOKOUHI_DIR / f"{prefix}_{year}_nokouhi.xlsx"
        else:
            path = src_dir / f"{prefix}_{year}.xlsx"
        if not path.exists():
            print(f"  skip (not found): {path.name}")
            continue
        wb = openpyxl.load_workbook(path, read_only=True)
        sheets = wb.sheetnames
        wb.close()
        for sheet in sheets:
            # 例: 「注射薬 外来 (院内)」→ 区分部分だけ残す
            kubun = clean(sheet).replace("注射薬", "")
            for meta, vals in load_sheet(path, sheet):
                for (top, sub), v in vals:
                    raw = clean(v)
                    if raw == "":  # 表に存在しない列（余白）はスキップ
                        continue
                    rec = {"年度": year, "区分": kubun, **meta}
                    if kind == "pref":
                        rec["都道府県"] = sub
                    else:
                        # 2014年度は「男性/女性」、2015年度以降は「男/女」表記
                        rec["性別"] = top.replace("性", "")
                        rec["年齢階級"] = sub
                    # 「-」は10未満のため非公開（秘匿）。行は残し、値を空欄＋フラグで示す。
                    hidden = raw in ("-", "－")
                    rec["処方数量"] = "" if hidden else raw
                    rec["秘匿フラグ"] = 1 if hidden else 0
                    rec["原表記"] = raw
                    records.append(rec)
        print(f"  {path.name}: cumulative {len(records)} rows")
    return records


def write_csv(records, path, cols):
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        w.writerows(records)
    print(f"wrote {path} ({len(records)} rows)")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--nokouhi", action="store_true",
                    help="2024年度を公費レセプトを含まない版にして「公費含まない」へ出力")
    args = ap.parse_args()
    out_dir = OUT_DIR / "公費含まない" if args.nokouhi else OUT_DIR
    out_dir.mkdir(parents=True, exist_ok=True)
    meta_cols = [
        "年度", "区分", "薬効分類", "薬効分類名称", "医薬品コード", "医薬品名",
        "単位", "薬価基準収載医薬品コード", "薬価", "後発品区分", "総計_処方数量",
    ]

    print("[都道府県版]")
    pref = extract(PREF_DIR, "ndb_chusha", "pref", args.nokouhi)
    write_csv(pref, out_dir / "ophthalmic_injection_prefecture.csv",
              meta_cols + ["都道府県", "処方数量", "秘匿フラグ", "原表記"])

    print("[年齢性別版]")
    agesex = extract(AGESEX_DIR, "ndb_chusha_agesex", "agesex", args.nokouhi)
    write_csv(agesex, out_dir / "ophthalmic_injection_agesex.csv",
              meta_cols + ["性別", "年齢階級", "処方数量", "秘匿フラグ", "原表記"])


if __name__ == "__main__":
    main()
