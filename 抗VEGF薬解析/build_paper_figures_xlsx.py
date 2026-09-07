"""
build_paper_figures_xlsx.py
抗VEGF薬論文用のFigureデータシート＋Excelチャートを生成する。

    python 抗VEGF薬解析/build_paper_figures_xlsx.py --nokouhi   # 公費含まない版（論文用）
    python 抗VEGF薬解析/build_paper_figures_xlsx.py             # 親フォルダのデータ

出力先: <データフォルダ>/04_図表/ に**用途別の5ファイル**を書き出す。
1ファイルにまとめると開くのが重くなるため、章立てに沿って分割している。

    antivegf_fig_main.xlsx     Fig 1A〜Suppl Fig（論文本体のFigure）
    antivegf_fig_plots.xlsx    P01〜P10（plots/*.png のExcel再現）
    antivegf_fig_sexdiff.xlsx  S1〜S5（年齢階級別の男女差と秘匿状況）
    antivegf_fig_cost.xlsx     C01〜C06（医療費の要因分解・反実仮想・薬価改定）
    antivegf_fig_methods.xlsx  M01〜M07（変化点・収束・感度分析・突合の検証）

入力CSVが揃っていないグループはスキップする（親フォルダで analyze_advanced.py を
未実行なら methods だけが飛ぶ、といった挙動になる）。

依存: openpyxl, pandas
"""

import argparse
import os

import pandas as pd
from pathlib import Path
from openpyxl import Workbook
from openpyxl.chart import (
    BarChart, LineChart, AreaChart, Reference
)
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import Marker
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import (CharacterProperties, Paragraph,
                                   ParagraphProperties)
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.chart.layout import Layout, ManualLayout

BASE = Path(__file__).parent

# main() で解析対象フォルダに応じて差し替える（--nokouhi の有無で切り替わる）
ROOT = BASE / "公費含まない"
FIG_DIR = ROOT / "04_図表"

OKU = 1e8  # 円 → 億円

# ---------------------------------------------------------------------------
# グラフの共通書式
#
# antivegf_fig_sexdiff.xlsx の S1 シートでExcel上で調整された書式を採寸し、
# 全ワークブックの全グラフに同じ設定を適用する。値の根拠（S1のchart XML）:
#   タイトル sz=1400 / 軸ラベル sz=1000 / 目盛・凡例 sz=900
#   系列線 w=28575EMU(2.25pt) cap=rnd、マーカー=circle、smooth=0
#   目盛線なし、凡例は下(b)、軸は表示(delete=0)
# フォント書体は指定せず、テーマの本文フォント（+mn-lt/+mn-ea）を継承する。
# ---------------------------------------------------------------------------
FONT_TITLE = 1400        # グラフタイトル 14pt
FONT_AXIS_TITLE = 1000   # 軸ラベル 10pt
FONT_TICK = 900          # 目盛 9pt
FONT_LEGEND = 900        # 凡例 9pt
SERIES_LINE_WIDTH = 28575  # 系列線 2.25pt（EMU）
MARKER_SYMBOL = "circle"
MARKER_SIZE = 5

# 既定のグラフサイズ（cm）。S1 と同じ。
CHART_W, CHART_H = 28, 14
# 47都道府県の横棒だけは縦長を保つ（14cmに潰すと県名が読めない）。
# builder が TALL_THRESHOLD 以上の高さを指定したものを「縦長」とみなす。
TALL_W, TALL_H = 20, 34
TALL_THRESHOLD = 25


def _rich(sz, bold=False):
    """指定サイズの既定文字書式を持つ RichText を作る。"""
    cp = CharacterProperties(sz=sz, b=bold)
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def _set_title_font(title, sz):
    """Title オブジェクト内の各段落・各ランに文字サイズを流し込む。"""
    if title is None or title.tx is None or title.tx.rich is None:
        return
    cp = CharacterProperties(sz=sz)
    for p in title.tx.rich.p:
        p.pPr = ParagraphProperties(defRPr=cp)
        for r in (p.r or []):
            r.rPr = cp


def _style_axis(ax):
    ax.delete = False           # 軸を必ず表示する
    ax.majorGridlines = None    # 目盛線は出さない（S1に合わせる）
    ax.txPr = _rich(FONT_TICK)
    _set_title_font(ax.title, FONT_AXIS_TITLE)


def _style_chart(chart):
    """S1 と同じ見た目になるようグラフ全体の書式を揃える。

    複合グラフ（chart += line）は chart._charts に各要素が入るため、
    系列の書式はそちらを回して設定する。
    """
    chart.roundedCorners = False
    _set_title_font(chart.title, FONT_TITLE)

    if chart.legend is not None:
        chart.legend.position = "b"
        chart.legend.overlay = False
        chart.legend.txPr = _rich(FONT_LEGEND)

    for sub in getattr(chart, "_charts", [chart]) or [chart]:
        for ax in (getattr(sub, "x_axis", None), getattr(sub, "y_axis", None)):
            if ax is not None:
                _style_axis(ax)
        # 線の太さ・マーカー・スムージングは折れ線にのみ適用する。
        # 棒・面グラフに太い線を入れると輪郭が潰れて見えるため。
        if not isinstance(sub, LineChart):
            continue
        for s in sub.series:
            s.smooth = False
            s.marker = Marker(symbol=MARKER_SYMBOL, size=MARKER_SIZE)
            gp = s.graphicalProperties
            if gp is None:
                s.graphicalProperties = gp = GraphicalProperties()
            gp.line.width = SERIES_LINE_WIDTH
            gp.line.cap = "rnd"


# 走査から除外するフォルダ。src/paths.py の find() と同じ方針。
# 「公費含まない」は同名・別内容のCSVを持つ独立したデータセットなので、
# 親フォルダからの走査では必ず除外する（混ざると図表が静かに壊れる）。
EXCLUDE_DIRS = {"公費含まない", "__pycache__", ".git"}


class _CsvLocator:
    """フォルダ再編後もファイル名だけでCSVを引けるようにする。"""

    def __init__(self, root):
        self.root = Path(root)
        self._index = {}
        for dirpath, dirnames, filenames in os.walk(self.root):
            dirnames[:] = [d for d in dirnames if d not in EXCLUDE_DIRS]
            for f in filenames:
                if f.endswith(".csv"):
                    self._index.setdefault(f, Path(dirpath) / f)

    def __truediv__(self, name):
        try:
            return self._index[name]
        except KeyError:
            raise FileNotFoundError(f"{name} が {self.root} 配下に見つかりません")


# main() が対象フォルダを確定してから構築する。ここで既定値を入れてしまうと、
# --nokouhi の指定に関わらず「公費含まない」のCSVを読みかねない（静かに壊れる）
PROC = None

# 年齢群の正しいソート順
AGE_ORDER = [
    "0～4歳", "5～9歳", "10～14歳", "15～19歳", "20～24歳",
    "25～29歳", "30～34歳", "35～39歳", "40～44歳", "45～49歳",
    "50～54歳", "55～59歳", "60～64歳", "65～69歳", "70～74歳",
    "75～79歳", "80～84歳", "85～89歳", "90歳以上",
]

# 年齢群（公開粒度21区分用）
AGE_ORDER_21 = [
    "0～4歳", "5～9歳", "10～14歳", "15～19歳", "20～24歳",
    "25～29歳", "30～34歳", "35～39歳", "40～44歳", "45～49歳",
    "50～54歳", "55～59歳", "60～64歳", "65～69歳", "70～74歳",
    "75～79歳", "80～84歳", "85～89歳", "90～94歳", "95～99歳", "100歳以上",
]


def write_df(ws, df, start_row=1):
    """DataFrameをワークシートに書き込む"""
    for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), start_row):
        for c_idx, val in enumerate(row, 1):
            ws.cell(row=r_idx, column=c_idx, value=val)
    return start_row + len(df)  # 次の空行


def build_fig1a(wb):
    """Fig 1A: 年齢群別 処方数量（2024年度、男女合算）"""
    ws = wb.create_sheet("Fig1A_data")
    df = pd.read_csv(PROC / "agesex_distribution_published_antivegf.csv")
    df2024 = df[df["year"] == 2024].copy()

    # 男女合算
    both = df2024.groupby("age_group_detail")["quantity"].sum().reset_index()
    both.rename(columns={"age_group_detail": "age_group"}, inplace=True)

    # age_orderに基づいてソート（21区分）
    order = [a for a in AGE_ORDER_21 if a in both["age_group"].values]
    both["age_group"] = pd.Categorical(both["age_group"], categories=order, ordered=True)
    both = both.sort_values("age_group").reset_index(drop=True)

    write_df(ws, both)

    # チャート: 棒グラフ
    chart = BarChart()
    chart.type = "col"
    chart.title = "Fig 1A: 年齢群別 処方数量（2024年度）"
    chart.y_axis.title = "処方数量（本）"
    chart.x_axis.title = "年齢群"
    chart.style = 10
    chart.width = 28
    chart.height = 15

    data = Reference(ws, min_col=2, min_row=1, max_row=len(both) + 1)
    cats = Reference(ws, min_col=1, min_row=2, max_row=len(both) + 1)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.shape = 4
    chart.legend = None

    ws.add_chart(chart, "D2")


def build_fig1b(wb):
    """Fig 1B: 性別×年齢群（2024年度、男女2系列）"""
    ws = wb.create_sheet("Fig1B_data")
    df = pd.read_csv(PROC / "agesex_distribution_published_antivegf.csv")
    df2024 = df[df["year"] == 2024].copy()

    # ピボット
    pivot = df2024.pivot_table(index="age_group_detail", columns="sex", values="quantity", aggfunc="sum").reset_index()
    pivot.rename(columns={"age_group_detail": "age_group"}, inplace=True)

    order = [a for a in AGE_ORDER_21 if a in pivot["age_group"].values]
    pivot["age_group"] = pd.Categorical(pivot["age_group"], categories=order, ordered=True)
    pivot = pivot.sort_values("age_group").reset_index(drop=True)

    # 列名整理
    cols = ["age_group"]
    if "男" in pivot.columns:
        cols.append("男")
    if "女" in pivot.columns:
        cols.append("女")
    pivot = pivot[cols]
    pivot["age_group"] = pivot["age_group"].astype(str)
    pivot = pivot.fillna(0)

    write_df(ws, pivot)

    # チャート: 折れ線
    chart = LineChart()
    chart.title = "Fig 1B: 性別×年齢群 処方数量（2024年度）"
    chart.y_axis.title = "処方数量（本）"
    chart.x_axis.title = "年齢群"
    chart.style = 10
    chart.width = 28
    chart.height = 15

    for i, col_name in enumerate(cols[1:], 2):
        data = Reference(ws, min_col=i, min_row=1, max_row=len(pivot) + 1)
        chart.add_data(data, titles_from_data=True)

    cats = Reference(ws, min_col=1, min_row=2, max_row=len(pivot) + 1)
    chart.set_categories(cats)

    ws.add_chart(chart, "D2")


def build_fig1c(wb):
    """Fig 1C: 年齢分布の経年シフト（2014 vs 2024）"""
    ws = wb.create_sheet("Fig1C_data")
    df = pd.read_csv(PROC / "agesex_distribution_comparable_antivegf.csv")

    # 男女合算
    both = df.groupby(["year", "age_group"])["quantity"].sum().reset_index()
    # 各年度の合計に対するshare
    totals = both.groupby("year")["quantity"].transform("sum")
    both["share_pct"] = both["quantity"] / totals * 100

    # 2014と2024のみ
    pivot = both[both["year"].isin([2014, 2024])].pivot_table(
        index="age_group", columns="year", values="share_pct"
    ).reset_index()

    order = [a for a in AGE_ORDER if a in pivot["age_group"].values]
    pivot["age_group"] = pd.Categorical(pivot["age_group"], categories=order, ordered=True)
    pivot = pivot.sort_values("age_group").reset_index(drop=True)
    pivot.columns = ["age_group", "2014年度_share_pct", "2024年度_share_pct"]

    write_df(ws, pivot)

    # チャート: 棒グラフ（2系列並列）
    chart = BarChart()
    chart.type = "col"
    chart.title = "Fig 1C: 年齢分布の経年シフト（2014→2024）"
    chart.y_axis.title = "構成比（%）"
    chart.style = 10
    chart.width = 28
    chart.height = 15

    data = Reference(ws, min_col=2, max_col=3, min_row=1, max_row=len(pivot) + 1)
    cats = Reference(ws, min_col=1, min_row=2, max_row=len(pivot) + 1)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)

    ws.add_chart(chart, "E2")


def build_fig2a(wb):
    """Fig 2A: 近似平均年齢・75歳以上比率・男性比率の推移"""
    ws = wb.create_sheet("Fig2A_data")
    df = pd.read_csv(PROC / "agesex_summary_antivegf.csv")
    df = df.rename(columns={
        "mean_age_approx": "近似平均年齢",
        "share_75plus_pct": "75歳以上比率(%)",
        "male_share_pct": "男性比率(%)"
    })
    df["year"] = df["year"].astype(int)

    write_df(ws, df)

    # チャート: 2軸折れ線（左軸=年齢、右軸=%）
    chart = LineChart()
    chart.title = "Fig 2A: 近似平均年齢・75歳以上比率・男性比率の推移"
    chart.style = 10
    chart.width = 28
    chart.height = 15
    chart.y_axis.title = "年齢（歳）"

    n = len(df) + 1

    # 近似平均年齢（左軸）
    data1 = Reference(ws, min_col=2, min_row=1, max_row=n)
    chart.add_data(data1, titles_from_data=True)

    # 75歳以上比率（右軸）
    data2 = Reference(ws, min_col=3, min_row=1, max_row=n)
    chart.add_data(data2, titles_from_data=True)

    # 男性比率（右軸）
    data3 = Reference(ws, min_col=4, min_row=1, max_row=n)
    chart.add_data(data3, titles_from_data=True)

    cats = Reference(ws, min_col=1, min_row=2, max_row=n)
    chart.set_categories(cats)

    # 2軸設定
    from openpyxl.chart import LineChart as LC2
    chart2 = LC2()
    chart2.y_axis.title = "割合（%）"
    chart2.y_axis.axId = 200

    # 75歳以上比率と男性比率を右軸に
    s1 = chart.series[1]
    s2 = chart.series[2]
    chart.series.pop(2)
    chart.series.pop(1)
    chart2.series.append(s1)
    chart2.series.append(s2)

    chart.y_axis.crosses = "min"
    chart2.y_axis.crosses = "max"
    chart += chart2

    ws.add_chart(chart, "F2")


def build_fig2b(wb):
    """Fig 2B: 外来/入院比率の推移"""
    ws = wb.create_sheet("Fig2B_data")
    df = pd.read_csv(PROC / "product_by_setting_antivegf.csv")

    # 年度×setting で合算
    setting_total = df.groupby(["year", "setting"])["quantity"].sum().reset_index()
    pivot = setting_total.pivot_table(index="year", columns="setting", values="quantity").reset_index().fillna(0)

    # 入院比率を計算
    if "入院" in pivot.columns and "外来(院内)" in pivot.columns:
        pivot["合計"] = pivot["入院"] + pivot["外来(院内)"]
        pivot["入院比率(%)"] = pivot["入院"] / pivot["合計"] * 100
    else:
        pivot["合計"] = 0
        pivot["入院比率(%)"] = 0

    out = pivot[["year", "外来(院内)", "入院", "合計", "入院比率(%)"]].copy()
    out.columns = ["年度", "外来（院内）", "入院", "合計", "入院比率(%)"]
    out["年度"] = out["年度"].astype(int)

    write_df(ws, out)

    # チャート: 2軸（左=数量棒、右=入院比率%折れ線）
    chart = BarChart()
    chart.type = "col"
    chart.title = "Fig 2B: 外来/入院の推移"
    chart.y_axis.title = "処方数量（本）"
    chart.style = 10
    chart.width = 28
    chart.height = 15

    n = len(out) + 1

    # 外来
    d1 = Reference(ws, min_col=2, min_row=1, max_row=n)
    chart.add_data(d1, titles_from_data=True)
    # 入院
    d2 = Reference(ws, min_col=3, min_row=1, max_row=n)
    chart.add_data(d2, titles_from_data=True)

    cats = Reference(ws, min_col=1, min_row=2, max_row=n)
    chart.set_categories(cats)

    # 入院比率を右軸折れ線
    line = LineChart()
    line.y_axis.title = "入院比率（%）"
    line.y_axis.axId = 200

    d3 = Reference(ws, min_col=5, min_row=1, max_row=n)
    line.add_data(d3, titles_from_data=True)
    line.set_categories(cats)

    chart.y_axis.crosses = "min"
    line.y_axis.crosses = "max"
    chart += line

    ws.add_chart(chart, "G2")


def build_fig3a(wb):
    """Fig 3A: 都道府県別ランキング（2024年度、横棒）"""
    ws = wb.create_sheet("Fig3A_data")
    df = pd.read_csv(PROC / "prefecture_per_capita_ranking_antivegf.csv")

    # 2024年度のデータ
    out = df[["prefecture", "2024年度_人口10万対", "2024年度_順位"]].copy()
    out = out.sort_values("2024年度_順位", ascending=True).reset_index(drop=True)
    out.columns = ["都道府県", "人口10万対", "順位"]

    write_df(ws, out)

    # 横棒グラフ
    chart = BarChart()
    chart.type = "bar"
    chart.title = "Fig 3A: 都道府県別 人口10万対処方数量（2024年度）"
    chart.x_axis.title = "人口10万対（本）"
    chart.style = 10
    chart.width = 20
    chart.height = 35  # 47県分なので縦長

    n = len(out) + 1
    data = Reference(ws, min_col=2, min_row=1, max_row=n)
    cats = Reference(ws, min_col=1, min_row=2, max_row=n)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.legend = None

    # 順位の低い県が上に来るように反転
    chart.x_axis.scaling.orientation = "maxMin"

    ws.add_chart(chart, "E2")


def build_fig3b(wb):
    """Fig 3B: Gini・CVの経年推移"""
    ws = wb.create_sheet("Fig3B_data")
    df = pd.read_csv(PROC / "geographic_disparity_antivegf.csv")
    total = df[df["code"] == "ANTI_VEGF_TOTAL"][["year", "gini", "cv", "max_to_min_ratio", "p90_to_p10_ratio"]].copy()
    total = total.sort_values("year").reset_index(drop=True)
    total.columns = ["年度", "Gini", "CV", "最大/最小比", "P90/P10比"]
    total["年度"] = total["年度"].astype(int)

    write_df(ws, total)

    # 折れ線
    chart = LineChart()
    chart.title = "Fig 3B: 地域格差指標の推移（抗VEGF薬合計）"
    chart.y_axis.title = "指標値"
    chart.style = 10
    chart.width = 28
    chart.height = 15

    n = len(total) + 1
    for col_idx in range(2, 6):
        data = Reference(ws, min_col=col_idx, min_row=1, max_row=n)
        chart.add_data(data, titles_from_data=True)

    cats = Reference(ws, min_col=1, min_row=2, max_row=n)
    chart.set_categories(cats)

    ws.add_chart(chart, "G2")


def build_fig4a(wb):
    """Fig 4A: 成分別 処方数量推移（人口10万対）"""
    ws = wb.create_sheet("Fig4A_data")
    df = pd.read_csv(PROC / "national_trends_antivegf.csv")

    # 成分コードをグループ化（AFLIBERCEPT_ALL相当にまとめる）
    # year × code の count_per_100k をピボット
    # 主要4成分+合計
    codes_of_interest = {
        "AFLIBERCEPT": "アフリベルセプト2mg",
        "AFLIBERCEPT_8MG": "アフリベルセプト8mg",
        "BROLUCIZUMAB": "ブロルシズマブ",
        "FARICIMAB": "ファリシマブ",
        "RANIBIZUMAB_ORIG": "ラニビズマブ先発",
        "RANIBIZUMAB_BS": "ラニビズマブBS",
        "PEGAPTANIB": "ペガプタニブ",
    }

    df_filtered = df[df["code"].isin(codes_of_interest.keys())].copy()
    df_filtered["label"] = df_filtered["code"].map(codes_of_interest)

    pivot = df_filtered.pivot_table(
        index="year", columns="label", values="count_per_100k", aggfunc="sum"
    ).reset_index().fillna("")

    # 合計行を追加
    total_by_year = df_filtered.groupby("year")["count_per_100k"].sum().reset_index()
    total_by_year.columns = ["year", "抗VEGF薬合計"]
    pivot = pivot.merge(total_by_year, on="year", how="left")

    # 列順を整理
    col_order = ["year", "抗VEGF薬合計", "アフリベルセプト2mg", "アフリベルセプト8mg",
                 "ファリシマブ", "ラニビズマブ先発", "ラニビズマブBS", "ブロルシズマブ", "ペガプタニブ"]
    col_order = [c for c in col_order if c in pivot.columns]
    pivot = pivot[col_order]
    pivot.rename(columns={"year": "年度"}, inplace=True)
    pivot["年度"] = pivot["年度"].astype(int)

    write_df(ws, pivot)

    # 折れ線チャート
    chart = LineChart()
    chart.title = "Fig 4A: 成分別 処方数量推移（人口10万対）"
    chart.y_axis.title = "人口10万対（本）"
    chart.style = 10
    chart.width = 28
    chart.height = 15

    n = len(pivot) + 1
    for col_idx in range(2, len(col_order) + 1):
        data = Reference(ws, min_col=col_idx, min_row=1, max_row=n)
        chart.add_data(data, titles_from_data=True)

    cats = Reference(ws, min_col=1, min_row=2, max_row=n)
    chart.set_categories(cats)

    ws.add_chart(chart, "K2")


def build_fig4b(wb):
    """Fig 4B: 成分別シェア推移（積み上げ面）"""
    ws = wb.create_sheet("Fig4B_data")
    df = pd.read_csv(PROC / "national_trends_antivegf.csv")

    codes_of_interest = {
        "AFLIBERCEPT": "アフリベルセプト2mg",
        "AFLIBERCEPT_8MG": "アフリベルセプト8mg",
        "BROLUCIZUMAB": "ブロルシズマブ",
        "FARICIMAB": "ファリシマブ",
        "RANIBIZUMAB_ORIG": "ラニビズマブ先発",
        "RANIBIZUMAB_BS": "ラニビズマブBS",
        "PEGAPTANIB": "ペガプタニブ",
    }

    df_filtered = df[df["code"].isin(codes_of_interest.keys())].copy()
    df_filtered["label"] = df_filtered["code"].map(codes_of_interest)

    pivot = df_filtered.pivot_table(
        index="year", columns="label", values="share_of_antivegf_pct", aggfunc="sum"
    ).reset_index().fillna(0)

    col_order = ["year", "アフリベルセプト2mg", "アフリベルセプト8mg",
                 "ファリシマブ", "ラニビズマブ先発", "ラニビズマブBS", "ブロルシズマブ", "ペガプタニブ"]
    col_order = [c for c in col_order if c in pivot.columns]
    pivot = pivot[col_order]
    pivot.rename(columns={"year": "年度"}, inplace=True)
    pivot["年度"] = pivot["年度"].astype(int)

    write_df(ws, pivot)

    # 積み上げ面グラフ
    chart = AreaChart()
    chart.grouping = "stacked"
    chart.title = "Fig 4B: 成分別シェア推移"
    chart.y_axis.title = "シェア（%）"
    chart.style = 10
    chart.width = 28
    chart.height = 15

    n = len(pivot) + 1
    for col_idx in range(2, len(col_order) + 1):
        data = Reference(ws, min_col=col_idx, min_row=1, max_row=n)
        chart.add_data(data, titles_from_data=True)

    cats = Reference(ws, min_col=1, min_row=2, max_row=n)
    chart.set_categories(cats)

    ws.add_chart(chart, "K2")


def build_fig5a(wb):
    """Fig 5A: 剤形比率の推移（注射液 vs キット）"""
    ws = wb.create_sheet("Fig5A_data")
    df = pd.read_csv(PROC / "formulation_total_antivegf.csv")
    pivot = df.pivot_table(
        index="year", columns="formulation", values="share_pct"
    ).reset_index().fillna(0)
    pivot.rename(columns={"year": "年度"}, inplace=True)
    pivot["年度"] = pivot["年度"].astype(int)

    # 列順: 注射液, キット
    cols = ["年度"]
    if "注射液" in pivot.columns:
        cols.append("注射液")
    if "キット" in pivot.columns:
        cols.append("キット")
    pivot = pivot[cols]

    write_df(ws, pivot)

    # 積み上げ棒
    chart = BarChart()
    chart.type = "col"
    chart.grouping = "stacked"
    chart.title = "Fig 5A: 剤形比率の推移（注射液 vs キット）"
    chart.y_axis.title = "シェア（%）"
    chart.style = 10
    chart.width = 28
    chart.height = 15

    n = len(pivot) + 1
    for col_idx in range(2, len(cols) + 1):
        data = Reference(ws, min_col=col_idx, min_row=1, max_row=n)
        chart.add_data(data, titles_from_data=True)

    cats = Reference(ws, min_col=1, min_row=2, max_row=n)
    chart.set_categories(cats)

    ws.add_chart(chart, "E2")


def build_fig5b(wb):
    """Fig 5B: バイオシミラーシェア推移"""
    ws = wb.create_sheet("Fig5B_data")
    df = pd.read_csv(PROC / "biosimilar_national_share.csv")
    pivot = df.pivot_table(
        index="year", columns="brand_type", values="share_pct"
    ).reset_index().fillna(0)
    pivot.rename(columns={"year": "年度"}, inplace=True)
    pivot["年度"] = pivot["年度"].astype(int)

    cols = ["年度"]
    if "先発" in pivot.columns:
        cols.append("先発")
    if "バイオシミラー" in pivot.columns:
        cols.append("バイオシミラー")
    pivot = pivot[cols]

    write_df(ws, pivot)

    # 積み上げ面グラフ
    chart = AreaChart()
    chart.grouping = "stacked"
    chart.title = "Fig 5B: ラニビズマブ 先発 vs バイオシミラー シェア推移"
    chart.y_axis.title = "シェア（%）"
    chart.style = 10
    chart.width = 28
    chart.height = 15

    n = len(pivot) + 1
    for col_idx in range(2, len(cols) + 1):
        data = Reference(ws, min_col=col_idx, min_row=1, max_row=n)
        chart.add_data(data, titles_from_data=True)

    cats = Reference(ws, min_col=1, min_row=2, max_row=n)
    chart.set_categories(cats)

    ws.add_chart(chart, "E2")


def build_fig5c(wb):
    """Fig 5C: 薬剤費推移と構成（製品別積み上げ棒）"""
    ws = wb.create_sheet("Fig5C_data")
    df = pd.read_csv(PROC / "product_trends_antivegf.csv")

    # 製品名×年度で薬剤費を億円に
    df["cost_oku"] = df["cost"] / 1e8

    pivot = df.pivot_table(
        index="year", columns="product_name", values="cost_oku", aggfunc="sum"
    ).reset_index().fillna(0)
    pivot.rename(columns={"year": "年度"}, inplace=True)
    pivot["年度"] = pivot["年度"].astype(int)

    write_df(ws, pivot)

    # 積み上げ棒グラフ
    chart = BarChart()
    chart.type = "col"
    chart.grouping = "stacked"
    chart.title = "Fig 5C: 薬剤費推移と構成（億円）"
    chart.y_axis.title = "薬剤費（億円）"
    chart.style = 10
    chart.width = 28
    chart.height = 15

    n = len(pivot) + 1
    ncols = len(pivot.columns)
    for col_idx in range(2, ncols + 1):
        data = Reference(ws, min_col=col_idx, min_row=1, max_row=n)
        chart.add_data(data, titles_from_data=True)

    cats = Reference(ws, min_col=1, min_row=2, max_row=n)
    chart.set_categories(cats)

    ws.add_chart(chart, f"{chr(64 + ncols + 2)}2")


def build_suppl(wb):
    """Suppl Fig: G016算定回数 vs 抗VEGF薬数量"""
    ws = wb.create_sheet("SuppFig_data")
    df = pd.read_csv(PROC / "g016_vs_drug_validation.csv")
    out = df[["year", "g016_procedures", "antivegf_vials", "antivegf_per_g016_pct"]].copy()
    out.columns = ["年度", "G016算定回数", "抗VEGF薬数量", "抗VEGF/G016(%)"]
    out["年度"] = out["年度"].astype(int)

    write_df(ws, out)

    # 2軸：左=数量（棒2本）、右=比率%（折れ線）
    chart = BarChart()
    chart.type = "col"
    chart.title = "Suppl Fig: G016算定回数 vs 抗VEGF薬数量"
    chart.y_axis.title = "件数/本数"
    chart.style = 10
    chart.width = 28
    chart.height = 15

    n = len(out) + 1
    d1 = Reference(ws, min_col=2, min_row=1, max_row=n)
    d2 = Reference(ws, min_col=3, min_row=1, max_row=n)
    chart.add_data(d1, titles_from_data=True)
    chart.add_data(d2, titles_from_data=True)

    cats = Reference(ws, min_col=1, min_row=2, max_row=n)
    chart.set_categories(cats)

    # 右軸に比率
    line = LineChart()
    line.y_axis.title = "抗VEGF/G016（%）"
    line.y_axis.axId = 200
    d3 = Reference(ws, min_col=4, min_row=1, max_row=n)
    line.add_data(d3, titles_from_data=True)
    line.set_categories(cats)

    chart.y_axis.crosses = "min"
    line.y_axis.crosses = "max"
    chart += line

    ws.add_chart(chart, "F2")


# ---------------------------------------------------------------------------
# 04_図表/plots/*.png の再現（P01〜P10）
# ---------------------------------------------------------------------------

# 公開粒度（2015年度以降は90歳以上が3区分に細分化）— visualization_antivegf.py と同一
AGE_ORDER_PUBLISHED = AGE_ORDER[:-1] + ["90～94歳", "95～99歳", "100歳以上", "90歳以上"]


def _note(ws, text, row=1):
    """シート先頭に元PNGの対応を書く。データは row+2 行目から。"""
    c = ws.cell(row=row, column=1, value=text)
    c.font = Font(bold=True)


def _year_pivot(df, columns, values, aggfunc="sum"):
    """year × columns のピボットを「年度」始まりの表に整える。"""
    piv = df.pivot_table(index="year", columns=columns, values=values,
                         aggfunc=aggfunc).fillna(0).reset_index()
    piv.rename(columns={"year": "年度"}, inplace=True)
    piv["年度"] = piv["年度"].astype(int)
    piv.columns = [str(c) for c in piv.columns]
    return piv


def _series_chart(ws, chart, df, first_row, anchor):
    """1列目=カテゴリ、2列目以降=系列としてチャートに流し込む。"""
    n_last = first_row + len(df)
    for col_idx in range(2, len(df.columns) + 1):
        chart.add_data(Reference(ws, min_col=col_idx, min_row=first_row, max_row=n_last),
                       titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=first_row + 1, max_row=n_last))
    chart.style = 10
    ws.add_chart(chart, anchor)


def build_p01(wb):
    """P01: product_trends.png — 製品別 処方数量の推移（折れ線）"""
    ws = wb.create_sheet("P01_product_trends")
    _note(ws, "plots/product_trends.png の再現（product_trends_antivegf.csv / year × product_name × quantity）")
    df = pd.read_csv(PROC / "product_trends_antivegf.csv")
    piv = _year_pivot(df, "product_name", "quantity")
    write_df(ws, piv, start_row=3)

    chart = LineChart()
    chart.title = "抗VEGF薬 製品別 処方数量の推移（全国・外来＋入院）"
    chart.x_axis.title = "年度"
    chart.y_axis.title = "処方数量（バイアル・シリンジ本数）"
    chart.width = 28
    chart.height = 15
    _series_chart(ws, chart, piv, 3, f"{get_column_letter(len(piv.columns) + 2)}3")


def build_p02(wb):
    """P02: product_share_stacked.png — 製品別シェアの積み上げ面"""
    ws = wb.create_sheet("P02_product_share_stacked")
    _note(ws, "plots/product_share_stacked.png の再現（product_trends_antivegf.csv / share_pct）")
    df = pd.read_csv(PROC / "product_trends_antivegf.csv")
    piv = _year_pivot(df, "product_name", "share_pct", aggfunc="mean")
    write_df(ws, piv, start_row=3)

    chart = AreaChart()
    chart.grouping = "stacked"
    chart.overlap = 100
    chart.title = "抗VEGF薬 製品別シェアの推移（%）"
    chart.x_axis.title = "年度"
    chart.y_axis.title = "シェア（%）"
    chart.y_axis.scaling.min = 0
    chart.y_axis.scaling.max = 100
    chart.width = 28
    chart.height = 15
    _series_chart(ws, chart, piv, 3, f"{get_column_letter(len(piv.columns) + 2)}3")


def build_p03(wb):
    """P03: formulation_share_total.png — 剤形比率（積み上げ棒）"""
    ws = wb.create_sheet("P03_formulation_share_total")
    _note(ws, "plots/formulation_share_total.png の再現（formulation_total_antivegf.csv / share_pct）")
    df = pd.read_csv(PROC / "formulation_total_antivegf.csv")
    piv = _year_pivot(df, "formulation", "share_pct")
    cols = ["年度"] + [c for c in ["注射液", "キット"] if c in piv.columns]
    piv = piv[cols]
    write_df(ws, piv, start_row=3)

    chart = BarChart()
    chart.type = "col"
    chart.grouping = "stacked"
    chart.overlap = 100
    chart.title = "抗VEGF薬全体の剤形比率（注射液 vs キット）"
    chart.x_axis.title = "年度"
    chart.y_axis.title = "シェア（%）"
    chart.width = 24
    chart.height = 13
    _series_chart(ws, chart, piv, 3, "F3")


def build_p04(wb):
    """P04: formulation_kit_share_by_molecule.png — 成分別キット比率（折れ線）"""
    ws = wb.create_sheet("P04_formulation_kit_share_mol")
    _note(ws, "plots/formulation_kit_share_by_molecule.png の再現"
              "（formulation_by_molecule_antivegf.csv / has_both_formulations=True かつ formulation=キット）")
    df = pd.read_csv(PROC / "formulation_by_molecule_antivegf.csv")
    df = df[df["has_both_formulations"] & (df["formulation"] == "キット")]
    piv = _year_pivot(df, "molecule_name", "share_within_molecule_pct", aggfunc="mean")
    write_df(ws, piv, start_row=3)

    chart = LineChart()
    chart.title = "成分別 キット製剤の比率"
    chart.x_axis.title = "年度"
    chart.y_axis.title = "キット比率（%）"
    chart.y_axis.scaling.min = 0
    chart.y_axis.scaling.max = 100
    chart.width = 24
    chart.height = 13
    _series_chart(ws, chart, piv, 3, f"{get_column_letter(len(piv.columns) + 2)}3")


def build_p05(wb):
    """P05: biosimilar_share.png — 左:数量積み上げ棒 / 右:シェア折れ線（2パネル）"""
    ws = wb.create_sheet("P05_biosimilar_share")
    _note(ws, "plots/biosimilar_share.png の再現（biosimilar_national_share.csv / 左=quantity, 右=share_pct）")
    df = pd.read_csv(PROC / "biosimilar_national_share.csv")

    qty = _year_pivot(df, "brand_type", "quantity")
    shr = _year_pivot(df, "brand_type", "share_pct")
    order = [c for c in ["先発", "バイオシミラー"] if c in qty.columns]
    qty = qty[["年度"] + order]
    shr = shr[["年度"] + order]

    ws.cell(row=3, column=1, value="【左パネル】処方数量").font = Font(bold=True)
    write_df(ws, qty, start_row=4)
    start2 = 4 + len(qty) + 3
    ws.cell(row=start2 - 1, column=1, value="【右パネル】シェア（%）").font = Font(bold=True)
    write_df(ws, shr, start_row=start2)

    c1 = BarChart()
    c1.type = "col"
    c1.grouping = "stacked"
    c1.overlap = 100
    c1.title = "ラニビズマブ 処方数量（先発 vs BS）"
    c1.x_axis.title = "年度"
    c1.y_axis.title = "処方数量"
    c1.width = 16
    c1.height = 11
    _series_chart(ws, c1, qty, 4, "F3")

    c2 = LineChart()
    c2.title = "ラニビズマブ シェア推移（%）"
    c2.x_axis.title = "年度"
    c2.y_axis.title = "シェア（%）"
    c2.y_axis.scaling.min = 0
    c2.y_axis.scaling.max = 100
    c2.width = 16
    c2.height = 11
    # 共通書式で幅28cmに揃うため、横並び（R3）では重なる。縦に積む
    _series_chart(ws, c2, shr, start2, "F32")


def build_p06(wb):
    """P06: agesex_pyramid_latest.png — 年齢階級・性別ピラミッド（最新年度）"""
    ws = wb.create_sheet("P06_agesex_pyramid_latest")
    df = pd.read_csv(PROC / "agesex_distribution_published_antivegf.csv")
    latest_year = int(df["year"].max())
    _note(ws, f"plots/agesex_pyramid_latest.png の再現"
              f"（agesex_distribution_published_antivegf.csv / year={latest_year}）"
              "　※「男（描画用）」は左向き描画のため符号を反転した列で、実数は「男」列")

    latest = df[df["year"] == latest_year]
    piv = latest.pivot_table(index="age_group_detail", columns="sex",
                             values="quantity", aggfunc="sum").fillna(0)
    order = [a for a in AGE_ORDER_PUBLISHED if a in piv.index]
    piv = piv.reindex(order).reset_index()
    piv.rename(columns={"age_group_detail": "年齢階級"}, inplace=True)
    for s in ["男", "女"]:
        if s not in piv.columns:
            piv[s] = 0
    out = piv[["年齢階級", "男", "女"]].copy()
    out["男（描画用）"] = -out["男"]
    write_df(ws, out, start_row=3)

    chart = BarChart()
    chart.type = "bar"
    chart.grouping = "stacked"
    chart.overlap = 100
    chart.title = f"抗VEGF薬 年齢階級・性別分布（{latest_year}年度）"
    chart.x_axis.title = "処方数量（左：男 / 右：女）"
    chart.width = 20
    chart.height = 16
    n_last = 3 + len(out)
    for col_idx in (4, 3):  # 男（描画用）→ 左、女 → 右
        chart.add_data(Reference(ws, min_col=col_idx, min_row=3, max_row=n_last),
                       titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=4, max_row=n_last))
    chart.x_axis.tickLblPos = "low"  # 目盛ラベルを軸の外側へ（棒と重ねない）
    chart.style = 10
    ws.add_chart(chart, "F3")


def build_p07(wb):
    """P07: age_distribution_heatmap.png — 年齢構成比の推移（条件付き書式のカラースケール）"""
    ws = wb.create_sheet("P07_age_distribution_heatmap")
    _note(ws, "plots/age_distribution_heatmap.png の再現"
              "（agesex_distribution_comparable_antivegf.csv / 男女計の年度内構成比%）"
              "　※Excelにヒートマップのグラフ種別が無いため、条件付き書式のカラースケールで再現")
    df = pd.read_csv(PROC / "agesex_distribution_comparable_antivegf.csv")
    age = df.groupby(["year", "age_group"], as_index=False)["quantity"].sum()
    age["share"] = age.groupby("year")["quantity"].transform(lambda s: s / s.sum() * 100)
    piv = age.pivot_table(index="age_group", columns="year", values="share")
    piv = piv.reindex([a for a in AGE_ORDER if a in piv.index]).reset_index()
    piv.rename(columns={"age_group": "年齢階級"}, inplace=True)
    piv.columns = ["年齢階級"] + [f"{int(c)}年度" for c in piv.columns[1:]]
    write_df(ws, piv, start_row=3)

    last_col = get_column_letter(len(piv.columns))
    ws.conditional_formatting.add(
        f"B4:{last_col}{3 + len(piv)}",
        ColorScaleRule(start_type="min", start_color="FFFFE5",
                       mid_type="percentile", mid_value=50, mid_color="FEB24C",
                       end_type="max", end_color="BD0026"),
    )
    for col in range(2, len(piv.columns) + 1):
        for row in range(4, 4 + len(piv)):
            ws.cell(row=row, column=col).number_format = "0.0"

    # 同じ行列を転置し、Excelのグラフ（100%積み上げ縦棒）でも構成比の推移を描く
    trans = piv.set_index("年齢階級").T.reset_index()
    trans.rename(columns={"index": "年度"}, inplace=True)
    start2 = 3 + len(piv) + 3
    ws.cell(row=start2 - 1, column=1,
            value="【グラフ用】上表を転置（年度×年齢階級）").font = Font(bold=True)
    write_df(ws, trans, start_row=start2)

    chart = BarChart()
    chart.type = "col"
    chart.grouping = "percentStacked"
    chart.overlap = 100
    chart.title = "抗VEGF薬 年齢構成比の推移（%）"
    chart.x_axis.title = "年度"
    chart.y_axis.title = "構成比（%）"
    chart.width = 28
    chart.height = 15
    _series_chart(ws, chart, trans, start2, f"{get_column_letter(len(piv.columns) + 2)}3")


def build_p08(wb):
    """P08: prefecture_ranking_latest.png — 都道府県別 人口10万対（最新年度、横棒）"""
    ws = wb.create_sheet("P08_prefecture_ranking_latest")
    panel = pd.read_csv(PROC / "panel_zero.csv")
    latest_year = int(panel["year"].max())
    _note(ws, f"plots/prefecture_ranking_latest.png の再現"
              f"（panel_zero.csv / code=ANTI_VEGF_TOTAL, year={latest_year}）")

    total = panel[(panel["code"] == "ANTI_VEGF_TOTAL") & (panel["year"] == latest_year)]
    out = total[["prefecture", "count_per_100k"]].sort_values(
        "count_per_100k", ascending=False).reset_index(drop=True)
    out.columns = ["都道府県", "人口10万対 処方数量"]
    write_df(ws, out, start_row=3)

    chart = BarChart()
    chart.type = "bar"
    chart.title = f"抗VEGF薬 都道府県別 人口10万対処方数量（{latest_year}年度）"
    chart.x_axis.title = "人口10万対 処方数量"
    chart.width = 18
    chart.height = 32
    chart.legend = None
    _series_chart(ws, chart, out, 3, "D3")
    chart.x_axis.scaling.orientation = "maxMin"  # 上位が上に来るよう反転


def build_p08b(wb):
    """P08b: P08 の65歳以上人口分母版（M02【2】と同じ数値を最新年度だけ抜き出した横棒）"""
    # 入力は analyze_advanced.py の出力。未実行でも plots ファイル全体を落とさず
    # このシートだけ飛ばす（_build は FileNotFoundError でファイルごとスキップするため）。
    try:
        df = pd.read_csv(PROC / "prefecture_per_65plus_ranking_antivegf.csv")
    except FileNotFoundError as e:
        print(f"    → P08b をスキップ（{e}）")
        return

    rate_cols = [c for c in df.columns if c.endswith("_65歳以上10万対")]
    latest_col = rate_cols[-1]
    latest_year = int(latest_col.split("年度")[0])

    ws = wb.create_sheet("P08b_prefecture_ranking_65plus")
    _note(ws, "P08 の65歳以上人口分母版"
              "（prefecture_per_65plus_ranking_antivegf.csv / M02【2】の最新年度列）。"
              "分母は総務省統計局「人口推計」都道府県別 年齢3区分別人口（各年10月1日現在）の"
              "65歳以上人口。※都道府県×年齢のクロス表が非公表のため真の年齢調整率ではない")

    out = df[["prefecture", latest_col]].sort_values(
        latest_col, ascending=False).reset_index(drop=True)
    out.columns = ["都道府県", "65歳以上人口10万対 処方数量"]
    write_df(ws, out, start_row=3)

    chart = BarChart()
    chart.type = "bar"
    chart.title = f"抗VEGF薬 都道府県別 65歳以上人口10万対処方数量（{latest_year}年度）"
    chart.x_axis.title = "65歳以上人口10万対 処方数量"
    chart.width = 18
    chart.height = 32
    chart.legend = None
    _series_chart(ws, chart, out, 3, "D3")
    chart.x_axis.scaling.orientation = "maxMin"  # 上位が上に来るよう反転


def build_p09(wb):
    """P09: gini_trends.png — 主要解析単位のGini係数推移（折れ線）"""
    ws = wb.create_sheet("P09_gini_trends")
    _note(ws, "plots/gini_trends.png の再現（geographic_disparity_antivegf.csv / "
              "code=ANTI_VEGF_TOTAL, AFLIBERCEPT, RANIBIZUMAB_ALL, FARICIMAB, RANIBIZUMAB_BS）")
    df = pd.read_csv(PROC / "geographic_disparity_antivegf.csv")
    codes = ["ANTI_VEGF_TOTAL", "AFLIBERCEPT", "RANIBIZUMAB_ALL", "FARICIMAB", "RANIBIZUMAB_BS"]
    sub = df[df["code"].isin(codes)]
    piv = _year_pivot(sub, "name", "gini", aggfunc="mean")
    write_df(ws, piv, start_row=3)

    chart = LineChart()
    chart.title = "都道府県間格差（Gini係数）の推移"
    chart.x_axis.title = "年度"
    chart.y_axis.title = "Gini係数"
    chart.width = 24
    chart.height = 13
    _series_chart(ws, chart, piv, 3, f"{get_column_letter(len(piv.columns) + 2)}3")


def build_p10(wb):
    """P10: cost_trends.png — 製品別 薬剤費の推移（積み上げ面、億円）"""
    ws = wb.create_sheet("P10_cost_trends")
    _note(ws, "plots/cost_trends.png の再現（product_trends_antivegf.csv / cost を億円換算）")
    df = pd.read_csv(PROC / "product_trends_antivegf.csv")
    df = df.assign(cost_oku=df["cost"] / 1e8)
    piv = _year_pivot(df, "product_name", "cost_oku")
    write_df(ws, piv, start_row=3)

    chart = AreaChart()
    chart.grouping = "stacked"
    chart.overlap = 100
    chart.title = "抗VEGF薬 薬剤費の推移（薬価×処方数量）"
    chart.x_axis.title = "年度"
    chart.y_axis.title = "薬剤費（億円）"
    chart.width = 28
    chart.height = 15
    _series_chart(ws, chart, piv, 3, f"{get_column_letter(len(piv.columns) + 2)}3")


# ---------------------------------------------------------------------------
# 年齢階級別 男女差（S1〜S5）
# ---------------------------------------------------------------------------

SETTINGS = {
    "外来のみ": ["外来(院内)"],
    "外来＋入院": ["外来(院内)", "入院"],
}


def _load_agesex(settings):
    """秘匿セルを0で補完した年齢性別データ（抗VEGF薬のみ）を読む。"""
    df = pd.read_csv(PROC / "processed_agesex_zero.csv")
    df = df[(df["category"] == "ANTI_VEGF") & (df["setting"].isin(settings))]
    return df.copy()


def _sex_diff_table(df, keys):
    """keys（年度・成分など）×年齢階級で男女差の指標を組み立てる。

    quantity は秘匿セルを0で補完した値。masked/cells で秘匿の程度を併記する。
    """
    g = df.groupby(keys + ["age_group", "sex"]).agg(
        qty=("quantity", "sum"), masked=("masked", "sum"), cells=("masked", "size")
    ).reset_index()

    wide = g.pivot_table(index=keys + ["age_group"], columns="sex",
                         values=["qty", "masked", "cells"], fill_value=0)
    wide.columns = [f"{a}_{b}" for a, b in wide.columns]
    wide = wide.reset_index()
    for col in ["qty_男", "qty_女", "masked_男", "masked_女", "cells_男", "cells_女"]:
        if col not in wide.columns:
            wide[col] = 0

    wide["男−女"] = wide["qty_男"] - wide["qty_女"]
    wide["男/女比"] = [m / f if f > 0 else None
                    for m, f in zip(wide["qty_男"], wide["qty_女"])]
    # 性別内構成比: その集計単位・その性別の総量に占める当該年齢階級の割合
    for s in ["男", "女"]:
        tot = wide.groupby(keys)[f"qty_{s}"].transform("sum")
        wide[f"{s}_構成比%"] = [q / t * 100 if t > 0 else None
                             for q, t in zip(wide[f"qty_{s}"], tot)]
    wide["男性比率%"] = [m / (m + f) * 100 if (m + f) > 0 else None
                     for m, f in zip(wide["qty_男"], wide["qty_女"])]
    wide["秘匿セル数"] = wide["masked_男"] + wide["masked_女"]
    wide["全セル数"] = wide["cells_男"] + wide["cells_女"]
    wide["秘匿率%"] = wide["秘匿セル数"] / wide["全セル数"] * 100

    order = {a: i for i, a in enumerate(AGE_ORDER)}
    wide["_o"] = wide["age_group"].map(order)
    wide = wide.sort_values(keys + ["_o"]).drop(columns=["_o"])

    out = wide[keys + ["age_group", "qty_男", "qty_女", "男−女", "男/女比",
                       "男_構成比%", "女_構成比%", "男性比率%",
                       "秘匿セル数", "全セル数", "秘匿率%"]].copy()
    out.columns = keys + ["年齢階級", "男", "女", "男−女", "男/女比",
                          "男_構成比%", "女_構成比%", "男性比率%",
                          "秘匿セル数", "全セル数", "秘匿率%"]
    return out


def _write_block(ws, df, start_row, caption):
    ws.cell(row=start_row - 1, column=1, value=caption).font = Font(bold=True)
    write_df(ws, df, start_row=start_row)
    return start_row + len(df) + 3


def _fmt(ws, first_row, n_rows, col_formats):
    for col, fmt in col_formats.items():
        for r in range(first_row + 1, first_row + 1 + n_rows):
            ws.cell(row=r, column=col).number_format = fmt


def _build_sexdiff_total(wb, sheet_name, label, settings):
    """S1/S2: 抗VEGF薬合計の年齢階級別 男女差。"""
    ws = wb.create_sheet(sheet_name)
    _note(ws, f"抗VEGF薬合計 × 年齢階級 × 性別（{label}）"
              "　※秘匿セル（NDB原表記「-」＝10未満）は0で補完済み。"
              "「秘匿セル数／全セル数」列で各行の信頼度を確認すること")
    df = _load_agesex(settings)
    tbl = _sex_diff_table(df, ["year"])
    tbl = tbl.rename(columns={"year": "年度"})
    tbl["年度"] = tbl["年度"].astype(int)

    row = _write_block(ws, tbl, 4, f"【明細】年度×年齢階級（{label}）")
    _fmt(ws, 4, len(tbl), {4: "#,##0.0", 5: "#,##0.0", 6: "#,##0.0",
                           7: "0.000", 8: "0.00", 9: "0.00", 10: "0.00", 13: "0.0"})

    # グラフ用ピボット（年齢階級 × 年度）
    def _piv(value):
        p = tbl.pivot_table(index="年齢階級", columns="年度", values=value)
        p = p.reindex([a for a in AGE_ORDER if a in p.index]).reset_index()
        p.columns = ["年齢階級"] + [f"{int(c)}年度" for c in p.columns[1:]]
        return p

    blocks = [
        ("男性比率%", "【グラフ1】年齢階級別 男性比率（%）の推移",
         "年齢階級別 男性比率（%）の推移", "男性比率（%）"),
        ("男−女", "【グラフ2】年齢階級別 男女差（男−女、本）の推移",
         "年齢階級別 男女差（男−女、本）の推移", "男−女（本）"),
    ]
    anchor_col = get_column_letter(len(tbl.columns) + 2)
    anchor_row = 3
    for value, caption, title, ytitle in blocks:
        p = _piv(value)
        start = row
        row = _write_block(ws, p, start, caption)
        chart = LineChart()
        chart.title = f"{title}（{label}）"
        chart.x_axis.title = "年齢階級"
        chart.y_axis.title = ytitle
        chart.width = 28
        chart.height = 14
        _series_chart(ws, chart, p, start, f"{anchor_col}{anchor_row}")
        anchor_row += 29

    # 性別内構成比（最新年度）
    latest = int(tbl["年度"].max())
    comp = tbl[tbl["年度"] == latest][["年齢階級", "男_構成比%", "女_構成比%"]].reset_index(drop=True)
    start = row
    row = _write_block(ws, comp, start, f"【グラフ3】{latest}年度 性別内 年齢構成比（%）")
    chart = LineChart()
    chart.title = f"{latest}年度 性別内 年齢構成比（%）（{label}）"
    chart.x_axis.title = "年齢階級"
    chart.y_axis.title = "構成比（%）"
    chart.width = 28
    chart.height = 14
    _series_chart(ws, chart, comp, start, f"{anchor_col}{anchor_row}")


def _build_sexdiff_molecule(wb, sheet_name, label, settings):
    """S3/S4: 成分別の年齢階級別 男女差。"""
    ws = wb.create_sheet(sheet_name)
    _note(ws, f"成分別 × 年齢階級 × 性別（{label}）"
              "　※秘匿セルは0で補完済み。ペガプタニブは全年度・全セル秘匿のため実質データなし")
    df = _load_agesex(settings)
    tbl = _sex_diff_table(df, ["year", "molecule_name"])
    tbl = tbl.rename(columns={"year": "年度", "molecule_name": "成分"})
    tbl["年度"] = tbl["年度"].astype(int)

    row = _write_block(ws, tbl, 4, f"【明細】年度×成分×年齢階級（{label}）")
    _fmt(ws, 4, len(tbl), {5: "#,##0.0", 6: "#,##0.0", 7: "#,##0.0",
                           8: "0.000", 9: "0.00", 10: "0.00", 11: "0.00", 14: "0.0"})

    anchor_col = get_column_letter(len(tbl.columns) + 2)

    # グラフ1: 成分別 男性比率（%）の推移（全年齢計）
    tot = df.groupby(["year", "molecule_name", "sex"])["quantity"].sum().reset_index()
    tot = tot.pivot_table(index=["year", "molecule_name"], columns="sex",
                          values="quantity", fill_value=0).reset_index()
    for s in ["男", "女"]:
        if s not in tot.columns:
            tot[s] = 0
    tot["男性比率%"] = [m / (m + f) * 100 if (m + f) > 0 else None
                    for m, f in zip(tot["男"], tot["女"])]
    p1 = tot.pivot_table(index="year", columns="molecule_name",
                         values="男性比率%").reset_index()
    p1.rename(columns={"year": "年度"}, inplace=True)
    p1["年度"] = p1["年度"].astype(int)
    p1.columns = [str(c) for c in p1.columns]
    start = row
    row = _write_block(ws, p1, start, "【グラフ1】成分別 男性比率（%）の推移（全年齢計）")
    chart = LineChart()
    chart.title = f"成分別 男性比率（%）の推移（{label}）"
    chart.x_axis.title = "年度"
    chart.y_axis.title = "男性比率（%）"
    chart.width = 26
    chart.height = 14
    _series_chart(ws, chart, p1, start, f"{anchor_col}3")

    # グラフ2: 最新年度の成分別・年齢階級別 男性比率（%）
    latest = int(tbl["年度"].max())
    p2 = tbl[tbl["年度"] == latest].pivot_table(index="年齢階級", columns="成分",
                                              values="男性比率%")
    p2 = p2.reindex([a for a in AGE_ORDER if a in p2.index]).reset_index()
    p2.columns = [str(c) for c in p2.columns]
    start = row
    row = _write_block(ws, p2, start, f"【グラフ2】{latest}年度 成分別・年齢階級別 男性比率（%）")
    chart = LineChart()
    chart.title = f"{latest}年度 成分別・年齢階級別 男性比率（%）（{label}）"
    chart.x_axis.title = "年齢階級"
    chart.y_axis.title = "男性比率（%）"
    chart.width = 26
    chart.height = 14
    _series_chart(ws, chart, p2, start, f"{anchor_col}32")


def build_s5_masking(wb):
    """S5: S1〜S4で使った年齢性別データの秘匿状況サマリ。"""
    ws = wb.create_sheet("S5_秘匿状況")
    _note(ws, "S1〜S4の元データ（processed_agesex_zero.csv）の秘匿状況。"
              "捕捉率＝年齢性別内訳の合計÷公表総計。公表総計は秘匿なしの値")

    raw = pd.read_csv(PROC / "ophthalmic_injection_agesex.csv", dtype=str)
    raw["医薬品コード"] = raw["医薬品コード"].astype(str)
    zero = pd.read_csv(PROC / "processed_agesex_zero.csv")
    zero = zero[zero["category"] == "ANTI_VEGF"]  # マキュエイド・ビスダインを除く
    av_codes = set(zero["drug_code"].astype(str))
    raw = raw[raw["医薬品コード"].isin(av_codes)].copy()
    raw["masked"] = raw["秘匿フラグ"].astype(int)
    raw["val"] = pd.to_numeric(raw["処方数量"], errors="coerce").fillna(0.0)
    raw["総計"] = pd.to_numeric(raw["総計_処方数量"], errors="coerce")
    raw["年度"] = raw["年度"].astype(int)
    name = dict(zip(zero["drug_code"].astype(str), zero["molecule_name"]))
    raw["成分"] = raw["医薬品コード"].map(name)

    row = 4
    for label, settings in SETTINGS.items():
        sub = raw[raw["区分"].isin(settings)]
        num = sub.groupby("年度")["val"].sum()
        den = sub.drop_duplicates(["年度", "医薬品コード", "区分"]).groupby("年度")["総計"].sum()
        cell = sub.groupby("年度")["masked"].agg(["size", "sum"])
        t = pd.DataFrame({
            "年度": num.index.astype(int),
            "内訳合計": num.values,
            "公表総計": den.values,
            "欠落": (den - num).values,
            "捕捉率%": (num / den * 100).values,
            "全セル数": cell["size"].values,
            "秘匿セル数": cell["sum"].values,
            "セル秘匿率%": (cell["sum"] / cell["size"] * 100).values,
        })
        start = row
        row = _write_block(ws, t, start, f"【1】抗VEGF薬合計の捕捉率（{label}）")
        _fmt(ws, start, len(t), {2: "#,##0.0", 3: "#,##0.0", 4: "#,##0.0",
                                 5: "0.00", 8: "0.0"})
        c = LineChart()
        c.add_data(Reference(ws, min_col=5, min_row=start,
                             max_row=start + len(t)), titles_from_data=True)
        c.set_categories(Reference(ws, min_col=1, min_row=start + 1,
                                   max_row=start + len(t)))
        c.title = f"抗VEGF薬合計 年齢性別内訳の捕捉率（%）（{label}）"
        c.x_axis.title = "年度"
        c.y_axis.title = "捕捉率（%）"
        c.width, c.height, c.style = 22, 12, 10
        ws.add_chart(c, f"K{start}")

    # 成分別の捕捉率
    for label, settings in SETTINGS.items():
        sub = raw[raw["区分"].isin(settings)]
        num = sub.groupby(["成分", "年度"])["val"].sum()
        den = sub.drop_duplicates(["年度", "医薬品コード", "区分"]).groupby(
            ["成分", "年度"])["総計"].sum()
        t = (num / den * 100).unstack().reset_index()
        t.columns = ["成分"] + [f"{int(c)}年度" for c in t.columns[1:]]
        start = row
        row = _write_block(ws, t, start, f"【2】成分別 捕捉率%（{label}）")
        _fmt(ws, start, len(t), {c: "0.0" for c in range(2, len(t.columns) + 1)})


# ---------------------------------------------------------------------------
# 医療費（C01〜C06）— analyze_cost.py / analyze_advanced.py の出力
# ---------------------------------------------------------------------------

def build_c01_decomposition(wb):
    """C01: 薬剤費変化の要因分解（数量・品目構成・薬価）。"""
    ws = wb.create_sheet("C01_費用要因分解")
    _note(ws, "cost_decomposition_yearly.csv / cost_decomposition_cumulative.csv。"
              "ΔC = 数量効果 + 構成効果 + 薬価効果 に厳密分解（残差0）。単位は億円")
    df = pd.read_csv(PROC / "cost_decomposition_yearly.csv")
    out = pd.DataFrame({
        "年度": df["year"].astype(int),
        "薬剤費（前年）": df["cost_prev"] / OKU,
        "薬剤費（当年）": df["cost_current"] / OKU,
        "変化額": df["cost_change"] / OKU,
        "数量効果": df["volume_effect"] / OKU,
        "構成効果": df["mix_effect"] / OKU,
        "薬価効果": df["price_effect"] / OKU,
        "新規収載品目数": df["n_new_products"].astype(int),
        "平均単価（前年・円）": df["mean_price_prev"],
        "平均単価（当年・円）": df["mean_price_current"],
    })
    row = _write_block(ws, out, 4, "【明細】年度別の要因分解（億円）")
    _fmt(ws, 4, len(out), {c: "#,##0.1" for c in range(2, 8)})

    cum = pd.read_csv(PROC / "cost_decomposition_cumulative.csv")
    cum_out = pd.DataFrame({
        "期間": cum["period"],
        "変化額": cum["cost_change"] / OKU,
        "数量効果": cum["volume_effect"] / OKU,
        "構成効果": cum["mix_effect"] / OKU,
        "薬価効果": cum["price_effect"] / OKU,
    })
    row = _write_block(ws, cum_out, row, "【累計】期間全体の要因分解（億円）")

    chart = BarChart()
    chart.type = "col"
    chart.grouping = "stacked"
    chart.overlap = 100
    chart.title = "薬剤費変化の要因分解（億円）"
    chart.x_axis.title = "年度"
    chart.y_axis.title = "変化額（億円）"
    chart.width = 28
    chart.height = 15
    n_last = 4 + len(out)
    for col in (5, 6, 7):  # 数量・構成・薬価
        chart.add_data(Reference(ws, min_col=col, min_row=4, max_row=n_last),
                       titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=5, max_row=n_last))
    chart.style = 10
    ws.add_chart(chart, "M3")


def build_c02_total_cost(wb):
    """C02: 薬剤費＋手技料の治療総額と薬剤費シェア。"""
    ws = wb.create_sheet("C02_治療総額")
    _note(ws, "cost_total_with_procedure.csv。手技料 = G016算定回数 × 点数 × 10円。"
              "薬剤費は薬価×数量（薬価基準額であり実際の償還額ではない）")
    df = pd.read_csv(PROC / "cost_total_with_procedure.csv")
    out = pd.DataFrame({
        "年度": df["year"].astype(int),
        "薬剤費": df["drug_cost"] / OKU,
        "手技料": df["procedure_cost"] / OKU,
        "治療総額": df["total_cost"] / OKU,
        "薬剤費シェア(%)": df["drug_share_pct"],
        "G016算定回数": df["g016_procedures"].astype(int),
        "点数": df["points"].astype(int),
    })
    write_df(ws, out, start_row=3)
    _fmt(ws, 3, len(out), {2: "#,##0.1", 3: "#,##0.1", 4: "#,##0.1", 5: "0.00"})

    chart = BarChart()
    chart.type = "col"
    chart.grouping = "stacked"
    chart.overlap = 100
    chart.title = "抗VEGF治療の総額（薬剤費＋手技料、億円）"
    chart.x_axis.title = "年度"
    chart.y_axis.title = "費用（億円）"
    chart.width = 26
    chart.height = 14
    n_last = 3 + len(out)
    for col in (2, 3):
        chart.add_data(Reference(ws, min_col=col, min_row=3, max_row=n_last),
                       titles_from_data=True)
    cats = Reference(ws, min_col=1, min_row=4, max_row=n_last)
    chart.set_categories(cats)

    line = LineChart()
    line.y_axis.title = "薬剤費シェア（%）"
    line.y_axis.axId = 200
    line.y_axis.scaling.min = 90
    line.y_axis.scaling.max = 100
    line.add_data(Reference(ws, min_col=5, min_row=3, max_row=n_last),
                  titles_from_data=True)
    line.set_categories(cats)
    chart.y_axis.crosses = "min"
    line.y_axis.crosses = "max"
    chart += line
    chart.style = 10
    ws.add_chart(chart, "I3")


def build_c03_counterfactual(wb):
    """C03: 反実仮想シナリオ（BS非参入／薬価据置）。"""
    ws = wb.create_sheet("C03_反実仮想")
    _note(ws, "cost_counterfactual_scenarios.csv。単位は億円。"
              "「BS非参入」はBSの数量を先発（注射液）薬価で換算した場合、"
              "「薬価据置」は各製品を収載時薬価のまま据え置いた場合")
    df = pd.read_csv(PROC / "cost_counterfactual_scenarios.csv")
    df["year"] = df["year"].astype(int)

    row = 4
    for scenario, g in df.groupby("scenario", sort=False):
        out = pd.DataFrame({
            "年度": g["year"].values,
            "実績": g["actual_cost"].values / OKU,
            "反実仮想": g["counterfactual_cost"].values / OKU,
            "差額（抑制額）": g["difference"].values / OKU,
        })
        out.loc[len(out)] = ["累計", out["実績"].sum(), out["反実仮想"].sum(),
                             out["差額（抑制額）"].sum()]
        start = row
        row = _write_block(ws, out, start, f"【{scenario}】")
        _fmt(ws, start, len(out), {2: "#,##0.1", 3: "#,##0.1", 4: "#,##0.1"})

        chart = LineChart()
        chart.title = f"{scenario}（億円）"
        chart.x_axis.title = "年度"
        chart.y_axis.title = "薬剤費（億円）"
        chart.width = 24
        chart.height = 13
        # 累計行はグラフに含めない
        n_last = start + len(out) - 1
        for col in (2, 3):
            chart.add_data(Reference(ws, min_col=col, min_row=start, max_row=n_last),
                           titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=start + 1, max_row=n_last))
        chart.style = 10
        ws.add_chart(chart, f"G{start}")


def build_c04_bs_potential(wb):
    """C04: BSシェアの地域差を解消した場合の追加削減余地。"""
    ws = wb.create_sheet("C04_BS地域削減余地")
    _note(ws, "cost_biosimilar_regional_potential.csv。"
              "全国のBSシェアを上位10%の県の水準（P90）まで引き上げた場合の追加削減額。"
              "分母は都道府県内訳の合計（公表総計より秘匿分だけ小さい）")
    df = pd.read_csv(PROC / "cost_biosimilar_regional_potential.csv")
    out = pd.DataFrame({
        "年度": df["year"].astype(int),
        "全国BSシェア(%)": df["current_national_bs_share_pct"],
        "P90県のBSシェア(%)": df["target_bs_share_p90"],
        "追加切替可能数量（本）": df["additional_bs_switch_quantity"],
        "1本あたり削減額（円）": df["unit_saving_yen"].astype(int),
        "追加削減余地（億円）": df["additional_saving"] / OKU,
    })
    write_df(ws, out, start_row=3)
    _fmt(ws, 3, len(out), {2: "0.00", 3: "0.00", 4: "#,##0.0", 6: "#,##0.00"})

    chart = BarChart()
    chart.type = "col"
    chart.title = "BSシェアの地域差解消による追加削減余地"
    chart.x_axis.title = "年度"
    chart.y_axis.title = "追加削減余地（億円）"
    chart.width = 24
    chart.height = 13
    n_last = 3 + len(out)
    chart.add_data(Reference(ws, min_col=6, min_row=3, max_row=n_last),
                   titles_from_data=True)
    cats = Reference(ws, min_col=1, min_row=4, max_row=n_last)
    chart.set_categories(cats)

    line = LineChart()
    line.y_axis.title = "BSシェア（%）"
    line.y_axis.axId = 200
    for col in (2, 3):
        line.add_data(Reference(ws, min_col=col, min_row=3, max_row=n_last),
                      titles_from_data=True)
    line.set_categories(cats)
    chart.y_axis.crosses = "min"
    line.y_axis.crosses = "max"
    chart += line
    chart.style = 10
    ws.add_chart(chart, "H3")


def build_c05_price_gap(wb):
    """C05: 実勢価格乖離を仮定した感度分析。"""
    ws = wb.create_sheet("C05_薬価差感度")
    _note(ws, "cost_price_gap_sensitivity.csv。薬価基準額は実勢価格を上回るため、"
              "乖離0/5/10%を仮定した治療総額の幅を示す。単位は億円")
    df = pd.read_csv(PROC / "cost_price_gap_sensitivity.csv")
    piv = df.pivot_table(index="year", columns="price_gap_pct",
                         values="total_cost_adjusted").reset_index()
    piv.columns = ["年度"] + [f"乖離{int(c)}%" for c in piv.columns[1:]]
    piv["年度"] = piv["年度"].astype(int)
    for c in piv.columns[1:]:
        piv[c] = piv[c] / OKU
    write_df(ws, piv, start_row=3)
    _fmt(ws, 3, len(piv), {c: "#,##0.1" for c in range(2, len(piv.columns) + 1)})

    chart = LineChart()
    chart.title = "実勢価格乖離を仮定した治療総額（億円）"
    chart.x_axis.title = "年度"
    chart.y_axis.title = "治療総額（億円）"
    chart.width = 24
    chart.height = 13
    _series_chart(ws, chart, piv, 3, "G3")


def build_c06_price_revision(wb):
    """C06: 薬価改定の年表（製品別の改定率）。"""
    ws = wb.create_sheet("C06_薬価改定年表")
    _note(ws, "product_price_revision_timeline.csv / product_price_revision_summary.csv。"
              "改定区分は中医協の薬価改定スケジュールに基づく注記で、改定率は"
              "NDBオープンデータの薬価欄から算出。薬価は年度あたり1値のため、"
              "年度途中の改定（2019年10月の消費税対応）は分離できない")
    df = pd.read_csv(PROC / "product_price_revision_timeline.csv")

    summary = pd.read_csv(PROC / "product_price_revision_summary.csv")
    s_out = pd.DataFrame({
        "年度": summary["year"].astype(int),
        "改定区分": summary["revision_type"],
        "収載品目数": summary["n_products"].astype(int),
        "薬価が動いた品目数": summary["n_revised"].astype(int),
        "平均改定率(%)": summary["mean_change_pct"],
        "最大下落(%)": summary["min_change_pct"],
        "最大上昇(%)": summary["max_change_pct"],
    })
    row = _write_block(ws, s_out, 5, "【1】年度別サマリ")
    _fmt(ws, 5, len(s_out), {5: "0.00", 6: "0.00", 7: "0.00"})

    piv = df.pivot_table(index="product_name", columns="year",
                         values="change_pct")
    piv = piv.reset_index()
    piv.columns = ["製品"] + [f"{int(c)}年度" for c in piv.columns[1:]]
    start = row
    row = _write_block(ws, piv, start, "【2】製品別 改定率（%）")
    _fmt(ws, start, len(piv), {c: "0.0" for c in range(2, len(piv.columns) + 1)})

    chart = LineChart()
    chart.title = "製品別 薬価改定率（%）"
    chart.x_axis.title = "製品"
    chart.y_axis.title = "改定率（%）"
    chart.width = 26
    chart.height = 14
    _series_chart(ws, chart, piv, start, "J5")


# ---------------------------------------------------------------------------
# 統計手法・検証（M01〜M07）— analyze_advanced.py の出力
# ---------------------------------------------------------------------------

def build_m01_joinpoint(wb):
    """M01: 変化点解析（分節対数線形回帰）。"""
    ws = wb.create_sheet("M01_変化点解析")
    _note(ws, "trend_joinpoint_antivegf.csv。変化点は観測年上で探索し、"
              "各区間3点以上・最大2点の条件でBIC最小のモデルを採用。"
              "AAPCは区間長で重みづけした平均傾き")
    df = pd.read_csv(PROC / "trend_joinpoint_antivegf.csv")
    out = pd.DataFrame({
        "解析単位": df["name"],
        "変化点数": df["n_joinpoints"].astype(int),
        "変化点": df["joinpoints"].fillna(""),
        "区間": df["segment"].astype(int),
        "開始年度": df["segment_start"].astype(int),
        "終了年度": df["segment_end"].astype(int),
        "区間APC(%)": df["apc"],
        "95%CI下限": df["apc_low"],
        "95%CI上限": df["apc_high"],
        "有意": df["significant"],
        "AAPC(%)": df["aapc"],
        "AAPC下限": df["aapc_low"],
        "AAPC上限": df["aapc_high"],
        "R²": df["r2"],
    })
    write_df(ws, out, start_row=3)
    _fmt(ws, 3, len(out), {c: "0.00" for c in range(7, 14)})
    _fmt(ws, 3, len(out), {14: "0.000"})

    # 区間APCを解析単位×区間で並べた棒グラフ
    piv = df.pivot_table(index="name", columns="segment", values="apc").reset_index()
    piv.columns = ["解析単位"] + [f"区間{int(c)}" for c in piv.columns[1:]]
    start = 3 + len(out) + 3
    ws.cell(row=start - 1, column=1, value="【グラフ用】区間別APC（%）").font = Font(bold=True)
    write_df(ws, piv, start_row=start)

    chart = BarChart()
    chart.type = "col"
    chart.title = "区間別 年平均変化率（APC, %）"
    chart.y_axis.title = "APC（%）"
    chart.width = 26
    chart.height = 14
    _series_chart(ws, chart, piv, start, "Q3")


def build_m02_65plus(wb):
    """M02: 65歳以上人口を分母にした率と格差。"""
    ws = wb.create_sheet("M02_65歳以上分母")
    _note(ws, "geographic_disparity_65plus_antivegf.csv / "
              "prefecture_per_65plus_ranking_antivegf.csv。"
              "※真の年齢調整率は都道府県×年齢のクロス表が非公表のため算出できない。"
              "高齢人口を分母に置くことで年齢構成の県差を部分的にのみ調整している")
    disp = pd.read_csv(PROC / "geographic_disparity_65plus_antivegf.csv")
    tot = disp[disp["code"] == "ANTI_VEGF_TOTAL"].sort_values("year")
    out = pd.DataFrame({
        "年度": tot["year"].astype(int).values,
        "平均（65歳以上10万対）": tot["mean_rate_per_100k_65plus"].values,
        "CV（65歳以上）": tot["cv_65plus"].values,
        "Gini（65歳以上）": tot["gini_65plus"].values,
        "Gini（総人口）": tot["gini_total_pop"].values,
        "Gini低下率(%)": tot["gini_reduction_pct"].values,
        "最大/最小（65歳以上）": tot["max_to_min_ratio_65plus"].values,
        "最多": tot["max_prefecture"].values,
        "最少": tot["min_prefecture"].values,
    })
    row = _write_block(ws, out, 5, "【1】抗VEGF薬合計 格差指標（分母の比較）")
    _fmt(ws, 5, len(out), {2: "#,##0.0", 3: "0.000", 4: "0.000", 5: "0.000",
                           6: "0.0", 7: "0.00"})

    chart = LineChart()
    chart.title = "分母を変えたときのGini係数（総人口 vs 65歳以上人口）"
    chart.x_axis.title = "年度"
    chart.y_axis.title = "Gini係数"
    chart.width = 24
    chart.height = 13
    n_last = 5 + len(out)
    for col in (4, 5):
        chart.add_data(Reference(ws, min_col=col, min_row=5, max_row=n_last),
                       titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=6, max_row=n_last))
    chart.style = 10
    ws.add_chart(chart, "L5")

    rank = pd.read_csv(PROC / "prefecture_per_65plus_ranking_antivegf.csv")
    _write_block(ws, rank, row, "【2】都道府県別 65歳以上10万対（最新年度の降順）")


def build_m03_convergence(wb):
    """M03: β収束・σ収束。"""
    ws = wb.create_sheet("M03_収束")
    _note(ws, "convergence_beta_antivegf.csv。β収束＝初期水準の高い県ほど"
              "その後の伸びが低いか（β<0で格差縮小）。σ収束＝対数率の標準偏差の推移。"
              "初期年に0（全セル秘匿）の県は対数がとれないため除外している")
    df = pd.read_csv(PROC / "convergence_beta_antivegf.csv")

    beta = df[df["section"] == "beta収束"]
    b_out = pd.DataFrame({
        "期間": beta["period"].values,
        "対象県数": beta["n_prefectures"].astype(int).values,
        "β": beta["beta"].values,
        "SE": beta["se"].values,
        "p値": beta["p_value"].values,
        "R²": beta["r2"].values,
        "収束（β<0かつp<0.05）": beta["converging"].values,
        "収束速度λ(/年)": beta["lambda_per_year"].values,
        "半減期（年）": beta["half_life_years"].values,
    })
    row = _write_block(ws, b_out, 5, "【1】β収束")
    _fmt(ws, 5, len(b_out), {3: "0.0000", 4: "0.0000", 5: "0.000E+00",
                             6: "0.000", 8: "0.0000", 9: "0.0"})

    sigma = df[df["section"] == "sigma収束"]
    s_out = pd.DataFrame({
        "年度": sigma["year"].astype(int).values,
        "対象県数": sigma["n_prefectures"].astype(int).values,
        "対数率のSD": sigma["sd_log_rate"].values,
        "CV": sigma["cv"].values,
    })
    start = row
    row = _write_block(ws, s_out, start, "【2】σ収束")
    _fmt(ws, start, len(s_out), {3: "0.000", 4: "0.000"})

    chart = LineChart()
    chart.title = "σ収束：都道府県別 人口10万対の散らばり"
    chart.x_axis.title = "年度"
    chart.y_axis.title = "対数率のSD / CV"
    chart.width = 24
    chart.height = 13
    n_last = start + len(s_out)
    for col in (3, 4):
        chart.add_data(Reference(ws, min_col=col, min_row=start, max_row=n_last),
                       titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=start + 1, max_row=n_last))
    chart.style = 10
    ws.add_chart(chart, "L5")


def build_m04_sensitivity_apc(wb):
    """M04: 補完戦略別のAPC。"""
    ws = wb.create_sheet("M04_APC感度")
    _note(ws, "sensitivity_apc_antivegf.csv。県別内訳を全国に足し上げた率でAPCを再計算し、"
              "秘匿補完の方法（zero/five/random）によって傾向推定が変わらないことを確認する。"
              "§2-3の全国APCは公表総計ベースなので、そもそも補完の影響を受けない")
    df = pd.read_csv(PROC / "sensitivity_apc_antivegf.csv")
    out = pd.DataFrame({
        "解析単位": df["name"], "コード": df["code"],
        "補完": df["imputation"],
        "期間": df["start_year"].astype(str) + "-" + df["end_year"].astype(str),
        "APC(%)": df["apc"], "95%CI下限": df["apc_low"], "95%CI上限": df["apc_high"],
        "p値": df["p_value"], "R²": df["r2"],
        "zeroとの差(pt)": df["apc_diff_vs_zero"],
    })
    row = _write_block(ws, out, 4, "【明細】補完戦略別のAPC")
    _fmt(ws, 4, len(out), {5: "0.000", 6: "0.000", 7: "0.000", 8: "0.000E+00",
                           9: "0.000", 10: "0.0000"})

    piv = df.pivot_table(index="name", columns="imputation", values="apc").reset_index()
    piv.columns = ["解析単位"] + [str(c) for c in piv.columns[1:]]
    start = row
    _write_block(ws, piv, start, "【グラフ用】解析単位×補完戦略のAPC（%）")

    chart = BarChart()
    chart.type = "col"
    chart.title = "補完戦略別のAPC（%）"
    chart.y_axis.title = "APC（%）"
    chart.width = 26
    chart.height = 14
    _series_chart(ws, chart, piv, start, "M4")


def build_m05_sex_tests(wb):
    """M05: 男性比率の傾向検定（参考値）。"""
    ws = wb.create_sheet("M05_性差検定")
    _note(ws, "agesex_sex_ratio_tests.csv。"
              "★観測単位はバイアル本数であって患者ではない。同一患者が年間6〜12回"
              "投与を受けるため実質的な標本サイズは本数よりはるかに小さく、p値は"
              "極端に小さく出る。効果量（男性比率・男/女比）で解釈し、p値は推論に用いないこと")
    df = pd.read_csv(PROC / "agesex_sex_ratio_tests.csv")
    out = pd.DataFrame({
        "検定": df["test"],
        "年度": df["year"],
        "男": df["male"], "女": df["female"],
        "男性比率(%)": df["male_share_pct"],
        "男/女比": df["male_female_ratio"],
        "統計量z": df["statistic"],
        "p値（参考）": df["p_value"],
        "χ²（性別×年齢の独立性）": df.get("chi2_independence"),
        "年齢階級数": df["n_age_groups"],
    })
    write_df(ws, out, start_row=5)
    _fmt(ws, 5, len(out), {3: "#,##0.0", 4: "#,##0.0", 5: "0.00", 6: "0.000",
                           7: "0.00", 9: "#,##0.0"})

    year_rows = df[df["year"].notna()]
    piv = pd.DataFrame({
        "年度": year_rows["year"].astype(int).values,
        "男性比率(%)": year_rows["male_share_pct"].values,
        "男/女比": year_rows["male_female_ratio"].values,
    })
    start = 5 + len(out) + 3
    ws.cell(row=start - 1, column=1, value="【グラフ用】男性比率の推移").font = Font(bold=True)
    write_df(ws, piv, start_row=start)

    chart = LineChart()
    chart.title = "男性比率と男/女比の推移（外来のみ・全年齢計）"
    chart.x_axis.title = "年度"
    chart.y_axis.title = "男性比率（%）"
    chart.width = 24
    chart.height = 13
    _series_chart(ws, chart, piv, start, "M5")


def build_m06_g016_prefecture(wb):
    """M06: G016算定回数と薬剤数量の県別突合。"""
    ws = wb.create_sheet("M06_G016県別突合")
    _note(ws, "g016_vs_drug_prefecture_summary.csv / g016_vs_drug_by_prefecture.csv。"
              "G016の県別データには秘匿セルが1件もない。両者の比は"
              "**県ごとの秘匿による欠落量の推定値**になる")
    summary = pd.read_csv(PROC / "g016_vs_drug_prefecture_summary.csv")
    s_out = pd.DataFrame({
        "年度": summary["year"].astype(int),
        "県数": summary["n_prefectures"].astype(int),
        "平均比(%)": summary["mean_ratio_pct"],
        "中央値(%)": summary["median_ratio_pct"],
        "最小(%)": summary["min_ratio_pct"],
        "最大(%)": summary["max_ratio_pct"],
        "100%超の県数": summary["n_over_100"].astype(int),
        "最小の県": summary["min_prefecture"],
    })
    row = _write_block(ws, s_out, 5, "【1】年度別サマリ（抗VEGF薬数量 ÷ G016算定回数）")
    _fmt(ws, 5, len(s_out), {3: "0.00", 4: "0.00", 5: "0.00", 6: "0.00"})

    chart = LineChart()
    chart.title = "抗VEGF薬数量 / G016算定回数（県別の分布）"
    chart.x_axis.title = "年度"
    chart.y_axis.title = "比（%）"
    chart.width = 24
    chart.height = 13
    n_last = 5 + len(s_out)
    for col in (3, 4, 5, 6):
        chart.add_data(Reference(ws, min_col=col, min_row=5, max_row=n_last),
                       titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=6, max_row=n_last))
    chart.style = 10
    ws.add_chart(chart, "K5")

    detail = pd.read_csv(PROC / "g016_vs_drug_by_prefecture.csv")
    piv = detail.pivot_table(index="prefecture", columns="year",
                             values="ratio_pct").reset_index()
    piv.columns = ["都道府県"] + [f"{int(c)}年度" for c in piv.columns[1:]]
    last = piv.columns[-1]
    piv = piv.sort_values(last).reset_index(drop=True)
    start = row
    _write_block(ws, piv, start, "【2】都道府県別の比（%、最新年度の昇順）")
    _fmt(ws, start, len(piv), {c: "0.0" for c in range(2, len(piv.columns) + 1)})

    last_col = get_column_letter(len(piv.columns))
    ws.conditional_formatting.add(
        f"B{start + 1}:{last_col}{start + len(piv)}",
        ColorScaleRule(start_type="min", start_color="F03B20",
                       mid_type="percentile", mid_value=50, mid_color="FEB24C",
                       end_type="max", end_color="FFFFB2"))


def build_m07_imputation(wb):
    """M07: 補完戦略によるGini・順位相関の感度分析（既存の出力）。"""
    ws = wb.create_sheet("M07_補完感度")
    _note(ws, "sensitivity_imputation_antivegf.csv。"
              "秘匿セルを 0 / 5 / 1〜9の乱数 で補完したときのGini係数と、"
              "zero補完との都道府県順位のSpearman相関")
    df = pd.read_csv(PROC / "sensitivity_imputation_antivegf.csv")
    df = df.sort_values(["code", "year"])
    write_df(ws, df, start_row=3)
    num_cols = [i + 1 for i, c in enumerate(df.columns)
                if c.startswith(("gini_", "rank_corr_"))]
    _fmt(ws, 3, len(df), {c: "0.000" for c in num_cols})


# ---------------------------------------------------------------------------

def _build(filename, title, builders):
    """1ファイル分のワークブックを組み立てて保存する。

    入力CSVが1つでも欠けていればファイルを作らずスキップする。中途半端な
    ワークブックを残すより、何が足りないかを表示して飛ばすほうが安全なため。
    """
    wb = Workbook()
    wb.remove(wb.active)
    for name, fn in builders:
        print(f"  {name}")
        try:
            fn(wb)
        except FileNotFoundError as e:
            print(f"  → スキップ: {filename} （{e}）\n")
            return None

    # 各builderが作り終えたあとで、全グラフに共通書式をまとめて適用する。
    # 個々のbuilderに書式指定を散らかさないための一元化。
    n_charts = 0
    for ws in wb.worksheets:
        for chart in ws._charts:
            tall = chart.height >= TALL_THRESHOLD
            chart.width, chart.height = (TALL_W, TALL_H) if tall else (CHART_W, CHART_H)
            _style_chart(chart)
            n_charts += 1

    path = FIG_DIR / filename
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))
    size_kb = path.stat().st_size / 1024
    print(f"→ {filename}  ({len(wb.sheetnames)} シート / {n_charts} グラフ / "
          f"{size_kb:,.0f} KB)  {title}\n")
    return path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--nokouhi", action="store_true",
                    help="公費レセプトを含まないデータ（抗VEGF薬解析/公費含まない）で作図する")
    args = ap.parse_args()

    global ROOT, FIG_DIR, PROC
    ROOT = BASE / "公費含まない" if args.nokouhi else BASE
    FIG_DIR = ROOT / "04_図表"
    PROC = _CsvLocator(ROOT)
    print(f"データフォルダ: {ROOT}\n出力先: {FIG_DIR}\n")

    workbooks = [
        ("antivegf_fig_main.xlsx", "論文本体のFigure", [
            ("Fig 1A", build_fig1a), ("Fig 1B", build_fig1b), ("Fig 1C", build_fig1c),
            ("Fig 2A", build_fig2a), ("Fig 2B", build_fig2b),
            ("Fig 3A", build_fig3a), ("Fig 3B", build_fig3b),
            ("Fig 4A", build_fig4a), ("Fig 4B", build_fig4b),
            ("Fig 5A", build_fig5a), ("Fig 5B", build_fig5b), ("Fig 5C", build_fig5c),
            ("Suppl Fig", build_suppl),
        ]),
        ("antivegf_fig_plots.xlsx", "plots/*.png のExcel再現", [
            ("P01 product_trends", build_p01),
            ("P02 product_share_stacked", build_p02),
            ("P03 formulation_share_total", build_p03),
            ("P04 formulation_kit_share_by_molecule", build_p04),
            ("P05 biosimilar_share", build_p05),
            ("P06 agesex_pyramid_latest", build_p06),
            ("P07 age_distribution_heatmap", build_p07),
            ("P08 prefecture_ranking_latest", build_p08),
            ("P08b prefecture_ranking_65plus", build_p08b),
            ("P09 gini_trends", build_p09),
            ("P10 cost_trends", build_p10),
        ]),
        ("antivegf_fig_sexdiff.xlsx", "年齢階級別の男女差", [
            ("S1 男女差 合計 外来のみ",
             lambda wb: _build_sexdiff_total(wb, "S1_男女差_合計_外来のみ",
                                             "外来のみ", SETTINGS["外来のみ"])),
            ("S2 男女差 合計 外来＋入院",
             lambda wb: _build_sexdiff_total(wb, "S2_男女差_合計_外来入院",
                                             "外来＋入院", SETTINGS["外来＋入院"])),
            ("S3 男女差 成分別 外来のみ",
             lambda wb: _build_sexdiff_molecule(wb, "S3_男女差_成分別_外来のみ",
                                                "外来のみ", SETTINGS["外来のみ"])),
            ("S4 男女差 成分別 外来＋入院",
             lambda wb: _build_sexdiff_molecule(wb, "S4_男女差_成分別_外来入院",
                                                "外来＋入院", SETTINGS["外来＋入院"])),
            ("S5 秘匿状況", build_s5_masking),
        ]),
        ("antivegf_fig_cost.xlsx", "医療費", [
            ("C01 費用要因分解", build_c01_decomposition),
            ("C02 治療総額", build_c02_total_cost),
            ("C03 反実仮想", build_c03_counterfactual),
            ("C04 BS地域削減余地", build_c04_bs_potential),
            ("C05 薬価差感度", build_c05_price_gap),
            ("C06 薬価改定年表", build_c06_price_revision),
        ]),
        ("antivegf_fig_methods.xlsx", "統計手法・感度分析・検証", [
            ("M01 変化点解析", build_m01_joinpoint),
            ("M02 65歳以上分母", build_m02_65plus),
            ("M03 収束", build_m03_convergence),
            ("M04 APC感度", build_m04_sensitivity_apc),
            ("M05 性差検定", build_m05_sex_tests),
            ("M06 G016県別突合", build_m06_g016_prefecture),
            ("M07 補完感度", build_m07_imputation),
        ]),
    ]

    built = 0
    for filename, title, builders in workbooks:
        print(f"=== {filename} — {title}")
        if _build(filename, title, builders):
            built += 1
    print(f"完了: {built}/{len(workbooks)} ファイルを生成")


if __name__ == "__main__":
    main()
