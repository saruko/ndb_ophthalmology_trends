# -*- coding: utf-8 -*-
"""追加図表（FigureA1/A2）の xlsx と 編集可能 pptx（解説ノート付き）を生成する。

入力は build_additional_figures_antivegf.py が書き出した 追加図表/FigureA1_data.csv・
FigureA2A_coefficients.csv・FigureA2B_residuals_2024.csv。
本編 Figure1〜6.pptx と同じ規格（縦長 7.5×11 in、Excel ネイティブグラフ）。
FigureA2A の 95%CI は Excel のカスタム誤差範囲で表現する（python-pptx の制約により
横向きの誤差バーではなく縦棒グラフ＋誤差バーとする）。

実行: .venv\\Scripts\\python.exe 抗VEGF薬解析/06_論文投稿/build_additional_pptx_xlsx_antivegf.py
"""
import sys
from pathlib import Path

import pandas as pd
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import build_pptx_figures_antivegf as B  # noqa: E402  helpers 共用（実行はされない）
from pptx_bounds import add_error_bars  # noqa: E402  (B が sys.path に 抗アレルギー点眼解析 を追加済み)

OUT = HERE / "追加図表"
LM, CW = B.LM, B.CW

NOTES_A1 = """FigureA1. 都道府県別 65歳以上人口10万対処方数量（2024年度）— Figure 5A の65歳以上分母版
Figure 5A（総人口分母）と同じ構成で、分母を65歳以上人口に置き換えたランキング。高齢化率の県差を除いた利用率を示す。
・棒: ゼロ補完した県別内訳による 65歳以上10万対（本）。上位: 群馬・香川・東京。総人口分母では下位でない東京などの都市部が相対的に上昇し、順位が入れ替わる。
・グレー（PNG版ではハッチ）の棒: 県別内訳が G016 算定回数の90%未満しか捕捉していない6県（山梨・高知・鳥取・福井・徳島・岩手）。秘匿による過小評価の可能性がある。
・全国値 2,938（Table 1 の 2024年度 65歳以上10万対）は注記テキストに記載（必要なら PowerPoint 上で縦線・複合グラフとして追加）。
データ: FigureA1_data.csv / _data.xlsx（prefecture_per_65plus_ranking_antivegf.csv, g016_vs_drug_by_prefecture.csv 由来、公費レセプト含まない集計）。"""

NOTES_A2 = """FigureA2. 二元固定効果パネル回帰（抗VEGF合計、2014–2024年度）の係数と補正後残差
A. 県・年の二元固定効果、県クラスターSEのパネル回帰係数と95%CI（高齢化率・眼科医密度・眼科施設密度）。PNG版は横向きの点＋CI、pptx版は python-pptx の制約により縦棒＋カスタム誤差範囲で表現。3共変量ともCIが0をまたぎ、県間差は供給側要因では説明できない（原稿 §4、Suppl Table S5）。係数はパイプライン出力 panel_regression_summary_antivegf.csv と一致することを生成時に照合済み。
B. 固定効果と共変量を除去した後の2024年度の県別残差（人口10万対）。正（赤）= 補正後もなお高い県（茨城・香川・和歌山など）、負（青）= 低い県（高知・沖縄・山梨など）。山梨・高知は秘匿による過小評価県（FigureA1 のグレー対象）のため、負の残差は秘匿の影響を含む点に注意。
データ: FigureA2A_coefficients.csv, FigureA2B_residuals_2024.csv / FigureA2_data.xlsx。残差は linearmodels PanelOLS を同一仕様で再推定して取得。"""


def build_a1():
    df = pd.read_csv(OUT / "FigureA1_data.csv")  # rate 降順で保存されている
    df = df.sort_values("per_100k_65plus_2024", ascending=True)  # 横棒は下から積むため昇順
    prs = B.new_prs()
    s = B.blank(prs)
    B.fig_label(s, "Fig. A1")
    B.textbox(s, LM, Inches(0.5), CW, Inches(0.3),
              "Prefecture-level utilisation per 100,000 population aged 65+, FY2024", 11, True)
    w = pd.DataFrame({"Vials / PFS per 100,000 aged 65+": df["per_100k_65plus_2024"].to_numpy()},
                     index=df["prefecture_en"])
    ch = B.add_chart(s, XL_CHART_TYPE.BAR_CLUSTERED, LM, Inches(0.8), CW, Inches(9.0), w, None,
                     ytitle="Vials / PFS per 100,000 population aged 65+, FY2024",
                     numfmt="#,##0", legend=False, gap_width=30, font_size=7)
    ser = ch.plots[0].series[0]
    for i, cap in enumerate(df["antivegf_per_g016_pct_2024"]):
        pt = ser.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = B.rgb("9E9E9E" if cap < 90 else "3182BD")
    grey = ", ".join(df.loc[df["antivegf_per_g016_pct_2024"] < 90, "prefecture_en"])
    B.textbox(s, LM, Inches(9.9), CW, Inches(0.6),
              "Grey bars: prefecture breakdown captures <90% of G016 claims (under-estimated by cell "
              f"suppression): {grey}. National value 2,938 per 100,000 aged 65+ "
              "(add as a line if needed).", 7, color="595959")
    s.notes_slide.notes_text_frame.text = NOTES_A1
    prs.save(OUT / "FigureA1.pptx")
    print("saved FigureA1.pptx")

    with pd.ExcelWriter(OUT / "FigureA1_data.xlsx") as xw:
        out = df.sort_values("per_100k_65plus_2024", ascending=False)
        out = out[["prefecture", "prefecture_en", "per_100k_65plus_2024", "antivegf_per_g016_pct_2024"]]
        out.to_excel(xw, sheet_name="FigA1_per_65plus_2024", index=False)
    print("saved FigureA1_data.xlsx")


def build_a2():
    coef = pd.read_csv(OUT / "FigureA2A_coefficients.csv")
    resid = pd.read_csv(OUT / "FigureA2B_residuals_2024.csv")
    labels = {"aging_rate": "Ageing rate (%)",
              "docs_per_100k": "Ophthalmologists per 100,000",
              "facilities_per_100k": "Facilities per 100,000"}
    coef["label"] = coef["variable"].map(labels)

    prs = B.new_prs()
    s = B.blank(prs)
    B.fig_label(s, "Fig. A2")
    B.textbox(s, LM, Inches(0.5), CW, Inches(0.3),
              "A. Two-way fixed-effects panel regression, total anti-VEGF, FY2014–2024", 11, True)
    w = pd.DataFrame({"Coefficient (vials per 100,000)": coef["coefficient"].to_numpy()},
                     index=coef["label"])
    ch = B.add_chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, LM, Inches(0.8), CW, Inches(3.2), w,
                     {"Coefficient (vials per 100,000)": "3182BD"},
                     ytitle="Coefficient (95% CI, prefecture-clustered SE)",
                     numfmt="0", legend=False, gap_width=120, font_size=8)
    add_error_bars(ch.plots[0].series[0],
                   plus=(coef["ci_upper"] - coef["coefficient"]).to_numpy(),
                   minus=(coef["coefficient"] - coef["ci_lower"]).to_numpy())
    B.textbox(s, LM, Inches(4.05), CW, Inches(0.35),
              "Whiskers: 95% CI (custom error bars). All CIs cross zero.", 7, color="595959")

    B.textbox(s, LM, Inches(4.5), CW, Inches(0.3),
              "B. Regression-adjusted prefecture residuals, FY2024", 11, True)
    r = resid.sort_values("residual", ascending=True)
    w = pd.DataFrame({"Residual per 100,000": r["residual"].to_numpy()}, index=r["prefecture_en"])
    ch = B.add_chart(s, XL_CHART_TYPE.BAR_CLUSTERED, LM, Inches(4.8), CW, Inches(5.6), w, None,
                     ytitle="Residual utilisation per 100,000, FY2024",
                     numfmt="0", legend=False, gap_width=30, font_size=6)
    ser = ch.plots[0].series[0]
    for i, v in enumerate(r["residual"]):
        pt = ser.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = B.rgb("B2182B" if v > 0 else "3182BD")
    B.textbox(s, LM, Inches(10.45), CW, Inches(0.45),
              "Residuals after removing prefecture and year fixed effects and covariates "
              "(ageing rate, ophthalmologists, facilities). Red = higher than adjusted expectation, "
              "blue = lower.", 7, color="595959")
    s.notes_slide.notes_text_frame.text = NOTES_A2
    prs.save(OUT / "FigureA2.pptx")
    print("saved FigureA2.pptx")

    with pd.ExcelWriter(OUT / "FigureA2_data.xlsx") as xw:
        coef[["variable", "label", "coefficient", "ci_lower", "ci_upper"]].to_excel(
            xw, sheet_name="FigA2A_coefficients", index=False)
        resid.sort_values("residual", ascending=False)[
            ["prefecture", "prefecture_en", "year", "residual"]].to_excel(
            xw, sheet_name="FigA2B_residuals_2024", index=False)
    print("saved FigureA2_data.xlsx")


if __name__ == "__main__":
    build_a1()
    build_a2()
    print("done")
