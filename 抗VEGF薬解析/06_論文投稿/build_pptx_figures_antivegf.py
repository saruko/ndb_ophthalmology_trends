# -*- coding: utf-8 -*-
"""
抗VEGF 論文 投稿用 pptx（Excel ネイティブグラフ・編集可能）の生成

`build_submission_files_antivegf.py` が書き出した `投稿用/Figure*_data.csv`
（図に描いた値そのもの）を読み、1 図 = 1 pptx（1 スライド、縦長）として
python-pptx のネイティブグラフで再構成する。PNG 版と同じ値・同じ panel 構成。

  Figure1.pptx 〜 Figure6.pptx
  SupplFigureS1.pptx   都道府県コロプレス地図（県ごとのフリーフォーム図形、2 スライド:
                       A 総人口分母 / B 65歳以上分母）。抗アレルギー点眼解析/data/japan_prefectures.geojson
  SupplFigureS2.pptx   G016 vs 抗VEGF

python-pptx は複合グラフ（第2軸）を作れないため、PNG 版で第2軸にした系列
（Fig 3B 薬価、Fig 4C 平均年齢、Fig 5B P90/P10）は同じ panel 内の別グラフとして並べる。
Fig 4A のヒートマップは値入りの表（セル塗り）として出力する。

実行:  .venv\\Scripts\\python.exe 抗VEGF薬解析/06_論文投稿/build_pptx_figures_antivegf.py
"""
from __future__ import annotations

import json
import math
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from matplotlib import colormaps
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_TICK_LABEL_POSITION
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt, Emu

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
OUT = HERE / "投稿用"
REPO = ROOT.parent
sys.path.insert(0, str(REPO / "抗アレルギー点眼解析"))
from pptx_bounds import add_error_bars  # noqa: E402  (allergy 側の補助モジュールを共用)

GEOJSON = REPO / "抗アレルギー点眼解析" / "data" / "japan_prefectures.geojson"
SRC = ROOT / "公費含まない" / "03_解析結果"

SLIDE_W, SLIDE_H = Inches(7.5), Inches(11)
LM = Inches(0.4)
CW = SLIDE_W - 2 * LM
FONT = "Arial"

PRODUCT_COLOR = {
    "Aflibercept 2 mg vial": "1F77B4", "Aflibercept 2 mg PFS": "6BAED6", "Aflibercept 8 mg vial": "C6DBEF",
    "Faricimab vial": "2CA02C", "Brolucizumab PFS": "9467BD",
    "Ranibizumab originator vial (2.3 mg/0.23 mL)": "A50F15", "Ranibizumab originator vial (10 mg/mL)": "D62728",
    "Ranibizumab originator PFS": "FB6A4A", "Ranibizumab biosimilar PFS": "FCBBA1", "Pegaptanib PFS": "7F7F7F",
    "All anti-VEGF agents (total)": "000000", "Total": "000000",
}
UNIT_COLOR = {
    "Aflibercept 2 mg": "1F77B4", "Aflibercept 8 mg": "AEC7E8", "Faricimab": "2CA02C",
    "Brolucizumab": "9467BD", "Ranibizumab originator": "D62728", "Ranibizumab biosimilar": "FF9896",
    "Pegaptanib": "7F7F7F",
}


def rgb(hexstr):
    return RGBColor.from_string(hexstr.upper())


def load(fig):
    return pd.read_csv(OUT / f"{fig}_data.csv")


def wide(df, panel, series_order=None):
    g = df[df["panel"] == panel]
    w = g.pivot_table(index="category", columns="series", values="value_plotted", aggfunc="first", sort=False)
    cats = list(dict.fromkeys(g["category"]))
    w = w.reindex(cats)
    if series_order:
        w = w[[s for s in series_order if s in w.columns]]
    return w


def new_prs():
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    return prs


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def textbox(slide, x, y, w, h, text, size=10, bold=False, color="000000", align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = FONT
        r.font.color.rgb = rgb(color)
    return tb


def fig_label(slide, text):
    textbox(slide, Inches(0.2), Inches(0.1), Inches(2), Inches(0.4), text, size=18)


def style_chart(chart, ytitle=None, xtitle=None, legend=True, legend_pos=XL_LEGEND_POSITION.TOP, numfmt=None,
                font_size=8, ymin=None, ymax=None):
    chart.font.size = Pt(font_size)
    chart.font.name = FONT
    chart.has_legend = legend
    if legend:
        chart.legend.position = legend_pos
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(font_size)
    va = chart.value_axis
    va.has_major_gridlines = False
    va.format.line.color.rgb = rgb("000000")
    va.tick_labels.font.size = Pt(font_size)
    if numfmt:
        va.tick_labels.number_format = numfmt
        va.tick_labels.number_format_is_linked = False
    if ymin is not None:
        va.minimum_scale = ymin
    if ymax is not None:
        va.maximum_scale = ymax
    if ytitle:
        va.has_title = True
        va.axis_title.text_frame.text = ytitle
        va.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(font_size + 1)
        va.axis_title.text_frame.paragraphs[0].runs[0].font.bold = True
    ca = chart.category_axis
    ca.format.line.color.rgb = rgb("000000")
    ca.tick_labels.font.size = Pt(font_size)
    ca.has_major_gridlines = False
    if xtitle:
        ca.has_title = True
        ca.axis_title.text_frame.text = xtitle
        ca.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(font_size + 1)
        ca.axis_title.text_frame.paragraphs[0].runs[0].font.bold = True


def color_series(chart, colors: dict, line=False, marker=True, dashed=()):
    for s in chart.plots[0].series:
        c = colors.get(s.name)
        if c is None:
            continue
        if line:
            s.format.line.color.rgb = rgb(c)
            s.format.line.width = Pt(1.5)
            s.smooth = False
            if s.name in dashed:
                s.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
            if marker:
                s.marker.style = 8  # circle
                s.marker.size = 5
                s.marker.format.fill.solid()
                s.marker.format.fill.fore_color.rgb = rgb(c)
                s.marker.format.line.color.rgb = rgb(c)
        else:
            s.format.fill.solid()
            s.format.fill.fore_color.rgb = rgb(c)
            s.format.line.color.rgb = rgb("FFFFFF")
            s.format.line.width = Pt(0.25)


def add_chart(slide, kind, x, y, w, h, wide_df, colors=None, title=None, line=False, gap_width=40, overlap=None,
              dashed=(), **style):
    cd = CategoryChartData()
    cd.categories = [str(c) for c in wide_df.index]
    for col in wide_df.columns:
        vals = [None if (v is None or (isinstance(v, float) and math.isnan(v))) else float(v) for v in wide_df[col]]
        cd.add_series(col, vals)
    gf = slide.shapes.add_chart(kind, x, y, w, h, cd)
    ch = gf.chart
    ch.has_title = bool(title)
    if title:
        ch.chart_title.text_frame.text = title
        r = ch.chart_title.text_frame.paragraphs[0].runs[0]
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.name = FONT
    style_chart(ch, **style)
    if colors:
        color_series(ch, colors, line=line, dashed=dashed)
    if not line:
        try:
            ch.plots[0].gap_width = gap_width
            if overlap is not None:
                ch.plots[0].overlap = overlap
        except Exception:
            pass
    return ch


def save(prs, name):
    prs.save(OUT / f"{name}.pptx")
    print("saved", name + ".pptx")


# ---------------------------------------------------------------------------
# Figure 1
# ---------------------------------------------------------------------------
def build_fig1():
    df = load("Figure1")
    prs = new_prs()
    s = blank(prs)
    fig_label(s, "Fig. 1")
    w = wide(df, "Fig1A_quantity_by_product_thousand_vials")
    total = w.pop("All anti-VEGF agents (total)")
    textbox(s, LM, Inches(0.5), CW, Inches(0.3), "A. Nationwide quantity of anti-VEGF intravitreal agents by product, FY2014–2024", 11, True)
    ch = add_chart(s, XL_CHART_TYPE.COLUMN_STACKED, LM, Inches(0.8), CW, Inches(4.6), w, PRODUCT_COLOR,
                   ytitle="Vials / prefilled syringes (thousands)", xtitle="Fiscal year", numfmt="#,##0", overlap=100)
    # total as note (line on stacked column needs combo chart; listed in data sheet)
    textbox(s, LM, Inches(5.4), CW, Inches(0.3),
            "Total (thousand vials): " + ", ".join(f"FY{c} {v:,.0f}" for c, v in total.items()), 7, color="595959")
    w = wide(df, "Fig1B_share_pct_by_unit")
    textbox(s, LM, Inches(5.8), CW, Inches(0.3), "B. Composition by agent (analysis unit), FY2014–2024", 11, True)
    add_chart(s, XL_CHART_TYPE.COLUMN_STACKED_100, LM, Inches(6.1), CW, Inches(4.6), w, UNIT_COLOR,
              ytitle="Share of anti-VEGF quantity (%)", xtitle="Fiscal year", numfmt="0", overlap=100,
              legend_pos=XL_LEGEND_POSITION.BOTTOM)
    save(prs, "Figure1")


# ---------------------------------------------------------------------------
# Figure 2
# ---------------------------------------------------------------------------
def build_fig2():
    df = load("Figure2")
    prs = new_prs()
    s = blank(prs)
    fig_label(s, "Fig. 2")
    w = wide(df, "Fig2A_formulation_share_pct", ["Vial", "Prefilled syringe"])
    textbox(s, LM, Inches(0.5), CW, Inches(0.3), "A. Formulation share of all anti-VEGF agents (%)", 11, True)
    ch = add_chart(s, XL_CHART_TYPE.COLUMN_STACKED_100, LM, Inches(0.8), CW, Inches(4.6), w,
                   {"Vial": "9ECAE1", "Prefilled syringe": "3182BD"},
                   ytitle="Share of all anti-VEGF quantity (%)", xtitle="Fiscal year", numfmt="0", overlap=100)
    pl = ch.plots[0]
    pl.has_data_labels = True
    pl.data_labels.number_format = "0.0"
    pl.data_labels.number_format_is_linked = False
    pl.data_labels.font.size = Pt(7)
    w = wide(df, "Fig2B_pfs_share_within_molecule_pct")
    textbox(s, LM, Inches(5.8), CW, Inches(0.3), "B. PFS share within agents marketed in both formulations (%)", 11, True)
    add_chart(s, XL_CHART_TYPE.LINE_MARKERS, LM, Inches(6.1), CW, Inches(4.6), w,
              {"Aflibercept 2 mg (PFS launched FY2020)": "1F77B4", "Ranibizumab originator": "D62728",
               "Ranibizumab all (originator + biosimilar)": "D62728", "All anti-VEGF agents": "000000"}, line=True,
              dashed=("Ranibizumab all (originator + biosimilar)", "All anti-VEGF agents"),
              ytitle="PFS share within the agent (%)", xtitle="Fiscal year", numfmt="0", ymin=0, ymax=100)
    save(prs, "Figure2")


# ---------------------------------------------------------------------------
# Figure 3
# ---------------------------------------------------------------------------
def build_fig3():
    df = load("Figure3")
    prs = new_prs()
    s = blank(prs)
    fig_label(s, "Fig. 3")
    w = wide(df, "Fig3A_ranibizumab_thousand_vials", ["Originator", "Biosimilar"])
    textbox(s, LM, Inches(0.5), CW, Inches(0.3), "A. Ranibizumab: originator versus biosimilar quantity", 11, True)
    add_chart(s, XL_CHART_TYPE.COLUMN_STACKED, LM, Inches(0.8), CW, Inches(4.3), w,
              {"Originator": "D62728", "Biosimilar": "FCBBA1"},
              ytitle="Vials / PFS (thousands)", xtitle="Fiscal year", numfmt="#,##0", overlap=100)
    wb = wide(df, "Fig3B_bs_share_and_prices")
    textbox(s, LM, Inches(5.3), CW, Inches(0.3), "B. Biosimilar share (left) and NHI drug prices of ranibizumab products (right)", 11, True)
    half = int((CW - Inches(0.2)) / 2)
    ch = add_chart(s, XL_CHART_TYPE.LINE_MARKERS, LM, Inches(5.6), half, Inches(4.6), wb[["Biosimilar share (%)"]],
                   {"Biosimilar share (%)": "B2182B"}, line=True, ytitle="Biosimilar share (%)", xtitle="Fiscal year",
                   numfmt="0", ymin=0, ymax=100, legend=False)
    pl = ch.plots[0]
    pl.has_data_labels = True
    pl.data_labels.number_format = '0.0"%"'
    pl.data_labels.number_format_is_linked = False
    pl.data_labels.font.size = Pt(7)
    pl.data_labels.position = 0  # above
    prices = wb[["Originator PFS price (thousand JPY)", "Originator vial price (thousand JPY)", "Biosimilar price (thousand JPY)"]]
    prices.columns = ["Originator PFS", "Originator vial", "Biosimilar"]
    add_chart(s, XL_CHART_TYPE.LINE_MARKERS, LM + half + Inches(0.2), Inches(5.6), half, Inches(4.6), prices,
              {"Originator PFS": "D62728", "Originator vial": "A50F15", "Biosimilar": "FB6A4A"}, line=True,
              ytitle="NHI drug price (thousand JPY per unit)", xtitle="Fiscal year", numfmt="0", ymin=0, ymax=200)
    save(prs, "Figure3")


# ---------------------------------------------------------------------------
# Figure 4
# ---------------------------------------------------------------------------
def build_fig4():
    df = load("Figure4")
    prs = new_prs()
    s = blank(prs)
    fig_label(s, "Fig. 4")
    # A heatmap as table
    w = wide(df, "Fig4A_age_share_pct_heatmap")  # index: series? -> need years as columns, age rows
    g = df[df["panel"] == "Fig4A_age_share_pct_heatmap"]
    ages = list(dict.fromkeys(g["series"]))
    years = list(dict.fromkeys(g["category"]))
    mat = g.pivot_table(index="series", columns="category", values="value_plotted").reindex(ages)[years]
    textbox(s, LM, Inches(0.5), CW, Inches(0.3), "A. Age distribution of anti-VEGF quantity by fiscal year (%, 90+ pooled)", 11, True)
    rows, cols = len(ages) + 1, len(years) + 1
    tbl = s.shapes.add_table(rows, cols, LM, Inches(0.85), CW, Inches(3.4)).table
    cmap = colormaps["YlOrRd"]
    vmax = float(np.nanmax(mat.to_numpy()))
    def setcell(c, text, fill=None, bold=False, fc="000000"):
        c.text = text if text else " "  # non-breaking space keeps the row height small
        p = c.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(6.5)
        for r in p.runs:
            r.font.size = Pt(6.5)
            r.font.bold = bold
            r.font.name = FONT
            r.font.color.rgb = rgb(fc)
        c.margin_left = c.margin_right = Emu(0)
        c.margin_top = c.margin_bottom = Emu(0)
        if fill is not None:
            c.fill.solid()
            c.fill.fore_color.rgb = fill
    setcell(tbl.cell(0, 0), "Age", fill=rgb("FFFFFF"), bold=True)
    for j, y in enumerate(years):
        setcell(tbl.cell(0, j + 1), str(y), fill=rgb("FFFFFF"), bold=True)
    for i, a in enumerate(reversed(ages)):
        setcell(tbl.cell(i + 1, 0), a, fill=rgb("FFFFFF"), bold=True)
        for j, y in enumerate(years):
            v = float(mat.loc[a, y])
            r_, g_, b_, _ = cmap(v / vmax)
            fill = RGBColor(int(r_ * 255), int(g_ * 255), int(b_ * 255))
            setcell(tbl.cell(i + 1, j + 1), f"{v:.0f}" if v >= 1 else "", fill=fill,
                    fc="FFFFFF" if v / vmax > 0.65 else "000000")
    tbl.columns[0].width = Inches(0.7)
    for j in range(1, cols):
        tbl.columns[j].width = int((CW - Inches(0.7)) / len(years))
    for i in range(rows):
        tbl.rows[i].height = Inches(0.16)
    # B pyramid
    g = df[df["panel"] == "Fig4B_pyramid_2024_thousand_vials"]
    agesB = list(dict.fromkeys(g["category"]))
    male = g[g["series"] == "Male"].set_index("category").reindex(agesB)
    fem = g[g["series"] == "Female"].set_index("category").reindex(agesB)
    wb = pd.DataFrame({"Male": -male["value_plotted"].to_numpy(), "Female": fem["value_plotted"].to_numpy()}, index=agesB)
    textbox(s, LM, Inches(4.4), CW, Inches(0.3), "B. Age–sex distribution, FY2024 (bars = zero-imputed; whiskers = upper bound of the identification interval)", 10, True)
    ch = add_chart(s, XL_CHART_TYPE.BAR_STACKED, LM, Inches(4.7), CW, Inches(3.6), wb,
                   {"Male": "3182BD", "Female": "E6550D"}, ytitle="Vials / PFS (thousands; male shown as negative)",
                   xtitle="Age group (years)", numfmt="#,##0;#,##0", overlap=100, gap_width=30)
    ch.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    ser_m, ser_f = ch.plots[0].series
    ser_m.invert_if_negative = False
    ser_f.invert_if_negative = False
    dm = (male["value_upper"] - male["value_plotted"]).clip(lower=0).fillna(0).to_numpy()
    dfm = (fem["value_upper"] - fem["value_plotted"]).clip(lower=0).fillna(0).to_numpy()
    add_error_bars(ser_m, plus=np.zeros(len(dm)), minus=dm, color="000000", width_pt=0.75)
    add_error_bars(ser_f, plus=dfm, minus=np.zeros(len(dfm)), color="000000", width_pt=0.75)
    # C trends (two charts)
    g = wide(df, "Fig4C_patient_profile_trends")
    textbox(s, LM, Inches(8.4), CW, Inches(0.3), "C. Ageing and sex composition, FY2014–2024", 11, True)
    half = int((CW - Inches(0.2)) / 2)
    add_chart(s, XL_CHART_TYPE.LINE_MARKERS, LM, Inches(8.7), half, Inches(2.2), g[["Share aged 75+ (%)", "Male share (%)"]],
              {"Share aged 75+ (%)": "E6550D", "Male share (%)": "3182BD"}, line=True, ytitle="Percent (%)",
              xtitle="Fiscal year", numfmt="0", ymin=40, ymax=70, font_size=7)
    add_chart(s, XL_CHART_TYPE.LINE_MARKERS, LM + half + Inches(0.2), Inches(8.7), half, Inches(2.2),
              g[["Approximate mean age (years)"]], {"Approximate mean age (years)": "000000"}, line=True,
              ytitle="Approximate mean age (years)", xtitle="Fiscal year", numfmt="0.0", ymin=72, ymax=78, font_size=7)
    save(prs, "Figure4")


# ---------------------------------------------------------------------------
# Figure 5
# ---------------------------------------------------------------------------
def build_fig5():
    df = load("Figure5")
    prs = new_prs()
    s = blank(prs)
    fig_label(s, "Fig. 5")
    g = df[df["panel"] == "Fig5A_prefecture_per_100k_2024"]
    prefs = list(dict.fromkeys(g["category"]))
    rate = g[g["series"].str.startswith("Rate")].set_index("category").reindex(prefs)["value_plotted"]
    cap = g[g["series"].str.startswith("Anti-VEGF")].set_index("category").reindex(prefs)["value_plotted"]
    w = pd.DataFrame({"Vials / PFS per 100,000 population": rate.to_numpy()}, index=prefs)
    textbox(s, LM, Inches(0.5), CW, Inches(0.3), "A. Prefecture-level utilisation per 100,000 population, FY2024", 11, True)
    ch = add_chart(s, XL_CHART_TYPE.BAR_CLUSTERED, LM, Inches(0.8), CW, Inches(6.9), w, None,
                   ytitle="Vials / PFS per 100,000 population, FY2024", numfmt="#,##0", legend=False, gap_width=30, font_size=7)
    ser = ch.plots[0].series[0]
    for i, p in enumerate(prefs):
        pt = ser.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = rgb("9E9E9E" if cap.iloc[i] < 90 else "3182BD")
    textbox(s, LM, Inches(7.7), CW, Inches(0.4),
            "Grey bars: prefecture breakdown captures <90% of G016 claims (under-estimated by cell suppression): "
            + ", ".join(p for p, c in zip(prefs, cap) if c < 90) + ". Dashed national value 860 per 100,000 (add as a line if needed).",
            7, color="595959")
    g = wide(df, "Fig5B_disparity_indices")
    textbox(s, LM, Inches(8.2), CW, Inches(0.3), "B. Inter-prefecture disparity in utilisation per 100,000, FY2014–2024", 11, True)
    half = int((CW - Inches(0.2)) / 2)
    cols = [c for c in ["Gini", "CV", "Gini (65+ denominator)"] if c in g.columns]
    add_chart(s, XL_CHART_TYPE.LINE_MARKERS, LM, Inches(8.5), half, Inches(2.4), g[cols],
              {"Gini": "B2182B", "CV": "3182BD", "Gini (65+ denominator)": "E08080"}, line=True, ytitle="Gini / CV",
              xtitle="Fiscal year", numfmt="0.00", ymin=0, ymax=0.5, font_size=7)
    add_chart(s, XL_CHART_TYPE.LINE_MARKERS, LM + half + Inches(0.2), Inches(8.5), half, Inches(2.4), g[["P90/P10"]],
              {"P90/P10": "000000"}, line=True, ytitle="P90/P10 ratio", xtitle="Fiscal year", numfmt="0.0", ymin=1, ymax=3.5,
              font_size=7)
    save(prs, "Figure5")


# ---------------------------------------------------------------------------
# Figure 6
# ---------------------------------------------------------------------------
def build_fig6():
    df = load("Figure6")
    prs = new_prs()
    s = blank(prs)
    fig_label(s, "Fig. 6")
    w = wide(df, "Fig6A_drug_cost_100M_jpy_by_product")
    total = w.pop("Total")
    textbox(s, LM, Inches(0.5), CW, Inches(0.3), "A. Estimated anti-VEGF drug cost (NHI price × quantity) by product", 11, True)
    add_chart(s, XL_CHART_TYPE.COLUMN_STACKED, LM, Inches(0.8), CW, Inches(4.6), w, PRODUCT_COLOR,
              ytitle="Estimated drug cost (100 million JPY)", xtitle="Fiscal year", numfmt="#,##0", overlap=100)
    textbox(s, LM, Inches(5.4), CW, Inches(0.3),
            "Total (100 million JPY): " + ", ".join(f"FY{c} {v:,.0f}" for c, v in total.items()), 7, color="595959")
    w = wide(df, "Fig6B_cost_decomposition_100M_jpy", ["Volume effect", "Product-mix effect", "Price-revision effect", "Net change"])
    net = w.pop("Net change")
    textbox(s, LM, Inches(5.8), CW, Inches(0.3), "B. Decomposition of the year-on-year change in drug cost", 11, True)
    add_chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, LM, Inches(6.1), CW, Inches(4.3), w,
              {"Volume effect": "3182BD", "Product-mix effect": "FD8D3C", "Price-revision effect": "756BB1"},
              ytitle="Change from previous year (100 million JPY)", xtitle="Fiscal year", numfmt="#,##0;-#,##0", gap_width=60, overlap=0)
    textbox(s, LM, Inches(10.4), CW, Inches(0.4),
            "Net change (100 million JPY): " + ", ".join(f"FY{c} {v:+,.0f}" for c, v in net.items()), 7, color="595959")
    save(prs, "Figure6")


# ---------------------------------------------------------------------------
# Suppl Fig S2
# ---------------------------------------------------------------------------
def build_s2():
    d = pd.read_csv(OUT / "SupplFigureS2_data.csv").set_index("year")
    prs = new_prs()
    s = blank(prs)
    fig_label(s, "Suppl Fig. S2")
    w = pd.DataFrame({"G016 intravitreal injection claims": d["g016_procedures"] / 1000,
                      "Anti-VEGF vials / PFS": d["antivegf_vials"] / 1000})
    textbox(s, LM, Inches(0.5), CW, Inches(0.3), "Anti-VEGF quantity versus G016 claims (published national totals)", 11, True)
    add_chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, LM, Inches(0.8), CW, Inches(4.2), w,
              {"G016 intravitreal injection claims": "BDBDBD", "Anti-VEGF vials / PFS": "3182BD"},
              ytitle="Thousands", xtitle="Fiscal year", numfmt="#,##0", gap_width=60, overlap=-10)
    add_chart(s, XL_CHART_TYPE.LINE_MARKERS, LM, Inches(5.2), CW, Inches(3.0),
              pd.DataFrame({"Anti-VEGF vials per G016 claim (%)": d["antivegf_per_g016_pct"]}),
              {"Anti-VEGF vials per G016 claim (%)": "000000"}, line=True,
              ytitle="Anti-VEGF vials per G016 claim (%)", xtitle="Fiscal year", numfmt="0.0", ymin=90, ymax=101, legend=False)
    save(prs, "SupplFigureS2")


# ---------------------------------------------------------------------------
# Suppl Fig S1: choropleth as editable freeform shapes
# ---------------------------------------------------------------------------
PREF_EN = {
    "北海道": "Hokkaido", "青森県": "Aomori", "岩手県": "Iwate", "宮城県": "Miyagi", "秋田県": "Akita", "山形県": "Yamagata",
    "福島県": "Fukushima", "茨城県": "Ibaraki", "栃木県": "Tochigi", "群馬県": "Gunma", "埼玉県": "Saitama", "千葉県": "Chiba",
    "東京都": "Tokyo", "神奈川県": "Kanagawa", "新潟県": "Niigata", "富山県": "Toyama", "石川県": "Ishikawa", "福井県": "Fukui",
    "山梨県": "Yamanashi", "長野県": "Nagano", "岐阜県": "Gifu", "静岡県": "Shizuoka", "愛知県": "Aichi", "三重県": "Mie",
    "滋賀県": "Shiga", "京都府": "Kyoto", "大阪府": "Osaka", "兵庫県": "Hyogo", "奈良県": "Nara", "和歌山県": "Wakayama",
    "鳥取県": "Tottori", "島根県": "Shimane", "岡山県": "Okayama", "広島県": "Hiroshima", "山口県": "Yamaguchi",
    "徳島県": "Tokushima", "香川県": "Kagawa", "愛媛県": "Ehime", "高知県": "Kochi", "福岡県": "Fukuoka", "佐賀県": "Saga",
    "長崎県": "Nagasaki", "熊本県": "Kumamoto", "大分県": "Oita", "宮崎県": "Miyazaki", "鹿児島県": "Kagoshima", "沖縄県": "Okinawa",
}


def ring_area(ring):
    x = np.array([p[0] for p in ring]); y = np.array([p[1] for p in ring])
    return 0.5 * abs(np.dot(x, np.roll(y, 1)) - np.dot(y, np.roll(x, 1)))


def load_polys():
    gj = json.load(open(GEOJSON, encoding="utf-8"))
    out = {}
    for feat in gj["features"]:
        name = feat["properties"]["nam_ja"]
        geom = feat["geometry"]
        polys = [geom["coordinates"]] if geom["type"] == "Polygon" else geom["coordinates"]
        rings = [poly[0] for poly in polys]  # outer rings only
        out[name] = rings
    return out


def build_s1():
    polys = load_polys()
    pref_rank = pd.read_csv(SRC / "都道府県_地域格差/prefecture_per_capita_ranking_antivegf.csv")
    pref_65 = pd.read_csv(SRC / "都道府県_地域格差/prefecture_per_65plus_ranking_antivegf.csv")
    vals_total = pref_rank.set_index("prefecture")["2024年度_人口10万対"]
    vals_65 = pref_65.set_index("prefecture")["2024年度_65歳以上10万対"]
    rank_total = pref_rank.set_index("prefecture")["2024年度_順位"]
    rank_65 = pref_65.set_index("prefecture")["2024年度_順位"]

    # projection: shift Okinawa to the upper-left inset, simple equirectangular with cos(36deg)
    def proj(lon, lat, name):
        if name == "沖縄県":
            lon, lat = lon + 4.0, lat + 12.0
        return lon * math.cos(math.radians(36)), lat

    # extent: main islands only (drop remote islands south of 30N / east of 146E and tiny rings)
    all_pts = [proj(x, y, n) for n, rings in polys.items() for r in rings
               if (ring_area(r) >= 0.004 or r is max(rings, key=ring_area))
               for x, y in r if (n == "沖縄県" or (y >= 30 and x <= 146))]
    xs = [p[0] for p in all_pts]; ys = [p[1] for p in all_pts]
    x0, x1, y0, y1 = min(xs), max(xs), min(ys), max(ys)
    map_w = CW
    map_h = Inches(8.2)
    scale = min(map_w / (x1 - x0), map_h / (y1 - y0))
    ox, oy = LM + Inches(0.0), Inches(1.2)

    def to_emu(x, y):
        return int(ox + (x - x0) * scale), int(oy + (y1 - y) * scale)

    cmap = colormaps["YlOrRd"]
    min_area = 0.004  # deg^2: drop tiny islands to keep shape count manageable

    prs = new_prs()
    for title, vals, ranks, unit in [
        ("A. Anti-VEGF vials / PFS per 100,000 population, FY2024", vals_total, rank_total, "per 100,000 population"),
        ("B. Anti-VEGF vials / PFS per 100,000 population aged 65+, FY2024", vals_65, rank_65, "per 100,000 aged 65+"),
    ]:
        s = blank(prs)
        fig_label(s, "Suppl Fig. S1")
        textbox(s, LM, Inches(0.5), CW, Inches(0.3), title, 11, True)
        vmin, vmax = float(vals.min()), float(vals.max())
        for name, rings in polys.items():
            v = float(vals[name])
            r_, g_, b_, _ = cmap((v - vmin) / (vmax - vmin))
            fill = RGBColor(int(r_ * 255), int(g_ * 255), int(b_ * 255))
            rings_sorted = sorted(rings, key=ring_area, reverse=True)
            shapes = []
            for ring in rings_sorted:
                if ring_area(ring) < min_area and ring is not rings_sorted[0]:
                    continue
                if name != "沖縄県" and (max(y for _, y in ring) < 30 or min(x for x, _ in ring) > 146):
                    continue  # remote islands outside the map extent
                pts = [to_emu(*proj(x, y, name)) for x, y in ring]
                fb = s.shapes.build_freeform(pts[0][0], pts[0][1], scale=1.0)
                fb.add_line_segments(pts[1:], close=True)
                shp = fb.convert_to_shape()
                shp.fill.solid()
                shp.fill.fore_color.rgb = fill
                shp.line.color.rgb = rgb("FFFFFF")
                shp.line.width = Pt(0.5)
                shp.name = f"{PREF_EN[name]} ({name})"
                shapes.append(shp)
            # label at centroid of the largest ring
            big = rings_sorted[0]
            cx = np.mean([proj(x, y, name)[0] for x, y in big]); cy = np.mean([proj(x, y, name)[1] for x, y in big])
            ex, ey = to_emu(cx, cy)
            lab = textbox(s, ex - Inches(0.45), ey - Inches(0.12), Inches(0.9), Inches(0.25),
                          f"{PREF_EN[name]}\n{v:,.0f} (#{int(ranks[name])})", 5, color="000000", align=PP_ALIGN.CENTER)
            lab.name = f"label {PREF_EN[name]}"
        # Okinawa inset box
        ok_rings = [r for r in polys["沖縄県"] if ring_area(r) >= min_area]
        okx = [proj(x, y, "沖縄県")[0] for r in ok_rings for x, y in r]
        oky = [proj(x, y, "沖縄県")[1] for r in ok_rings for x, y in r]
        bx0, by0 = to_emu(min(okx) - 0.3, max(oky) + 0.3)
        bx1, by1 = to_emu(max(okx) + 0.3, min(oky) - 0.3)
        box = s.shapes.add_shape(1, bx0, by0, bx1 - bx0, by1 - by0)
        box.fill.background()
        box.line.color.rgb = rgb("808080")
        box.line.width = Pt(0.5)
        box.name = "Okinawa inset"
        # colour bar
        n = 6
        cbx, cby, cbw, cbh = LM + Inches(4.6), Inches(7.9), Inches(0.35), Inches(0.3)
        textbox(s, cbx, cby - Inches(0.45), Inches(2.1), Inches(0.45), f"Vials / PFS {unit}\n(colour = zero-imputed value)", 7, True)
        for i in range(n):
            lo = vmin + (vmax - vmin) * i / n
            hi = vmin + (vmax - vmin) * (i + 1) / n
            r_, g_, b_, _ = cmap((i + 0.5) / n)
            rect = s.shapes.add_shape(1, cbx, cby + cbh * i, cbw, cbh)
            rect.fill.solid(); rect.fill.fore_color.rgb = RGBColor(int(r_ * 255), int(g_ * 255), int(b_ * 255))
            rect.line.color.rgb = rgb("FFFFFF")
            textbox(s, cbx + cbw + Inches(0.05), cby + cbh * i, Inches(2), cbh, f"{lo:,.0f} – {hi:,.0f}", 7)
        textbox(s, LM, Inches(10.1), CW, Inches(0.6),
                "Colour = zero-imputed prefecture breakdown (lower bound). Labels: prefecture, value (#rank). "
                "Okinawa is shown in the inset box. Six prefectures (Yamanashi, Kochi, Tottori, Fukui, Tokushima, Iwate) "
                "capture <90% of G016 claims and are under-estimated (see Fig 5A, Suppl Table S3/S7). "
                "Geometry: dataofjapan/land (public domain), small islands omitted.", 7, color="595959")
    save(prs, "SupplFigureS1")


if __name__ == "__main__":
    build_fig1(); build_fig2(); build_fig3(); build_fig4(); build_fig5(); build_fig6()
    build_s2(); build_s1()
    print("done")
