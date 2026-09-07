"""論文用の Table / Figure 成果物（CSV・Excel）を 05_論文成果物/ に生成する。

allergy解析/05_論文成果物/ と同じ考え方で、原稿にそのまま載せられる表と、
作図に使えるワイド形式のデータシート（Excelグラフ埋め込み済み）を出力する。

    python 緑内障点眼解析/build_paper_outputs.py

前提: run_glaucoma_pipeline.py / build_product_inventory.py / verify_coverage.py /
      run_paper_analyses.py / run_censoring_sensitivity.py の実行済み出力。

出力先: 緑内障点眼解析/05_論文成果物/
  Table1_drug_summary_2024.csv          主解析期間の代表年（2024年度）の薬剤・薬効群別サマリー
  Table2_group_change_2022_2024.csv     薬効群の2022→2024変化
  Table3_mf_ratio_2024.csv              年齢群別の男女比
  Table4_prefecture_ranking.csv         都道府県ランキング（2022–2024平均）
  Table5_balanced_panel_apc.csv         均衡パネルの長期APC
  Table6_within_agent_share_2024.csv    同一成分内の品目別シェア
  Table7_age_adjusted_prefecture_2024.csv 年齢調整した都道府県比較
  Tables_1to7.xlsx                      Table1〜7を1ブックにまとめたもの
  Fig1_age_sex_profile.xlsx             年齢×性別プロファイル
  Fig2_group_comparison.xlsx            薬効群の量・シェア（2022–2024）
  Fig3_prefecture.xlsx                  都道府県分布
  Fig4_balanced_panel_trends.xlsx       均衡パネルの長期トレンド
  Fig5_brand_generic.xlsx               先発品／後発品の推移
  Fig6_ingredient_exposure.xlsx         単剤 vs 成分ベース曝露
  Fig7_within_agent_share.xlsx          同一成分内の品目別シェア
  Fig8_age_adjustment.xlsx              粗率 vs 年齢調整率
  Fig9_group_share_trend.xlsx           薬効群シェアの年次推移（2014-2024）
  SupplTable1_annual_totals_by_group.xlsx  年度×薬効群の全国数量（NR表記つき）
  SupplTable2_product_master.xlsx          対象品目マスタ（全品目）
  SupplTable3_censoring_coverage.xlsx      秘匿・収載カバレッジ
  SupplTable4_regressions.xlsx             パネル回帰・代替性回帰・収束分析
  SupplTable5_censoring_sensitivity.xlsx   秘匿セルの識別区間（zero / upper）
  SupplTable6_agent_series.xlsx            成分別の年次系列と追跡可能性
"""

import argparse
import os
import sys

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.marker import Marker
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Alignment, Font
from openpyxl.utils.dataframe import dataframe_to_rows

BASE = os.path.dirname(os.path.abspath(__file__))
sys.path.append(os.path.join(BASE, "src"))

from labels import (DRUG_EN, DRUG_TYPE_EN, GROUP_EN, INGREDIENT_EN,  # noqa: E402
                    PREFECTURE_EN, SUBGROUP_EN, en)
from paths import add_nokouhi_arg, data_root, find  # noqa: E402
from preprocess_glaucoma import GROUP_DEFS, SUBGROUP_DEFS  # noqa: E402

PAPER_SUBDIR = "05_論文成果物"
FULL_COV_START = 2022
LATEST = 2024
GROUP_ORDER = [c for c, _, _ in GROUP_DEFS]
SUBGROUP_ORDER = [c for c, _, _ in SUBGROUP_DEFS]
NOTE_COVERAGE = (
    "NDB Open Data expanded the listed products of topical drugs from the 10th release "
    "(FY2022). Volumes before FY2022 are therefore not comparable with FY2022 onwards; "
    "the primary between-drug comparison is restricted to FY2022-2024."
)
NOTE_GAP = (
    "FY2021 is understated: nine products are missing from the listing that year "
    "although the same products are listed before and after (Suppl Table 3E). The "
    "estimated shortfall is 14.8 million mL, 7.5% of the annual total, concentrated "
    "in carbonic anhydrase inhibitors (4.6), beta-blockers (4.5), prostanoid "
    "receptor-related drugs (3.7) and fixed-dose combinations (2.0). The FY2021 dip "
    "is an artefact, not a fall in prescribing. FY2017-2020 are understated by "
    "0.6-1.9% for the same reason."
)
NOTE_UNIT = (
    "Volumes are dispensed quantities converted to millilitres. Nine products are "
    "published in container units (bottle or piece) and were converted using the volume "
    "stated in the product name; all other products are published in millilitres."
)
# 系列色。明度と色相を離した12色を順に割り当て、13系列目以降は循環する。
SERIES_COLORS = [
    "1F77B4", "D62728", "2CA02C", "FF7F0E", "9467BD", "8C564B",
    "17BECF", "E377C2", "7F7F7F", "BCBD22", "393B79", "8C6D31",
]


# ─────────────────────────────────────────────
# Excel 書き出しの共通ヘルパー
# ─────────────────────────────────────────────

def _write_df(ws, df, start_row=1, index=False, number_format=None):
    """DataFrame をワークシートへ書き出し、次に空く行番号を返す。"""
    rows = dataframe_to_rows(df, index=index, header=True)
    r = start_row
    for i, row in enumerate(rows):
        if index and i == 1 and all(v is None for v in row):
            continue          # dataframe_to_rows が挟む index 名の空行
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            if r == start_row:
                cell.font = Font(bold=True)
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            elif number_format and isinstance(v, (int, float)):
                cell.number_format = number_format
        r += 1
    for col_cells in ws.iter_cols(min_row=start_row, max_row=start_row):
        for cell in col_cells:
            width = max(12, min(38, len(str(cell.value or "")) + 2))
            ws.column_dimensions[cell.column_letter].width = width
    return r


def _titled_sheet(wb, name, title, notes=()):
    ws = wb.create_sheet(name) if wb.sheetnames != ["Sheet"] else wb.active
    if ws.title == "Sheet":
        ws.title = name
    ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=12)
    r = 2
    for n in notes:
        ws.cell(row=r, column=1, value=n).font = Font(size=9, italic=True)
        r += 1
    return ws, r + 1


def _style_chart(chart, kind="line"):
    """全グラフ共通の体裁。折れ線は直線・丸マーカー（サイズ5）、系列色は固定パレット。"""
    for i, s in enumerate(chart.series):
        color = SERIES_COLORS[i % len(SERIES_COLORS)]
        if kind == "line":
            s.smooth = False
            s.graphicalProperties = GraphicalProperties(
                ln=LineProperties(solidFill=color, w=22000))
            s.marker = Marker(
                symbol="circle", size=5,
                spPr=GraphicalProperties(solidFill=color,
                                         ln=LineProperties(solidFill=color)))
        else:
            s.graphicalProperties = GraphicalProperties(
                solidFill=color, ln=LineProperties(solidFill="FFFFFF", w=6350))
    return chart


def _add_chart(ws, kind, df_rows, df_cols, anchor, title, y_title, x_title,
               data_start_row, stacked=False):
    """データ範囲からグラフを作り、シートに貼る。"""
    chart = LineChart() if kind == "line" else BarChart()
    if kind == "bar":
        chart.type = "col"
        if stacked:
            chart.grouping = "stacked"
            chart.overlap = 100
    chart.title = title
    chart.y_axis.title = y_title
    chart.x_axis.title = x_title
    chart.height, chart.width = 9, 18
    data = Reference(ws, min_col=2, max_col=df_cols,
                     min_row=data_start_row, max_row=data_start_row + df_rows)
    cats = Reference(ws, min_col=1, min_row=data_start_row + 1,
                     max_row=data_start_row + df_rows)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    _style_chart(chart, kind)
    ws.add_chart(chart, anchor)
    return chart


def _save(wb, out_dir, name):
    if "Sheet" in wb.sheetnames and wb["Sheet"].max_row == 1 and wb["Sheet"].max_column == 1:
        del wb["Sheet"]
    path = os.path.join(out_dir, name)
    wb.save(path)
    print(f"  wrote {name}")
    return path


# ─────────────────────────────────────────────
# Tables
# ─────────────────────────────────────────────

def build_table1(src, out_dir):
    """Table 1: 2024年度の薬効群別・薬剤別サマリー（量・人口10万対・シェア・後発品率）。"""
    summary = pd.read_csv(src["fullcov_summary"])
    inv = pd.read_csv(src["inventory"])
    s = summary[summary["year"] == LATEST].copy()

    ge = inv[(inv["year"] == LATEST) & (inv["group_code"] != "EXCLUDED")].copy()
    ge["kind"] = np.where(ge["drug_type"].str.startswith("後発"), "GE", "BR")
    by_code = ge.groupby(["code", "kind"])["quantity_ml"].sum().unstack(fill_value=0)
    for c in ("GE", "BR"):
        if c not in by_code:
            by_code[c] = 0.0
    by_code["generic_share_pct"] = (
        by_code["GE"] / (by_code["GE"] + by_code["BR"]) * 100)
    n_products = ge.groupby("code")["product_name"].nunique()

    # 薬効群・全体の行にも、構成する薬剤の品目数・後発品率を集計して入れる
    ge_g = ge.groupby(["group_code", "kind"])["quantity_ml"].sum().unstack(fill_value=0)
    for c in ("GE", "BR"):
        if c not in ge_g:
            ge_g[c] = 0.0
    group_ge = ge_g["GE"] / (ge_g["GE"] + ge_g["BR"]) * 100
    group_np = ge.groupby("group_code")["product_name"].nunique()
    total_ge = ge["quantity_ml"][ge["kind"] == "GE"].sum() / ge["quantity_ml"].sum() * 100
    total_np = ge["product_name"].nunique()

    def _ge(code):
        if code == "GLAUCOMA_EYE_TOTAL":
            return total_ge
        return group_ge.get(code, by_code["generic_share_pct"].get(code))

    def _np(code):
        if code == "GLAUCOMA_EYE_TOTAL":
            return total_np
        return group_np.get(code, n_products.get(code))

    s["generic_share_pct"] = s["code"].map(_ge)
    s["n_products"] = s["code"].map(_np)
    s["level"] = np.where(s["code"] == "GLAUCOMA_EYE_TOTAL", "Total",
                          np.where(s["code"].isin(GROUP_ORDER), "Drug class", "Agent"))
    s["name_en"] = s["code"].map(lambda c: en(c))
    s["share_pct"] = s["share"] * 100

    order = {"Total": 0, "Drug class": 1, "Agent": 2}
    s["_o"] = s["level"].map(order)
    s = s.sort_values(["_o", "count_ml"], ascending=[True, False]).drop(columns="_o")

    cols = ["level", "code", "name_en", "procedure_name", "n_products",
            "count_ml", "count_per_100k", "share_pct", "generic_share_pct",
            "cv", "gini", "max_to_min_ratio", "min_prefecture", "max_prefecture"]
    t1 = s[cols].rename(columns={
        "level": "Level", "code": "Code", "name_en": "Drug / class (EN)",
        "procedure_name": "薬剤・薬効群", "n_products": "Products (n)",
        "count_ml": "Volume (mL)", "count_per_100k": "Per 100,000 population (mL)",
        "share_pct": "Share (%)", "generic_share_pct": "Generic share (%)",
        "cv": "CV", "gini": "Gini", "max_to_min_ratio": "Max/min ratio",
        "min_prefecture": "Lowest prefecture", "max_prefecture": "Highest prefecture"})
    t1.to_csv(os.path.join(out_dir, "Table1_drug_summary_2024.csv"),
              index=False, encoding="utf-8-sig")
    return t1


def build_table2(src, out_dir):
    """Table 2: 薬効群の2022→2024変化。"""
    ch = pd.read_csv(src["fullcov_change"])
    ch = ch[ch["code"].isin(GROUP_ORDER + ["GLAUCOMA_EYE_TOTAL"])].copy()
    ch["name_en"] = ch["code"].map(lambda c: en(c))
    ch["_o"] = ch["code"].map(
        {c: i for i, c in enumerate(["GLAUCOMA_EYE_TOTAL"] + GROUP_ORDER)})
    ch = ch.sort_values("_o").drop(columns="_o")
    ch["share_start_pct"] = ch["share_start"] * 100
    ch["share_end_pct"] = ch["share_end"] * 100
    t2 = ch[["code", "name_en", "procedure_name",
             "count_per_100k_start", "count_per_100k_end",
             "share_start_pct", "share_end_pct", "share_change_pt",
             "pct_change", "cagr_pct"]].rename(columns={
        "code": "Code", "name_en": "Drug class (EN)", "procedure_name": "薬効群",
        "count_per_100k_start": "FY2022 per 100,000 (mL)",
        "count_per_100k_end": "FY2024 per 100,000 (mL)",
        "share_start_pct": "FY2022 share (%)", "share_end_pct": "FY2024 share (%)",
        "share_change_pt": "Share change (pt)",
        "pct_change": "Change FY2022-2024 (%)",
        "cagr_pct": "Annual change (%/year)"})
    t2.to_csv(os.path.join(out_dir, "Table2_group_change_2022_2024.csv"),
              index=False, encoding="utf-8-sig")
    return t2


def build_table3(src, out_dir):
    """Table 3: 年齢群別の男女比（薬効群および全体）。"""
    mf = pd.read_csv(src["mf_ratio"])
    mf = mf[(mf["year"] == LATEST) &
            (mf["code"].isin(GROUP_ORDER + ["GLAUCOMA_EYE_TOTAL"]))].copy()
    mf["name_en"] = mf["code"].map(lambda c: en(c))
    t3 = mf[["code", "name_en", "procedure_name", "age_group",
             "count_ml_male", "count_ml_female",
             "count_per_100k_male", "count_per_100k_female",
             "mf_ratio_count", "mf_ratio_rate"]].rename(columns={
        "code": "Code", "name_en": "Drug class (EN)", "procedure_name": "薬効群",
        "age_group": "Age group",
        "count_ml_male": "Volume, male (mL)", "count_ml_female": "Volume, female (mL)",
        "count_per_100k_male": "Per 100,000, male (mL)",
        "count_per_100k_female": "Per 100,000, female (mL)",
        "mf_ratio_count": "M:F ratio (volume)", "mf_ratio_rate": "M:F ratio (rate)"})
    t3.to_csv(os.path.join(out_dir, "Table3_mf_ratio_2024.csv"),
              index=False, encoding="utf-8-sig")
    return t3


def build_table4(src, out_dir):
    """Table 4: 都道府県ランキング（2022–2024平均）。"""
    pr = pd.read_csv(src["fullcov_pref"])
    pr["name_en"] = pr["code"].map(lambda c: en(c))
    pr["prefecture_en"] = pr["prefecture"].map(PREFECTURE_EN)
    t4 = pr[["code", "name_en", "procedure_name", "rank",
             "prefecture", "prefecture_en", "count_per_100k", "count_ml"]].rename(
        columns={"code": "Code", "name_en": "Drug class (EN)",
                 "procedure_name": "薬効群", "rank": "Rank",
                 "prefecture": "都道府県", "prefecture_en": "Prefecture",
                 "count_per_100k": "Per 100,000 population (mL), FY2022-2024 mean",
                 "count_ml": "Volume (mL), FY2022-2024 total"})
    t4.to_csv(os.path.join(out_dir, "Table4_prefecture_ranking.csv"),
              index=False, encoding="utf-8-sig")
    return t4


def build_table5(src, out_dir):
    """Table 5: 均衡パネルの長期APC。"""
    apc = pd.read_csv(src["panel_apc"])
    prods = pd.read_csv(src["panel_products"])
    lab = prods.drop_duplicates(["panel", "product_name"]).set_index(
        ["panel", "product_name"])[["category_name", "drug_type"]]
    apc = apc[apc["level"].isin(["product", "panel_total"])].copy()
    apc = apc.join(lab, on=["panel", "series"])
    apc["drug_type_en"] = apc["drug_type"].map(DRUG_TYPE_EN)
    t5 = apc[["panel", "level", "series", "category_name", "drug_type_en",
              "start_year", "end_year", "n_years",
              "apc", "apc_low", "apc_high", "p_value", "r2"]].rename(columns={
        "panel": "Panel", "level": "Level", "series": "品目 / 系列",
        "category_name": "成分", "drug_type_en": "Type",
        "start_year": "From", "end_year": "To", "n_years": "Years (n)",
        "apc": "APC (%/year)", "apc_low": "95%CI lower", "apc_high": "95%CI upper",
        "p_value": "p", "r2": "R2"})
    t5.to_csv(os.path.join(out_dir, "Table5_balanced_panel_apc.csv"),
              index=False, encoding="utf-8-sig")
    return t5


def build_table6(src, out_dir):
    """Table 6: 同一成分内の品目別シェア（先発品 vs 後発品の分散度）。"""
    s = pd.read_csv(src["within_agent_summary"])
    s = s[s["year"] == LATEST].copy()
    s["Agent (EN)"] = s["code"].map(lambda c: DRUG_EN.get(c, c))
    s = s.sort_values("agent_total_ml", ascending=False)
    s["originator_share_pct"] = s["originator_share"] * 100
    s["generic_share_pct"] = s["generic_share"] * 100
    s["top_product_share_pct"] = s["top_product_share"] * 100
    s["top3_share_pct"] = s["top3_share"] * 100
    t6 = s[["Agent (EN)", "category_name", "n_products", "n_generic_products",
            "agent_total_ml", "originator_share_pct", "generic_share_pct",
            "top_product", "top_product_share_pct", "top3_share_pct",
            "hhi_within_agent"]].rename(columns={
        "category_name": "成分", "n_products": "Products (n)",
        "n_generic_products": "Generic products (n)",
        "agent_total_ml": "Volume (mL)",
        "originator_share_pct": "Originator share (%)",
        "generic_share_pct": "Generic share (%)",
        "top_product": "Leading product",
        "top_product_share_pct": "Leading product share (%)",
        "top3_share_pct": "Top-3 share (%)",
        "hhi_within_agent": "HHI within agent"})
    t6.to_csv(os.path.join(out_dir, "Table6_within_agent_share_2024.csv"),
              index=False, encoding="utf-8-sig")
    return t6


def build_table7(src, out_dir):
    """Table 7: 年齢調整した都道府県比較（間接標準化）。"""
    a = pd.read_csv(src["age_adjusted_pref"])
    a = a[(a["year"] == LATEST) & (a["code"] == "GLAUCOMA_EYE_TOTAL")].copy()
    a["Prefecture"] = a["prefecture"].map(PREFECTURE_EN)
    a["aging_rate_pct"] = a["aging_rate"] * 100
    a = a.sort_values("age_adjusted_rank")
    t7 = a[["age_adjusted_rank", "Prefecture", "prefecture", "aging_rate_pct",
            "crude_rate_per_100k", "crude_rank",
            "age_adjusted_rate_per_100k", "spr", "rank_change"]].rename(columns={
        "age_adjusted_rank": "Rank (age-adjusted)", "prefecture": "都道府県",
        "aging_rate_pct": "Population aged 65+ (%)",
        "crude_rate_per_100k": "Crude rate (mL per 100,000)",
        "crude_rank": "Rank (crude)",
        "age_adjusted_rate_per_100k": "Age-adjusted rate (mL per 100,000)",
        "spr": "SPR", "rank_change": "Rank change (crude - adjusted)"})
    t7.to_csv(os.path.join(out_dir, "Table7_age_adjusted_prefecture_2024.csv"),
              index=False, encoding="utf-8-sig")
    return t7


def build_tables_workbook(tables, out_dir):
    wb = Workbook()
    titles = {
        "Table1": ("Table 1. Prescribed volume of glaucoma eye drops in Japan, FY2024",
                   (NOTE_UNIT, NOTE_COVERAGE,
                    "Volumes are the sum of prefecture cells; suppressed cells (<1,000) "
                    "were imputed as zero.")),
        "Table2": ("Table 2. Change in prescribed volume and share by drug class, FY2022-2024",
                   (NOTE_COVERAGE,)),
        "Table3": ("Table 3. Male-to-female ratio by age group, FY2024", ()),
        "Table4": ("Table 4. Prefecture ranking of prescribed volume, FY2022-2024 mean", ()),
        "Table5": ("Table 5. Long-term annual percent change in the balanced product panels",
                   ("Panel A: products listed in all 11 fiscal years (FY2014-2024). "
                    "Panel B: products listed in all 10 fiscal years (FY2015-2024).",
                    "Volumes are the published national totals of each product. "
                    "Declines reflect substitution away from these specific products, "
                    "not a decline in treatment of the molecule as a whole.")),
        "Table6": ("Table 6. Distribution of volume across products within each agent, FY2024",
                   ("Share of each agent's total volume held by the leading product and "
                    "by originator versus generic products.",
                    "HHI = sum of squared product shares within the agent; 1.0 means a "
                    "single product accounts for all volume.",
                    "Products split into a regular and a '（選）' row from October 2024 "
                    "(selective medical care for long-listed originators) were merged.")),
        "Table7": ("Table 7. Age-adjusted prefecture comparison of all glaucoma eye drops, "
                   "FY2024",
                   ("Indirect standardisation with two age strata (0-64 and 65 years and "
                    "over), because NDB Open Data does not publish prefecture-level age "
                    "breakdowns.",
                    "SPR = observed volume / expected volume; 1.0 equals the national "
                    "average. Age-adjusted rate = SPR x national crude rate.",
                    "A positive rank change means the prefecture ranks higher after "
                    "age adjustment.")),
    }
    for key, df in tables.items():
        title, notes = titles[key]
        ws, r = _titled_sheet(wb, key, title, notes)
        _write_df(ws, df, start_row=r)
    _save(wb, out_dir, "Tables_1to7.xlsx")


# ─────────────────────────────────────────────
# Figures
# ─────────────────────────────────────────────

def build_fig1(src, out_dir):
    """Fig 1: 年齢×性別の人口10万対処方量（2024年度）。"""
    rates = pd.read_csv(src["age_sex"])
    r = rates[(rates["year"] == LATEST) & (rates["code"] == "GLAUCOMA_EYE_TOTAL")]
    age_order = [a for a in r[r["sex"] == "both"]["age_group"]]

    wb = Workbook()
    # Fig1A: 全体（男女計・男・女）
    a = r.pivot_table(index="age_group", columns="sex", values="count_per_100k")
    a = a.reindex(age_order)[["both", "male", "female"]]
    a.columns = ["Both sexes", "Male", "Female"]
    a.index.name = "Age group"
    ws, hdr = _titled_sheet(
        wb, "Fig1A_data",
        "Figure 1A. Prescribed volume of all glaucoma eye drops by age group and sex, FY2024",
        ("Values are millilitres per 100,000 population of the same age and sex.",))
    _write_df(ws, a.reset_index(), start_row=hdr)
    _add_chart(ws, "line", len(a), 4, "F3",
               "Glaucoma eye drops by age and sex (FY2024)",
               "mL per 100,000 population", "Age group", hdr)

    # Fig1B: 薬効群別（男女計）
    rb = rates[(rates["year"] == LATEST) & (rates["sex"] == "both") &
               (rates["code"].isin(GROUP_ORDER))]
    b = rb.pivot_table(index="age_group", columns="code", values="count_per_100k")
    b = b.reindex(age_order)
    b = b[[c for c in GROUP_ORDER if c in b.columns]]
    b.columns = [GROUP_EN[c] for c in b.columns]
    b.index.name = "Age group"
    ws2, hdr2 = _titled_sheet(
        wb, "Fig1B_data",
        "Figure 1B. Prescribed volume by age group and drug class, FY2024",
        ("Both sexes combined. Millilitres per 100,000 population of the same age.",))
    _write_df(ws2, b.reset_index(), start_row=hdr2)
    _add_chart(ws2, "line", len(b), b.shape[1] + 1, "N3",
               "Drug classes by age (FY2024)", "mL per 100,000 population",
               "Age group", hdr2)

    # Fig1C: M:F比
    c = r.pivot_table(index="age_group", columns="sex", values="count_per_100k")
    c = c.reindex(age_order)
    c["M:F ratio"] = c["male"] / c["female"]
    c.index.name = "Age group"
    ws3, hdr3 = _titled_sheet(
        wb, "Fig1C_data", "Figure 1C. Male-to-female ratio by age group, FY2024", ())
    _write_df(ws3, c.reset_index()[["Age group", "M:F ratio"]], start_row=hdr3)
    _add_chart(ws3, "line", len(c), 2, "E3", "M:F ratio by age (FY2024)",
               "Male : female", "Age group", hdr3)
    _save(wb, out_dir, "Fig1_age_sex_profile.xlsx")


def build_fig2(src, out_dir):
    """Fig 2: 薬効群の処方量とシェア（2022–2024）。"""
    s = pd.read_csv(src["fullcov_summary"])
    s = s[s["code"].isin(GROUP_ORDER)]

    vol = s.pivot(index="year", columns="code", values="count_per_100k")
    vol = vol[[c for c in GROUP_ORDER if c in vol.columns]]
    vol.columns = [GROUP_EN[c] for c in vol.columns]
    vol.index.name = "Fiscal year"

    sh = s.pivot(index="year", columns="code", values="share") * 100
    sh = sh[[c for c in GROUP_ORDER if c in sh.columns]]
    sh.columns = [GROUP_EN[c] for c in sh.columns]
    sh.index.name = "Fiscal year"

    wb = Workbook()
    ws, hdr = _titled_sheet(
        wb, "Fig2A_data",
        "Figure 2A. Prescribed volume by drug class, FY2022-2024",
        (NOTE_COVERAGE, "Millilitres per 100,000 population."))
    _write_df(ws, vol.reset_index(), start_row=hdr)
    _add_chart(ws, "bar", len(vol), vol.shape[1] + 1, "N3",
               "Volume by drug class", "mL per 100,000 population",
               "Fiscal year", hdr)

    ws2, hdr2 = _titled_sheet(
        wb, "Fig2B_data",
        "Figure 2B. Share of total glaucoma eye-drop volume by drug class, FY2022-2024",
        (NOTE_COVERAGE,))
    _write_df(ws2, sh.reset_index(), start_row=hdr2)
    _add_chart(ws2, "bar", len(sh), sh.shape[1] + 1, "N3",
               "Share by drug class", "Share (%)", "Fiscal year", hdr2, stacked=True)

    # Fig2C: 個別薬剤（2024年度・降順）
    agents = pd.read_csv(src["fullcov_summary"])
    agents = agents[(agents["year"] == LATEST) &
                    (~agents["code"].isin(GROUP_ORDER + ["GLAUCOMA_EYE_TOTAL"]))]
    agents = agents.sort_values("count_per_100k", ascending=False)
    c = agents[["code", "procedure_name", "count_per_100k"]].copy()
    c["Agent"] = c["code"].map(lambda x: DRUG_EN.get(x, x))
    c = c[["Agent", "procedure_name", "count_per_100k"]].rename(
        columns={"procedure_name": "薬剤", "count_per_100k": "mL per 100,000"})
    ws3, hdr3 = _titled_sheet(
        wb, "Fig2C_data", "Figure 2C. Prescribed volume by agent, FY2024",
        (NOTE_UNIT,))
    _write_df(ws3, c, start_row=hdr3)
    chart = BarChart()
    chart.type = "bar"
    chart.title = "Volume by agent (FY2024)"
    chart.x_axis.title = "Agent"
    chart.y_axis.title = "mL per 100,000 population"
    chart.height, chart.width = 14, 18
    chart.add_data(Reference(ws3, min_col=3, min_row=hdr3, max_row=hdr3 + len(c)),
                   titles_from_data=True)
    chart.set_categories(Reference(ws3, min_col=1, min_row=hdr3 + 1,
                                   max_row=hdr3 + len(c)))
    _style_chart(chart, "bar")
    ws3.add_chart(chart, "F3")
    _save(wb, out_dir, "Fig2_group_comparison.xlsx")


def build_fig3(src, out_dir):
    """Fig 3: 都道府県別の処方量（2022–2024平均）。"""
    pr = pd.read_csv(src["fullcov_pref"])
    total = pr[pr["code"] == "GLAUCOMA_EYE_TOTAL"].sort_values("rank")
    a = total[["rank", "prefecture", "count_per_100k"]].copy()
    a["Prefecture"] = a["prefecture"].map(PREFECTURE_EN)
    a = a[["rank", "Prefecture", "prefecture", "count_per_100k"]].rename(
        columns={"rank": "Rank", "prefecture": "都道府県",
                 "count_per_100k": "mL per 100,000"})

    wb = Workbook()
    ws, hdr = _titled_sheet(
        wb, "Fig3A_data",
        "Figure 3A. Prescribed volume of all glaucoma eye drops by prefecture, "
        "FY2022-2024 mean",
        ("Millilitres per 100,000 population, mean of FY2022-2024.",))
    _write_df(ws, a, start_row=hdr)
    chart = BarChart()
    chart.type = "bar"
    chart.title = "All glaucoma eye drops by prefecture"
    chart.x_axis.title = "Prefecture"
    chart.y_axis.title = "mL per 100,000 population"
    chart.height, chart.width = 18, 16
    chart.add_data(Reference(ws, min_col=4, min_row=hdr, max_row=hdr + len(a)),
                   titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=2, min_row=hdr + 1, max_row=hdr + len(a)))
    _style_chart(chart, "bar")
    ws.add_chart(chart, "G3")

    # Fig3B: 群×都道府県
    b = pr[pr["code"].isin(GROUP_ORDER)].pivot_table(
        index="prefecture", columns="code", values="count_per_100k")
    b = b[[c for c in GROUP_ORDER if c in b.columns]]
    b.columns = [GROUP_EN[c] for c in b.columns]
    b = b.loc[[p for p in total["prefecture"] if p in b.index]]
    b.insert(0, "Prefecture", [PREFECTURE_EN[p] for p in b.index])
    b.index.name = "都道府県"
    ws2, hdr2 = _titled_sheet(
        wb, "Fig3B_data",
        "Figure 3B. Prescribed volume by prefecture and drug class, FY2022-2024 mean",
        ("Prefectures are ordered by total volume (highest first).",))
    _write_df(ws2, b.reset_index(), start_row=hdr2)
    _save(wb, out_dir, "Fig3_prefecture.xlsx")


def build_fig4(src, out_dir):
    """Fig 4: 均衡パネルの長期トレンド。"""
    tr = pd.read_csv(src["panel_trends"])
    wb = Workbook()
    for panel, label, note in [
        ("A", "Figure 4A. Balanced panel A (products listed in all 11 fiscal years, "
              "FY2014-2024)",
         "Six originator products. Volumes are the published national totals."),
        ("B", "Figure 4B. Balanced panel B (products listed in all 10 fiscal years, "
              "FY2015-2024)",
         "Twenty products. Volumes are the published national totals."),
    ]:
        p = tr[(tr["panel"] == panel) & (tr["level"] == "product")]
        w = p.pivot(index="year", columns="series", values="count_per_100k")
        w.index.name = "Fiscal year"
        ws, hdr = _titled_sheet(
            wb, f"Fig4{panel}_data", label,
            (note,
             "Restricting to continuously listed products removes the effect of the "
             "FY2022 expansion of listed products.",
             "A decline means these particular products lost volume, mostly to "
             "generics; it is not a decline in treatment of the molecule."))
        _write_df(ws, w.reset_index(), start_row=hdr)
        _add_chart(ws, "line", len(w), w.shape[1] + 1, "N3",
                   f"Balanced panel {panel}: product-level trends",
                   "mL per 100,000 population", "Fiscal year", hdr)

        tot = tr[(tr["panel"] == panel) & (tr["level"] == "panel_total")][
            ["year", "count_per_100k"]].rename(
            columns={"year": "Fiscal year",
                     "count_per_100k": f"Panel {panel} total"})
        ws2, hdr2 = _titled_sheet(
            wb, f"Fig4{panel}_total", f"{label} - panel total", (note,))
        _write_df(ws2, tot, start_row=hdr2)
        _add_chart(ws2, "line", len(tot), 2, "E3",
                   f"Balanced panel {panel}: total", "mL per 100,000 population",
                   "Fiscal year", hdr2)
    _save(wb, out_dir, "Fig4_balanced_panel_trends.xlsx")


def build_fig5(src, out_dir):
    """Fig 5: 先発品／後発品の推移と成分別後発品率。"""
    inv = pd.read_csv(src["inventory"])
    inv = inv[inv["group_code"] != "EXCLUDED"].copy()
    inv["kind"] = np.where(inv["drug_type"].str.startswith("後発"), "Generic", "Originator")

    a = inv.groupby(["year", "kind"])["quantity_ml"].sum().unstack(fill_value=0)
    a["Generic share (%)"] = a["Generic"] / (a["Generic"] + a["Originator"]) * 100
    a.index.name = "Fiscal year"

    wb = Workbook()
    ws, hdr = _titled_sheet(
        wb, "Fig5A_data",
        "Figure 5A. Originator and generic volume of glaucoma eye drops, FY2014-2024",
        (NOTE_COVERAGE,
         "Volumes are the published national totals of each product. The jump between "
         "FY2021 and FY2022 mostly reflects the listing of generic products, not a "
         "sudden change in prescribing."))
    _write_df(ws, a.reset_index(), start_row=hdr)
    _add_chart(ws, "line", len(a), 4, "G3", "Originator vs generic",
               "Volume (mL) / share (%)", "Fiscal year", hdr)

    b = inv[inv["year"] >= FULL_COV_START].groupby(
        ["year", "code", "kind"])["quantity_ml"].sum().unstack(fill_value=0)
    b["Generic share (%)"] = b["Generic"] / (b["Generic"] + b["Originator"]) * 100
    b = b.reset_index()
    b["Agent"] = b["code"].map(lambda x: DRUG_EN.get(x, x))
    piv = b.pivot(index="Agent", columns="year", values="Generic share (%)")
    piv = piv.sort_values(LATEST, ascending=False)
    piv.columns = [f"FY{c}" for c in piv.columns]
    ws2, hdr2 = _titled_sheet(
        wb, "Fig5B_data",
        "Figure 5B. Generic share by agent, FY2022-2024",
        ("Per cent of volume dispensed as generic products. Agents with no listed "
         "generic are shown as 0.",))
    _write_df(ws2, piv.reset_index(), start_row=hdr2)
    _add_chart(ws2, "bar", len(piv), piv.shape[1] + 1, "H3",
               "Generic share by agent", "Generic share (%)", "Agent", hdr2)
    _save(wb, out_dir, "Fig5_brand_generic.xlsx")


def build_fig6(src, out_dir):
    """Fig 6: 単剤ベース vs 成分ベース曝露（配合剤の寄与）。"""
    ing = pd.read_csv(src["ingredient"])
    ing = ing[ing["year"] >= FULL_COV_START]
    n = ing.groupby(["year", "code"]).agg(
        ml=("count_ml", "sum"), pop=("population_total", "sum")).reset_index()
    n["per100k"] = n["ml"] / n["pop"] * 100000
    w = n.pivot(index="year", columns="code", values="per100k")
    w.columns = [INGREDIENT_EN.get(c, c) for c in w.columns]
    w.index.name = "Fiscal year"

    mono = pd.read_csv(src["fullcov_summary"])
    pair = {"ING_PGA": "PGA", "ING_BETA": "BETA", "ING_CAI": "CAI",
            "ING_ALPHA2": "ALPHA2", "ING_ROCK": "ROCK"}
    rows = []
    for ing_code, grp in pair.items():
        m = mono[mono["code"] == grp].set_index("year")["count_per_100k"]
        i = n[n["code"] == ing_code].set_index("year")["per100k"]
        for y in sorted(set(m.index) & set(i.index)):
            rows.append({"Fiscal year": y, "Class": GROUP_EN[grp],
                         "Monotherapy only (mL per 100,000)": m[y],
                         "Ingredient basis incl. FDC (mL per 100,000)": i[y],
                         "FDC contribution (%)": (1 - m[y] / i[y]) * 100})
    comp = pd.DataFrame(rows)

    wb = Workbook()
    ws, hdr = _titled_sheet(
        wb, "Fig6A_data",
        "Figure 6A. Ingredient-level exposure including fixed-dose combinations, "
        "FY2022-2024",
        ("Fixed-dose combinations are decomposed into their component classes, so the "
         "series overlap and do not sum to the total.",))
    _write_df(ws, w.reset_index(), start_row=hdr)
    _add_chart(ws, "line", len(w), w.shape[1] + 1, "J3",
               "Ingredient-level exposure", "mL per 100,000 population",
               "Fiscal year", hdr)

    ws2, hdr2 = _titled_sheet(
        wb, "Fig6B_data",
        "Figure 6B. Monotherapy volume versus ingredient-level exposure by class",
        ("The gap between the two columns is the volume delivered through fixed-dose "
         "combinations.",))
    _write_df(ws2, comp, start_row=hdr2)
    _save(wb, out_dir, "Fig6_ingredient_exposure.xlsx")


def build_fig7(src, out_dir):
    """Fig 7: 同一成分内の品目別シェア。"""
    p = pd.read_csv(src["within_agent_products"])
    s = pd.read_csv(src["within_agent_summary"])
    p = p[p["year"] == LATEST]
    s = s[s["year"] == LATEST].sort_values("agent_total_ml", ascending=False)

    wb = Workbook()
    # 7A: 成分ごとの先発／後発比と集中度
    a = s.copy()
    a["Agent"] = a["code"].map(lambda c: DRUG_EN.get(c, c))
    a = a[["Agent", "category_name", "n_products", "originator_share",
           "generic_share", "top_product_share", "hhi_within_agent"]]
    a[["originator_share", "generic_share", "top_product_share"]] *= 100
    a.columns = ["Agent", "成分", "Products (n)", "Originator share (%)",
                 "Generic share (%)", "Leading product share (%)", "HHI"]
    ws, hdr = _titled_sheet(
        wb, "Fig7A_data",
        "Figure 7A. Originator/generic split and concentration within each agent, FY2024",
        ("HHI = sum of squared product shares within the agent. A low HHI means the "
         "agent's volume is spread across many products.",))
    _write_df(ws, a, start_row=hdr)
    chart = BarChart()
    chart.type = "bar"
    chart.grouping = "stacked"
    chart.overlap = 100
    chart.title = "Originator vs generic share within agent (FY2024)"
    chart.x_axis.title = "Agent"
    chart.y_axis.title = "Share (%)"
    chart.height, chart.width = 14, 18
    chart.add_data(Reference(ws, min_col=4, max_col=5, min_row=hdr,
                             max_row=hdr + len(a)), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=hdr + 1,
                                   max_row=hdr + len(a)))
    _style_chart(chart, "bar")
    ws.add_chart(chart, "K3")

    # 7B: 主要成分の品目別シェア（縦持ち）
    top_agents = s.head(8)["code"].tolist()
    b = p[p["code"].isin(top_agents)].copy()
    b["Agent"] = b["code"].map(lambda c: DRUG_EN.get(c, c))
    b["share_pct"] = b["share_within_agent"] * 100
    b = b[["Agent", "category_name", "rank_within_agent", "product_name",
           "drug_type", "quantity_ml", "share_pct"]].rename(columns={
        "category_name": "成分", "rank_within_agent": "Rank",
        "product_name": "医薬品名", "drug_type": "Type",
        "quantity_ml": "Volume (mL)", "share_pct": "Share within agent (%)"})
    ws2, hdr2 = _titled_sheet(
        wb, "Fig7B_data",
        "Figure 7B. Product-level share within the eight highest-volume agents, FY2024",
        ("Products are ranked by volume within each agent.",))
    _write_df(ws2, b, start_row=hdr2)
    _save(wb, out_dir, "Fig7_within_agent_share.xlsx")


def build_fig8(src, out_dir):
    """Fig 8: 粗率 vs 年齢調整率。"""
    a = pd.read_csv(src["age_adjusted_pref"])
    d = pd.read_csv(src["age_adjusted_disparity"])
    t = a[(a["year"] == LATEST) & (a["code"] == "GLAUCOMA_EYE_TOTAL")].copy()
    t["Prefecture"] = t["prefecture"].map(PREFECTURE_EN)
    t["aging_rate_pct"] = t["aging_rate"] * 100
    t = t.sort_values("age_adjusted_rate_per_100k", ascending=False)

    wb = Workbook()
    x = t[["Prefecture", "prefecture", "aging_rate_pct",
           "crude_rate_per_100k", "age_adjusted_rate_per_100k", "spr",
           "crude_rank", "age_adjusted_rank", "rank_change"]].rename(columns={
        "prefecture": "都道府県", "aging_rate_pct": "Aged 65+ (%)",
        "crude_rate_per_100k": "Crude rate", "age_adjusted_rate_per_100k": "Age-adjusted rate",
        "spr": "SPR", "crude_rank": "Rank (crude)",
        "age_adjusted_rank": "Rank (adjusted)", "rank_change": "Rank change"})
    ws, hdr = _titled_sheet(
        wb, "Fig8A_data",
        "Figure 8A. Crude versus age-adjusted prescribed volume by prefecture, FY2024",
        ("Indirect standardisation with two age strata (0-64, 65+).",
         "Rates are millilitres per 100,000 population. A positive rank change means "
         "the prefecture ranks higher after adjustment."))
    _write_df(ws, x, start_row=hdr)
    chart = BarChart()
    chart.type = "bar"
    chart.title = "Crude vs age-adjusted rate by prefecture (FY2024)"
    chart.x_axis.title = "Prefecture"
    chart.y_axis.title = "mL per 100,000 population"
    chart.height, chart.width = 18, 18
    chart.add_data(Reference(ws, min_col=4, max_col=5, min_row=hdr,
                             max_row=hdr + len(x)), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=hdr + 1,
                                   max_row=hdr + len(x)))
    _style_chart(chart, "bar")
    ws.add_chart(chart, "L3")

    # 8B: 群別に見た年齢調整の効果
    b = d[d["year"] == LATEST].copy()
    b["Class (EN)"] = b["code"].map(lambda c: en(c))
    b = b[["Class (EN)", "procedure_name", "crude_cv", "age_adjusted_cv",
           "cv_reduction_pct", "spearman_rho_crude_vs_adjusted",
           "max_rank_change", "spr_min", "spr_min_prefecture",
           "spr_max", "spr_max_prefecture"]].rename(columns={
        "procedure_name": "薬効群", "crude_cv": "CV (crude)",
        "age_adjusted_cv": "CV (age-adjusted)",
        "cv_reduction_pct": "CV reduction (%)",
        "spearman_rho_crude_vs_adjusted": "Spearman rho (crude vs adjusted)",
        "max_rank_change": "Max rank change", "spr_min": "SPR min",
        "spr_min_prefecture": "SPR min prefecture", "spr_max": "SPR max",
        "spr_max_prefecture": "SPR max prefecture"})
    ws2, hdr2 = _titled_sheet(
        wb, "Fig8B_data",
        "Figure 8B. Effect of age adjustment on between-prefecture disparity by drug "
        "class, FY2024",
        ("A small CV reduction means the disparity is not explained by the age "
         "structure of the prefecture.",))
    _write_df(ws2, b, start_row=hdr2)
    _save(wb, out_dir, "Fig8_age_adjustment.xlsx")


# ─────────────────────────────────────────────
# Supplementary tables
# ─────────────────────────────────────────────

def build_suppl1(src, out_dir):
    """Suppl Table 1: 年度×薬効群の全国数量（未収載はNR）。"""
    tr = pd.read_csv(src["national_trends"])
    cov = pd.read_csv(src["coverage_by_year"])
    listed = cov.groupby("year")["code"].apply(set).to_dict()
    members = {g: set(m) for g, _, m in GROUP_DEFS}

    w = tr[tr["code"].isin(GROUP_ORDER + ["GLAUCOMA_EYE_TOTAL"])].pivot(
        index="year", columns="code", values="count_ml")
    out = pd.DataFrame(index=w.index)
    for c in ["GLAUCOMA_EYE_TOTAL"] + GROUP_ORDER:
        col = GROUP_EN[c]
        vals = []
        for y in w.index:
            if c != "GLAUCOMA_EYE_TOTAL" and not (members[c] & listed.get(y, set())):
                vals.append("NR")
            elif c in w.columns and pd.notna(w.loc[y, c]):
                vals.append(round(float(w.loc[y, c])))
            else:
                vals.append("NR")
        out[col] = vals
    out.index.name = "Fiscal year"

    n_listed = cov.groupby("year")["n_products"].sum()
    out["Products listed (n)"] = [int(n_listed.get(y, 0)) for y in out.index]

    wb = Workbook()
    ws, hdr = _titled_sheet(
        wb, "SupplTable1",
        "Supplementary Table 1. Annual national volume of glaucoma eye drops by drug "
        "class, FY2014-2024 (mL)",
        ("NR = not reported. NDB Open Data listed only the highest-volume products of "
         "each therapeutic category before the 10th release (FY2022); a class shown as "
         "NR had no product listed in that fiscal year and was not necessarily unused.",
         NOTE_UNIT, NOTE_GAP,
         "Because the set of listed products changes in FY2022, volumes before and from "
         "FY2022 must not be compared directly."))
    _write_df(ws, out.reset_index(), start_row=hdr)
    _save(wb, out_dir, "SupplTable1_annual_totals_by_group.xlsx")
    out.reset_index().to_csv(
        os.path.join(out_dir, "SupplTable1_annual_totals_by_group.csv"),
        index=False, encoding="utf-8-sig")


def build_suppl2(src, out_dir):
    """Suppl Table 2: 対象品目マスタ。"""
    m = pd.read_csv(src["master"])
    m["Class (EN)"] = m["group_code"].map(lambda c: GROUP_EN.get(c, "Listed separately"))
    m["Agent (EN)"] = m["code"].map(lambda c: DRUG_EN.get(c, c))
    m["Type"] = m["drug_type"].map(DRUG_TYPE_EN)
    out = m[["Class (EN)", "group_name", "Agent (EN)", "category_name",
             "product_name", "Type", "unit_published", "ml_per_unit",
             "n_years", "years_listed", "quantity_ml_total"]].rename(columns={
        "group_name": "薬効群", "category_name": "成分", "product_name": "医薬品名",
        "unit_published": "Published unit", "ml_per_unit": "mL per unit",
        "n_years": "Fiscal years listed (n)", "years_listed": "Fiscal years listed",
        "quantity_ml_total": "Total volume FY2014-2024 (mL)"})
    wb = Workbook()
    ws, hdr = _titled_sheet(
        wb, "SupplTable2",
        "Supplementary Table 2. All products included in the analysis",
        ("Products were selected as therapeutic category 131 (ophthalmic agents) with "
         "'点眼' (eye drops) in the product name and an intraocular-pressure-lowering "
         "indication.",
         NOTE_UNIT,
         "Products marked '（選）' (selective medical care for long-listed originators, "
         "from October 2024) were merged into the parent product.",
         "Originator / generic status is inferred from the product name; NDB Open Data "
         "does not publish this attribute."))
    _write_df(ws, out, start_row=hdr)
    _save(wb, out_dir, "SupplTable2_product_master.xlsx")


def build_suppl3(src, out_dir):
    """Suppl Table 3: 秘匿と収載カバレッジ。"""
    cen = pd.read_csv(src["censoring"])
    cov = pd.read_csv(src["coverage_by_year"])

    by_year = cen.groupby("year").agg(
        cells=("n_pref_cells", "sum"), censored=("n_pref_censored", "sum"),
        volume=("quantity_ml", "sum"), pref_sum=("pref_sum_ml", "sum")).reset_index()
    by_year["Suppressed cells (%)"] = by_year["censored"] / by_year["cells"] * 100
    by_year["Volume lost to suppression (%)"] = (
        1 - by_year["pref_sum"] / by_year["volume"]) * 100
    by_year = by_year.rename(columns={
        "year": "Fiscal year", "cells": "Prefecture cells (n)",
        "censored": "Suppressed cells (n)", "volume": "National volume (mL)",
        "pref_sum": "Sum of prefecture cells (mL)"})

    by_code = cen.groupby(["code", "category_name"]).agg(
        cells=("n_pref_cells", "sum"), censored=("n_pref_censored", "sum"),
        volume=("quantity_ml", "sum"), pref_sum=("pref_sum_ml", "sum")).reset_index()
    by_code["Agent (EN)"] = by_code["code"].map(lambda c: DRUG_EN.get(c, c))
    by_code["Suppressed cells (%)"] = by_code["censored"] / by_code["cells"] * 100
    by_code["Volume lost to suppression (%)"] = (
        1 - by_code["pref_sum"] / by_code["volume"]) * 100
    by_code = by_code.sort_values("Volume lost to suppression (%)", ascending=False)
    by_code = by_code[["Agent (EN)", "category_name", "cells", "censored",
                       "Suppressed cells (%)", "Volume lost to suppression (%)"]].rename(
        columns={"category_name": "成分", "cells": "Prefecture cells (n)",
                 "censored": "Suppressed cells (n)"})

    npr = cov.pivot(index="year", columns="code", values="n_products").fillna(0).astype(int)
    npr.columns = [DRUG_EN.get(c, c) for c in npr.columns]
    npr.index.name = "Fiscal year"

    wb = Workbook()
    ws, hdr = _titled_sheet(
        wb, "S3A_suppression_by_year",
        "Supplementary Table 3A. Cell suppression by fiscal year",
        ("Prefecture-level cells with a dispensed quantity below 1,000 are suppressed "
         "in NDB Open Data and were imputed as zero in the main analysis.",))
    _write_df(ws, by_year, start_row=hdr)

    ws2, hdr2 = _titled_sheet(
        wb, "S3B_suppression_by_agent",
        "Supplementary Table 3B. Cell suppression by agent, FY2014-2024",
        ("Agents with a high share of volume lost to suppression are not suitable for "
         "prefecture-level analysis.",))
    _write_df(ws2, by_code, start_row=hdr2)

    ws3, hdr3 = _titled_sheet(
        wb, "S3C_products_listed",
        "Supplementary Table 3C. Number of products listed by agent and fiscal year",
        ("The increase between FY2021 and FY2022 reflects the expansion of listed "
         "products in the 10th release of NDB Open Data.",))
    _write_df(ws3, npr.reset_index(), start_row=hdr3)

    # 3D: 年齢性別表の完全性（行単位の全面秘匿）
    ag = pd.read_csv(src["agesex_completeness"])
    by_y = ag.groupby("year").agg(
        rows=("product", "size"),
        fully=("fully_suppressed", "sum"),
        total=("total_published_ml", "sum"),
        cells=("agesex_sum_ml", "sum")).reset_index()
    by_y["Recovered (%)"] = by_y["cells"] / by_y["total"] * 100
    by_y["Volume lost (mL)"] = by_y["total"] - by_y["cells"]
    top = (ag[ag["fully_suppressed"]].sort_values("total_published_ml", ascending=False)
           .groupby("year")["product"].first())
    by_y["Largest affected product"] = by_y["year"].map(top).fillna("")
    by_y = by_y.rename(columns={
        "year": "Fiscal year", "rows": "Product rows (n)",
        "fully": "Rows with every age-sex cell suppressed (n)",
        "total": "Published row total (mL)",
        "cells": "Sum of age-sex cells (mL)"})
    ws4, hdr4 = _titled_sheet(
        wb, "S3D_agesex_completeness",
        "Supplementary Table 3D. Completeness of the age-and-sex tables, FY2014-2024",
        ("NDB Open Data publishes a row total for every listed product, but some rows "
         "have every age-and-sex cell suppressed even though the row total is in the "
         "millions of millilitres. Such rows are lost entirely from age-stratified "
         "analysis, and the loss is far larger than ordinary cell suppression "
         "(quantities below 1,000) can explain.",
         "Affected products are brimonidine (FY2014, FY2015), ripasudil (FY2016, "
         "FY2017), latanoprost (FY2017) and dorzolamide (FY2015).",
         "From FY2018 onwards 99% or more of the published totals are recovered. "
         "Age-stratified comparisons across fiscal years should be restricted to "
         "FY2018-2024; the main analysis of this study is restricted to FY2022-2024 "
         "and is therefore unaffected."))
    _write_df(ws4, by_y, start_row=hdr4)

    # 3E: 系列途中の収載欠落（前後の年度には収載されている品目）
    gap = pd.read_csv(src["listing_gaps"])
    gap["Agent (EN)"] = gap["code"].map(lambda c: DRUG_EN.get(c, c))
    ge = gap[["year", "product", "Agent (EN)", "prev_year", "prev_ml",
              "next_year", "next_ml", "expected_missing_ml"]].rename(columns={
        "year": "Fiscal year", "product": "品目",
        "prev_year": "Previous listed year", "prev_ml": "Volume then (mL)",
        "next_year": "Next listed year", "next_ml": "Volume then (mL) ",
        "expected_missing_ml": "Expected missing volume (mL)"})
    ws5, hdr5 = _titled_sheet(
        wb, "S3E_listing_gaps",
        "Supplementary Table 3E. Products absent from the listing in a fiscal year but "
        "listed both before and after, FY2014-2024",
        ("These are not suppressed cells: the whole product row is missing from NDB "
         "Open Data in that fiscal year, while the same product is listed in the "
         "preceding and following years at a comparable volume.",
         "The expected missing volume is the mean of the preceding and following "
         "fiscal years and estimates how much the class total is understated.",
         "FY2021 is by far the most affected year: nine products, an estimated 14.8 "
         "million mL, 7.5% of the annual total. The apparent FY2021 dip in Figure 9 "
         "and Supplementary Table 1 is therefore an artefact, not a fall in prescribing.",
         "The main analysis of this study is restricted to FY2022-2024 and is "
         "unaffected; the balanced panels (Table 5, Figure 4) exclude these products "
         "because they are not continuously listed."))
    _write_df(ws5, ge, start_row=hdr5, number_format="#,##0")
    _save(wb, out_dir, "SupplTable3_censoring_coverage.xlsx")


def build_suppl4(src, out_dir):
    """Suppl Table 4: パネル回帰・代替性・収束分析。"""
    wb = Workbook()
    panel = pd.read_csv(src["panel_reg"])
    panel["Name (EN)"] = panel["code"].map(lambda c: en(c))
    p = panel[["code", "Name (EN)", "procedure_name", "variable", "coefficient",
               "std_err", "t_stat", "p_value", "r2_within", "n_obs"]].rename(columns={
        "code": "Code", "procedure_name": "薬剤・薬効群", "variable": "Covariate",
        "coefficient": "Coefficient", "std_err": "SE", "t_stat": "t",
        "p_value": "p", "r2_within": "Within R2", "n_obs": "Observations"})
    ws, hdr = _titled_sheet(
        wb, "S4A_panel_regression",
        "Supplementary Table 4A. Two-way fixed-effects panel regression, FY2014-2024",
        ("Outcome: prescribed volume per 100,000 population (mL). Covariates: "
         "proportion aged 65 years and over, ophthalmologists per 100,000, and "
         "ophthalmic facilities per 100,000.",
         "Prefecture and fiscal-year fixed effects; standard errors clustered by "
         "prefecture."))
    _write_df(ws, p, start_row=hdr)

    sub = pd.read_csv(src["substitution"])
    s = sub.rename(columns={
        "model": "Model", "dependent": "Dependent variable", "variable": "Covariate",
        "coefficient": "Coefficient", "std_err": "SE", "t_stat": "t",
        "p_value": "p", "r2_within": "Within R2", "n_obs": "Observations"})
    ws2, hdr2 = _titled_sheet(
        wb, "S4B_substitution",
        "Supplementary Table 4B. First-difference panel regression of therapeutic "
        "substitution",
        ("Changes in within-class share are regressed on each other with prefecture and "
         "fiscal-year fixed effects. A negative coefficient indicates substitution.",))
    _write_df(ws2, s, start_row=hdr2)

    conv = pd.read_csv(src["convergence"])
    conv["Name (EN)"] = conv["code"].map(lambda c: en(c))
    ws3, hdr3 = _titled_sheet(
        wb, "S4C_convergence",
        "Supplementary Table 4C. Sigma- and beta-convergence of prefecture-level volume",
        ("Sigma-convergence: regression of the coefficient of variation on fiscal year. "
         "Beta-convergence: regression of log growth on log initial level.",))
    _write_df(ws3, conv, start_row=hdr3)

    disp = pd.read_csv(src["disparity"])
    disp["Name (EN)"] = disp["code"].map(lambda c: en(c))
    ws4, hdr4 = _titled_sheet(
        wb, "S4D_disparity",
        "Supplementary Table 4D. Prefecture-level disparity indices by fiscal year", ())
    _write_df(ws4, disp, start_row=hdr4)
    _save(wb, out_dir, "SupplTable4_regressions.xlsx")


def build_suppl5(src, out_dir):
    """Suppl Table 5: 秘匿セルの感度分析（識別区間）。"""
    c = pd.read_csv(src["sens_categories"])
    p = pd.read_csv(src["sens_prefecture"])
    c = c[c["year"] == LATEST].copy()
    c["Name (EN)"] = c["code"].map(lambda x: en(x))
    c["share_zero_pct"] = c["share_zero"] * 100
    c["share_upper_pct"] = c["share_upper"] * 100
    a = c[["Name (EN)", "procedure_name",
           "count_per_100k_zero", "count_per_100k_upper",
           "count_per_100k_rel_diff_pct", "share_zero_pct", "share_upper_pct",
           "cv_zero", "cv_upper", "gini_zero", "gini_upper"]].rename(columns={
        "procedure_name": "薬剤・薬効群",
        "count_per_100k_zero": "Per 100,000, lower bound (zero)",
        "count_per_100k_upper": "Per 100,000, upper bound (999)",
        "count_per_100k_rel_diff_pct": "Width of identified interval (%)",
        "share_zero_pct": "Share, lower bound (%)",
        "share_upper_pct": "Share, upper bound (%)",
        "cv_zero": "CV, lower bound", "cv_upper": "CV, upper bound",
        "gini_zero": "Gini, lower bound", "gini_upper": "Gini, upper bound"})

    p = p[p["year"] == LATEST].copy()
    p["Name (EN)"] = p["code"].map(lambda x: en(x))
    b = p[["Name (EN)", "code", "spearman_rho", "max_rank_change",
           "n_prefectures_rank_changed", "top1_zero", "top1_upper",
           "bottom1_zero", "bottom1_upper"]].rename(columns={
        "code": "Code", "spearman_rho": "Spearman rho (zero vs upper)",
        "max_rank_change": "Max rank change",
        "n_prefectures_rank_changed": "Prefectures with a changed rank (n)",
        "top1_zero": "Highest, lower bound", "top1_upper": "Highest, upper bound",
        "bottom1_zero": "Lowest, lower bound", "bottom1_upper": "Lowest, upper bound"})

    wb = Workbook()
    ws, hdr = _titled_sheet(
        wb, "S5A_identified_interval",
        "Supplementary Table 5A. Identified interval of the main estimates under cell "
        "suppression, FY2024",
        ("Prefecture-level cells with a dispensed quantity below 1,000 are suppressed. "
         "The true value lies in [0, 1000), so imputing 0 gives the lower bound of every "
         "estimate and imputing 999 gives the upper bound.",
         "The main analysis uses the lower bound (zero imputation)."))
    _write_df(ws, a, start_row=hdr)

    ws2, hdr2 = _titled_sheet(
        wb, "S5B_prefecture_robustness",
        "Supplementary Table 5B. Robustness of the prefecture ranking to the imputation "
        "of suppressed cells, FY2024",
        ("A Spearman rho close to 1 means the ranking is unaffected by the choice of "
         "imputation.",))
    _write_df(ws2, b, start_row=hdr2)
    _save(wb, out_dir, "SupplTable5_censoring_sensitivity.xlsx")


# ─────────────────────────────────────────────

def main():
    ap = argparse.ArgumentParser()
    add_nokouhi_arg(ap)
    args = ap.parse_args()

    out_dir = os.path.join(data_root(args.nokouhi), PAPER_SUBDIR)
    os.makedirs(out_dir, exist_ok=True)

    names = {
        "fullcov_summary": "fullcov_summary_glaucoma.csv",
        "fullcov_change": "fullcov_change_2022_2024_glaucoma.csv",
        "fullcov_pref": "fullcov_prefecture_ranking_glaucoma.csv",
        "panel_trends": "balanced_panel_trends_glaucoma.csv",
        "panel_apc": "balanced_panel_apc_glaucoma.csv",
        "panel_products": "balanced_panel_products_glaucoma.csv",
        "inventory": "product_inventory_glaucoma.csv",
        "master": "product_master_glaucoma.csv",
        "censoring": "censoring_report_glaucoma.csv",
        "coverage_by_year": "coverage_by_year_glaucoma.csv",
        "agesex_completeness": "agesex_completeness_glaucoma.csv",
        "listing_gaps": "listing_gaps_glaucoma.csv",
        "age_sex": "age_sex_rates_glaucoma.csv",
        "mf_ratio": "mf_ratio_by_age_glaucoma.csv",
        "national_trends": "national_trends_glaucoma.csv",
        "ingredient": "ingredient_exposure_glaucoma_zero.csv",
        "panel_reg": "panel_regression_summary_glaucoma.csv",
        "substitution": "substitution_regression_glaucoma.csv",
        "convergence": "convergence_analysis_glaucoma.csv",
        "disparity": "geographic_disparity_glaucoma.csv",
        "within_agent_products": "within_agent_product_share_glaucoma.csv",
        "within_agent_summary": "within_agent_summary_glaucoma.csv",
        "age_adjusted_pref": "age_adjusted_prefecture_glaucoma.csv",
        "age_adjusted_disparity": "age_adjusted_disparity_glaucoma.csv",
        "sens_categories": "censoring_sensitivity_glaucoma.csv",
        "sens_prefecture": "censoring_sensitivity_prefecture.csv",
        "agent_series": "agent_annual_series_glaucoma.csv",
        "agent_trace": "agent_traceability_glaucoma.csv",
        "group_share_by_year": "group_share_by_year_glaucoma.csv",
    }
    src = {}
    for k, v in names.items():
        p = find(v, args.nokouhi)
        if not os.path.exists(p):
            raise FileNotFoundError(
                f"{v} が見つかりません。README の実行順どおりに先行スクリプトを"
                f"実行してください。")
        src[k] = p

    print(f"出力先: {out_dir}")
    print("Tables:")
    tables = {
        "Table1": build_table1(src, out_dir),
        "Table2": build_table2(src, out_dir),
        "Table3": build_table3(src, out_dir),
        "Table4": build_table4(src, out_dir),
        "Table5": build_table5(src, out_dir),
        "Table6": build_table6(src, out_dir),
        "Table7": build_table7(src, out_dir),
    }
    build_tables_workbook(tables, out_dir)

    print("Figures:")
    build_fig1(src, out_dir)
    build_fig2(src, out_dir)
    build_fig3(src, out_dir)
    build_fig4(src, out_dir)
    build_fig5(src, out_dir)
    build_fig6(src, out_dir)
    build_fig7(src, out_dir)
    build_fig8(src, out_dir)
    build_fig9(src, out_dir)

    print("Supplementary:")
    build_suppl1(src, out_dir)
    build_suppl2(src, out_dir)
    build_suppl3(src, out_dir)
    build_suppl4(src, out_dir)
    build_suppl5(src, out_dir)
    build_suppl6(src, out_dir)

    print("\n完了。")


def build_fig9(src, out_dir):
    """Fig 9: 薬効群シェアの年次推移（2014–2024、下位分類を含む）。"""
    g = pd.read_csv(src["group_share_by_year"])

    a = g[g["level"] == "薬効群"].pivot(
        index="year", columns="code", values="share_of_total_pct")
    a = a[[c for c in GROUP_ORDER if c in a.columns]]
    a.columns = [GROUP_EN[c] for c in a.columns]
    a.index.name = "Fiscal year"

    wb = Workbook()
    ws, hdr = _titled_sheet(
        wb, "Fig9A_data",
        "Figure 9A. Share of total glaucoma eye-drop volume by drug class, FY2014-2024",
        (NOTE_COVERAGE, NOTE_GAP,
         "Blank cells mean the class had no product listed in that fiscal year; they "
         "do not mean the class was unused.",
         "Shares before FY2022 are computed within the products listed in that year, "
         "so they are not directly comparable with FY2022 onwards."))
    _write_df(ws, a.reset_index(), start_row=hdr)
    _add_chart(ws, "line", len(a), a.shape[1] + 1, "N3",
               "Drug-class share, FY2014-2024", "Share (%)", "Fiscal year", hdr)

    b = g[g["level"] == "下位分類"].pivot(
        index="year", columns="code", values="share_of_total_pct")
    b = b[[c for c in SUBGROUP_ORDER if c in b.columns]]
    b.columns = [SUBGROUP_EN[c] for c in b.columns]
    b.index.name = "Fiscal year"
    ws2, hdr2 = _titled_sheet(
        wb, "Fig9B_data",
        "Figure 9B. Share of total volume by subclass (prostanoid FP and EP2 receptor "
        "agonists; non-selective, beta1-selective and alpha1-beta blockers), FY2014-2024",
        (NOTE_COVERAGE, NOTE_GAP,
         "Subclasses overlap with their parent class and are shown as a share of the "
         "overall total.",
         "Blank cells mean the subclass had no product listed in that fiscal year; they "
         "do not mean the subclass was unused. Alpha1-beta blockers had no listed "
         "product in FY2014 and in FY2019-2021, and reappear in FY2022 at a higher "
         "volume than in FY2018.",
         "A value of 0 means every prefecture cell was suppressed and imputed as zero, "
         "not that nothing was prescribed: alpha1-beta blockers in FY2015 have a "
         "published national total of 2.0 million mL.",
         "Alpha1-beta blockers (nipradilol) are a subclass of the beta-blockers and are "
         "distinct from the sympathetic alpha1-receptor blockers (bunazosin) shown in "
         "Figure 9A."))
    _write_df(ws2, b.reset_index(), start_row=hdr2)
    _add_chart(ws2, "line", len(b), b.shape[1] + 1, "N3",
               "Subclass share, FY2014-2024", "Share (%)", "Fiscal year", hdr2)

    c = g[g["level"] == "薬効群"].pivot(
        index="year", columns="code", values="count_per_100k")
    c = c[[x for x in GROUP_ORDER if x in c.columns]]
    c.columns = [GROUP_EN[x] for x in c.columns]
    c.index.name = "Fiscal year"
    ws3, hdr3 = _titled_sheet(
        wb, "Fig9C_data",
        "Figure 9C. Prescribed volume by drug class, FY2014-2024",
        (NOTE_COVERAGE, NOTE_GAP, "Millilitres per 100,000 population."))
    _write_df(ws3, c.reset_index(), start_row=hdr3)
    _add_chart(ws3, "line", len(c), c.shape[1] + 1, "N3",
               "Volume by drug class, FY2014-2024", "mL per 100,000 population",
               "Fiscal year", hdr3)
    _save(wb, out_dir, "Fig9_group_share_trend.xlsx")


def build_suppl6(src, out_dir):
    """Suppl Table 6: 成分別の年次系列と追跡可能性。"""
    s = pd.read_csv(src["agent_series"])
    t = pd.read_csv(src["agent_trace"])

    s["Class (EN)"] = s["group_code"].map(lambda c: GROUP_EN.get(c, c))
    s["Agent (EN)"] = s["code"].map(lambda c: DRUG_EN.get(c, c))
    piv = s.pivot_table(index=["Class (EN)", "Agent (EN)", "category_name"],
                        columns="year", values="count_per_100k")
    piv.columns = [f"FY{int(c)}" for c in piv.columns]
    piv = piv.reset_index()
    order = {GROUP_EN[c]: i for i, c in enumerate(GROUP_ORDER)}
    piv["_o"] = piv["Class (EN)"].map(order)
    piv = piv.sort_values(["_o", "Agent (EN)"]).drop(columns="_o")
    piv = piv.rename(columns={"category_name": "成分"})

    t["Class (EN)"] = t["group_code"].map(lambda c: GROUP_EN.get(c, c))
    t["Agent (EN)"] = t["code"].map(lambda c: DRUG_EN.get(c, c))
    tt = t[["Class (EN)", "Agent (EN)", "category_name", "first_listed_year",
            "continuous_since", "n_years_continuous", "missing_years",
            "traceability", "start_pattern", "first_year_per_100k",
            "last_year_per_100k", "apc_over_continuous_period",
            "apc_low", "apc_high", "apc_p_value"]].rename(columns={
        "category_name": "成分", "first_listed_year": "First listed (FY)",
        "continuous_since": "Continuously listed since (FY)",
        "n_years_continuous": "Continuous years (n)",
        "missing_years": "Missing fiscal years", "traceability": "追跡可能性",
        "start_pattern": "初収載の解釈",
        "first_year_per_100k": "First year, per 100,000 (mL)",
        "last_year_per_100k": "FY2024, per 100,000 (mL)",
        "apc_over_continuous_period": "APC over continuous period (%/year)",
        "apc_low": "95%CI lower", "apc_high": "95%CI upper", "apc_p_value": "p"})

    wb = Workbook()
    ws, hdr = _titled_sheet(
        wb, "S6A_annual_series",
        "Supplementary Table 6A. Prescribed volume of each agent by fiscal year, "
        "FY2014-2024 (mL per 100,000 population)",
        ("Blank cells mean the agent had no product listed in NDB Open Data that "
         "fiscal year; they do not mean the agent was unused.",
         NOTE_COVERAGE, NOTE_UNIT))
    _write_df(ws, piv, start_row=hdr)

    ws2, hdr2 = _titled_sheet(
        wb, "S6B_traceability",
        "Supplementary Table 6B. From which fiscal year each agent can be tracked",
        ("'Continuously listed since' is the earliest fiscal year from which the agent "
         "is listed in every year up to FY2024.",
         "For agents first listed in FY2022, the ratio of FY2024 to the first year is "
         "given as evidence of whether the agent was newly launched (steep rise from a "
         "low base) or simply became visible when NDB Open Data expanded its product "
         "list (flat from the first year).",
         "The APC is calculated over the continuously listed period only."))
    _write_df(ws2, tt, start_row=hdr2)
    _save(wb, out_dir, "SupplTable6_agent_series.xlsx")


if __name__ == "__main__":
    main()
