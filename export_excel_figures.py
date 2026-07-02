# -*- coding: utf-8 -*-
"""
export_excel_figures.py
解析結果をExcelグラフ作成用に整形し、openpyxlでグラフオブジェクトを埋め込んで出力する。

出力: data/processed/figure_data_for_excel.xlsx
  Sheet1:  全国トレンド（経年）        -- データ＋折れ線グラフ
  Sheet2:  APC経年変化率               -- データ＋棒グラフ
  Sheet3:  地域格差指標推移            -- データ＋折れ線グラフ(Gini/CV)
  Sheet4:  都道府県別データ(2023)      -- データ＋横棒グラフ
  Sheet5:  Spearman相関                -- データ
  Sheet6:  パネル回帰係数              -- データ＋棒グラフ
"""
import sys
import pandas as pd
import numpy as np
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import SeriesLabel
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')

proc = pd.read_csv('data/processed/ndb_processed_zero.csv', encoding='utf-8-sig')
apc_lin = pd.read_csv('data/processed/national_apc_linear.csv', encoding='utf-8-sig')
apc_jp = pd.read_csv('data/processed/national_apc_joinpoint.csv', encoding='utf-8-sig')
disp = pd.read_csv('data/processed/geographic_disparity.csv', encoding='utf-8-sig')
corr = pd.read_csv('data/processed/covariates_correlation.csv', encoding='utf-8-sig')
reg = pd.read_csv('data/processed/panel_regression_summary.csv', encoding='utf-8-sig')

CODE_LABEL = {
    'J039-2': '抗VEGF薬注射',
    'K282':   '白内障手術',
    'K280':   '硝子体手術',
    'K268':   '緑内障手術',
    'K259':   '角膜移植術',
}
CODES = ['J039-2', 'K282', 'K280', 'K268', 'K259']

nat = proc.groupby(['year', 'code', 'procedure_name']).agg(
    count=('count', 'sum'),
    population_total=('population_total', 'sum'),
    population_65plus=('population_65plus', 'sum')
).reset_index()
nat['count_per_100k'] = (nat['count'] / nat['population_total']) * 100000
nat['count_per_100k_65plus'] = (nat['count'] / nat['population_65plus']) * 100000

# ─── Sheet1: 全国トレンド ───
rows_trend = []
for code in CODES:
    d = nat[nat['code'] == code].sort_values('year')
    for _, r in d.iterrows():
        rows_trend.append({
            '診療行為コード': code,
            '診療行為名': CODE_LABEL.get(code, code),
            '年度': int(r['year']),
            '全国施行件数': int(r['count']),
            '人口10万対施行件数': round(r['count_per_100k'], 2),
            '65歳以上人口10万対施行件数': round(r['count_per_100k_65plus'], 2),
        })
df_trend = pd.DataFrame(rows_trend)

df_trend_wide = df_trend.pivot_table(
    index='年度', columns='診療行為名', values='人口10万対施行件数'
).reset_index()
df_trend_wide.columns.name = None

df_trend_wide_count = df_trend.pivot_table(
    index='年度', columns='診療行為名', values='全国施行件数'
).reset_index()
df_trend_wide_count.columns.name = None

# ─── Sheet2: APC ───
apc_rows = []
for _, a in apc_lin.iterrows():
    if a['code'] in CODES:
        apc_rows.append({
            '診療行為コード': a['code'],
            '診療行為名': CODE_LABEL.get(a['code'], a['code']),
            '解析期間': f"{int(a['start_year'])}-{int(a['end_year'])}",
            'APC(%)': round(a['apc'], 2),
            '95%CI下限(%)': round(a['apc_low'], 2),
            '95%CI上限(%)': round(a['apc_high'], 2),
            'p値': round(a['p_value'], 6),
            'R2': round(a['r2'], 4),
        })

jp_rows = []
for _, j in apc_jp.iterrows():
    if j['code'] in CODES:
        jp_rows.append({
            '診療行為コード': j['code'],
            '診療行為名': CODE_LABEL.get(j['code'], j['code']),
            'セグメント': int(j['segment']),
            '開始年': int(j['start_year']),
            '終了年': int(j['end_year']),
            'APC(%)': round(j['apc'], 2),
        })

df_apc = pd.DataFrame(apc_rows)
df_jp = pd.DataFrame(jp_rows)

# ─── Sheet3: 地域格差指標推移 ───
disp_rows = []
for code in CODES:
    d = disp[disp['code'] == code].sort_values('year')
    for _, r in d.iterrows():
        disp_rows.append({
            '診療行為コード': code,
            '診療行為名': CODE_LABEL.get(code, code),
            '年度': int(r['year']),
            '変動係数(CV)': round(r['cv'], 4),
            'Gini係数': round(r['gini'], 4),
            '平均施行率(per100k)': round(r['mean_rate_per_100k'], 2),
            '最小施行率': round(r['min_rate'], 2),
            '最大施行率': round(r['max_rate'], 2),
            '最大/最小比': round(r['max_to_min_ratio'], 2) if pd.notna(r['max_to_min_ratio']) else None,
            '最小都道府県': r['min_prefecture'],
            '最大都道府県': r['max_prefecture'],
        })
df_disp = pd.DataFrame(disp_rows)

df_gini_wide = df_disp.pivot_table(
    index='年度', columns='診療行為名', values='Gini係数'
).reset_index()
df_gini_wide.columns.name = None

df_cv_wide = df_disp.pivot_table(
    index='年度', columns='診療行為名', values='変動係数(CV)'
).reset_index()
df_cv_wide.columns.name = None

# ─── Sheet4: 都道府県別データ(2023) ───
pref_rows = []
for code in CODES:
    d2023 = proc[(proc['code'] == code) & (proc['year'] == 2023)][
        ['prefecture', 'count', 'count_per_100k', 'count_per_100k_65plus',
         'aging_rate', 'docs_per_100k', 'facilities_per_100k']
    ].copy()
    d2023['診療行為コード'] = code
    d2023['診療行為名'] = CODE_LABEL.get(code, code)
    d2023 = d2023.rename(columns={
        'prefecture': '都道府県',
        'count': '施行件数',
        'count_per_100k': '人口10万対施行件数',
        'count_per_100k_65plus': '65歳以上10万対施行件数',
        'aging_rate': '高齢化率',
        'docs_per_100k': '眼科医数(per10万人)',
        'facilities_per_100k': '眼科施設数(per10万人)',
    })
    d2023['高齢化率(%)'] = (d2023['高齢化率'] * 100).round(1)
    d2023 = d2023.drop(columns=['高齢化率'])
    d2023 = d2023.sort_values('人口10万対施行件数', ascending=False).reset_index(drop=True)
    d2023.index = d2023.index + 1
    d2023.index.name = '順位'
    pref_rows.append((code, d2023.reset_index()))

# ─── Sheet5: Spearman相関 ───
corr_rows = []
for code in CODES:
    for yr in sorted(corr['year'].unique()):
        c = corr[(corr['code'] == code) & (corr['year'] == yr)]
        if not c.empty:
            r = c.iloc[0]
            corr_rows.append({
                '診療行為コード': code,
                '診療行為名': CODE_LABEL.get(code, code),
                '年度': int(yr),
                '高齢化率_ρ': round(r['spearman_rho_aging'], 4),
                '高齢化率_p値': round(r['p_value_aging'], 6),
                '眼科医数_ρ': round(r['spearman_rho_docs'], 4),
                '眼科医数_p値': round(r['p_value_docs'], 6),
                '施設数_ρ': round(r['spearman_rho_facilities'], 4),
                '施設数_p値': round(r['p_value_facilities'], 6),
            })
df_corr = pd.DataFrame(corr_rows)

# ─── Sheet6: パネル回帰係数 ───
reg_rows = []
for code in CODES:
    rr = reg[reg['code'] == code]
    if rr.empty:
        continue
    r2 = rr.iloc[0]['r2_within']
    n = int(rr.iloc[0]['n_obs'])
    for _, rv in rr.iterrows():
        var_label = {
            'aging_rate': '高齢化率',
            'docs_per_100k': '眼科医数(per10万人)',
            'facilities_per_100k': '眼科施設数(per10万人)',
        }.get(rv['variable'], rv['variable'])
        reg_rows.append({
            '診療行為コード': code,
            '診療行為名': CODE_LABEL.get(code, code),
            '説明変数': var_label,
            '回帰係数': round(rv['coefficient'], 2),
            '標準誤差': round(rv['std_err'], 2),
            't値': round(rv['t_stat'], 2),
            'p値': round(rv['p_value'], 6),
            'R2_within': round(r2, 4),
            '観測数': n,
        })
df_reg = pd.DataFrame(reg_rows)

# ─── Excel出力（データ＋グラフ埋め込み） ───
out_path = 'data/processed/figure_data_for_excel.xlsx'
with pd.ExcelWriter(out_path, engine='openpyxl') as writer:
    # --- Sheet1: 全国トレンド ---
    df_trend.to_excel(writer, sheet_name='1_全国トレンド(縦)', index=False)
    df_trend_wide.to_excel(writer, sheet_name='1b_トレンド_per100k', index=False)
    df_trend_wide_count.to_excel(writer, sheet_name='1c_トレンド_件数', index=False)

    # --- Sheet2: APC ---
    df_apc.to_excel(writer, sheet_name='2_APC経年変化率', index=False)
    df_jp.to_excel(writer, sheet_name='2b_Joinpoint回帰', index=False)

    # --- Sheet3: 地域格差 ---
    df_disp.to_excel(writer, sheet_name='3_地域格差(縦)', index=False)
    df_gini_wide.to_excel(writer, sheet_name='3b_Gini推移', index=False)
    df_cv_wide.to_excel(writer, sheet_name='3c_CV推移', index=False)

    # --- Sheet4: 都道府県別 ---
    for code, df_pref in pref_rows:
        df_pref.to_excel(writer, sheet_name=f'4_{code}_都道府県2023', index=False)

    # --- Sheet5: Spearman ---
    df_corr.to_excel(writer, sheet_name='5_Spearman相関', index=False)

    # --- Sheet6: パネル回帰 ---
    df_reg.to_excel(writer, sheet_name='6_パネル回帰係数', index=False)

    wb = writer.book

    # ========================================================
    # グラフ1: 全国トレンド（人口10万対施行件数）折れ線グラフ
    # ========================================================
    ws1 = wb['1b_トレンド_per100k']
    n_rows = len(df_trend_wide)
    n_cols = len(df_trend_wide.columns)

    chart1 = LineChart()
    chart1.title = '人口10万対施行件数の推移'
    chart1.y_axis.title = '人口10万対施行件数'
    chart1.x_axis.title = '年度'
    chart1.style = 10
    chart1.width = 28
    chart1.height = 16

    cats = Reference(ws1, min_col=1, min_row=2, max_row=n_rows + 1)
    for col_idx in range(2, n_cols + 1):
        vals = Reference(ws1, min_col=col_idx, min_row=1, max_row=n_rows + 1)
        chart1.add_data(vals, titles_from_data=True)
    chart1.set_categories(cats)

    for s in chart1.series:
        s.graphicalProperties.line.width = 25000

    ws1.add_chart(chart1, f"A{n_rows + 4}")

    # ========================================================
    # グラフ1c: 全国トレンド（施行件数）折れ線グラフ
    # ========================================================
    ws1c = wb['1c_トレンド_件数']
    n_rows_c = len(df_trend_wide_count)
    n_cols_c = len(df_trend_wide_count.columns)

    chart1c = LineChart()
    chart1c.title = '全国施行件数の推移'
    chart1c.y_axis.title = '施行件数'
    chart1c.x_axis.title = '年度'
    chart1c.style = 10
    chart1c.width = 28
    chart1c.height = 16

    cats_c = Reference(ws1c, min_col=1, min_row=2, max_row=n_rows_c + 1)
    for col_idx in range(2, n_cols_c + 1):
        vals = Reference(ws1c, min_col=col_idx, min_row=1, max_row=n_rows_c + 1)
        chart1c.add_data(vals, titles_from_data=True)
    chart1c.set_categories(cats_c)

    for s in chart1c.series:
        s.graphicalProperties.line.width = 25000

    ws1c.add_chart(chart1c, f"A{n_rows_c + 4}")

    # ========================================================
    # グラフ2: APC棒グラフ
    # ========================================================
    ws2 = wb['2_APC経年変化率']
    n_apc = len(df_apc)

    chart2 = BarChart()
    chart2.type = "col"
    chart2.title = 'APC（年平均変化率）'
    chart2.y_axis.title = 'APC (%)'
    chart2.style = 10
    chart2.width = 22
    chart2.height = 14

    cats2 = Reference(ws2, min_col=2, min_row=2, max_row=n_apc + 1)
    vals2 = Reference(ws2, min_col=4, min_row=1, max_row=n_apc + 1)
    chart2.add_data(vals2, titles_from_data=True)
    chart2.set_categories(cats2)
    chart2.shape = 4

    ws2.add_chart(chart2, f"A{n_apc + 4}")

    # ========================================================
    # グラフ3b: Gini係数推移 折れ線グラフ
    # ========================================================
    ws3b = wb['3b_Gini推移']
    n_gini = len(df_gini_wide)
    n_gini_cols = len(df_gini_wide.columns)

    chart3b = LineChart()
    chart3b.title = 'Gini係数の推移（地域格差）'
    chart3b.y_axis.title = 'Gini係数'
    chart3b.x_axis.title = '年度'
    chart3b.style = 10
    chart3b.width = 28
    chart3b.height = 16

    cats3b = Reference(ws3b, min_col=1, min_row=2, max_row=n_gini + 1)
    for col_idx in range(2, n_gini_cols + 1):
        vals = Reference(ws3b, min_col=col_idx, min_row=1, max_row=n_gini + 1)
        chart3b.add_data(vals, titles_from_data=True)
    chart3b.set_categories(cats3b)

    for s in chart3b.series:
        s.graphicalProperties.line.width = 25000

    ws3b.add_chart(chart3b, f"A{n_gini + 4}")

    # ========================================================
    # グラフ3c: CV推移 折れ線グラフ
    # ========================================================
    ws3c = wb['3c_CV推移']
    n_cv = len(df_cv_wide)
    n_cv_cols = len(df_cv_wide.columns)

    chart3c = LineChart()
    chart3c.title = '変動係数(CV)の推移（地域格差）'
    chart3c.y_axis.title = '変動係数(CV)'
    chart3c.x_axis.title = '年度'
    chart3c.style = 10
    chart3c.width = 28
    chart3c.height = 16

    cats3c = Reference(ws3c, min_col=1, min_row=2, max_row=n_cv + 1)
    for col_idx in range(2, n_cv_cols + 1):
        vals = Reference(ws3c, min_col=col_idx, min_row=1, max_row=n_cv + 1)
        chart3c.add_data(vals, titles_from_data=True)
    chart3c.set_categories(cats3c)

    for s in chart3c.series:
        s.graphicalProperties.line.width = 25000

    ws3c.add_chart(chart3c, f"A{n_cv + 4}")

    # ========================================================
    # グラフ4: 都道府県別ランキング（2023, per100k）横棒グラフ
    # ========================================================
    for code, df_pref in pref_rows:
        sheet_name = f'4_{code}_都道府県2023'
        ws4 = wb[sheet_name]
        n_pref = len(df_pref)

        chart4 = BarChart()
        chart4.type = "bar"
        chart4.title = f'{CODE_LABEL.get(code, code)} 都道府県別施行率(2023)'
        chart4.x_axis.title = '人口10万対施行件数'
        chart4.style = 10
        chart4.width = 24
        chart4.height = max(20, n_pref * 0.5)

        # 都道府県名=col2, 人口10万対施行件数=col4
        cats4 = Reference(ws4, min_col=2, min_row=2, max_row=n_pref + 1)
        vals4 = Reference(ws4, min_col=4, min_row=1, max_row=n_pref + 1)
        chart4.add_data(vals4, titles_from_data=True)
        chart4.set_categories(cats4)
        chart4.legend = None

        ws4.add_chart(chart4, f"K1")

    # ========================================================
    # グラフ6: パネル回帰係数 棒グラフ（コード別×変数別）
    # ========================================================
    # 回帰係数をコード別にピボットして棒グラフにする
    if not df_reg.empty:
        df_reg_pivot = df_reg.pivot_table(
            index='診療行為名', columns='説明変数', values='回帰係数'
        ).reset_index()
        df_reg_pivot.columns.name = None

        reg_sheet_name = '6b_回帰係数グラフ用'
        df_reg_pivot.to_excel(writer, sheet_name=reg_sheet_name, index=False)
        ws6 = wb[reg_sheet_name]
        n_reg = len(df_reg_pivot)
        n_reg_cols = len(df_reg_pivot.columns)

        chart6 = BarChart()
        chart6.type = "col"
        chart6.title = 'パネル回帰係数（Two-way FE）'
        chart6.y_axis.title = '回帰係数'
        chart6.style = 10
        chart6.width = 22
        chart6.height = 14

        cats6 = Reference(ws6, min_col=1, min_row=2, max_row=n_reg + 1)
        for col_idx in range(2, n_reg_cols + 1):
            vals = Reference(ws6, min_col=col_idx, min_row=1, max_row=n_reg + 1)
            chart6.add_data(vals, titles_from_data=True)
        chart6.set_categories(cats6)

        ws6.add_chart(chart6, f"A{n_reg + 4}")

print(f'Excel出力完了（グラフ埋め込み済み）: {out_path}')

# CSVも個別出力（必要な人向け）
df_trend.to_csv('data/processed/fig1_national_trends.csv', index=False, encoding='utf-8-sig')
df_apc.to_csv('data/processed/fig2_apc_results.csv', index=False, encoding='utf-8-sig')
df_disp.to_csv('data/processed/fig3_geographic_disparity.csv', index=False, encoding='utf-8-sig')
df_corr.to_csv('data/processed/fig5_spearman_correlation.csv', index=False, encoding='utf-8-sig')
df_reg.to_csv('data/processed/fig6_panel_regression.csv', index=False, encoding='utf-8-sig')
print('CSV個別出力完了 (fig1_*.csv - fig6_*.csv)')
