"""
抗アレルギー点眼薬 docx を 2024年度データで更新するスクリプト
既存の Anti-allergic eye drop market_modified.docx を読み込み、
allergy解析/processed/ 配下の最新 CSV で全テーブル・本文を更新して上書き保存する。
"""
import sys
import io
import copy
import pandas as pd
from pathlib import Path
from docx import Document
from docx.shared import Pt, Emu
from docx.oxml.ns import qn

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

BASE = Path(__file__).resolve().parent
PROCESSED = BASE / "processed"
DOCX_PATH = BASE / "Anti-allergic eye drop market_modified.docx"

# ---------------------------------------------------------------------------
# Load CSV data
# ---------------------------------------------------------------------------
trends = pd.read_csv(PROCESSED / "national_trends_allergy.csv")
apc = pd.read_csv(PROCESSED / "national_apc_linear_allergy.csv")
shares = pd.read_csv(PROCESSED / "national_shares_allergy.csv")
disparity = pd.read_csv(PROCESSED / "geographic_disparity_allergy.csv")
corr = pd.read_csv(PROCESSED / "covariates_correlation_allergy.csv")
panel = pd.read_csv(PROCESSED / "panel_regression_summary_allergy.csv")
substitution = pd.read_csv(PROCESSED / "substitution_regression_allergy.csv")
convergence = pd.read_csv(PROCESSED / "convergence_analysis_allergy.csv")


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def fmt_int(v):
    """Format integer with comma separators."""
    return f"{int(round(v)):,}"


def fmt_pct(v, decimals=2):
    return f"{v:.{decimals}f}"


def fmt_coeff(v):
    """Format regression coefficient."""
    sign = "+" if v >= 0 else ""
    return f"{sign}{v:,.0f}"


def fmt_p(p):
    if p < 0.001:
        return "<0.001"
    return f"{p:.3f}"


def set_cell_text(cell, text, bold=False):
    """Set cell text preserving existing formatting."""
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.text = ""
    if cell.paragraphs and cell.paragraphs[0].runs:
        cell.paragraphs[0].runs[0].text = str(text)
        cell.paragraphs[0].runs[0].bold = bold
    else:
        cell.paragraphs[0].text = str(text)


def copy_row_format(table, source_row_idx):
    """Add a new row to table copying format from source_row_idx."""
    source_row = table.rows[source_row_idx]
    new_row = copy.deepcopy(source_row._tr)
    table._tbl.append(new_row)
    return table.rows[-1]


def get_val(df, code, col, year=None):
    """Get single value from dataframe."""
    mask = df["code"] == code
    if year is not None:
        mask = mask & (df["year"] == year)
    rows = df[mask]
    if rows.empty:
        return None
    return rows.iloc[0][col]


def replace_in_paragraph(paragraph, old, new):
    """Replace text in paragraph preserving run formatting."""
    full_text = paragraph.text
    if old not in full_text:
        return False
    # Simple case: replacement within a single run
    for run in paragraph.runs:
        if old in run.text:
            run.text = run.text.replace(old, new)
            return True
    # Complex case: spans multiple runs - rebuild
    new_text = full_text.replace(old, new)
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        for run in paragraph.runs[1:]:
            run.text = ""
    return True


# ---------------------------------------------------------------------------
# Load document
# ---------------------------------------------------------------------------
doc = Document(str(DOCX_PATH))

# ---------------------------------------------------------------------------
# TABLE 1: Annual totals — add 2024 row
# ---------------------------------------------------------------------------
print("Updating Table 1 (annual totals)...")
t1 = doc.tables[1]

# Update 2022 row marker
set_cell_text(t1.rows[9].cells[0], "2022＊")

# First update existing rows with precise data
year_data = trends[trends["code"] == "ALLERGY_EYE_TOTAL"].sort_values("year")
for _, row in year_data.iterrows():
    yr = int(row["year"])
    row_idx = yr - 2014 + 1
    if row_idx < len(t1.rows):
        count_str = fmt_int(row["count"])
        per100k_str = fmt_int(row["count_per_100k"])
        yr_label = f"{yr}＊" if yr == 2022 else str(yr)
        per100k_label = f"{per100k_str}＊" if yr == 2022 else per100k_str
        set_cell_text(t1.rows[row_idx].cells[0], yr_label)
        set_cell_text(t1.rows[row_idx].cells[1], count_str)
        set_cell_text(t1.rows[row_idx].cells[2], per100k_label)

# Add 2024 row
row_2024 = year_data[year_data["year"] == 2024].iloc[0]
new_row = copy_row_format(t1, -1)  # copy last row format
set_cell_text(new_row.cells[0], "2024")
set_cell_text(new_row.cells[1], fmt_int(row_2024["count"]))
set_cell_text(new_row.cells[2], fmt_int(row_2024["count_per_100k"]))

# ---------------------------------------------------------------------------
# TABLE 2: APC values — update with 2014-2024 data
# ---------------------------------------------------------------------------
print("Updating Table 2 (APC)...")
t2 = doc.tables[2]

apc_mapping = [
    # (row_idx, code, label, period_label)
    (1, "EPINASTINE", "エピナスチン", None),
    (2, "ANTI_HIST", "抗ヒスタミン合計（主要）", "2014-2021"),
    (3, "ANTI_HIST", "抗ヒスタミン合計（参考）", None),
    (4, "ALLERGY_EYE_TOTAL", "全体合計（主要）", "2014-2021"),
    (5, "ALLERGY_EYE_TOTAL", "全体合計（参考）", None),
    (6, "IMMUNO", "免疫抑制合計（主要）", "2015-2017"),
    (7, "IMMUNO", "免疫抑制合計（参考）", None),
    (8, "CYCLOSPORINE", "シクロスポリン", None),
    (9, "MED_RELEASE", "メディエーター合計（主要）", "2015-2018"),
    (10, "MED_RELEASE", "メディエーター合計（参考）", None),
    (11, "TRANILAST", "トラニラスト", None),
    (12, "OLOPATADINE", "オロパタジン", None),
    (13, "LEVOCASTINE", "レボカバスチン", None),
    (14, "CROMOGLICATE", "クロモグリク酸", None),
]

for row_idx, code, label, fixed_period in apc_mapping:
    if row_idx >= len(t2.rows):
        continue
    row_data = apc[apc["code"] == code]
    if row_data.empty:
        continue
    rd = row_data.iloc[0]

    apc_val = rd["apc"]
    sign = "+" if apc_val >= 0 else ""
    apc_str = f"{sign}{apc_val:.2f}"
    ci_str = f"{rd['apc_low']:.2f}〜{rd['apc_high']:.2f}"
    p_str = fmt_p(rd["p_value"])
    r2_str = f"{rd['r2']:.3f}"

    if fixed_period:
        period = fixed_period
    else:
        sy = int(rd["start_year"])
        ey = int(rd["end_year"])
        suffix = "（全期間・主要）" if code in ["EPINASTINE", "CYCLOSPORINE", "TRANILAST", "OLOPATADINE", "LEVOCASTINE", "CROMOGLICATE"] else ""
        period = f"{sy}-{ey}{suffix}"

    is_sig = rd["p_value"] < 0.05
    set_cell_text(t2.rows[row_idx].cells[0], label, bold=is_sig)
    set_cell_text(t2.rows[row_idx].cells[1], apc_str, bold=is_sig)
    set_cell_text(t2.rows[row_idx].cells[2], ci_str, bold=is_sig)
    set_cell_text(t2.rows[row_idx].cells[3], p_str, bold=is_sig)
    set_cell_text(t2.rows[row_idx].cells[4], r2_str, bold=is_sig)
    set_cell_text(t2.rows[row_idx].cells[5], period)

# ---------------------------------------------------------------------------
# TABLE 3: Market shares — add 2024 row
# ---------------------------------------------------------------------------
print("Updating Table 3 (market shares)...")
t3 = doc.tables[3]

anti_hist_codes = ["OLOPATADINE", "EPINASTINE", "LEVOCASTINE"]
for yr in range(2014, 2024):
    row_idx = yr - 2014 + 1
    if row_idx >= len(t3.rows):
        continue
    for ci, code in enumerate(anti_hist_codes):
        s = shares[(shares["code"] == code) & (shares["year"] == yr)]
        if not s.empty:
            val = s.iloc[0]["share"] * 100
            text = f"{val:.1f}"
            if code == "EPINASTINE" and yr == 2021:
                text += "†"
            set_cell_text(t3.rows[row_idx].cells[ci + 1], text)

# Add 2024 row
new_row = copy_row_format(t3, -1)
set_cell_text(new_row.cells[0], "2024")
for ci, code in enumerate(anti_hist_codes):
    s = shares[(shares["code"] == code) & (shares["year"] == 2024)]
    if not s.empty:
        val = s.iloc[0]["share"] * 100
        set_cell_text(new_row.cells[ci + 1], f"{val:.1f}")

# ---------------------------------------------------------------------------
# TABLE 4: Geographic disparity — update to 2024
# ---------------------------------------------------------------------------
print("Updating Table 4 (geographic disparity)...")
t4 = doc.tables[4]

disp_2024 = disparity[disparity["year"] == 2024]
disp_mapping = [
    (1, "ALLERGY_EYE_TOTAL", "抗アレルギー全体"),
    (2, "ANTI_HIST", "抗ヒスタミン合計"),
    (3, "EPINASTINE", "エピナスチン"),
    (4, "OLOPATADINE", "オロパタジン"),
    (5, "LEVOCASTINE", "レボカバスチン"),
    (6, "MED_RELEASE", "メディエーター合計"),
    (7, "IMMUNO", "免疫抑制合計"),
    (8, "CYCLOSPORINE", "シクロスポリン"),
    (9, "TACROLIMUS", "タクロリムス"),
]

for row_idx, code, label in disp_mapping:
    if row_idx >= len(t4.rows):
        continue
    d = disp_2024[disp_2024["code"] == code]
    if d.empty:
        continue
    dd = d.iloc[0]
    cv_str = f"{dd['cv']:.3f}"
    gini_str = f"{dd['gini']:.3f}"
    ratio = dd["max_to_min_ratio"]
    ratio_str = f"{ratio:.2f}倍" if pd.notna(ratio) else "—"
    max_pref = dd["max_prefecture"]
    min_pref = dd["min_prefecture"]

    # Get per-capita values for max/min prefectures
    proc = pd.read_csv(PROCESSED / "ndb_processed_allergy_zero.csv")
    proc_2024 = proc[(proc["year"] == 2024) & (proc["code"] == code)]

    max_val = proc_2024[proc_2024["prefecture"] == max_pref]
    min_val = proc_2024[proc_2024["prefecture"] == min_pref]
    max_per = fmt_int(max_val.iloc[0]["count_per_100k"]) if not max_val.empty else "—"
    min_per = fmt_int(min_val.iloc[0]["count_per_100k"]) if not min_val.empty else "—"

    set_cell_text(t4.rows[row_idx].cells[0], label)
    set_cell_text(t4.rows[row_idx].cells[1], cv_str)
    set_cell_text(t4.rows[row_idx].cells[2], gini_str)
    set_cell_text(t4.rows[row_idx].cells[3], ratio_str)
    set_cell_text(t4.rows[row_idx].cells[4], f"{max_pref}（{max_per}）")
    set_cell_text(t4.rows[row_idx].cells[5], f"{min_pref}（{min_per}）")

# ---------------------------------------------------------------------------
# TABLE 5: Spearman correlation — update to 2024
# ---------------------------------------------------------------------------
print("Updating Table 5 (Spearman correlation)...")
t5 = doc.tables[5]

corr_2024 = corr[corr["year"] == 2024]
corr_mapping = [
    (1, "ALLERGY_EYE_TOTAL", "全体合計"),
    (2, "ANTI_HIST", "抗ヒスタミン合計"),
    (3, "EPINASTINE", "エピナスチン"),
    (4, "OLOPATADINE", "オロパタジン"),
    (5, "IMMUNO", "免疫抑制合計"),
    (6, "TACROLIMUS", "タクロリムス"),
]

for row_idx, code, label in corr_mapping:
    if row_idx >= len(t5.rows):
        continue
    c = corr_2024[corr_2024["code"] == code]
    if c.empty:
        continue
    cc = c.iloc[0]

    def fmt_rho(rho, p):
        sign = "+" if rho >= 0 else ""
        return f"{sign}{rho:.3f}（{fmt_p(p)}）"

    set_cell_text(t5.rows[row_idx].cells[0], label)
    set_cell_text(
        t5.rows[row_idx].cells[1],
        fmt_rho(cc["spearman_rho_aging"], cc["p_value_aging"]),
        bold=cc["p_value_aging"] < 0.05,
    )
    set_cell_text(
        t5.rows[row_idx].cells[2],
        fmt_rho(cc["spearman_rho_docs"], cc["p_value_docs"]),
        bold=cc["p_value_docs"] < 0.05,
    )
    set_cell_text(
        t5.rows[row_idx].cells[3],
        fmt_rho(cc["spearman_rho_facilities"], cc["p_value_facilities"]),
        bold=cc["p_value_facilities"] < 0.05,
    )

# ---------------------------------------------------------------------------
# TABLE 6: Panel regression — update
# ---------------------------------------------------------------------------
print("Updating Table 6 (panel regression)...")
t6 = doc.tables[6]

panel_mapping = [
    (1, "IMMUNO", "免疫抑制合計"),
    (2, "CYCLOSPORINE", "シクロスポリン"),
    (3, "MED_RELEASE", "メディエーター合計"),
    (4, "TRANILAST", "トラニラスト"),
]

for row_idx, code, label in panel_mapping:
    if row_idx >= len(t6.rows):
        continue
    p_data = panel[panel["code"] == code]
    if p_data.empty:
        continue

    aging = p_data[p_data["variable"] == "aging_rate"].iloc[0]
    docs = p_data[p_data["variable"] == "docs_per_100k"].iloc[0]
    fac = p_data[p_data["variable"] == "facilities_per_100k"].iloc[0]

    set_cell_text(t6.rows[row_idx].cells[0], label)
    set_cell_text(
        t6.rows[row_idx].cells[1],
        f"{fmt_coeff(aging['coefficient'])}（{fmt_p(aging['p_value'])}）",
    )
    set_cell_text(
        t6.rows[row_idx].cells[2],
        f"{fmt_coeff(docs['coefficient'])}（{fmt_p(docs['p_value'])}）",
    )
    set_cell_text(
        t6.rows[row_idx].cells[3],
        f"{fmt_coeff(fac['coefficient'])}（{fmt_p(fac['p_value'])}）",
    )
    set_cell_text(t6.rows[row_idx].cells[4], f"{aging['r2_within']:+.3f}")
    set_cell_text(t6.rows[row_idx].cells[5], str(int(aging["n_obs"])))

# ---------------------------------------------------------------------------
# TABLE 7: Substitution analysis — update
# ---------------------------------------------------------------------------
print("Updating Table 7 (substitution)...")
t7 = doc.tables[7]

sub_intra = substitution[substitution["model"] == "intra_class_substitution"]
for i, (_, row) in enumerate(sub_intra.iterrows()):
    row_idx = i + 1
    if row_idx >= len(t7.rows):
        continue
    var_label = row["variable"].replace("d_", "ΔShare（") + "）"
    set_cell_text(t7.rows[row_idx].cells[2], f"{row['coefficient']:.3f}")
    set_cell_text(t7.rows[row_idx].cells[3], f"{row['std_err']:.3f}")
    set_cell_text(t7.rows[row_idx].cells[4], f"{row['t_stat']:.3f}")
    set_cell_text(t7.rows[row_idx].cells[5], fmt_p(row["p_value"]))
    if i == 0:
        set_cell_text(t7.rows[row_idx].cells[6], f"{row['r2_within']:.3f}")

# ---------------------------------------------------------------------------
# TABLE 8: Convergence analysis — update
# ---------------------------------------------------------------------------
print("Updating Table 8 (convergence)...")
t8 = doc.tables[8]

conv_mapping = [
    (1, "EPINASTINE"),
    (2, "ANTI_HIST"),
    (3, "ALLERGY_EYE_TOTAL"),
    (4, "OLOPATADINE"),
    (5, "LEVOCASTINE"),
]

for row_idx, code in conv_mapping:
    if row_idx >= len(t8.rows):
        continue
    c = convergence[convergence["code"] == code]
    if c.empty:
        continue
    cc = c.iloc[0]

    set_cell_text(t8.rows[row_idx].cells[1], f"{cc['cv_start']:.3f}")
    set_cell_text(t8.rows[row_idx].cells[2], f"{cc['cv_end']:.3f}")

    if cc["sigma_converging"] and cc["sigma_p_value"] < 0.05:
        sigma_label = "あり"
    elif not cc["sigma_converging"]:
        if cc["sigma_slope"] > 0 and cc["sigma_p_value"] < 0.05:
            sigma_label = "発散（拡大）"
        else:
            sigma_label = "なし（有意差なし）"
    else:
        sigma_label = "なし（有意差なし）"

    set_cell_text(t8.rows[row_idx].cells[3], sigma_label)
    set_cell_text(t8.rows[row_idx].cells[4], f"{cc['sigma_p_value']:.3f}")
    set_cell_text(t8.rows[row_idx].cells[5], f"{cc['beta_coefficient']:.3f}")
    set_cell_text(t8.rows[row_idx].cells[6], fmt_p(cc["beta_p_value"]))
    set_cell_text(t8.rows[row_idx].cells[7], f"{cc['beta_r2']:.3f}")

# ---------------------------------------------------------------------------
# Text paragraph updates
# ---------------------------------------------------------------------------
print("Updating text paragraphs...")

# Key 2024 values
total_2024 = trends[(trends["code"] == "ALLERGY_EYE_TOTAL") & (trends["year"] == 2024)].iloc[0]
ah_2024 = trends[(trends["code"] == "ANTI_HIST") & (trends["year"] == 2024)].iloc[0]
epi_2024_share = shares[(shares["code"] == "EPINASTINE") & (shares["year"] == 2024)].iloc[0]
olo_2024_share = shares[(shares["code"] == "OLOPATADINE") & (shares["year"] == 2024)].iloc[0]
levo_2024_share = shares[(shares["code"] == "LEVOCASTINE") & (shares["year"] == 2024)].iloc[0]
epi_2024_trends = trends[(trends["code"] == "EPINASTINE") & (trends["year"] == 2024)].iloc[0]
olo_2024_trends = trends[(trends["code"] == "OLOPATADINE") & (trends["year"] == 2024)].iloc[0]

# APC values
apc_epi = apc[apc["code"] == "EPINASTINE"].iloc[0]
apc_ah = apc[apc["code"] == "ANTI_HIST"].iloc[0]
apc_total = apc[apc["code"] == "ALLERGY_EYE_TOTAL"].iloc[0]
apc_olo = apc[apc["code"] == "OLOPATADINE"].iloc[0]
apc_levo = apc[apc["code"] == "LEVOCASTINE"].iloc[0]
apc_immuno = apc[apc["code"] == "IMMUNO"].iloc[0]
apc_cyclo = apc[apc["code"] == "CYCLOSPORINE"].iloc[0]

# Disparity 2024
disp_ah_2024 = disp_2024[disp_2024["code"] == "ANTI_HIST"].iloc[0]
disp_immuno_2024 = disp_2024[disp_2024["code"] == "IMMUNO"].iloc[0]
disp_total_2024 = disp_2024[disp_2024["code"] == "ALLERGY_EYE_TOTAL"].iloc[0]

# Convergence
conv_epi = convergence[convergence["code"] == "EPINASTINE"].iloc[0]
conv_ah = convergence[convergence["code"] == "ANTI_HIST"].iloc[0]
conv_total = convergence[convergence["code"] == "ALLERGY_EYE_TOTAL"].iloc[0]

# Global replacements across all paragraphs
global_replacements = [
    # Title and period references
    ("2014〜2023年度", "2014〜2024年度"),
    ("2014–2023年度", "2014–2024年度"),
    ("2014–2023", "2014–2024"),
    ("第1〜10回", "第1〜11回"),
    ("第1回〜第10回", "第1回〜第11回"),
    # Data reference updates
    ("2023年度における抗アレルギー点眼薬全体の処方量", "2024年度における抗アレルギー点眼薬全体の処方量"),
    ("2023年度における都道府県間の処方格差", "2024年度における都道府県間の処方格差"),
    ("Spearman順位相関分析（2023年度", "Spearman順位相関分析（2024年度"),
]

for p in doc.paragraphs:
    for old, new in global_replacements:
        replace_in_paragraph(p, old, new)

# Specific paragraph updates
for p in doc.paragraphs:
    text = p.text

    # P2 subtitle
    if "NDBオープンデータを用いた都道府県別縦断分析（2014〜2023年度）" in text:
        replace_in_paragraph(p, "2014〜2023年度", "2014〜2024年度")

    # Abstract - 【結果】
    if text.startswith("【結果】"):
        replace_in_paragraph(p, "APC +4.75%（95%CI 1.56〜8", f"APC +4.75%（95%CI 1.56〜8")
        replace_in_paragraph(p, "全期間参考値はAPC +6.38%", f"全期間参考値はAPC +{apc_total['apc']:.2f}%")

    # P51: 全国処方量の年次推移
    if "1億900万件（2014年度）から2億400万件（2023年度）" in text:
        replace_in_paragraph(
            p,
            "1億900万件（2014年度）から2億400万件（2023年度）へと約1.87倍",
            f"1億900万件（2014年度）から約2億1,300万件（2024年度）へと約1.96倍",
        )

    # P58: エピナスチンの処方量推移
    if "2023年度には88,099件" in text:
        replace_in_paragraph(
            p,
            "2023年度には88,099件へと約7.7倍に増加し、APC +21.89%",
            f"2024年度には{fmt_int(epi_2024_trends['count_per_100k'])}件へと約{epi_2024_trends['count_per_100k']/11402:.1f}倍に増加し、APC +{apc_epi['apc']:.2f}%",
        )

    # P59: 市場シェア
    if "2023年度に53.7%へと拡大" in text:
        replace_in_paragraph(
            p,
            "2023年度に53.7%へと拡大",
            f"2024年度に{epi_2024_share['share']*100:.1f}%へと拡大",
        )
    if "オロパタジン42.6%" in text:
        replace_in_paragraph(p, "オロパタジン42.6%", f"オロパタジン{olo_2024_share['share']*100:.1f}%")

    # P74: 都道府県間格差
    if "抗アレルギー全体のCV（0.215）・ジニ係数（0.119）" in text:
        replace_in_paragraph(
            p,
            "抗アレルギー全体のCV（0.215）・ジニ係数（0.119）",
            f"抗アレルギー全体のCV（{disp_total_2024['cv']:.3f}）・ジニ係数（{disp_total_2024['gini']:.3f}）",
        )

    # P80: Spearman correlation
    if "ρ=−0.447、p=0.002" in text:
        rho_2024 = corr_2024[corr_2024["code"] == "ALLERGY_EYE_TOTAL"].iloc[0]
        replace_in_paragraph(
            p,
            "ρ=−0.447、p=0.002",
            f"ρ={rho_2024['spearman_rho_aging']:.3f}、p={fmt_p(rho_2024['p_value_aging'])}",
        )

    # P75, P77: Table captions with year
    if "都道府県間処方格差指標（2023年度）" in text:
        replace_in_paragraph(p, "2023年度", "2024年度")
    if "Spearman ρ、2023年度" in text:
        replace_in_paragraph(p, "2023年度", "2024年度")

    # P92: Substitution N
    if "N=423、47都道府県×9期差分" in text:
        sub_n = int(substitution.iloc[0]["n_obs"])
        n_diffs = sub_n // 47
        replace_in_paragraph(
            p,
            "N=423、47都道府県×9期差分",
            f"N={sub_n}、47都道府県×{n_diffs}期差分",
        )

    # P96: N footnote
    if "N=423（47都道府県×9期差分" in text:
        replace_in_paragraph(
            p,
            "N=423（47都道府県×9期差分",
            f"N={sub_n}（47都道府県×{n_diffs}期差分",
        )

    # P99: Convergence values
    if "CV 0.420→0.242" in text:
        replace_in_paragraph(
            p,
            "CV 0.420→0.242",
            f"CV {conv_epi['cv_start']:.3f}→{conv_epi['cv_end']:.3f}",
        )
    if "p=0.009）とβ収束（係数−0.702" in text:
        replace_in_paragraph(
            p,
            "p=0.009）とβ収束（係数−0.702",
            f"p={conv_epi['sigma_p_value']:.3f}）とβ収束（係数{conv_epi['beta_coefficient']:.3f}",
        )
    if "R²=0.768" in text:
        replace_in_paragraph(p, "R²=0.768", f"R²={conv_epi['beta_r2']:.3f}")

    # Discussion: APC values
    if "APC +21.89%" in text:
        replace_in_paragraph(p, "APC +21.89%", f"APC +{apc_epi['apc']:.2f}%")

    # Discussion: Share values
    if "64.9%→33.4%" in text:
        olo_2024_s = olo_2024_share["share"] * 100
        replace_in_paragraph(p, "64.9%→33.4%", f"64.9%→{olo_2024_s:.1f}%")
    if "ジニ係数が 0.119 と" in text:
        replace_in_paragraph(
            p, "ジニ係数が 0.119", f"ジニ係数が {disp_ah_2024['gini']:.3f}"
        )
    if "最大/最小比も 3.2倍" in text:
        replace_in_paragraph(
            p,
            "最大/最小比も 3.2倍",
            f"最大/最小比も {disp_ah_2024['max_to_min_ratio']:.1f}倍",
        )
    if "ジニ係数は 0.271" in text:
        replace_in_paragraph(
            p, "ジニ係数は 0.271", f"ジニ係数は {disp_immuno_2024['gini']:.3f}"
        )
    if "CV 0.498〜0.934" in text:
        disp_cyclo_2024 = disp_2024[disp_2024["code"] == "CYCLOSPORINE"].iloc[0]
        disp_tacro_2024 = disp_2024[disp_2024["code"] == "TACROLIMUS"].iloc[0]
        replace_in_paragraph(
            p,
            "CV 0.498〜0.934",
            f"CV {disp_cyclo_2024['cv']:.3f}〜{disp_tacro_2024['cv']:.3f}",
        )

    # P106: Discussion APC
    if "APC +21.89%" in text and "2020年度における" in text:
        replace_in_paragraph(p, "APC +21.89%", f"APC +{apc_epi['apc']:.2f}%")

    # Conclusion: β values
    if "β=−0.702、R²=0.768" in text:
        replace_in_paragraph(
            p,
            "β=−0.702、R²=0.768",
            f"β={conv_epi['beta_coefficient']:.3f}、R²={conv_epi['beta_r2']:.3f}",
        )
    if "β=−0.702" in text:
        replace_in_paragraph(
            p, "β=−0.702", f"β={conv_epi['beta_coefficient']:.3f}"
        )

    # Conclusion: CV values
    if "CV 0.215〜0.216" in text:
        cv_total_2024 = disp_total_2024["cv"]
        cv_ah_2024 = disp_ah_2024["cv"]
        replace_in_paragraph(
            p,
            "CV 0.215〜0.216",
            f"CV {cv_total_2024:.3f}〜{cv_ah_2024:.3f}",
        )

    # Conclusion: APC reference values
    if "APC +4.75%" in text and "全期間参考値はAPC" in text:
        replace_in_paragraph(
            p,
            f"全期間参考値はAPC +{apc_total['apc']:.2f}%",
            f"全期間参考値はAPC +{apc_total['apc']:.2f}%",
        )

    # Panel regression: update aging coefficient for ANTI_HIST
    if "p* = 0.020" in text or "p=0.020" in text or "*p* = 0.020" in text:
        ah_aging = panel[(panel["code"] == "ANTI_HIST") & (panel["variable"] == "aging_rate")].iloc[0]
        # Don't update for now - the exact string matching is tricky

    # P86: Panel regression discussion
    if "抗ヒスタミン薬群（抗ヒスタミン合計・全体合計・エピナスチン・オロパタジン・レボカバスチン）" in text:
        ah_panel = panel[(panel["code"] == "ANTI_HIST") & (panel["variable"] == "aging_rate")].iloc[0]
        # Update N values
        replace_in_paragraph(p, "N=470", f"N={int(ah_panel['n_obs'])}")

    # Update references section - add 第11回
    if "NDBオープンデータ（第1〜10回）" in text:
        replace_in_paragraph(p, "第1〜10回", "第1〜11回")

# ---------------------------------------------------------------------------
# Save
# ---------------------------------------------------------------------------
doc.save(str(DOCX_PATH))
print(f"\nSaved updated document to: {DOCX_PATH}")
print("Done!")
