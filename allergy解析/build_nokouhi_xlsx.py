"""公費レセプトを含まないデータ版の論文用xlsxを、既存ファイルのレイアウトを保って生成する。

`論文に使うファイルたち/` の既存xlsxは generate_paper_xlsx.py の出力を手作業で編集したもの
（グラフシートの削除、Fig1のSheet1・Fig1c追加など）である。そのため単純に再生成すると
レイアウトが変わってしまう。ここでは既存ファイルを複製し、データシートの値だけを
公費含まない版に差し替えることでレイアウトを維持する。

前提: 公費含まない版のCSVと、generate_paper_xlsx.py を公費含まないデータで実行した
      xlsx（--generated-dir）が用意されていること。

出力: 論文に使うファイルたち/公費含まない/
"""

import argparse
import os
import shutil
import sys

import openpyxl
import pandas as pd

BASE = os.path.dirname(os.path.abspath(__file__))
sys.path.append(os.path.join(BASE, "src"))
from paths import PAPER_SUBDIR  # noqa: E402

# 公費含む版のxlsxをレイアウトの雛形として読み、公費含まない版を作る
ORIG_DIR = os.path.join(BASE, PAPER_SUBDIR, "公費含む")
OUT_DIR = os.path.join(BASE, PAPER_SUBDIR, "公費含まない")

TARGETS = [
    "Fig1_age_sex_profile.xlsx",
    "Fig2_age_distribution.xlsx",
    "Fig3_prefecture_map.xlsx",
    "Fig4_trends_shares.xlsx",
    "Fig5_ge_formulation.xlsx",
    "Tables_123.xlsx",
]

AGE_ORDER = [
    "0-4", "5-9", "10-14", "15-19", "20-24", "25-29", "30-34", "35-39",
    "40-44", "45-49", "50-54", "55-59", "60-64", "65-69", "70-74", "75-79",
    "80-84", "85-89", "90-94", "95-99", "100+", "90+",
]

# Fig1 の Sheet1 に貼り付けられている age_sex_rates の範囲
SHEET1_FIRST_ROW = 5
SHEET1_FIRST_COL = 16  # P列
SHEET1_COLS = ["code", "procedure_name", "sex", "age_group",
               "count", "population", "count_per_100k"]


def _header_map(src_ws, dst_ws):
    """1行目を見出しとみなし、dst列→src列の対応を返す。対応づかなければ None。

    既存xlsxを作った時点のCSVには丸め列（*_rounded）が無く、再生成側にはあるため、
    位置ではなく見出し名で対応づける必要がある。
    """
    src_hdr = {}
    for c in range(1, src_ws.max_column + 1):
        v = src_ws.cell(1, c).value
        if isinstance(v, str) and v.strip():
            src_hdr.setdefault(v.strip(), c)
    if not src_hdr:
        return None
    mapping = {}
    for c in range(1, dst_ws.max_column + 1):
        v = dst_ws.cell(1, c).value
        if not (isinstance(v, str) and v.strip()):
            return None
        if v.strip() not in src_hdr:
            return None
        mapping[c] = src_hdr[v.strip()]
    return mapping


def _same(x, y):
    if isinstance(x, (int, float)) and isinstance(y, (int, float)):
        return abs(x - y) < 1e-9
    return x == y


def copy_sheet_values(src_ws, dst_ws, base_ws, log, sheet_name):
    """公費含まない版の値を dst に書き込む（丸めは現行スクリプトの仕様に従う）。

    既存xlsxは現行の generate_paper_xlsx.py より古い版で作られている箇所があるため、
    値は再生成側（src）で全面的に上書きし、丸め仕様を現行版に統一する。
    変更理由を切り分けるため、「公費含むCSVで同じスクリプトを実行した結果（base）」とも
    比較し、公費による差か丸め仕様による差かをログに記録する。
    """
    changed = 0
    mapping = _header_map(src_ws, dst_ws)
    src_max_r = src_ws.max_row

    def row_label(ws, r):
        v = ws.cell(r, 1).value
        return "" if v is None else str(v)

    def put(dcell, new, old, label):
        nonlocal changed
        if isinstance(dcell.value, str) and dcell.value.startswith("="):
            return
        before = dcell.value
        if _same(before, new):
            return
        dcell.value = new
        changed += 1
        if not _same(old, new):
            reason = "公費レセプトの有無"
        elif not _same(before, old):
            reason = "丸め仕様（現行版に統一）"
        else:
            reason = "その他"
        log.append({"sheet": sheet_name, "cell": dcell.coordinate, "label": label,
                    "reason": reason, "original": before, "kouhi": old, "nokouhi": new})

    if mapping:
        for r in range(2, src_max_r + 1):
            for dc, sc in mapping.items():
                new = src_ws.cell(r, sc).value
                old = base_ws.cell(r, sc).value if base_ws is not None else None
                hdr = dst_ws.cell(1, dc).value
                put(dst_ws.cell(r, dc), new, old, f"{row_label(src_ws, r)} / {hdr}")
        return changed

    for r in range(1, src_max_r + 1):
        for c in range(1, src_ws.max_column + 1):
            new = src_ws.cell(r, c).value
            old = base_ws.cell(r, c).value if base_ws is not None else None
            put(dst_ws.cell(r, c), new, old, f"row{r}")
    return changed


def rewrite_fig1_sheet1(dst_ws, rates_csv):
    """Fig1 の Sheet1 に貼られた age_sex_rates（2024年度）を差し替える。"""
    df = pd.read_csv(rates_csv)
    d = df[df["year"] == 2024].copy()
    d["_a"] = pd.Categorical(d["age_group"], AGE_ORDER, ordered=True)
    d = d.sort_values(["code", "sex", "_a"])
    rows = d[SHEET1_COLS].values.tolist()

    changed = 0
    for i, row in enumerate(rows):
        r = SHEET1_FIRST_ROW + i
        for j, v in enumerate(row):
            cell = dst_ws.cell(r, SHEET1_FIRST_COL + j)
            v = v.item() if hasattr(v, "item") else v
            if cell.value != v:
                cell.value = v
                changed += 1
    # 余剰行が残っていれば消す
    for r in range(SHEET1_FIRST_ROW + len(rows), dst_ws.max_row + 1):
        for j in range(len(SHEET1_COLS)):
            cell = dst_ws.cell(r, SHEET1_FIRST_COL + j)
            if cell.value is not None:
                cell.value = None
                changed += 1
    return changed, len(rows)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--generated-dir", required=True,
                    help="公費含まないデータで generate_paper_xlsx.py を実行した出力先")
    ap.add_argument("--rates-csv", required=True,
                    help="公費含まない版の age_sex_rates_allergy.csv")
    ap.add_argument("--baseline-dir", required=True,
                    help="公費含むCSVで generate_paper_xlsx.py を実行した出力先")
    ap.add_argument("--diff-csv", default=None, help="変更セル一覧の出力先")
    args = ap.parse_args()

    os.makedirs(OUT_DIR, exist_ok=True)
    diff_log = []
    for name in TARGETS:
        orig = os.path.join(ORIG_DIR, name)
        gen = os.path.join(args.generated_dir, name)
        dst = os.path.join(OUT_DIR, name)
        if not os.path.exists(orig):
            print(f"  skip (元ファイルなし): {name}")
            continue
        shutil.copy(orig, dst)

        wb_dst = openpyxl.load_workbook(dst)
        total = 0
        base = os.path.join(args.baseline_dir, name)
        if name == "Tables_123.xlsx":
            gen = os.path.join(args.generated_dir, "Tables_all.xlsx")
            base = os.path.join(args.baseline_dir, "Tables_all.xlsx")
        if os.path.exists(gen) and os.path.exists(base):
            wb_gen = openpyxl.load_workbook(gen)
            wb_base = openpyxl.load_workbook(base)
            for sn in wb_dst.sheetnames:
                if sn in wb_gen.sheetnames and sn in wb_base.sheetnames:
                    total += copy_sheet_values(wb_gen[sn], wb_dst[sn],
                                               wb_base[sn], diff_log, f"{name}:{sn}")
            wb_gen.close(); wb_base.close()

        note = ""
        if name == "Fig1_age_sex_profile.xlsx" and "Sheet1" in wb_dst.sheetnames:
            n, rows = rewrite_fig1_sheet1(wb_dst["Sheet1"], args.rates_csv)
            total += n
            note = f" / Sheet1 {rows}行を差し替え"
        wb_dst.save(dst)
        wb_dst.close()
        print(f"  {name:32} 更新セル {total:>6,}{note}")
    if args.diff_csv and diff_log:
        pd.DataFrame(diff_log).to_csv(args.diff_csv, index=False, encoding="utf-8-sig")
        print(f"変更セル一覧: {args.diff_csv} ({len(diff_log)}件)")
    print(f"完了: {OUT_DIR}")


if __name__ == "__main__":
    main()
